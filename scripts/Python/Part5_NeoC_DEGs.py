# =============================================================================
# Part5_NeoC_DEGs.py
# SEA-AD CEP-IP/CMRF, Part 5, DEGs of M4-top Mic.13 cells, neocortex
#   Pseudobulk paired Wilcoxon signed-rank DEGs (donor level, BH-FDR) of
#   M4-top cells against non-M4-top cells of the same supertype: PVM_3 per
#   region (AnG, DFC, ITG, MTG, STG, V1C), PVM_2_3 and PVM_4 pooled across
#   regions. The M4 mask is read from the Part 2 joblib.
#
# Usage (from the SEA-AD folder, or set SEAAD_ROOT), after Part 2
#   python3 Part5_NeoC_DEGs.py
#
# Outputs ([1]_NeoC/Part5_NeoC_DEGs/): DEG workbook per comparison, master
#   workbook, volcano, dot and overlap figures (PNG, SVG).
# =============================================================================

import os, warnings, gc
warnings.filterwarnings("ignore")

import numpy as np
import pandas as pd
import scipy.sparse as sp
from scipy.stats import wilcoxon
from statsmodels.stats.multitest import multipletests
import matplotlib
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.cm as mcm
from matplotlib.lines import Line2D

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import random as _random_stdlib
GLOBAL_SEED = 42
np.random.seed(GLOBAL_SEED)
_random_stdlib.seed(GLOBAL_SEED)
os.environ.setdefault("PYTHONHASHSEED", str(GLOBAL_SEED))

matplotlib.use("Agg")   # no GUI

# == Output folder ============================================================
BASE_DIR   = os.environ.get("SEAAD_ROOT", os.getcwd())
ARM_FOLDER = r"[1]_NeoC"
ARM_TAG    = "NeoC"
ARM_NAME   = "neocortical"
PART_TAG   = "Part5"
DIR_TAG    = "DEGs"
ARM_ROOT   = os.path.join(BASE_DIR, ARM_FOLDER)
ARM_DIR    = os.path.join(ARM_ROOT, f"{PART_TAG}_{ARM_TAG}_{DIR_TAG}")
OUT_DIR    = ARM_DIR
os.makedirs(OUT_DIR, exist_ok=True)
print(f"[OUT] arm folder : {ARM_ROOT}")
print(f"[OUT] this part  : {ARM_DIR}")

print("=" * 70)
print("Part 5  -  DEG Analysis: M4-top Mic.13 vs same-supertype background")
print("(neocortical; PVM_2_3 and PVM_4 pooled across regions)")
print("=" * 70)


# =============================================================================
# SECTION 1: Constants, auto-loader, guardrails
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 1: Constants, auto-loader, guardrails")
print("=" * 70)

# == Raw h5ad file, for expression extraction =================================
FULL_H5AD   = ("SEAAD_Immune_HIP_MEC_LEC_ITG_MTG_FI_STG_DFC_AnG_V1C_"
               "RNAseq_final-nuclei.2026-06-22.h5ad")
PART2_DIR_TAG  = "M4"
PART2_DIR      = os.path.join(ARM_ROOT, f"Part2_{ARM_TAG}_{PART2_DIR_TAG}")
PART2_ENV_NAME = "Part2_neocortex_env.joblib"
_JOBLIB_PATH   = os.path.join(PART2_DIR, PART2_ENV_NAME)

CANDIDATE_SUPERTYPES = [
    "Micro-PVM_3-SEAAD",
    "Micro-PVM_2_3-SEAAD",
    "Micro-PVM_4-SEAAD",
]
SHORT_ST = {
    "Micro-PVM_3-SEAAD"  : "PVM_3",
    "Micro-PVM_2_3-SEAAD": "PVM_2_3",
    "Micro-PVM_4-SEAAD"  : "PVM_4",
}
NEOCORTEX_REGIONS = ["AnG", "DFC", "ITG", "MTG", "STG", "V1C"]
REGION_ORDER      = ["AnG", "DFC", "ITG", "MTG", "STG", "V1C"]

# == Analysis mode per supertype ==============================================
# "per_region" : separate DEG per region (PVM_3)
# "pooled"     : all regions combined into one DEG (PVM_2_3, PVM_4)
ST_MODE = {
    "PVM_3"  : "per_region",
    "PVM_2_3": "pooled",
    "PVM_4"  : "pooled",
}

PARENT_A = "M1_FC"
PARENT_B = "M2_Hybrid"

# DEG parameters
MIN_PCT        = 0.10
MIN_CELLS      = 3
MIN_CELLS_EXPR = 10
LOG2FC_THR     = 0.25
NORM_TARGET    = 1e4
PADJ_THR       = 0.05

# Expression gate for significant DEGs in the SIG sheet (pct_fg >= PCT_FG_GATE);
# volcano_rank = -log10(p_adj) x |log2FC| ranks the gated UP and DOWN lists.
PCT_FG_GATE = 0.15

# Region colours
SPECTRAL_6 = ["#D53E4F", "#FC8D59", "#FEE08B", "#99D594", "#3288BD", "#5E4FA2"]
REGION_PAL = {r: SPECTRAL_6[i] for i, r in enumerate(REGION_ORDER)}
ST_COLORS  = {
    "Micro-PVM_3-SEAAD"  : "#4393C3",
    "Micro-PVM_2_3-SEAAD": "#D6604D",
    "Micro-PVM_4-SEAAD"  : "#74C476",
}

# =============================================================================
# Automatic joblib loader (Part 2 variables)
# =============================================================================
_REQUIRED_VARS = [
    "obs_full", "M4_TOP_MASK", "N_M4_TOP",
    "M4_THR_M1", "M4_THR_M2", "TOP_PCT", "score_cols",
    "donor_col", "stype_col", "region_col",
]
_missing_vars = [v for v in _REQUIRED_VARS if v not in dir()]

if _missing_vars:
    print(f"  [AUTO-LOADER] {len(_missing_vars)} variables missing "
          f" -  loading from joblib ...")
    import joblib as _jl6
    if not os.path.exists(_JOBLIB_PATH):
        raise FileNotFoundError(
            f"\n  FATAL: joblib not found: {_JOBLIB_PATH}\n"
            f"  Run Part 2 first.")
    _env6 = _jl6.load(_JOBLIB_PATH)
    obs_full    = _env6["obs_full"]
    M4_TOP_MASK = _env6["M4_TOP_MASK"]
    N_M4_TOP    = _env6["N_M4_TOP"]
    M4_THR_M1   = _env6["M4_THR_M1"]
    M4_THR_M2   = _env6["M4_THR_M2"]
    TOP_PCT     = _env6["TOP_PCT"]
    score_cols  = _env6["score_cols"]
    donor_col   = _env6["donor_col"]
    stype_col   = _env6["stype_col"]
    region_col  = _env6["region_col"]
    del _env6; gc.collect()
    print("  [AUTO-LOADER] All variables loaded successfully.")
else:
    print(f"  [AUTO-LOADER] All {len(_REQUIRED_VARS)} variables already "
          f"in namespace  -  skipping load.")

TOP_PCT_STR = f"{TOP_PCT * 100:.1f}%"

# == M4_TOP_MASK integrity check ==============================================
print("\n  [M4_TOP_MASK INTEGRITY CHECK]")
assert isinstance(M4_TOP_MASK, np.ndarray) and M4_TOP_MASK.dtype == bool, \
    "M4_TOP_MASK must be a boolean numpy array"
assert len(M4_TOP_MASK) == len(obs_full), (
    f"M4_TOP_MASK length {len(M4_TOP_MASK)} != obs_full {len(obs_full)} "
    f" -  row alignment broken")
assert int(M4_TOP_MASK.sum()) == N_M4_TOP, (
    f"M4_TOP_MASK.sum()={M4_TOP_MASK.sum()} != N_M4_TOP={N_M4_TOP}")
print(f"    PASS: M4_TOP_MASK n={N_M4_TOP:,} cells, "
      f"M1_thr={M4_THR_M1:.6f}, M2_thr={M4_THR_M2:.6f}")

# == Arm guardrail ============================================================
print(f"\n  [{ARM_TAG} ARM GUARDRAIL]")
_nc_expected6 = set(NEOCORTEX_REGIONS)
_obs_regions6 = set(obs_full[region_col].astype(str).unique())
_unexpected6  = _obs_regions6 - _nc_expected6
if _unexpected6:
    raise AssertionError(
        f"\n  GUARDRAIL FAILED: obs_full carries regions that do not belong to "
        f"the {ARM_TAG} arm: {sorted(_unexpected6)}\n"
        f"  Expected exactly {sorted(_nc_expected6)}.")
_missing_nc6 = _nc_expected6 - _obs_regions6
if _missing_nc6:
    print(f"    WARNING: expected {ARM_NAME} regions absent: {sorted(_missing_nc6)}")
print(f"    PASS: obs_full regions = {sorted(_obs_regions6)}  "
      f"({ARM_NAME} only, n={len(obs_full):,} cells)")

# == Load adata (raw h5ad) ====================================================
print(f"\n  Loading adata from: {FULL_H5AD}")
import scanpy as sc
adata = sc.read_h5ad(FULL_H5AD, backed="r")
print(f"  adata: {adata.n_obs:,} cells × {adata.n_vars:,} genes")
assert "UMIs" in adata.layers, (
    "adata.layers['UMIs'] not found")
print(f"  adata.layers['UMIs'] confirmed present.")

# == Canonical M4 mask ========================================================
# The Part 2 mask is used as is.
m4_mask_global = M4_TOP_MASK
print(f"\n  m4_mask_global = canonical M4_TOP_MASK: n={m4_mask_global.sum():,} cells")
print(f"  TOP_PCT        : {TOP_PCT_STR}")
print(f"  obs_full cells : {len(obs_full):,}")
print(f"  adata cells    : {adata.n_obs:,}")
print(f"  adata genes    : {adata.n_vars:,}")


# =============================================================================
# SECTION 2: Expression matrix extraction helper
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 2: Expression extraction & normalisation helpers")
print("(Source: adata.layers['UMIs']  -  raw integer counts confirmed)")
print("=" * 70)

_OBS_NAME_TO_IDX = None

def _get_obs_idx_map(adata_obj):
    global _OBS_NAME_TO_IDX
    if _OBS_NAME_TO_IDX is None:
        _OBS_NAME_TO_IDX = pd.Series(
            np.arange(adata_obj.n_obs), index=adata_obj.obs_names)
    return _OBS_NAME_TO_IDX

def _load_raw_counts(adata_obj, obs_names_list):
    """
    Load raw UMI counts from adata.layers['UMIs'] for a list of obs_names.
    Three-tier approach: scanpy slice → h5py row-selective → full-matrix fallback.
    """
    n_genes  = adata_obj.n_vars
    n_obs    = adata_obj.n_obs
    obs_map  = _get_obs_idx_map(adata_obj)
    int_idx  = obs_map.loc[list(obs_names_list)].values
    sort_ord  = np.argsort(int_idx)
    unsort    = np.argsort(sort_ord)
    sorted_idx = int_idx[sort_ord]

    # Tier 1: scanpy layer slice
    try:
        layer_slice = adata_obj.layers["UMIs"][int_idx, :]
        if hasattr(layer_slice, "toarray"):
            return layer_slice.toarray().astype(np.float32)
        return np.array(layer_slice, dtype=np.float32)
    except Exception:
        pass

    # Tier 2: h5py row-selective read (memory efficient)
    try:
        import h5py
        fname = str(adata_obj.filename)
        with h5py.File(fname, "r") as f:
            grp         = f["layers"]["UMIs"]
            indptr_full = grp["indptr"][:]
            row_starts  = indptr_full[sorted_idx]
            row_ends    = indptr_full[sorted_idx + 1]
            n_req       = len(sorted_idx)
            X_dense     = np.zeros((n_req, n_genes), dtype=np.float32)
            i = 0
            while i < n_req:
                j = i + 1
                while j < n_req and sorted_idx[j] == sorted_idx[j-1] + 1:
                    j += 1
                start_nnz = int(row_starts[i])
                end_nnz   = int(row_ends[j - 1])
                if end_nnz > start_nnz:
                    chunk_data = grp["data"][start_nnz:end_nnz]
                    chunk_cols = grp["indices"][start_nnz:end_nnz]
                    for k in range(j - i):
                        r_start = int(row_starts[i + k]) - start_nnz
                        r_end   = int(row_ends[i + k])   - start_nnz
                        if r_end > r_start:
                            X_dense[i + k, chunk_cols[r_start:r_end]] = \
                                chunk_data[r_start:r_end]
                i = j
        return X_dense[unsort, :]
    except Exception as e2:
        pass

    # Tier 3: full-matrix fallback (last resort)
    try:
        import h5py
        print(f"  Tier 3: loading the full {n_obs}×{n_genes} UMI "
              f"matrix into RAM.")
        fname = str(adata_obj.filename)
        with h5py.File(fname, "r") as f:
            grp         = f["layers"]["UMIs"]
            data_arr    = grp["data"][:]
            indices_arr = grp["indices"][:]
            indptr_arr  = grp["indptr"][:]
        import scipy.sparse as _sp
        X_full  = _sp.csr_matrix(
            (data_arr, indices_arr, indptr_arr), shape=(n_obs, n_genes))
        X_dense = X_full[int_idx, :].toarray().astype(np.float32)
        del X_full, data_arr, indices_arr, indptr_arr; gc.collect()
        return X_dense
    except Exception as e3:
        raise RuntimeError(
            f"Cannot load layers['UMIs'] from adata.\n"
            f"Tier-2 failed: {e2}\nTier-3 failed: {e3}\n"
            f"File: {getattr(adata_obj, 'filename', 'unknown')}")

def _normalise_log1p(X_raw):
    row_sums = X_raw.sum(axis=1, keepdims=True)
    row_sums[row_sums == 0] = 1.0
    return np.log1p((X_raw / row_sums) * NORM_TARGET).astype(np.float32)

def _runtime_dtype_check(X_raw_sample):
    mx     = float(X_raw_sample.max())
    is_int = np.allclose(X_raw_sample, X_raw_sample.astype(int), atol=0.01)
    if not is_int or mx < 20:
        raise RuntimeError(
            f"\n  FATAL: layers['UMIs'] does not look like raw integer counts.\n"
            f"      max value = {mx:.4f}  |  integer-like = {is_int}\n"
            f"      Expected: integers, max >> 20 for UMI counts.\n")
    print(f"  [UMI DTYPE CHECK PASS] max={mx:.0f}, integer-like={is_int}"
          f" -> layers['UMIs'] confirmed as raw counts.")

print("  Expression helpers ready.")


# =============================================================================
# SECTION 3: Excel helpers
# =============================================================================
_tn  = Side(style="thin", color="BBBBBB")
_bdr = Border(left=_tn, right=_tn, top=_tn, bottom=_tn)
_bf  = Font(name="Calibri", size=9)
_af  = PatternFill("solid", fgColor="EFF3FB")
_ctr = Alignment(horizontal="center", vertical="center")
_lft = Alignment(horizontal="left",   vertical="center")

def _write_header(ws, headers, widths, row=1):
    hf  = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
    hfl = PatternFill("solid", fgColor="1F3864")
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = hf; c.fill = hfl; c.border = _bdr; c.alignment = _ctr
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[row].height = 24

def _safe_cell(ws, ri, ci, val):
    c = ws.cell(ri, ci)
    try:
        if isinstance(val, (bool, np.bool_)):      c.value = bool(val)
        elif isinstance(val, (float, np.floating)):
            c.value = None if np.isnan(val) else round(float(val), 6)
        elif isinstance(val, (int, np.integer)):   c.value = int(val)
        else:                                       c.value = str(val)
    except Exception:                               c.value = str(val)
    c.font = _bf; c.border = _bdr
    return c

def _pval_fill(p):
    if pd.isna(p): return PatternFill("solid", fgColor="F2F2F2")
    if p < 0.001:  return PatternFill("solid", fgColor="C6EFCE")
    if p < 0.01:   return PatternFill("solid", fgColor="A9D18E")
    if p < 0.05:   return PatternFill("solid", fgColor="E2EFDA")
    return PatternFill("solid", fgColor="F5F5F5")

print("  Excel helpers ready.")


# =============================================================================
# SECTION 4: Core DEG function
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 4: Core DEG function (pseudobulk paired Wilcoxon signed-rank, BH-FDR)")
print("=" * 70)

def run_deg(adata_obj, obs_full_df, fg_obs_names, bg_obs_names,
            gene_names, label="", min_pct=MIN_PCT,
            min_cells=MIN_CELLS, min_cells_expr=MIN_CELLS_EXPR,
            log2fc_thr=LOG2FC_THR):
    """
    Pseudobulk paired DEG: donor-level Wilcoxon signed-rank. For each gene:
      1. Average log-normalised expression per donor within fg and bg
      2. Restrict to donors present in BOTH fg and bg (paired design)
      3. Per-donor difference (fg_mean - bg_mean)
      4. One-sample Wilcoxon signed-rank on the differences

    pct_fg/pct_bg are cell level; log2FC is from the donor means.
    """
    n_fg = len(fg_obs_names)
    n_bg = len(bg_obs_names)

    if n_fg < min_cells or n_bg < min_cells:
        print(f"    [{label}] SKIP: fg={n_fg}, bg={n_bg} (< min_cells={min_cells})")
        return pd.DataFrame()

    print(f"    [{label}] fg={n_fg:,}, bg={n_bg:,}  -  loading expression ...")

    X_fg_raw = _load_raw_counts(adata_obj, list(fg_obs_names))
    X_bg_raw = _load_raw_counts(adata_obj, list(bg_obs_names))

    # Dtype check once
    if not getattr(run_deg, "_dtype_checked", False):
        _runtime_dtype_check(X_fg_raw[:min(50, n_fg), :])
        run_deg._dtype_checked = True

    X_fg = _normalise_log1p(X_fg_raw)
    X_bg = _normalise_log1p(X_bg_raw)
    del X_fg_raw, X_bg_raw; gc.collect()

    n_genes = X_fg.shape[1]

    # Cell-level pct (retained for reporting / gating)
    pct_fg  = (X_fg > 0).mean(axis=0)
    pct_bg  = (X_bg > 0).mean(axis=0)

    # Gene filters (cell level)
    pct_mask = (pct_fg >= min_pct) | (pct_bg >= min_pct)
    n_expr_combined = (X_fg > 0).sum(axis=0) + (X_bg > 0).sum(axis=0)
    expr_mask = n_expr_combined >= min_cells_expr
    gene_mask = pct_mask & expr_mask
    n_tested  = int(gene_mask.sum())
    print(f"    [{label}] genes passing pct filter : {int(pct_mask.sum()):,}")
    print(f"    [{label}] genes passing both filters: {n_tested:,}")

    if n_tested == 0:
        print(f"    [{label}] No genes pass filters  -  skipping.")
        return pd.DataFrame()

    X_fg_f  = X_fg[:, gene_mask]
    X_bg_f  = X_bg[:, gene_mask]
    genes_f = gene_names[gene_mask]
    del X_fg, X_bg; gc.collect()

    # == Pseudobulk: mean expression per donor ================================
    donors_fg = obs_full_df.loc[list(fg_obs_names), donor_col].values
    donors_bg = obs_full_df.loc[list(bg_obs_names), donor_col].values

    shared_donors = sorted(set(donors_fg) & set(donors_bg))
    n_shared = len(shared_donors)
    print(f"    [{label}] donors: fg={len(set(donors_fg))}, "
          f"bg={len(set(donors_bg))}, shared(paired)={n_shared}")

    MIN_DONORS_DEG = 5
    if n_shared < MIN_DONORS_DEG:
        print(f"    [{label}] SKIP: only {n_shared} shared donors "
              f"(< {MIN_DONORS_DEG})")
        return pd.DataFrame()

    # Donor-mean pseudobulk matrices (n_shared x n_tested)
    pb_fg = np.zeros((n_shared, n_tested), dtype=np.float64)
    pb_bg = np.zeros((n_shared, n_tested), dtype=np.float64)
    for di, d in enumerate(shared_donors):
        mask_fg_d = donors_fg == d
        mask_bg_d = donors_bg == d
        if mask_fg_d.sum() > 0:
            pb_fg[di, :] = X_fg_f[mask_fg_d, :].mean(axis=0)
        if mask_bg_d.sum() > 0:
            pb_bg[di, :] = X_bg_f[mask_bg_d, :].mean(axis=0)

    pb_diff = pb_fg - pb_bg

    # log2FC from donor-mean pseudobulk
    mean_fg_pb = pb_fg.mean(axis=0)
    mean_bg_pb = pb_bg.mean(axis=0)
    log2fc = (mean_fg_pb - mean_bg_pb) / np.log(2)

    # == Paired Wilcoxon signed-rank test per gene ============================
    p_vals = np.ones(n_tested, dtype=np.float64)
    for gi in range(n_tested):
        diffs = pb_diff[:, gi]
        if np.all(diffs == 0):
            continue
        try:
            _, p_vals[gi] = wilcoxon(diffs, alternative="two-sided")
        except Exception:
            p_vals[gi] = 1.0

    _, p_adj, _, _ = multipletests(p_vals, method="fdr_bh")

    result_df = pd.DataFrame({
        "gene"     : genes_f,
        "log2FC"   : np.round(log2fc.astype(np.float64), 4),
        "pct_fg"   : np.round(pct_fg[gene_mask].astype(np.float64), 4),
        "pct_bg"   : np.round(pct_bg[gene_mask].astype(np.float64), 4),
        "p_raw"    : np.round(p_vals, 8),
        "p_adj"    : np.round(p_adj,  8),
        "n_donors" : n_shared,
    })

    # == Significance flag (before the expression gate) =======================
    result_df["significant"] = (
        (result_df["p_adj"] < PADJ_THR) &
        (result_df["log2FC"].abs() >= LOG2FC_THR))
    result_df["direction"] = np.where(
        result_df["log2FC"] > 0, "UP_in_M4top", "DOWN_in_M4top")

    # == volcano_rank = -log10(p_adj) x |log2FC| ==============================
    _neg_log10_padj = -np.log10(np.clip(result_df["p_adj"].values, 1e-300, 1.0))
    _abs_lfc        = result_df["log2FC"].abs().values
    _sig_arr        = result_df["significant"].values
    _pct_fg_arr     = result_df["pct_fg"].values
    _gate_arr       = _pct_fg_arr >= PCT_FG_GATE
    _rank_score     = np.where(
        _sig_arr & _gate_arr,
        np.round(_neg_log10_padj * _abs_lfc, 6),
        np.nan)
    result_df["volcano_rank"] = _rank_score

    # == Sort: significant and gated first, by volcano_rank ===================
    result_df["abs_log2FC"] = _abs_lfc
    result_df = result_df.sort_values(
        ["significant", "volcano_rank", "abs_log2FC"],
        ascending=[False, False, False],
        na_position="last").reset_index(drop=True)
    result_df.drop(columns=["abs_log2FC"], inplace=True)

    # == Counts after the pct_fg gate (SIG sheet) =============================
    _sig_gated = result_df["significant"] & (result_df["pct_fg"] >= PCT_FG_GATE)
    n_sig      = int(result_df["significant"].sum())
    n_sig_g    = int(_sig_gated.sum())
    n_up       = int((_sig_gated & (result_df["log2FC"] > 0)).sum())
    n_down     = int((_sig_gated & (result_df["log2FC"] < 0)).sum())
    n_dropped  = n_sig - n_sig_g
    print(f"    [{label}] significant DEGs (p_adj<{PADJ_THR}, |log2FC|>={LOG2FC_THR}): "
          f"{n_sig}  \u2192 after pct_fg>={PCT_FG_GATE} gate: {n_sig_g} "
          f"(UP={n_up}, DOWN={n_down}, dropped={n_dropped})  "
          f"[n_donors={n_shared}]")

    del X_fg_f, X_bg_f, pb_fg, pb_bg, pb_diff; gc.collect()
    return result_df

print("  DEG function ready.")


# =============================================================================
# SECTION 5: Run all comparisons
# PVM_3 per region; PVM_2_3 and PVM_4 pooled across regions. FG = M4-top cells,
# BG = non-M4-top cells of the same supertype (and region).
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 5: Running all DEG comparisons")
print("  PVM_3   → per_region (one comparison per region)")
print("  PVM_2_3 → pooled across all regions (1 comparison)")
print("  PVM_4   → pooled across all regions (1 comparison)")
print("=" * 70)

cell_labels     = obs_full[stype_col].astype(str).values
region_vals     = obs_full[region_col].astype(str).values
gene_names      = np.array(adata.var_names.tolist())
obs_full_index  = obs_full.index.astype(str).values

all_results  = {}
summary_rows = []

for st_full in CANDIDATE_SUPERTYPES:
    short    = SHORT_ST[st_full]
    mode     = ST_MODE[short]
    st_mask  = cell_labels == st_full

    print(f"\n  ── {short} ({st_full}) | mode={mode} ──────────────────────────")
    print(f"     Total cells in supertype: {int(st_mask.sum()):,}")

    # == Per-region mode (PVM_3) ==============================================
    if mode == "per_region":
        for region in REGION_ORDER:
            reg_mask = region_vals == region

            fg_mask = st_mask & reg_mask & m4_mask_global
            bg_mask = st_mask & reg_mask & (~m4_mask_global)

            n_fg = int(fg_mask.sum())
            n_bg = int(bg_mask.sum())
            label = f"{short}_{region}"
            print(f"\n  {label}: fg={n_fg} (M4-top), "
                  f"bg={n_bg} (non-M4-top, same ST+region)")

            if n_fg < MIN_CELLS or n_bg < MIN_CELLS:
                print(f"    SKIP (insufficient cells)")
                summary_rows.append({
                    "supertype": short, "region": region,
                    "n_fg": n_fg, "n_bg": n_bg,
                    "n_genes_tested": 0, "n_sig": 0,
                    "n_up": 0, "n_down": 0,
                    "mode": mode, "status": "SKIPPED"})
                continue

            fg_obs = obs_full_index[fg_mask]
            bg_obs = obs_full_index[bg_mask]

            result_df = run_deg(
                adata_obj=adata,
                obs_full_df=obs_full,
                fg_obs_names=fg_obs,
                bg_obs_names=bg_obs,
                gene_names=gene_names,
                label=label)

            if len(result_df) == 0:
                summary_rows.append({
                    "supertype": short, "region": region,
                    "n_fg": n_fg, "n_bg": n_bg,
                    "n_genes_tested": 0, "n_sig": 0,
                    "n_up": 0, "n_down": 0,
                    "mode": mode, "status": "NO_GENES"})
                continue

            all_results[(short, region)] = result_df

            # SIG_DEGs: significant and gated, UP then DOWN by volcano_rank
            _sig_g   = result_df[result_df["significant"] & (result_df["pct_fg"] >= PCT_FG_GATE)]
            _sig_g_up = _sig_g[_sig_g["log2FC"] > 0].sort_values("volcano_rank", ascending=False)
            _sig_g_dn = _sig_g[_sig_g["log2FC"] < 0].sort_values("volcano_rank", ascending=False)
            _sig_g_out = pd.concat([_sig_g_up, _sig_g_dn], ignore_index=True)
            indiv_fn = os.path.join(OUT_DIR, f"DEG_{label}.xlsx")
            with pd.ExcelWriter(indiv_fn, engine="openpyxl") as _ew:
                _sig_g_out.to_excel(_ew, sheet_name="SIG_DEGs", index=False)
                result_df.to_excel(_ew, sheet_name="ALL_tested", index=False)
            print(f"    Saved {indiv_fn}")

            n_sig  = int(_sig_g.shape[0])
            n_up   = int((_sig_g["log2FC"] > 0).sum())
            n_down = int((_sig_g["log2FC"] < 0).sum())
            summary_rows.append({
                "supertype": short, "region": region,
                "n_fg": n_fg, "n_bg": n_bg,
                "n_genes_tested": int(len(result_df)),
                "n_sig": n_sig, "n_up": n_up, "n_down": n_down,
                "mode": mode, "status": "OK"})

    # == Pooled mode (PVM_2_3, PVM_4) =========================================
    elif mode == "pooled":
        # Pool across all regions
        all_reg_mask = np.zeros(len(st_mask), dtype=bool)
        for region in REGION_ORDER:
            all_reg_mask |= (region_vals == region)

        fg_mask = st_mask & all_reg_mask & m4_mask_global
        bg_mask = st_mask & all_reg_mask & (~m4_mask_global)

        n_fg = int(fg_mask.sum())
        n_bg = int(bg_mask.sum())
        # Region breakdown for reporting
        region_breakdown_fg = {r: int((st_mask & (region_vals == r) & m4_mask_global).sum())
                                for r in REGION_ORDER}
        region_breakdown_bg = {r: int((st_mask & (region_vals == r) & (~m4_mask_global)).sum())
                                for r in REGION_ORDER}

        label = f"{short}_ALL"
        print(f"\n  {label}: fg={n_fg} (M4-top, ALL regions pooled), "
              f"bg={n_bg} (non-M4-top, ALL regions pooled)")
        print(f"    fg per region: {region_breakdown_fg}")
        print(f"    bg per region: {region_breakdown_bg}")

        if n_fg < MIN_CELLS or n_bg < MIN_CELLS:
            print(f"    SKIP (insufficient cells)")
            summary_rows.append({
                "supertype": short, "region": "ALL",
                "n_fg": n_fg, "n_bg": n_bg,
                "n_genes_tested": 0, "n_sig": 0,
                "n_up": 0, "n_down": 0,
                "mode": mode, "status": "SKIPPED"})
        else:
            fg_obs = obs_full_index[fg_mask]
            bg_obs = obs_full_index[bg_mask]

            result_df = run_deg(
                adata_obj=adata,
                obs_full_df=obs_full,
                fg_obs_names=fg_obs,
                bg_obs_names=bg_obs,
                gene_names=gene_names,
                label=label)

            if len(result_df) == 0:
                summary_rows.append({
                    "supertype": short, "region": "ALL",
                    "n_fg": n_fg, "n_bg": n_bg,
                    "n_genes_tested": 0, "n_sig": 0,
                    "n_up": 0, "n_down": 0,
                    "mode": mode, "status": "NO_GENES"})
            else:
                all_results[(short, "ALL")] = result_df

                # SIG_DEGs: significant and gated, UP then DOWN by volcano_rank
                _sig_g_p   = result_df[result_df["significant"] & (result_df["pct_fg"] >= PCT_FG_GATE)]
                _sig_g_up_p = _sig_g_p[_sig_g_p["log2FC"] > 0].sort_values("volcano_rank", ascending=False)
                _sig_g_dn_p = _sig_g_p[_sig_g_p["log2FC"] < 0].sort_values("volcano_rank", ascending=False)
                _sig_g_out_p = pd.concat([_sig_g_up_p, _sig_g_dn_p], ignore_index=True)
                indiv_fn = os.path.join(OUT_DIR, f"DEG_{label}.xlsx")
                with pd.ExcelWriter(indiv_fn, engine="openpyxl") as _ew:
                    _sig_g_out_p.to_excel(_ew, sheet_name="SIG_DEGs", index=False)
                    result_df.to_excel(_ew, sheet_name="ALL_tested", index=False)
                print(f"    Saved {indiv_fn}")

                n_sig  = int(_sig_g_p.shape[0])
                n_up   = int((_sig_g_p["log2FC"] > 0).sum())
                n_down = int((_sig_g_p["log2FC"] < 0).sum())
                summary_rows.append({
                    "supertype": short, "region": "ALL",
                    "n_fg": n_fg, "n_bg": n_bg,
                    "n_genes_tested": int(len(result_df)),
                    "n_sig": n_sig, "n_up": n_up, "n_down": n_down,
                    "mode": mode, "status": "OK"})

summary_df = pd.DataFrame(summary_rows)
print("\n\n  ── DEG SUMMARY ──────────────────────────────────────────────")
print(f"  {'Comparison':<22} {'mode':<12} {'fg':>6} {'bg':>6} "
      f"{'n_sig':>7} {'UP':>5} {'DOWN':>5}  status")
print("  " + "-" * 75)
for _, row_s in summary_df.iterrows():
    lbl_s = f"{row_s['supertype']}_{row_s['region']}"
    print(f"  {lbl_s:<22} {row_s['mode']:<12} {row_s['n_fg']:>6} {row_s['n_bg']:>6} "
          f"{row_s['n_sig']:>7} {row_s['n_up']:>5} {row_s['n_down']:>5}  "
          f"{row_s['status']}")


# =============================================================================
# SECTION 6: Master Excel workbook
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 6: Master Excel workbook")
print("=" * 70)

fp_master_xl = os.path.join(OUT_DIR, "DEG_master_results.xlsx")
wb_master    = Workbook(); wb_master.remove(wb_master.active)

ws_sum = wb_master.create_sheet("00_summary")
_write_header(ws_sum,
              ["Supertype", "Region", "Mode",
               "N fg (M4-top)", "N bg (non-M4-top)",
               "Genes tested",
               f"N sig (p_adj<0.05, |log2FC|>={LOG2FC_THR}, pct_fg>={PCT_FG_GATE})",
               "N up (M4-top)", "N down (M4-top)", "Status"],
              [14, 10, 12, 16, 18, 14, 42, 14, 16, 12])
for ri_s, (_, row_s2) in enumerate(summary_df.iterrows(), 2):
    vals_s = [row_s2["supertype"], row_s2["region"], row_s2["mode"],
              row_s2["n_fg"], row_s2["n_bg"],
              row_s2["n_genes_tested"], row_s2["n_sig"],
              row_s2["n_up"], row_s2["n_down"], row_s2["status"]]
    for ci_s, v_s in enumerate(vals_s, 1):
        c_s = _safe_cell(ws_sum, ri_s, ci_s, v_s)
        c_s.alignment = _lft if ci_s <= 3 else _ctr
        if ri_s % 2 == 0: c_s.fill = _af
    if row_s2["n_sig"] > 0:
        ws_sum.cell(ri_s, 7).fill = PatternFill("solid", fgColor="C6EFCE")
        ws_sum.cell(ri_s, 7).font = Font(name="Calibri", size=9, bold=True)

# Columns of the SIG and ALL_tested sheets
DEG_HEADERS = ["gene", "log2FC", "pct_fg", "pct_bg",
               "p_raw", "p_adj", "significant", "direction", "volcano_rank"]
DEG_WIDTHS  = [18, 10, 10, 10, 14, 14, 12, 16, 14]

# volcano_rank cell fill
def _vrank_fill(vr):
    """Light gold fill scaled to volcano_rank magnitude (capped at 20)."""
    if pd.isna(vr) or vr <= 0:
        return None
    intensity = min(float(vr) / 20.0, 1.0)   # saturates at rank 20
    # white to gold
    g_hex = format(int(0xD7 + (0xFF - 0xD7) * (1 - intensity)), "02X")
    b_hex = format(int(0x00 + (0xFF - 0x00) * (1 - intensity)), "02X")
    return PatternFill("solid", fgColor=f"FF{g_hex}{b_hex}")

for (short, region), df_r in sorted(all_results.items()):
    # Sheet name from the region label ("ALL" when pooled)
    sheet_name = f"{short[:8]}_{region}"

    # SIG sheet: significant DEGs passing the pct_fg gate
    ws_sig = wb_master.create_sheet(f"{sheet_name}_SIG")
    _write_header(ws_sig, DEG_HEADERS, DEG_WIDTHS)
    sig_df = df_r[
        df_r["significant"] & (df_r["pct_fg"] >= PCT_FG_GATE)
    ].copy()
    # Sort UP and DOWN separately by volcano_rank desc, then concatenate
    _up_df  = sig_df[sig_df["log2FC"] > 0].sort_values("volcano_rank", ascending=False)
    _dn_df  = sig_df[sig_df["log2FC"] < 0].sort_values("volcano_rank", ascending=False)
    sig_df  = pd.concat([_up_df, _dn_df], ignore_index=True)
    for ri_d, (_, row_d) in enumerate(sig_df.iterrows(), 2):
        for ci_d, col_d in enumerate(DEG_HEADERS, 1):
            c_d = _safe_cell(ws_sig, ri_d, ci_d, row_d[col_d])
            c_d.alignment = _lft if ci_d == 1 else _ctr
            if col_d == "p_adj":
                try: c_d.fill = _pval_fill(float(row_d["p_adj"]))
                except: pass
            elif col_d == "volcano_rank":
                try:
                    vr_fill = _vrank_fill(row_d["volcano_rank"])
                    if vr_fill: c_d.fill = vr_fill
                except: pass
            elif ri_d % 2 == 0:
                c_d.fill = _af
            if col_d == "direction":
                if row_d["direction"] == "UP_in_M4top":
                    c_d.font = Font(name="Calibri", size=9, color="C00000", bold=True)
                else:
                    c_d.font = Font(name="Calibri", size=9, color="2B6CB0", bold=True)

    # ALL_tested sheet: every gene tested (no gate, no significance filter)
    ws_all = wb_master.create_sheet(f"{sheet_name}_ALL")
    _write_header(ws_all, DEG_HEADERS, DEG_WIDTHS)
    for ri_a, (_, row_a) in enumerate(df_r.iterrows(), 2):
        for ci_a, col_a in enumerate(DEG_HEADERS, 1):
            c_a = _safe_cell(ws_all, ri_a, ci_a, row_a[col_a])
            c_a.alignment = _lft if ci_a == 1 else _ctr
            if ri_a % 2 == 0: c_a.fill = _af

wb_master.save(fp_master_xl)
print(f"  Saved {fp_master_xl}")


# =============================================================================
# SECTION 7: Figure 1  -  Volcano plots
# Fig 1a: PVM_3 per region; Fig 1b: PVM_2_3 and PVM_4 pooled.
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 7: Figure 1  -  Volcano plots")
print("  Fig1a: PVM_3 per-region (1×6 grid)")
print("  Fig1b: PVM_2_3 and PVM_4 pooled (1×2 grid)")
print("=" * 70)

def _draw_volcano(ax_v, df_v, n_fg_v, n_bg_v, title_str):
    """Shared volcano rendering helper."""
    lfc_v  = df_v["log2FC"].values
    plog_v = -np.log10(np.clip(df_v["p_adj"].values, 1e-300, 1.0))
    sig_v  = df_v["significant"].values
    up_v   = sig_v & (lfc_v > 0)
    dn_v   = sig_v & (lfc_v < 0)
    ns_v   = ~sig_v

    ax_v.set_facecolor("white")
    ax_v.spines[["top", "right"]].set_visible(False)

    ax_v.scatter(lfc_v[ns_v], plog_v[ns_v], c="#CCCCCC", s=4,
                 alpha=0.40, linewidths=0, rasterized=True, zorder=2)
    if up_v.sum() > 0:
        ax_v.scatter(lfc_v[up_v], plog_v[up_v], c="#C00000", s=10,
                     alpha=0.80, linewidths=0, rasterized=True, zorder=4)
    if dn_v.sum() > 0:
        ax_v.scatter(lfc_v[dn_v], plog_v[dn_v], c="#2B6CB0", s=10,
                     alpha=0.80, linewidths=0, rasterized=True, zorder=4)

    ax_v.axhline(-np.log10(PADJ_THR), color="#555555", lw=0.8, ls="--", alpha=0.7)
    ax_v.axvline( LOG2FC_THR, color="#AAAAAA", lw=0.7, ls=":", alpha=0.7)
    ax_v.axvline(-LOG2FC_THR, color="#AAAAAA", lw=0.7, ls=":", alpha=0.7)

    _sig_sub = df_v[df_v["significant"]]
    top_genes_v = pd.concat([
        _sig_sub.nlargest(10,  "log2FC"),
        _sig_sub.nsmallest(10, "log2FC")
    ]).drop_duplicates(subset="gene")
    for _, gr in top_genes_v.iterrows():
        ax_v.text(gr["log2FC"],
                  -np.log10(max(gr["p_adj"], 1e-300)) + 0.15,
                  gr["gene"], ha="center", va="bottom",
                  fontsize=5.5, color="#111111", zorder=6, clip_on=True)

    n_sig_v = int(sig_v.sum())
    n_up_v  = int(up_v.sum())
    n_dn_v  = int(dn_v.sum())
    ax_v.set_title(
        f"{title_str}\n"
        f"fg={n_fg_v:,} vs bg={n_bg_v:,}  |  "
        f"DEG={n_sig_v} (↑{n_up_v} ↓{n_dn_v})",
        fontsize=9, fontweight="bold", pad=6)
    ax_v.set_xlabel("log₂FC (M4-top vs background)", fontsize=8)
    ax_v.set_ylabel("−log₁₀(p_adj)", fontsize=8)
    ax_v.tick_params(labelsize=7)


_leg_handles_v = [
    Line2D([0],[0], marker="o", color="none", markerfacecolor="#C00000",
           markersize=7, label=f"UP in M4-top  (p_adj<{PADJ_THR}, |log2FC|>={LOG2FC_THR})"),
    Line2D([0],[0], marker="o", color="none", markerfacecolor="#2B6CB0",
           markersize=7, label="DOWN in M4-top"),
    Line2D([0],[0], marker="o", color="none", markerfacecolor="#CCCCCC",
           markersize=7, label="Not significant"),
]

# == Fig 1a: PVM_3 per region =================================================
PVM3_SHORT = "PVM_3"
pvm3_regions = [r for r in REGION_ORDER if (PVM3_SHORT, r) in all_results or
                any((summary_df["supertype"] == PVM3_SHORT) &
                    (summary_df["region"] == r))]

np.random.seed(GLOBAL_SEED); _random_stdlib.seed(GLOBAL_SEED)
fig1a, axes1a = plt.subplots(
    1, len(REGION_ORDER),
    figsize=(4.8 * len(REGION_ORDER), 5.0),
    facecolor="white", squeeze=False)

for ci_v, region_v in enumerate(REGION_ORDER):
    ax_v = axes1a[0, ci_v]
    ax_v.set_facecolor("white")
    ax_v.spines[["top", "right"]].set_visible(False)
    row_meta = summary_df[
        (summary_df["supertype"] == PVM3_SHORT) &
        (summary_df["region"]    == region_v)]

    if len(row_meta) == 0 or row_meta.iloc[0]["status"] != "OK":
        status_txt = row_meta.iloc[0]["status"] if len(row_meta) > 0 else "NOT RUN"
        n_fg_txt   = int(row_meta.iloc[0]["n_fg"]) if len(row_meta) > 0 else 0
        n_bg_txt   = int(row_meta.iloc[0]["n_bg"]) if len(row_meta) > 0 else 0
        ax_v.text(0.5, 0.5,
                  f"{status_txt}\n(fg={n_fg_txt}, bg={n_bg_txt})",
                  ha="center", va="center", transform=ax_v.transAxes,
                  fontsize=9, color="#888888")
        ax_v.set_title(f"{PVM3_SHORT} | {region_v}", fontsize=10, fontweight="bold")
        continue

    _draw_volcano(ax_v,
                  all_results[(PVM3_SHORT, region_v)],
                  int(row_meta.iloc[0]["n_fg"]),
                  int(row_meta.iloc[0]["n_bg"]),
                  f"{PVM3_SHORT} | {region_v}")

fig1a.legend(handles=_leg_handles_v, fontsize=9, loc="lower center", ncol=3,
             framealpha=0.95, edgecolor="#CCCCCC", bbox_to_anchor=(0.5, -0.04))
fig1a.suptitle(
    f"Fig 1a  -  Volcano plots: PVM_3 (per-region, neocortical)\n"
    f"M4-top Mic.13 (top {TOP_PCT_STR}) vs same-supertype non-M4-top background  |  "
    f"pseudobulk Wilcoxon signed-rank + BH-FDR  |  log2FC>={LOG2FC_THR}, p_adj<{PADJ_THR}",
    fontsize=11, y=1.03)
plt.tight_layout(h_pad=3.5, w_pad=2.5)
fn1a = "Fig1a_volcano_PVM3_perregion_neocortex.png"
fig1a.savefig(os.path.join(OUT_DIR, fn1a), dpi=150,
              bbox_inches="tight", facecolor="white")
fig1a.savefig(os.path.join(OUT_DIR, fn1a[:-4] + ".svg"),
              format="svg", bbox_inches="tight", facecolor="white")
plt.show(); plt.pause(0.5)
print(f"  Saved {fn1a}")

# == Fig 1b: PVM_2_3 and PVM_4 pooled =========================================
POOLED_SHORTS = ["PVM_2_3", "PVM_4"]
n_pooled      = len(POOLED_SHORTS)

np.random.seed(GLOBAL_SEED); _random_stdlib.seed(GLOBAL_SEED)
fig1b, axes1b = plt.subplots(
    1, n_pooled,
    figsize=(5.5 * n_pooled, 5.0),
    facecolor="white", squeeze=False)

for ci_p, short_p in enumerate(POOLED_SHORTS):
    ax_p = axes1b[0, ci_p]
    ax_p.set_facecolor("white")
    ax_p.spines[["top", "right"]].set_visible(False)
    row_meta_p = summary_df[
        (summary_df["supertype"] == short_p) &
        (summary_df["region"]    == "ALL")]

    if len(row_meta_p) == 0 or row_meta_p.iloc[0]["status"] != "OK":
        status_txt = row_meta_p.iloc[0]["status"] if len(row_meta_p) > 0 else "NOT RUN"
        n_fg_txt   = int(row_meta_p.iloc[0]["n_fg"]) if len(row_meta_p) > 0 else 0
        n_bg_txt   = int(row_meta_p.iloc[0]["n_bg"]) if len(row_meta_p) > 0 else 0
        ax_p.text(0.5, 0.5,
                  f"{status_txt}\n(fg={n_fg_txt}, bg={n_bg_txt})",
                  ha="center", va="center", transform=ax_p.transAxes,
                  fontsize=9, color="#888888")
        ax_p.set_title(f"{short_p} | ALL regions (pooled)", fontsize=10, fontweight="bold")
        continue

    _draw_volcano(ax_p,
                  all_results[(short_p, "ALL")],
                  int(row_meta_p.iloc[0]["n_fg"]),
                  int(row_meta_p.iloc[0]["n_bg"]),
                  f"{short_p} | ALL regions (AnG+DFC+ITG+MTG+STG+V1C pooled)")

fig1b.legend(handles=_leg_handles_v, fontsize=9, loc="lower center", ncol=3,
             framealpha=0.95, edgecolor="#CCCCCC", bbox_to_anchor=(0.5, -0.04))
fig1b.suptitle(
    f"Fig 1b  -  Volcano plots: PVM_2_3 & PVM_4 (ALL neocortical regions pooled)\n"
    f"M4-top Mic.13 (top {TOP_PCT_STR}) vs same-supertype non-M4-top background  |  "
    f"pseudobulk Wilcoxon signed-rank + BH-FDR  |  log2FC>={LOG2FC_THR}, p_adj<{PADJ_THR}",
    fontsize=11, y=1.03)
plt.tight_layout(h_pad=3.5, w_pad=2.5)
fn1b = "Fig1b_volcano_PVM23_PVM4_pooled_neocortex.png"
fig1b.savefig(os.path.join(OUT_DIR, fn1b), dpi=150,
              bbox_inches="tight", facecolor="white")
fig1b.savefig(os.path.join(OUT_DIR, fn1b[:-4] + ".svg"),
              format="svg", bbox_inches="tight", facecolor="white")
plt.show(); plt.pause(0.5)
print(f"  Saved {fn1b}")


# =============================================================================
# SECTION 8: Figure 2  -  Top-20 DEG dot plots per supertype
# PVM_3 one row per region; PVM_2_3 and PVM_4 one pooled row.
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 8: Figure 2  -  Top-DEG dot plots per supertype")
print("=" * 70)

def _dot_plot(short_dp, keys_dp, y_labels):
    """
    Render a dot plot for a given supertype.
    keys_dp   : list of (short, region) tuples in row order
    y_labels  : list of strings for y-axis tick labels (same order)
    """
    gene_pool = []
    for key_dp in keys_dp:
        sig_dp = all_results[key_dp][all_results[key_dp]["significant"]]
        if len(sig_dp) == 0:
            continue
        gene_pool.extend(sig_dp.nlargest(10,  "log2FC")["gene"].tolist())
        gene_pool.extend(sig_dp.nsmallest(10, "log2FC")["gene"].tolist())

    if not gene_pool:
        print(f"  {short_dp}: no significant DEGs  -  skipping dot plot.")
        return

    from collections import Counter
    top_genes_dp = [g for g, _ in Counter(gene_pool).most_common(30)]
    n_genes_dp   = len(top_genes_dp)
    n_rows_dp    = len(keys_dp)

    np.random.seed(GLOBAL_SEED); _random_stdlib.seed(GLOBAL_SEED)
    fig2dp, ax2dp = plt.subplots(
        1, 1,
        figsize=(max(10, n_genes_dp * 0.55 + 2),
                 max(3.5, n_rows_dp * 0.85 + 2)),
        facecolor="white")
    ax2dp.set_facecolor("white")
    ax2dp.spines[["top", "right", "bottom", "left"]].set_visible(False)

    lfc_all = [all_results[k].set_index("gene").loc[g, "log2FC"]
               for k in keys_dp for g in top_genes_dp
               if g in all_results[k]["gene"].values]
    lfc_abs_max = max(0.5, max(abs(v) for v in lfc_all)) if lfc_all else 1.0
    norm_dp = matplotlib.colors.TwoSlopeNorm(
        vmin=-lfc_abs_max, vcenter=0, vmax=lfc_abs_max)
    cmap_dp = matplotlib.cm.RdBu_r

    for xi_dp, gene_dp in enumerate(top_genes_dp):
        for yi_dp, key_dp in enumerate(keys_dp):
            df_key = all_results[key_dp].set_index("gene")
            if gene_dp not in df_key.index:
                continue
            row_dp = df_key.loc[gene_dp]
            lfc_v2 = float(row_dp["log2FC"])
            pct_v  = float(row_dp["pct_fg"])
            sig_v2 = bool(row_dp["significant"])
            ax2dp.scatter(xi_dp, -yi_dp,
                          s=max(pct_v * 400, 2), c=[cmap_dp(norm_dp(lfc_v2))],
                          alpha=0.90,
                          linewidths=0.8 if sig_v2 else 0.0,
                          edgecolors="black" if sig_v2 else "none", zorder=3)

    ax2dp.set_xticks(range(n_genes_dp))
    ax2dp.set_xticklabels(top_genes_dp, rotation=45, ha="right",
                           fontsize=7.5, style="italic")
    ax2dp.set_yticks(range(0, -n_rows_dp, -1))
    ax2dp.set_yticklabels(y_labels, fontsize=9)
    ax2dp.set_xlim(-0.6, n_genes_dp - 0.4)
    ax2dp.set_ylim(-n_rows_dp + 0.4, 0.6)

    sm_dp = matplotlib.cm.ScalarMappable(cmap=cmap_dp, norm=norm_dp)
    sm_dp.set_array([])
    cbar_dp = fig2dp.colorbar(sm_dp, ax=ax2dp, fraction=0.025, pad=0.02,
                               shrink=0.7, aspect=20)
    cbar_dp.set_label("log₂FC (M4-top vs bg)", fontsize=9)
    cbar_dp.ax.tick_params(labelsize=8)

    for pct_leg, lbl_leg in [(0.1, "10%"), (0.5, "50%"), (1.0, "100%")]:
        ax2dp.scatter([], [], s=pct_leg * 400, c="#888888", alpha=0.7,
                      label=f"pct_fg = {lbl_leg}")
    ax2dp.legend(title="% expressed (fg)", fontsize=7.5, title_fontsize=8,
                 loc="upper right", framealpha=0.92, edgecolor="#CCCCCC",
                 bbox_to_anchor=(1.18, 1.0))

    mode_str = ST_MODE.get(short_dp, "")
    region_str = (", ".join(y_labels) if mode_str == "per_region"
                  else "ALL regions (AnG+DFC+ITG+MTG+STG+V1C pooled)")
    fig2dp.suptitle(
        f"{short_dp}  -  Top DEGs dot plot (top {TOP_PCT_STR} M4-top Mic.13)\n"
        f"Dot size = % expressed in M4-top cells  |  "
        f"Colour = log₂FC  |  Black edge = significant\n"
        f"{'Region: ' + region_str if mode_str == 'pooled' else 'Regions: ' + region_str}",
        fontsize=10, y=1.03)
    plt.tight_layout()
    fn2dp = f"Fig2_dotplot_{short_dp}.png"
    fig2dp.savefig(os.path.join(OUT_DIR, fn2dp), dpi=150,
                   bbox_inches="tight", facecolor="white")
    fig2dp.savefig(os.path.join(OUT_DIR, fn2dp[:-4] + ".svg"),
                   format="svg", bbox_inches="tight", facecolor="white")
    plt.show(); plt.pause(0.3)
    print(f"  Saved {fn2dp}")

# PVM_3: per-region dot plot
_pvm3_keys = [(PVM3_SHORT, r) for r in REGION_ORDER
              if (PVM3_SHORT, r) in all_results]
if _pvm3_keys:
    _dot_plot(PVM3_SHORT, _pvm3_keys,
              [f"{r}" for _, r in _pvm3_keys])
else:
    print(f"  PVM_3: no valid comparisons  -  skipping dot plot.")

# PVM_2_3 and PVM_4: pooled dot plots
for _short_p in ["PVM_2_3", "PVM_4"]:
    if (_short_p, "ALL") in all_results:
        _dot_plot(_short_p, [(_short_p, "ALL")],
                  ["ALL regions"])
    else:
        print(f"  {_short_p}: pooled comparison not available  -  skipping dot plot.")


# =============================================================================
# SECTION 9: Figure 3  -  DEG overlap across regions (UpSet-style)
#
# PVM_3 only; PVM_2_3 and PVM_4 have one pooled DEG list each.
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 9: Figure 3  -  DEG overlap across regions per supertype")
print("  PVM_3 only (per-region mode)")
print("=" * 70)

for st_up in [SHORT_ST[s] for s in CANDIDATE_SUPERTYPES]:
    mode_up = ST_MODE[st_up]

    if mode_up == "pooled":
        print(f"  {st_up}: pooled mode, "
              f"skipped.")
        continue

    # per-region mode
    keys_up = [(st_up, r) for r in REGION_ORDER if (st_up, r) in all_results]
    if len(keys_up) < 2:
        print(f"  {st_up}: fewer than 2 valid comparisons  -  skipping overlap.")
        continue

    all_sig_genes = set()
    for key_up in keys_up:
        all_sig_genes |= set(
            all_results[key_up].loc[all_results[key_up]["significant"], "gene"])
    all_sig_genes = sorted(all_sig_genes)
    n_sig_g = len(all_sig_genes)

    if n_sig_g == 0:
        print(f"  {st_up}: no significant DEGs in any region  -  skipping.")
        continue

    regions_up = [r for _, r in keys_up]
    pres_mat = pd.DataFrame(False, index=all_sig_genes, columns=regions_up)
    lfc_mat  = pd.DataFrame(0.0,   index=all_sig_genes, columns=regions_up)
    for key_up in keys_up:
        reg_up = key_up[1]
        sig_up = all_results[key_up][all_results[key_up]["significant"]].set_index("gene")
        for g_up in sig_up.index:
            if g_up in pres_mat.index:
                pres_mat.loc[g_up, reg_up] = True
                lfc_mat.loc[g_up, reg_up]  = float(sig_up.loc[g_up, "log2FC"])

    share_counts = pres_mat.sum(axis=1)
    share_freq   = share_counts.value_counts().sort_index()
    pres_mat["n_regions"] = share_counts
    pres_mat = pres_mat.sort_values(["n_regions"] + regions_up, ascending=False)
    pres_mat.drop(columns=["n_regions"], inplace=True)

    np.random.seed(GLOBAL_SEED); _random_stdlib.seed(GLOBAL_SEED)
    fig3up, (ax3a, ax3b) = plt.subplots(
        1, 2,
        figsize=(max(12, n_sig_g * 0.30 + 3), 5.5),
        facecolor="white",
        gridspec_kw={"width_ratios": [3, 1]})

    pres_arr = pres_mat.values.T.astype(float)
    lfc_arr  = np.zeros_like(pres_arr)
    for ri_h, reg_h in enumerate(regions_up):
        for ci_h, g_h in enumerate(pres_mat.index):
            if pres_mat.loc[g_h, reg_h]:
                lfc_arr[ri_h, ci_h] = lfc_mat.loc[g_h, reg_h]

    lfc_max3 = max(0.5, float(np.abs(lfc_arr[pres_arr > 0]).max())
                   if (pres_arr > 0).any() else 1.0)
    norm3 = matplotlib.colors.TwoSlopeNorm(vmin=-lfc_max3, vcenter=0, vmax=lfc_max3)
    cmap3 = matplotlib.cm.RdBu_r

    img_arr = np.full((*pres_arr.shape, 4), fill_value=[0.92, 0.92, 0.92, 1.0])
    for ri_h in range(len(regions_up)):
        for ci_h in range(n_sig_g):
            if pres_arr[ri_h, ci_h]:
                img_arr[ri_h, ci_h] = cmap3(norm3(lfc_arr[ri_h, ci_h]))

    ax3a.imshow(img_arr, aspect="auto", interpolation="nearest", origin="upper")
    ax3a.set_yticks(range(len(regions_up)))
    ax3a.set_yticklabels(regions_up, fontsize=9)
    ax3a.set_xlabel("Significant DEG genes", fontsize=9)
    ax3a.set_xticks([])
    ax3a.set_title(
        f"{st_up}  -  DEG overlap across regions\n"
        f"(n={n_sig_g} unique sig genes  |  colour = log₂FC  |  grey = absent)",
        fontsize=10, fontweight="bold", pad=6)
    sm3  = matplotlib.cm.ScalarMappable(cmap=cmap3, norm=norm3)
    sm3.set_array([])
    cbar3 = fig3up.colorbar(sm3, ax=ax3a, fraction=0.025, pad=0.01,
                             shrink=0.7, aspect=20, location="right")
    cbar3.set_label("log₂FC", fontsize=8); cbar3.ax.tick_params(labelsize=7)

    region_counts_list = sorted(share_freq.index.tolist())
    bar_vals = [int(share_freq.get(n, 0)) for n in region_counts_list]
    _bar_pal = ["#D53E4F", "#FC8D59", "#FEE08B", "#99D594", "#3288BD", "#5E4FA2"]
    bar_cols = _bar_pal[:len(region_counts_list)]
    ax3b.barh(range(len(region_counts_list)), bar_vals,
              color=bar_cols, edgecolor="white", linewidth=0.8, height=0.6)
    ax3b.set_yticks(range(len(region_counts_list)))
    ax3b.set_yticklabels([f"In {n} region{'s' if n>1 else ''}"
                          for n in region_counts_list], fontsize=9)
    ax3b.set_xlabel("N genes", fontsize=9)
    ax3b.spines[["top", "right"]].set_visible(False)
    for xi_b, v_b in enumerate(bar_vals):
        ax3b.text(v_b + 0.3, xi_b, str(v_b), va="center", fontsize=8.5)
    ax3b.set_title("Region sharing", fontsize=10, fontweight="bold", pad=6)

    plt.tight_layout(w_pad=1.5)
    fn3up = f"Fig3_overlap_{st_up}.png"
    fig3up.savefig(os.path.join(OUT_DIR, fn3up), dpi=150,
                   bbox_inches="tight", facecolor="white")
    fig3up.savefig(os.path.join(OUT_DIR, fn3up[:-4] + ".svg"),
                   format="svg", bbox_inches="tight", facecolor="white")
    plt.show(); plt.pause(0.3)
    print(f"  Saved {fn3up}")

    # Add overlap sheet to master workbook
    overlap_df = pres_mat.copy()
    overlap_df["n_regions"] = share_counts[pres_mat.index]
    for reg_lfc in regions_up:
        overlap_df[f"log2FC_{reg_lfc}"] = [
            lfc_mat.loc[g, reg_lfc] if pres_mat.loc[g, reg_lfc] else np.nan
            for g in overlap_df.index]
    ws_ov = wb_master.create_sheet(f"{st_up[:10]}_overlap")
    ov_cols = (["gene", "n_regions"] + regions_up +
               [f"log2FC_{r}" for r in regions_up])
    ov_widths = [18, 12] + [10]*len(regions_up) + [12]*len(regions_up)
    _write_header(ws_ov, ov_cols, ov_widths)
    for ri_ov, (gene_ov, row_ov) in enumerate(overlap_df.iterrows(), 2):
        vals_ov = ([gene_ov, int(row_ov["n_regions"])] +
                   [bool(row_ov[r]) for r in regions_up] +
                   [row_ov.get(f"log2FC_{r}", np.nan) for r in regions_up])
        for ci_ov, v_ov in enumerate(vals_ov, 1):
            c_ov = _safe_cell(ws_ov, ri_ov, ci_ov, v_ov)
            c_ov.alignment = _lft if ci_ov == 1 else _ctr
            if ri_ov % 2 == 0: c_ov.fill = _af

wb_master.save(fp_master_xl)
print(f"\n  Master workbook re-saved with overlap sheets: {fp_master_xl}")


# =============================================================================
# SECTION 10: Final report
# =============================================================================
print("\n" + "=" * 70)
print("SECTION 10: Part 5 final report (neocortical)")
print("=" * 70)

_n_run  = int((summary_df["status"] == "OK").sum())
_n_skip = int((summary_df["status"] == "SKIPPED").sum())
_total_sig = int(summary_df["n_sig"].sum())
_total_up  = int(summary_df["n_up"].sum())
_total_dn  = int(summary_df["n_down"].sum())

# Separate counts by mode
_df_pr   = summary_df[summary_df["mode"] == "per_region"]
_df_pool = summary_df[summary_df["mode"] == "pooled"]

print(f"""
Part 5 complete (neocortical)
{'='*70}
CANONICAL MASK SOURCE:
  M4_TOP_MASK from the Part 2 joblib.
  Definition: top-{TOP_PCT_STR} M1_FC AND top-{TOP_PCT_STR} M2_Hybrid simultaneously.
  n={N_M4_TOP:,} cells total  |  M1_thr={M4_THR_M1:.6f}  |  M2_thr={M4_THR_M2:.6f}

[M4_TOP_MASK GUARDRAIL]: PASS  -  n={N_M4_TOP:,} cells confirmed
[NEOCORTICAL GUARDRAIL]: PASS  -  regions = {sorted(_obs_regions6)}

ANALYSIS MODES:
  PVM_3   : per_region   -  {len([r for r in REGION_ORDER])} separate comparisons (one per region)
  PVM_2_3 : pooled       -  1 comparison (all regions combined)
  PVM_4   : pooled       -  1 comparison (all regions combined)

DEG DESIGN:
  per_region foreground : M4-top (M4_TOP_MASK=True) of supertype in that region
  per_region background : non-M4-top of SAME supertype in SAME region
  pooled foreground     : ALL M4-top of supertype across AnG+DFC+ITG+MTG+STG+V1C
  pooled background     : ALL non-M4-top of SAME supertype across AnG+DFC+ITG+MTG+STG+V1C

DEG PARAMETERS:
  Source     : adata.layers['UMIs'] (raw integer UMI counts)
  Normalise  : CP10K ({int(NORM_TARGET):,}) + log1p
  Test       : pseudobulk paired Wilcoxon signed-rank (donor-level)
  min_pct    : {MIN_PCT} (>= 10% in either group, applied inside run_deg)
  min_cells  : {MIN_CELLS} per group
  min_cells_expr: {MIN_CELLS_EXPR} cells across fg+bg combined
  |log2FC|   : >= {LOG2FC_THR} (Seurat default)
  p_adj_thr  : < {PADJ_THR} (BH-FDR per comparison)
  pct_fg gate: >= {PCT_FG_GATE} (SIG sheet only)
  volcano_rank: -log10(p_adj) × |log2FC| for sig+gated genes; NaN otherwise

TOTAL SIGNIFICANT DEGs (across all comparisons):
  Total sig  : {_total_sig}  (UP={_total_up}, DOWN={_total_dn})
  Run OK     : {_n_run}  |  Skipped: {_n_skip}

COMPARISON SUMMARY:
""")
print(f"  {'Comparison':<22} {'mode':<12} {'fg':>6} {'bg':>6} "
      f"{'n_sig':>7} {'UP':>5} {'DOWN':>5}  status")
print("  " + "-" * 75)
for _, row_s3 in summary_df.iterrows():
    print(f"  {row_s3['supertype']:<10} {row_s3['region']:<8}  "
          f"{row_s3['mode']:<12} "
          f"fg={row_s3['n_fg']:>4}  bg={row_s3['n_bg']:>5}  "
          f"sig={row_s3['n_sig']:>4}  "
          f"(UP={row_s3['n_up']:>3} DOWN={row_s3['n_down']:>3})  "
          f"{row_s3['status']}")

print(f"\nOUTPUTS in {OUT_DIR}/:")
for fn_out in sorted(os.listdir(OUT_DIR)):
    if not fn_out.startswith("."):
        kb = os.path.getsize(os.path.join(OUT_DIR, fn_out)) / 1024
        print(f"  {fn_out:<70} {kb:>7.1f} KB")

print("\nPart 5 (neocortical) complete.")


# =============================================================================
# DIAGNOSTICS: filter breakdowns
# =============================================================================
print("\n" + "─" * 70)
print("DIAGNOSTICS: filter breakdowns (canonical M4_TOP_MASK)")
print("─" * 70)

# == Diag 1: PVM_3 MTG, per-region filter breakdown ===========================
print("\n[Diag 1] PVM_3 MTG  -  per-region filter breakdown")
_st_d  = "Micro-PVM_3-SEAAD"
_reg_d = "MTG"
_st_mask_d   = (obs_full[stype_col] == _st_d).values
_reg_mask_d  = (obs_full[region_col] == _reg_d).values
_fg_mask_d   = _st_mask_d & _reg_mask_d & m4_mask_global
_bg_mask_d   = _st_mask_d & _reg_mask_d & (~m4_mask_global)
_fg_obs_d    = obs_full.index.astype(str).values[_fg_mask_d]
_bg_obs_d    = obs_full.index.astype(str).values[_bg_mask_d]

X_fg_d = _load_raw_counts(adata, list(_fg_obs_d))
X_bg_d = _load_raw_counts(adata, list(_bg_obs_d))
X_fg_n = _normalise_log1p(X_fg_d)
X_bg_n = _normalise_log1p(X_bg_d)

pct_fg_d = (X_fg_n > 0).mean(axis=0)
pct_bg_d = (X_bg_n > 0).mean(axis=0)
n_expr_d = (X_fg_n > 0).sum(axis=0) + (X_bg_n > 0).sum(axis=0)

pct_mask_d  = (pct_fg_d >= 0.10) | (pct_bg_d >= 0.10)
expr_mask_d = n_expr_d >= 10
both_mask_d = pct_mask_d & expr_mask_d

print(f"PVM_3 MTG: fg={len(_fg_obs_d):,} (M4-top), bg={len(_bg_obs_d):,} (non-M4-top)")
print(f"Total genes in panel          : {len(pct_fg_d):,}")
print(f"Pass pct>=10% (either group)  : {pct_mask_d.sum():,}")
print(f"Pass min_cells_expr>=10       : {expr_mask_d.sum():,}")
print(f"Pass BOTH filters             : {both_mask_d.sum():,}")
print(f"Never expressed (pct=0 both)  : {((pct_fg_d==0) & (pct_bg_d==0)).sum():,}")
del X_fg_d, X_bg_d, X_fg_n, X_bg_n; gc.collect()

# == Diag 2: PVM_2_3 pooled, filter breakdown =================================
print("\n[Diag 2] PVM_2_3 ALL (pooled)  -  filter breakdown")
_st_p  = "Micro-PVM_2_3-SEAAD"
_st_mask_p   = (obs_full[stype_col] == _st_p).values
_all_reg_mask_p = np.zeros(len(_st_mask_p), dtype=bool)
for _r_p in REGION_ORDER:
    _all_reg_mask_p |= (obs_full[region_col].astype(str).values == _r_p)
_fg_mask_p   = _st_mask_p & _all_reg_mask_p & m4_mask_global
_bg_mask_p   = _st_mask_p & _all_reg_mask_p & (~m4_mask_global)
_fg_obs_p    = obs_full.index.astype(str).values[_fg_mask_p]
_bg_obs_p    = obs_full.index.astype(str).values[_bg_mask_p]

if len(_fg_obs_p) >= MIN_CELLS and len(_bg_obs_p) >= MIN_CELLS:
    X_fg_p = _load_raw_counts(adata, list(_fg_obs_p))
    X_bg_p = _load_raw_counts(adata, list(_bg_obs_p))
    X_fg_pn = _normalise_log1p(X_fg_p)
    X_bg_pn = _normalise_log1p(X_bg_p)

    pct_fg_p = (X_fg_pn > 0).mean(axis=0)
    pct_bg_p = (X_bg_pn > 0).mean(axis=0)
    n_expr_p = (X_fg_pn > 0).sum(axis=0) + (X_bg_pn > 0).sum(axis=0)

    pct_mask_p  = (pct_fg_p >= 0.10) | (pct_bg_p >= 0.10)
    expr_mask_p = n_expr_p >= 10
    both_mask_p = pct_mask_p & expr_mask_p

    print(f"PVM_2_3 ALL: fg={len(_fg_obs_p):,} (M4-top, pooled), "
          f"bg={len(_bg_obs_p):,} (non-M4-top, pooled)")
    print(f"Total genes in panel          : {len(pct_fg_p):,}")
    print(f"Pass pct>=10% (either group)  : {pct_mask_p.sum():,}")
    print(f"Pass min_cells_expr>=10       : {expr_mask_p.sum():,}")
    print(f"Pass BOTH filters             : {both_mask_p.sum():,}")
    print(f"Never expressed (pct=0 both)  : {((pct_fg_p==0) & (pct_bg_p==0)).sum():,}")
    del X_fg_p, X_bg_p, X_fg_pn, X_bg_pn; gc.collect()
else:
    print(f"PVM_2_3 ALL: fg={len(_fg_obs_p)}, bg={len(_bg_obs_p)}  -  "
          f"below MIN_CELLS threshold, diagnostics skipped.")

print("\nDiagnostics complete.")
