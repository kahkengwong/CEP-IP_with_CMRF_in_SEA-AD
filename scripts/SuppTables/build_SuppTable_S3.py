#!/usr/bin/env python3
"""
build_SuppTable_S3.py

Builds the data sheets of Supplementary Table S3 (choice of MT-CO3 as the gene
of interest). Sheet A sets the 13 mitochondrial protein-coding genes side by
side: detection and mean counts in all Micro-PVM_3-SEAAD nuclei (Part 1b), the
pooled PVM_3 screen of each region set (DFG partners, best Spearman rs and the
Kendall tau of that pair from Part 4, rank in the Part 6b MasterDFG ranking),
and the 10 regional PVM_3 screens of Part 6 (regions with partners, median
partners, regions where the gene holds the largest DFG set with ties, and
regions where Part 6 lists it as the top gene). Sheet B gives the DFG partners
of every gene with partners in each regional screen. Sheets C and D list every
gene passing the expression gate in the pooled PVM_3 screen of each region set.
Sheet E is a composite rank of the 13 genes computed from the rounded values
of sheet A: the weighted mean rank over nine measures, with detection and mean
counts weighted twice, plus a leave-one-measure-out check and a version over
all columns of sheet A. finalize_supp_tables.py adds the Contents sheet, renames
column regions_listed_top_in_Part6 to regions_listed_top_S4B and saves the
published file.

Usage:
    python build_SuppTable_S3.py
    python build_SuppTable_S3.py --root /data/SEA-AD
Root: --root, else $SEAAD_ROOT, else the current folder.

Inputs (<arm> is NeoC in [1]_NeoC and NonNeoC in [2]_Non-NeoC):
    <root>/Part1_PreChecks/S1b_mito_gene_inventory.csv
        scripts/Python/Part1b_MT-audit_code.py (columns gene, detect_pct,
        mean_counts)
    <root>/<arm folder>/Part6_<arm>_GOI-DFG/Part6_cross_population_summary.xlsx
        scripts/Python/Part6_<arm>_GOI-DFG_code.py (sheets 01_dfg_counts and
        PVM_3_<region>)
    <root>/<arm folder>/Part6b_<arm>_Pooled_GOI-DFG/PVM_3_pooled/PVM_3_pooled_MasterDFG_Ranking.xlsx
        scripts/Python/Part6b_<arm>_Pooled_GOI-DFG.py (sheet Master_DFG_Ranking)
    <root>/<arm folder>/Part4_<arm>_ScreenScatter/PVM_3_pooled_screen_scatter.csv
        scripts/Python/Part4_<arm>_ScreenScatter.py
Output:
    <root>/Manuscript/Supp_Tables/build/SuppTable_S3_MT-CO3_selection.xlsx
        sheets A_mito_gene_summary, B_region_screen, C_pooled_screen_NeoC,
        D_pooled_screen_NonNeoC, E_GOI_composite
    <root>/Manuscript/Supp_Tables/build/build_SuppTable_S3_log.txt
        checks and row counts, ending with VERDICT: PASS or VERDICT: FAIL.
        On FAIL no workbook is written.

Formatting: bold header row, frozen top row, no fills, no borders.
"""

import argparse
import os
import sys

import pandas as pd
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ARMS = {"NeoC": ("[1]_NeoC", ["DFC", "MTG", "AnG", "STG", "V1C", "ITG"]),
        "NonNeoC": ("[2]_Non-NeoC", ["HIP", "MEC", "LEC", "FI"])}
MITO = ["MT-CO3", "MT-ATP6", "MT-CO2", "MT-ND3", "MT-ND4", "MT-CYB", "MT-ND1", "MT-ND2",
        "MT-CO1", "MT-ATP8", "MT-ND4L", "MT-ND5", "MT-ND6"]
MODULE_DFG = ["MT-ATP6", "MT-CO2", "MT-ND3", "MT-ND4", "MT-CYB", "MT-ND1", "MT-ND2", "MT-CO1"]
# Composite (sheet E): one measure per property; higher is better for all nine.
# Expression is weighted twice because the analysed cells are those with GOI > 0.
COMPOSITE = [("detect_pct_PVM3", 2), ("mean_counts_PVM3", 2),
             ("NeoC_pooled_partners", 1), ("NonNeoC_pooled_partners", 1),
             ("NeoC_pooled_best_rs", 1), ("NonNeoC_pooled_best_rs", 1),
             ("regions_with_partners_of_10", 1), ("median_partners_10_regions", 1),
             ("regions_largest_set_ties_included", 1)]
RANK_COLS = ("NeoC_pooled_rank", "NonNeoC_pooled_rank")   # lower is better (all-columns check)
# Reference values of the original run, used only in checks.
EXPECT_POOLED_CO3 = {"NeoC": 7, "NonNeoC": 5}
EXPECT_GATED = {"NeoC": 3904, "NonNeoC": 2723}
EXPECT_COMPOSITE = {"MT-CO3": 2.14, "MT-ATP6": 3.00}       # weighted mean rank, 2 decimals
OUT_NAME = "SuppTable_S3_MT-CO3_selection.xlsx"
LOG_NAME = "build_SuppTable_S3_log.txt"


class Log:
    def __init__(self):
        self.lines, self.fails = [], 0

    def say(self, s=""):
        print(s)
        self.lines.append(str(s))

    def check(self, ok, what):
        self.say(f"  [{'PASS' if ok else 'FAIL'}] {what}")
        if not ok:
            self.fails += 1
        return ok

    def write(self, path):
        self.say("")
        self.say(f"VERDICT: {'PASS' if self.fails == 0 else 'FAIL'}")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.lines) + "\n")


def inputs(root):
    p = {"inv": os.path.join(root, "Part1_PreChecks", "S1b_mito_gene_inventory.csv")}
    for arm, (folder, _) in ARMS.items():
        p[f"p6_{arm}"] = os.path.join(root, folder, f"Part6_{arm}_GOI-DFG",
                                      "Part6_cross_population_summary.xlsx")
        p[f"p6b_{arm}"] = os.path.join(root, folder, f"Part6b_{arm}_Pooled_GOI-DFG", "PVM_3_pooled",
                                       "PVM_3_pooled_MasterDFG_Ranking.xlsx")
        p[f"p4_{arm}"] = os.path.join(root, folder, f"Part4_{arm}_ScreenScatter",
                                      "PVM_3_pooled_screen_scatter.csv")
    return p


def ranks(frame, cols):
    """Rank 1 = best; ties share the average rank; a missing value ranks last."""
    out = pd.DataFrame(index=frame.index)
    for c in cols:
        v = pd.to_numeric(frame[c], errors="coerce").astype(float)
        if c in RANK_COLS:
            v = -v.fillna(v.max() + 1)        # lower rank number is better
        else:
            v = v.fillna(v.min() - 1)
        out[c] = v.rank(ascending=False, method="average")
    return out


def weighted_mean_rank(rk, weights):
    w = pd.Series(weights, dtype=float)
    return (rk[list(w.index)] * w).sum(axis=1) / w.sum()


def add_sheet(wb, name, df):
    bold = Font(bold=True)
    w = wb.create_sheet(name)
    w.append(list(df.columns))
    for c in w[1]:
        c.font = bold
    for rec in df.itertuples(index=False):
        w.append([None if (v is pd.NA or (isinstance(v, float) and pd.isna(v)))
                  else (v.item() if hasattr(v, "item") else v) for v in rec])
    for j, col in enumerate(df.columns, 1):
        w.column_dimensions[get_column_letter(j)].width = max(10, min(32, len(str(col)) + 2))
    w.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--root", default=os.environ.get("SEAAD_ROOT", os.getcwd()))
    root = os.path.abspath(ap.parse_args().root)
    out_dir = os.path.join(root, "Manuscript", "Supp_Tables", "build")
    os.makedirs(out_dir, exist_ok=True)
    out_path, log_path = os.path.join(out_dir, OUT_NAME), os.path.join(out_dir, LOG_NAME)
    L = Log()
    L.say("build_SuppTable_S3.py")
    L.say(f"root: {root}")

    IN = inputs(root)
    for k, v in IN.items():
        L.say(f"  {'found  ' if os.path.isfile(v) else 'MISSING'}  {os.path.relpath(v, root)}")
    if not L.check(all(os.path.isfile(v) for v in IN.values()), "all inputs present"):
        L.write(log_path)
        sys.exit(1)

    # ------------------------------------------------------------ read
    inv = pd.read_csv(IN["inv"]).set_index("gene")
    L.check(all(g in inv.index for g in MITO), "inventory holds the 13 mitochondrial genes")
    region = {}          # region -> Series(gene -> partners)
    top_listed = {}      # region -> gene listed as top in 01_dfg_counts
    for arm, (_, regs) in ARMS.items():
        dc = pd.read_excel(IN[f"p6_{arm}"], sheet_name="01_dfg_counts")
        for r in regs:
            d = pd.read_excel(IN[f"p6_{arm}"], sheet_name=f"PVM_3_{r}")
            region[r] = d.set_index("GOI")["DFG_Total"]
            top_listed[r] = dc.loc[dc["Population"] == f"PVM_3_{r}", "Top_GOI_by_DFG"].iloc[0]
    pooled, rank = {}, {}
    for arm in ARMS:
        sc = pd.read_csv(IN[f"p4_{arm}"])
        ref = pd.read_excel(IN[f"p6b_{arm}"], sheet_name="Master_DFG_Ranking")
        ref = ref[ref["DFG_Total"] >= 1]
        got = sc[sc["DFG_Total"] > 0].set_index("GOI")["DFG_Total"]
        L.check(len(sc) == EXPECT_GATED[arm],
                f"{arm}: {EXPECT_GATED[arm]:,} genes pass the expression gate (this run {len(sc):,})")
        L.check(set(got.index) == set(ref["GOI"]) and all(got[g] == n for g, n in zip(ref["GOI"], ref["DFG_Total"])),
                f"{arm}: pooled screen matches Part 6b (genes with partners and counts)")
        L.check(int(got.get("MT-CO3", 0)) == EXPECT_POOLED_CO3[arm],
                f"{arm}: MT-CO3 has {EXPECT_POOLED_CO3[arm]} partners in the pooled screen "
                f"(this run {int(got.get('MT-CO3', 0))})")
        pooled[arm] = sc
        rank[arm] = dict(zip(ref["GOI"], ref["Rank"]))

    # ------------------------------------------------------------ sheet A
    all_regions = [r for _, (_, regs) in ARMS.items() for r in regs]
    rows = []
    for g in MITO:
        r = {"gene": g,
             "role": "gene of interest" if g == "MT-CO3" else ("module DFG" if g in MODULE_DFG else "other mitochondrial"),
             "detect_pct_PVM3": round(float(inv.loc[g, "detect_pct"]), 2),
             "mean_counts_PVM3": round(float(inv.loc[g, "mean_counts"]), 3)}
        for arm in ARMS:
            sc = pooled[arm].set_index("GOI")
            has = g in sc.index
            r[f"{arm}_pooled_partners"] = int(sc.loc[g, "DFG_Total"]) if has else 0
            r[f"{arm}_pooled_best_rs"] = round(float(sc.loc[g, "top_spearman"]), 4) if has else None
            r[f"{arm}_pooled_tau_same_pair"] = round(float(sc.loc[g, "top_kendall"]), 4) if has else None
            r[f"{arm}_pooled_rank"] = rank[arm].get(g)
        parts = [int(region[x].get(g, 0)) for x in all_regions]
        r["regions_with_partners_of_10"] = sum(p > 0 for p in parts)
        r["median_partners_10_regions"] = float(pd.Series(parts).median())
        r["regions_largest_set_ties_included"] = sum(int(region[x].get(g, 0)) == int(region[x].max()) and
                                                     int(region[x].get(g, 0)) > 0 for x in all_regions)
        r["regions_listed_top_in_Part6"] = sum(top_listed[x] == g for x in all_regions)
        rows.append(r)
    A = pd.DataFrame(rows)
    for arm in ARMS:
        A[f"{arm}_pooled_rank"] = A[f"{arm}_pooled_rank"].astype("Int64")   # whole numbers, blank if none
    co3 = A[A.gene == "MT-CO3"].iloc[0]
    L.check(co3["regions_with_partners_of_10"] == 10, "MT-CO3 gains partners in all 10 regions")
    L.say(f"    MT-CO3: largest set in {co3['regions_largest_set_ties_included']} of 10 regions with ties, "
          f"listed top in {co3['regions_listed_top_in_Part6']}; pooled rank "
          f"{co3['NeoC_pooled_rank']} (NeoC), {co3['NonNeoC_pooled_rank']} (NonNeoC)")

    # ------------------------------------------------------------ sheet B
    genes_b = sorted(set().union(*[set(region[x].index) for x in all_regions]),
                     key=lambda g: (g not in MITO, MITO.index(g) if g in MITO else 0, g))
    B = pd.DataFrame({"gene": genes_b})
    for x in all_regions:
        B[x] = [int(region[x].get(g, 0)) for g in genes_b]
    B["regions_with_partners"] = (B[all_regions] > 0).sum(axis=1)

    # ------------------------------------------------------------ sheet E
    Ai = A.set_index("gene")
    wts = dict(COMPOSITE)
    rk = ranks(Ai, list(wts))
    score = weighted_mean_rank(rk, wts)
    loo = pd.DataFrame({c: weighted_mean_rank(rk, {k: v for k, v in wts.items() if k != c})
                        .rank(method="min") for c in wts})
    all_cols = [c for c in Ai.columns if c != "role"]
    wts_all = {c: wts.get(c, 1) for c in all_cols}
    score_all = weighted_mean_rank(ranks(Ai, all_cols), wts_all)
    E = pd.DataFrame({"gene": Ai.index, "role": Ai["role"].values})
    for c in wts:
        E[f"rank_{c}"] = rk[c].values
    E["composite_mean_rank"] = score.round(2).values
    E["composite_rank"] = score.rank(method="min").astype(int).values
    E["worst_rank_leave_one_measure_out"] = loo.max(axis=1).astype(int).values
    E["rank_all_columns_of_A"] = score_all.rank(method="min").astype(int).values
    E = E.sort_values(["composite_rank", "gene"]).reset_index(drop=True)
    top = E.iloc[0]
    L.check(top["gene"] == "MT-CO3", "MT-CO3 ranks first on the composite")
    L.check(all(abs(E.set_index("gene").loc[g, "composite_mean_rank"] - v) < 0.006
                for g, v in EXPECT_COMPOSITE.items()),
            "composite mean ranks as in the original run (MT-CO3 2.14, MT-ATP6 3.00)")
    L.check(int(top["worst_rank_leave_one_measure_out"]) == 1,
            "MT-CO3 stays first when any one measure is left out")
    L.check(int(top["rank_all_columns_of_A"]) == 1, "MT-CO3 ranks first over all columns of sheet A")
    for rec in E.head(5).itertuples(index=False):
        L.say(f"    {rec.gene:8s} composite mean rank {rec.composite_mean_rank:.2f} (rank {rec.composite_rank})")

    # ------------------------------------------------------------ sheets C, D
    cols = {"GOI": "gene", "pct_expressing_gate": "pct_nuclei_expressing", "top_partner": "best_partner",
            "top_spearman": "best_rs", "top_kendall": "tau_same_pair", "DFG_Total": "DFG_partners",
            "max_Spearman": "max_rs_among_DFGs", "max_Kendall": "max_tau_among_DFGs"}
    CD = {}
    for arm in ARMS:
        d = pooled[arm][list(cols)].rename(columns=cols)
        d = d.sort_values(["DFG_partners", "best_rs"], ascending=[False, False]).reset_index(drop=True)
        for c in ("max_rs_among_DFGs", "max_tau_among_DFGs"):
            d.loc[d["DFG_partners"] == 0, c] = None
        CD[arm] = d

    # ------------------------------------------------------------ write
    sheets = [("A_mito_gene_summary", A), ("B_region_screen", B),
              ("C_pooled_screen_NeoC", CD["NeoC"]), ("D_pooled_screen_NonNeoC", CD["NonNeoC"]),
              ("E_GOI_composite", E)]
    L.say("")
    for name, df in sheets:
        L.say(f"  {name}: {len(df)} rows x {df.shape[1]} columns")
    if L.fails:
        L.say("")
        L.say("Checks failed; workbook not written.")
        L.write(log_path)
        sys.exit(1)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, df in sheets:
        add_sheet(wb, name, df)
    wb.save(out_path)
    L.say(f"  saved: {os.path.relpath(out_path, root)}")
    L.write(log_path)


if __name__ == "__main__":
    main()
