#!/usr/bin/env python3
"""
build_SuppTable_S14.py

Builds SuppTable_Fig16.xlsx (Figure 16 and Supplementary Figure 12 of the Part 12
rerun) from the stats and
plot-data files written by Part12_Fig16_Analysis.r (sections 13 to 17), the
figure log written by Part12_Fig16_Plots.r (section 18), the config
(Part12_01_config.json) and the unit table (05_units.csv) with the step 05 log.
finalize_supp_tables.py turns this workbook into Supplementary Table S14.

Usage:
    python build_SuppTable_S14.py
    python build_SuppTable_S14.py --root /data/Part12_rerun

--root defaults to $SEAAD_ROOT/Part12_rerun when SEAAD_ROOT is set, otherwise
to Part12_rerun in the current folder.

--root is searched recursively; every input file must occur exactly once under
it, so the script works on the Part12_rerun folder tree or on a flat folder
holding copies of the files. Without --out the workbook goes to
<root>/08_Fig16/SuppTable_Fig16.xlsx. A build log is written beside the workbook.

The unit table gives the Module post-IP% columns (post_ip_pct_median) of the
16A, 16B, 16F and 16G sheets and the Fig16_units sheet. Each value is the
median Module post-IP% over the units or donors behind that row.

The script computes no new statistic. Before it writes anything it checks:
  B-01  every section log ends with VERDICT: PASS and 0 FAIL
  B-02  every input file is listed in the outputs of the log of the section that wrote it
  B-03  the five analysis logs name one script, one config version and one labels version
  B-04  every stats row carries the run ID, config version and labels version of its log
  B-05  stats and plot-data row counts equal the counts the logs report
  B-06  every fdr equals BH recomputed within its family
  B-07  16C shares of each row sum to 1, and kappa recomputed from the four shares
        equals the kappa row
  B-08  16D within-donor medians, quartiles and donor counts equal the plot data
  B-09  16E difference equals the post-IP minus the pre-IP offset share, and every
        16E verdict follows from the 90% interval and the FDR
  B-10  every 16F and 16G contrast equals the difference of its two rs rows
  B-11  the four 16G conditions recomputed from Stats_16G equal the section 17 log
  B-12  the figure log is PASS and was started after the last analysis section
  B-13  no text written to the workbook contains an en dash or an em dash
  B-14  the config exists and has the config version of the section logs, the unit
        table is listed in the outputs of its log, which is PASS, and the unit table
        carries the labels version of the section logs
  B-15  the unit and donor sets rebuilt from the unit table have the n_units and
        n_donors of every 16A, 16B, 16F and 16G row, and the library fields kept
        under the missing-value limit are the fields in Stats_16F and Stats_16G
  B-16  every 16A donor share and every observed 16F and 16G rs recomputed from the
        Fig16_units values equals its stats row (difference below 1e-9)
  B-17  Fig16_units holds every unit of the unit table
Any FAIL stops the build and no workbook is written.

Formatting: Arial throughout, bold header row, frozen top row. No fills, no
borders, no colours. The gap column of Fig16A_readout, the threshold column of
Fig16C_readout and the condition and verdict columns of Fig16G_conditions are
Excel formulas; the file is flagged to recalculate on open.
"""

import argparse
import datetime as dt
import json
import os
import pathlib
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
SECTIONS = [
    # (step, folder, files written by that section)
    ("13", "16AB_DonorStructure", ["Stats_16A.csv", "Stats_16B.csv"]),
    ("14", "16C_LabelAgreement", ["Stats_16C.csv"]),
    ("15", "16D_WithinDonorRho", ["Stats_16D.csv", "Plot_16D_donors.csv"]),
    ("16", "16E_Offset", ["Stats_16E.csv"]),
    ("17", "16FG_LibraryFields", ["Stats_16F.csv", "Stats_16G.csv"]),
]
FIG_LOG = "Part12_Fig16_log.txt"
CONFIG = "Part12_01_config.json"
UNITS = ("05_units.csv", "*05_Units*_log.txt")
ARMS = ["M4", "PVM_3", "PVM_1", "HKG8"]
PROV_COLS = ["run_id", "config_version", "labels_version"]
TOL = 1e-9
DASHES = ("\u2013", "\u2014")
## Unit-table column of each 16F/16G quantity, as section 17 of the analysis.
QS = {"module": "module_post_pct", "mito": "mito_log1p", "ncells": "log_n_cells",
      "depth": "log_median_umi"}
QA = ["module_post_pct", "module_resid_tq4", "mito_log1p", "contamination", "ctrl_pct_null"]
QB = ["module_post_pct", "module_resid_tq4", "mito_log1p"]
LEVELS = ["between donors", "within donor"]
MODULE_ROWS = {"module", "module_minus_mito", "module_minus_depth", "arm_minus_hkg8"}
MARGIN = 10.0
VERDICT_16E = {
    (True, False): "within the margin",
    (True, True): "non-zero but within the margin",
    (False, True): "non-zero and not shown to be within the margin",
    (False, False): "underpowered: neither non-zero nor shown to be within the margin",
}
QLAB = {"module_post_pct": "Module post-IP%", "module_resid_tq4": "Module post-IP%, TQR",
        "mito_log1p": "%MT (log1p)", "contamination": "Contamination",
        "ctrl_pct_null": "Matched non-MT share, null"}
FIELD_LAB = {"intronic": "intronic rate", "antisense": "antisense rate",
             "neun_fans": "NeuN-positive fraction", "cdna_400": "cDNA > 400 bp",
             "pcr_cycles": "PCR cycles", "saturation": "sequencing saturation"}

FONT = Font(name="Arial", size=10)
FONT_B = Font(name="Arial", size=10, bold=True)
FMT_P = "0.00E+00"


# ---------------------------------------------------------------------------
# Check ledger
# ---------------------------------------------------------------------------
class Ledger:
    def __init__(self):
        self.lines, self.n = [], {"PASS": 0, "FAIL": 0, "NOTE": 0}

    def chk(self, cid, ok, what, detail=""):
        tag = "PASS" if ok else "FAIL"
        self.n[tag] += 1
        self._say(f"[{tag}] {cid} {what} | {detail}")

    def note(self, what, detail=""):
        self.n["NOTE"] += 1
        self._say(f"[NOTE] {what} | {detail}")

    def _say(self, s):
        self.lines.append(s)
        print(s)


def find(root, name):
    hits = [p for p in root.rglob(name) if p.is_file()]
    if len(hits) != 1:
        sys.exit(f"ERROR: expected exactly one {name} under {root}, found {len(hits)}")
    return hits[0]


def find_glob(root, pattern):
    hits = [p for p in root.rglob(pattern) if p.is_file()]
    if len(hits) != 1:
        sys.exit(f"ERROR: expected exactly one file matching {pattern} under {root}, found {len(hits)}")
    return hits[0]


_NUM = re.compile(r"([+-]?)(\d*)(?:\.(\d*))?(?:[eE]([+-]?\d+))?$")


def r_strtod(s):
    """One number as R's read.csv parses it (R_strtod5): the digits are
    accumulated in long double, scaled by a power of ten built by repeated
    squaring in long double, then rounded to double. pandas' own parser and
    Python's float() both differ from read.csv in the last bit for some values,
    which changes the order of values that nearly tie after centring."""
    s = s.strip()
    if s in ("", "NA"):
        return np.nan
    low = s.lower()
    if low in ("inf", "+inf"):
        return np.inf
    if low == "-inf":
        return -np.inf
    if low == "nan":
        return np.nan
    m = _NUM.match(s)
    if not m or not (m.group(2) or m.group(3)):
        raise ValueError(s)
    sign, ip, fp, ex = m.group(1), m.group(2) or "", m.group(3) or "", m.group(4)
    digits = ip + fp
    if len(digits.lstrip("0")) >= 18:
        return float(s)
    expn = -len(fp) + (int(ex) if ex else 0)
    ld = np.longdouble
    ans = ld(0)
    for ch in digits:
        ans = ld(10) * ans + ld(ord(ch) - 48)
    if expn != 0 and ans != 0:
        n, p10, fac = abs(expn), ld(10), ld(1)
        while n:
            if n & 1:
                fac *= p10
            n >>= 1
            p10 *= p10
        ans = ans / fac if expn < 0 else ans * fac
    v = float(ans)
    return -v if sign == "-" else v


def r_mean(v):
    """R's mean() of a double vector: long double sum in order, divided by n,
    plus the long double correction pass, rounded to double."""
    ld = np.longdouble
    n = len(v)
    s = ld(0)
    for a in v:
        s += ld(a)
    s /= n
    if np.isfinite(float(s)):
        t = ld(0)
        for a in v:
            t += ld(a) - s
        s += t / n
    return float(s)


def read_csv_like_r(path, char_cols=(), usecols=None):
    """Read a CSV with R read.csv missing-value rules, so complete-case filters
    match the analysis: "NA" is missing in every column; an empty field is
    missing in a numeric column and an empty string in a character column."""
    d = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[], usecols=usecols)
    for c in d.columns:
        v = d[c].where(d[c] != "NA")
        if c in char_cols:
            d[c] = v
            continue
        try:
            d[c] = pd.Series([np.nan if x is None or (isinstance(x, float) and np.isnan(x)) else r_strtod(x)
                              for x in v], index=d.index, dtype=float)
        except ValueError:
            d[c] = v
    return d



# ---------------------------------------------------------------------------
# Recomputation from the unit table (section 13 and section 17 rules)
# ---------------------------------------------------------------------------
def rs(x, y):
    """Spearman rs as p12_rs(): Pearson correlation of average ranks; NaN when
    fewer than 3 values or when either side has no spread."""
    x, y = np.asarray(x, float), np.asarray(y, float)
    if len(x) < 3:
        return np.nan
    rx, ry = pd.Series(x).rank().values, pd.Series(y).rank().values
    if not (np.std(rx, ddof=1) > 0 and np.std(ry, ddof=1) > 0):
        return np.nan
    return float(np.corrcoef(rx, ry)[0, 1])


def centre_within(v, donor, min_n):
    """p12_centre_within(): each donor's mean removed, for donors with at least
    min_n finite values; NaN otherwise."""
    v = np.asarray(v, float)
    out = np.full(len(v), np.nan)
    donor = np.asarray(donor)
    for u in pd.unique(donor):
        i = np.where((donor == u) & np.isfinite(v))[0]
        if len(i) >= min_n:
            out[i] = v[i] - r_mean(v[i])
    return out


def rebuild_16a(U, min_d):
    """16A per arm and quantity: units, donors, donor share (donor entered
    first) and the median Module post-IP% over the units."""
    out = {}
    for a in ARMS:
        for q in QA:
            u = U[(U["arm"] == a) & np.isfinite(U[q].astype(float))]
            y = u[q].astype(float).values
            don = u["donor_key"].values
            nd = len(pd.unique(don))
            est = np.nan
            if nd >= min_d and len(y) > 1 and np.var(y, ddof=1) > 0:
                s = pd.Series(y).groupby(don).agg(["sum", "size"])
                ssb = float((s["sum"] ** 2 / s["size"]).sum() - len(y) * y.mean() ** 2)
                est = ssb / float(np.sum((y - y.mean()) ** 2))
            out[(a, q)] = {"n_units": len(u), "n_donors": nd, "estimate": est,
                           "post": float(np.median(u["module_post_pct"])) if len(u) else np.nan}
    return out


def rebuild_16b(U):
    """16B per arm and quantity: units of donors with at least two units, and
    the median Module post-IP% over them."""
    out = {}
    for a in ARMS:
        for q in QB:
            u = U[(U["arm"] == a) & np.isfinite(U[q].astype(float))]
            k = u["donor_key"].value_counts()
            dons = set(k[k >= 2].index)
            u = u[u["donor_key"].isin(dons)]
            out[(a, q)] = {"n_units": len(u), "n_donors": len(dons),
                           "post": float(np.median(u["module_post_pct"])) if len(u) else np.nan}
    return out


def rebuild_16fg(U, fields, min_d, min_u):
    """16F/16G per arm, field, level and quantity, as section 17: the values
    entering rs, their counts, the observed rs and, for the Module, the median
    Module post-IP% over those units (within donor) or donor means (between)."""
    out = {}
    for a in ARMS:
        ua = U[U["arm"] == a]
        for f in fields:
            y_all = ua[f"lib_{f}"].astype(float).values
            for lv in LEVELS:
                for q, col in QS.items():
                    x_all = ua[col].astype(float).values
                    ok = np.isfinite(x_all) & np.isfinite(y_all)
                    d_all = ua["donor_key"].values
                    if lv == "between donors":
                        ## tapply(x, donor, mean): one mean per donor over its units in table order
                        dk = d_all[ok]
                        d = np.array(sorted(pd.unique(dk)))
                        x = np.array([r_mean(x_all[ok][dk == k]) for k in d])
                        y = np.array([r_mean(y_all[ok][dk == k]) for k in d])
                        pm = ua["module_post_pct"].astype(float).values[ok]
                        post_vals = np.array([r_mean(pm[dk == k]) for k in d])
                    else:
                        d = d_all[ok]
                        x = centre_within(x_all[ok], d, min_u)
                        y = centre_within(y_all[ok], d, min_u)
                        k = np.isfinite(x) & np.isfinite(y)
                        x, y, d = x[k], y[k], d[k]
                        post_vals = ua["module_post_pct"].values[ok][k]
                    nd = len(pd.unique(d))
                    out[(a, f, lv, q)] = {"n_units": len(x), "n_donors": nd,
                                          "estimate": rs(x, y) if nd >= min_d else np.nan,
                                          "post": float(np.median(post_vals)) if len(post_vals) else np.nan}
    return out


def conditions_16g(G):
    """The four 16G conditions per arm and field, within donor, as section 17
    prints them. Returns {(arm, field): (c1, c2, c3, c4, verdict)}."""
    res = {}
    excl = lambda r: len(r) == 1 and np.isfinite(r["lo"].iloc[0]) and (r["lo"].iloc[0] > 0 or r["hi"].iloc[0] < 0)
    for a in [x for x in ARMS if x != "HKG8"]:
        for f in pd.unique(G["measure"]):
            g = G[(G.arm == a) & (G.level == "within donor") & (G.measure == f)]
            m = g[g.quantity == "module"]
            c1 = len(m) == 1 and np.isfinite(m["fdr"].iloc[0]) and m["fdr"].iloc[0] < 0.05
            c2 = excl(g[g.quantity == "module_minus_depth"])
            c3 = excl(g[g.quantity == "arm_minus_hkg8"])
            mm = g[g.quantity == "module_minus_mito"]
            c4 = len(mm) == 1 and np.isfinite(mm["hi"].iloc[0]) and mm["hi"].iloc[0] >= 0
            testable = len(m) == 1 and np.isfinite(m["estimate"].iloc[0])
            v = ("untestable (below the donor floor)" if not testable else
                 "orders Module post-IP%" if (c1 and c2 and c3 and c4) else "does not order Module post-IP%")
            res[(a, f)] = (bool(c1), bool(c2), bool(c3), bool(c4), v)
    return res


def log_conditions(txt):
    """The section 17 log's four-condition lines."""
    pat = (r"^\s+(\S+)\s+(\S+)\s+non-zero (TRUE|FALSE)\s+\| vs depth (TRUE|FALSE)\s+\| vs HKG8 (TRUE|FALSE)"
           r"\s+\| not matched by %MT (TRUE|FALSE)\s+-> (.+?)\s*$")
    return {(m[0], m[1]): tuple(v == "TRUE" for v in m[2:6]) + (m[6],) for m in re.findall(pat, txt, re.M)}


def kappa_from_shares(s):
    """Cohen's kappa from the four agreement shares (both, test only, ref only,
    neither), as kappa4() in section 14."""
    b, t, r, n = s
    po = b + n
    pe = (b + t) * (b + r) + (n + r) * (n + t)
    return np.nan if pe >= 1 else (po - pe) / (1 - pe)


def parse_log(path):
    txt = path.read_text(encoding="utf-8", errors="replace")
    g = lambda pat: (re.search(pat, txt, re.M).group(1) if re.search(pat, txt, re.M) else None)
    # Output lines: "  <folder>/<file> | <size> MB"; anything after the size is ignored.
    outs = set(re.findall(r"^\s+\S*/([^/\s|]+) \| [\d.]+ MB", txt, re.M))
    return {
        "path": path,
        "script": g(r"^Script: (\S+)"),
        "run_id": g(r"^Run ID: (\S+)"),
        "start": g(r"^Start: (.+)$"),
        "config_version": g(r"^Config: .* \| version (\S+?)(?: \||\s*$)"),
        "labels_version": g(r"^Labels version: (\S+)"),
        "checks": g(r"^Checks: (.+?) \|"),
        "n_fail": int(g(r"^Checks: \d+ PASS, (\d+) FAIL") or -1),
        "verdict": g(r"^VERDICT: (\S+)"),
        "outputs": outs,
        "rows_line": g(r"^\[PASS\] \S+ stats (?:and plot )?files? written(?: and read back)? \| (.+)$"),
        "text": txt,
    }


def bh(p):
    """Benjamini-Hochberg as R p.adjust(method = "BH"), NaN left as NaN."""
    p = np.asarray(p, dtype=float)
    out = np.full(p.shape, np.nan)
    ok = np.isfinite(p)
    m = ok.sum()
    if m == 0:
        return out
    pv = p[ok]
    order = np.argsort(pv)[::-1]
    ranks = np.arange(m, 0, -1)
    q = np.minimum.accumulate(pv[order] * m / ranks)
    res = np.empty(m)
    res[order] = np.minimum(q, 1.0)
    out[ok] = res
    return out


def close(a, b, tol=TOL):
    a, b = np.asarray(a, float), np.asarray(b, float)
    both_nan = np.isnan(a) & np.isnan(b)
    return bool(np.all(both_nan | (np.abs(a - b) <= tol)))


def clean(v):
    """Excel-safe cell value: NaN to empty, numpy scalars to Python."""
    if v is None:
        return None
    if isinstance(v, (float, np.floating)) and not np.isfinite(v):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    return v


def write_table(wb, name, header, rows, widths=None, p_cols=(), fmt_cols=None, wrap=False):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append([clean(v) for v in r])
    for c in ws[1]:
        c.font = FONT_B
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            if wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    fmt_cols = dict(fmt_cols or {})
    for col in p_cols:
        fmt_cols[col] = FMT_P
    for col, fmt in fmt_cols.items():
        j = header.index(col) + 1
        for i in range(2, ws.max_row + 1):
            ws.cell(i, j).number_format = fmt
    for j, h in enumerate(header, start=1):
        if widths and h in widths:
            w = widths[h]
        else:
            vals = [len(str(h))] + [len(str(ws.cell(i, j).value)) for i in range(2, min(ws.max_row, 200) + 1)
                                    if ws.cell(i, j).value is not None
                                    and not str(ws.cell(i, j).value).startswith("=")]
            w = min(max(vals) + 2, 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def stats_sheet(wb, name, df, post_ip=None):
    """A stats file as a sheet, provenance columns left out. post_ip, if given, is
    one value per row, written as post_ip_pct_median after n_donors."""
    cols = [c for c in df.columns if c not in PROV_COLS]
    d = df[cols].copy()
    fmt = {}
    if post_ip is not None:
        d.insert(cols.index("n_donors") + 1, "post_ip_pct_median", list(post_ip))
        fmt = {"post_ip_pct_median": "0.00"}
    write_table(wb, name, list(d.columns), d.itertuples(index=False, name=None),
                p_cols=("p", "fdr"), fmt_cols=fmt)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    default_root = pathlib.Path(os.environ.get("SEAAD_ROOT") or ".") / "Part12_rerun"
    ap.add_argument("--root", default=str(default_root),
                    help="folder holding the Figure 16 outputs and logs (default: $SEAAD_ROOT/Part12_rerun, "
                         "or ./Part12_rerun)")
    ap.add_argument("--out", default=None, help="output xlsx (default: <root>/08_Fig16/SuppTable_Fig16.xlsx)")
    args = ap.parse_args()
    root = pathlib.Path(args.root).resolve()
    out = pathlib.Path(args.out).resolve() if args.out else root / "08_Fig16" / "SuppTable_Fig16.xlsx"
    L = Ledger()
    L._say("build_SuppTable_S14.py")
    L._say(f"Start: {dt.datetime.now():%Y-%m-%d %H:%M:%S} | root {root}")
    L._say("")

    # ---- logs and input files ---------------------------------------------------------
    L._say("SECTION A -- logs and inputs")
    logs, files, D = {}, {}, {}
    for step, folder, names in SECTIONS:
        lg = parse_log(find(root, f"Part12_{folder}_log.txt"))
        logs[step] = lg
        L.chk(f"B-01.{step}", lg["verdict"] == "PASS" and lg["n_fail"] == 0,
              f"section {step} log is PASS with 0 FAIL", f"{lg['checks']} | {lg['verdict']}")
        for nm in names:
            p = find(root, nm)
            files[nm] = (step, p)
            L.chk(f"B-02.{nm}", nm in lg["outputs"],
                  f"{nm} is listed in the outputs of the section {step} log", lg["path"].name)
            D[nm] = read_csv_like_r(p, char_cols=("panel", "arm", "level", "quantity", "measure", "family",
                                                  "run_id", "labels_version", "donor_key", "pair"))
    one = lambda k: sorted({lg[k] for lg in logs.values()})
    L.chk("B-03", len(one("script")) == 1 and len(one("config_version")) == 1
          and len(one("labels_version")) == 1,
          "one analysis script, config and labels version across sections 13 to 17",
          f"{one('script')} | config {one('config_version')} | labels {one('labels_version')}")

    # ---- row provenance -------------------------------------------------------------
    for nm, (step, _) in files.items():
        if not nm.startswith("Stats_"):
            continue
        d, lg = D[nm], logs[step]
        ok = (set(d["run_id"]) == {lg["run_id"]}
              and set(d["config_version"].astype(int).astype(str)) == {lg["config_version"]}
              and set(d["labels_version"]) == {lg["labels_version"]})
        L.chk(f"B-04.{nm}", ok, f"{nm} rows carry the log's run ID, config and labels version",
              f"{lg['run_id']} | {lg['config_version']} | {lg['labels_version']}")

    # ---- row counts against the logs ----------------------------------------------------
    ## "16A 20 rows | 16B 12 rows", "30 rows", "10 stats rows | 166 donor rows"
    for step, folder, names in SECTIONS:
        rl = logs[step]["rows_line"] or ""
        pairs = dict(re.findall(r"(\w+) (\d+) rows", rl))
        kinds = {k: int(v) for v, k in re.findall(r"(\d+) (stats|donor) rows", rl)}
        stats = [n for n in names if n.startswith("Stats_")]
        if kinds:
            ok = (kinds.get("stats", -1) == sum(len(D[n]) for n in stats)
                  and all(kinds.get("donor", -1) == len(D[n]) for n in names if n.startswith("Plot_")))
        elif pairs:
            ok = all(int(pairs.get(n[6:-4], -1)) == len(D[n]) for n in stats)
        else:
            m = re.search(r"(\d+) rows", rl)
            ok = bool(m) and int(m.group(1)) == sum(len(D[n]) for n in stats)
        L.chk(f"B-05.{step}", ok, f"section {step} row counts equal the log", rl)

    # ---- BH within family -----------------------------------------------------------
    for nm in [n for n in D if n.startswith("Stats_")]:
        d = D[nm]
        ok = all(close(bh(g["p"].values), g["fdr"].values, 1e-12) for _, g in d.groupby("family", sort=False))
        L.chk(f"B-06.{nm}", ok, f"{nm} fdr equals BH within family", f"{d['family'].nunique()} families")
    L._say("")

    # ---- panel consistency -------------------------------------------------------------
    L._say("SECTION B -- panel consistency")
    A, Bt, C, Dd, Pd, E, F, G = (D[n] for n in ("Stats_16A.csv", "Stats_16B.csv", "Stats_16C.csv",
                                                 "Stats_16D.csv", "Plot_16D_donors.csv", "Stats_16E.csv",
                                                 "Stats_16F.csv", "Stats_16G.csv"))
    # 16C
    SH = ["share_both", "share_test_only", "share_ref_only", "share_neither"]
    ok, worst, nk = True, 0.0, 0
    for (a, lv), g in C.groupby(["arm", "level"], sort=False):
        s = [g[g.quantity == q]["estimate"].iloc[0] for q in SH]
        k = g[g.quantity == "kappa"]["estimate"].iloc[0]
        ok &= abs(sum(s) - 1) < 1e-9 and g["n_units"].nunique() == 1 and g["n_donors"].nunique() == 1
        kk = kappa_from_shares(s)
        ok &= close([kk], [k], 1e-9)
        worst = max(worst, abs(kk - k)); nk += 1
    L.chk("B-07", ok, "16C shares sum to 1 and kappa from the four shares equals the kappa row",
          f"{nk} rows | largest kappa difference {worst:.1e}")

    # 16D
    ok = True
    for pair, g in Pd.groupby("pair", sort=False):
        s = Dd[Dd.measure == pair].set_index("quantity")
        for col, qn in (("rs_raw", "within_median_raw"), ("rs_partial", "within_median_partial")):
            v = g[col].dropna().values
            ok &= close([np.median(v)], [s.loc[qn, "estimate"]]) and int(s.loc[qn, "n_donors"]) == len(v)
        v = g["rs_raw"].dropna().values
        ok &= close([np.percentile(v, 25), np.median(v), np.percentile(v, 75)],
                    [s.loc["within_iqr_raw", "lo"], s.loc["within_iqr_raw", "estimate"], s.loc["within_iqr_raw", "hi"]])
    L.chk("B-08", ok, "16D within-donor medians, quartiles and donor counts equal the plot data",
          f"{Pd['pair'].nunique()} pairs, {Pd['donor_key'].nunique()} donors")

    # 16E
    ok, n_e = True, 0
    for a in ARMS:
        e = E[E.arm == a].set_index(["level", "quantity"])
        d = e.loc[("branches", "d_offset_share_pp")]
        d9 = e.loc[("branches", "d_offset_share_pp_90")]
        sp, sq = e.loc[("post-IP", "offset_share_pct"), "estimate"], e.loc[("pre-IP", "offset_share_pct"), "estimate"]
        if np.isfinite(d["estimate"]):
            ok &= close([d["estimate"]], [sp - sq], 1e-9)
            within = bool(np.isfinite(d9["lo"]) and d9["lo"] > -MARGIN and d9["hi"] < MARGIN)
            nonzero = bool(np.isfinite(d["fdr"]) and d["fdr"] < 0.05)
            ok &= d9["measure"] == VERDICT_16E[(within, nonzero)]
        else:
            ok &= d9["measure"] == "untestable"
        n_e += 1
    L.chk("B-09", ok, "16E difference equals post-IP minus pre-IP share and every verdict follows from the 90% interval and FDR",
          f"{n_e} arms")

    # 16F/16G contrasts
    ok, n_c = True, 0
    for S in (F, G):
        for (a, lv, f), g in S.groupby(["arm", "level", "measure"], sort=False):
            v = g.set_index("quantity")["estimate"]
            exp = {"module_minus_mito": abs(v["module"]) - abs(v["mito"]),
                   "module_minus_depth": v["module"] - v["depth"]}
            if a != "HKG8":
                h = S[(S.arm == "HKG8") & (S.level == lv) & (S.measure == f) & (S.quantity == "module")]["estimate"].iloc[0]
                exp["arm_minus_hkg8"] = v["module"] - h
            for k, x in exp.items():
                ok &= close([x], [v[k]], 1e-12); n_c += 1
    L.chk("B-10", ok, "every 16F and 16G contrast equals the difference of its two rs rows", f"{n_c} contrasts")

    # 16G four conditions against the log
    cond = conditions_16g(G)
    logc = log_conditions(logs["17"]["text"])
    L.chk("B-11", cond == logc and len(cond) > 0, "the four 16G conditions recomputed from Stats_16G equal the section 17 log",
          f"{len(cond)} arm-field rows; orders: " + (", ".join(f"{a} {f}" for (a, f), c in cond.items()
                                                              if c[4] == "orders Module post-IP%") or "none"))

    # ---- figure log ----------------------------------------------------------------
    fl = parse_log(find(root, FIG_LOG))
    last = max(dt.datetime.strptime(lg["start"], "%Y-%m-%d %H:%M:%S") for lg in logs.values())
    fstart = dt.datetime.strptime(fl["start"], "%Y-%m-%d %H:%M:%S")
    L.chk("B-12", fl["verdict"] == "PASS" and fl["n_fail"] == 0 and fstart >= last,
          "figure log is PASS and started after the last analysis section",
          f"{fl['script']} | {fl['start']}")
    L._say("")

    # ---- config and unit table ---------------------------------------------------------
    L._say("SECTION C -- unit table")
    cfg_path = find(root, CONFIG)
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    L.chk("B-14.config", cfg_path.is_file() and str(cfg.get("config_version")) == one("config_version")[0],
          f"{CONFIG} exists and has the config version of the section logs",
          f"version {cfg.get('config_version')}")
    up_path, up_log = find(root, UNITS[0]), parse_log(find_glob(root, UNITS[1]))
    L.chk("B-14.units", UNITS[0] in up_log["outputs"] and up_log["verdict"] == "PASS",
          f"{UNITS[0]} is listed in the outputs of its log, which is PASS", up_log["path"].name)
    U = read_csv_like_r(up_path, char_cols=("arm", "dataset", "donor_id", "donor_key", "region", "libs",
                                            "labels_version", "run_id"))
    L.chk("B-14.labels", set(U["labels_version"]) == set(one("labels_version")),
          "unit table carries the labels version of the section logs", ", ".join(sorted(set(U["labels_version"]))))
    min_d = int(cfg["floors"]["donors_per_test"])
    min_u = int(cfg["floors"]["finite_units_within_donor"])
    max_miss = float(cfg["floors"]["max_missing_share"])
    den = cfg["quantities"]["module_post_pct_denominator"]

    # library fields kept under the missing-value limit, as section 17
    lib_all = ["intronic", "antisense", "neun_fans", "cdna_400", "pcr_cycles", "saturation"]
    miss = {f: (float(np.mean(~np.isfinite(pd.to_numeric(U[f"lib_{f}"], errors="coerce").astype(float))))
                if f"lib_{f}" in U.columns else 1.0) for f in lib_all}
    kept = [f for f in lib_all if f"lib_{f}" in U.columns and miss[f] <= max_miss]
    f_fields = list(pd.unique(F["measure"]))
    g_fields = list(pd.unique(G["measure"]))
    dropped = [f for f in lib_all if f not in kept]

    R16A = rebuild_16a(U, min_d)
    R16B = rebuild_16b(U)
    R16FG = rebuild_16fg(U, f_fields + g_fields, min_d, min_u)

    ok = True
    for r in A.itertuples():
        x = R16A[(r.arm, r.quantity)]
        ok &= x["n_units"] == r.n_units and x["n_donors"] == r.n_donors
    for r in Bt.itertuples():
        x = R16B[(r.arm, r.quantity)]
        ok &= x["n_units"] == r.n_units and x["n_donors"] == r.n_donors
    for S in (F, G):
        for r in S.itertuples():
            q = r.quantity if r.quantity in QS else "module"
            x = R16FG[(r.arm, r.measure, r.level, q)]
            ok &= x["n_units"] == r.n_units and x["n_donors"] == r.n_donors
    ok &= kept == f_fields + g_fields
    L.chk("B-15", ok, "rebuilt unit and donor sets have the n_units and n_donors of every 16A, 16B, 16F and 16G row, "
          "and the fields kept are the fields in the stats files",
          f"{len(A) + len(Bt) + len(F) + len(G)} rows | kept {', '.join(kept)} | dropped "
          + (", ".join(f"{f} ({100 * miss[f]:.1f}% missing)" for f in dropped) or "none"))

    # ---- Fig16_units: the values the recomputation reads -------------------------------
    ucols = [("arm", "arm"), ("dataset", "region_set"), ("donor_id", "donor_id"), ("donor_key", "donor_key"),
             ("region", "region"), ("n_cells", "n_labelled_cells"), ("n_post", "n_module_post_ip_cells"),
             ("n_super", "n_supertype_cells"), ("module_post_pct", "post_ip_pct"),
             ("module_resid_tq4", "post_ip_pct_resid_tq4"), ("mito_pct", "mito_pct"), ("mito_log1p", "mito_log1p"),
             ("contamination", "contamination"), ("ctrl_pct_null", "ctrl_pct_null"),
             ("log_n_cells", "log_n_cells"), ("median_umi", "median_umi"), ("log_median_umi", "log_median_umi")]
    ucols += [(f"lib_{f}", f"lib_{f}") for f in kept]
    US = U.copy()
    US["_a"] = US["arm"].map({a: i for i, a in enumerate(ARMS)})
    ## Stable sort by arm only: within an arm the rows keep the unit-table order,
    ## which is the order R summed them in (05 writes the units sorted by arm,
    ## dataset, donor and region).
    US = US[US["_a"].notna()].sort_values("_a", kind="stable").reset_index(drop=True)
    # the recomputation reads the sheet's own frame, renamed back to the unit-table names
    SHEET = US[[c for c, _ in ucols]].copy()
    R16A_s = rebuild_16a(SHEET, min_d)
    R16FG_s = rebuild_16fg(SHEET, f_fields + g_fields, min_d, min_u)
    worst, worst_at, n_fit, ok, n_bad = 0.0, "", 0, True, 0
    def cmp16(v, ref, where):
        nonlocal worst, worst_at, n_fit, ok, n_bad
        n_fit += 1
        if np.isnan(v) or np.isnan(ref):
            if not (np.isnan(v) and np.isnan(ref)):
                ok = False; n_bad += 1
            return
        if abs(v - ref) > worst:
            worst, worst_at = abs(v - ref), where
        if abs(v - ref) >= TOL:
            n_bad += 1
    for r in A.itertuples():
        cmp16(R16A_s[(r.arm, r.quantity)]["estimate"], r.estimate, f"16A {r.arm} {r.quantity}")
    for S in (F, G):
        for r in S[S.quantity.isin(list(QS))].itertuples():
            cmp16(R16FG_s[(r.arm, r.measure, r.level, r.quantity)]["estimate"], r.estimate,
                  f"{r.panel} {r.arm} {r.level} {r.measure} {r.quantity}")
    ok &= worst < TOL
    L.chk("B-16", ok, "every 16A donor share and every observed 16F and 16G rs recomputed from the Fig16_units values equals its stats row",
          f"{n_fit} values | {n_bad} differ | largest difference {worst:.1e}" + (f" ({worst_at})" if worst_at else ""))
    L.chk("B-17", len(SHEET) == len(U), "Fig16_units holds every unit of the unit table", f"{len(SHEET)} units")
    L._say("")

    # ---- Module post-IP% per stats row ---------------------------------------------------
    post_a = [R16A[(r.arm, r.quantity)]["post"] if r.quantity in ("module_post_pct", "module_resid_tq4") else np.nan
              for r in A.itertuples()]
    post_b = [R16B[(r.arm, r.quantity)]["post"] if r.quantity in ("module_post_pct", "module_resid_tq4") else np.nan
              for r in Bt.itertuples()]
    post_fg = lambda S: [R16FG[(r.arm, r.measure, r.level, "module")]["post"] if r.quantity in MODULE_ROWS else np.nan
                         for r in S.itertuples()]

    # ---- workbook ------------------------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)
    B_ = int(cfg["inference"]["bootstrap_B"])
    R_ = int(cfg["inference"]["permutation_R"])
    nd16d = int(Dd[(Dd.quantity == "within_median_raw") & (Dd.measure == "MT-CO3 vs cDFG")]["n_donors"].iloc[0])
    orders = [f"{a} {FIELD_LAB.get(f, f)}" for (a, f), c in cond.items() if c[4] == "orders Module post-IP%"]
    e_untest = [a for a in ARMS if E[(E.arm == a) & (E.quantity == "d_offset_share_pp_90")]["measure"].iloc[0] == "untestable"]

    ov = [
        ["Fig 16A", "Is Module post-IP% donor-structured?",
         "Donor share of variance across donor-region units (donor entered first) for five quantities, against donor labels shuffled within region",
         "13", "Stats_16A.csv", "Fig16A_readout; Fig16A_stats; Fig16_units", "Notes 2, 5 and 11"],
        ["Fig 16B", "Does a donor keep its value across regions?",
         "Split-half rs of donor means over 200 random half-splits of each donor's units; median over splits",
         "13", "Stats_16B.csv", "Fig16B_stats; Fig16_units", "Notes 6 and 11"],
        ["Fig 16C", "Do two arms label the same cells?",
         "Cohen's kappa for Module membership on cells shared with the PVM_3 reference, with the four agreement shares",
         "14", "Stats_16C.csv", "Fig16C_readout; Fig16C_stats", "Note 7"],
        ["Fig 16D", "Does the MT-CO3 to cDFG relationship hold inside single donors?",
         "Per-donor rs of MT-CO3 with the cDFG composite, raw and depth-partialled; median over donors; pooled and between-donor rs; ACTB with cHKG8 as the control pair",
         "15", "Stats_16D.csv; Plot_16D_donors.csv", "Fig16D_stats; Fig16D_donors", "Note 8"],
        ["Fig 16E", "Is the offset share equal between the branches?",
         "Offset share of MT UMIs (median intercept over median fitted MT UMIs at the reference depth), Module post-IP minus Module pre-IP",
         "16", "Stats_16E.csv", "Fig16E_readout; Fig16E_stats", "Note 9"],
        ["Fig 16F", "Does Module post-IP% track the library intronic rate?",
         "Spearman rs with the library intronic rate between donors and within donor, with comparators and contrasts",
         "17", "Stats_16F.csv", "Fig16F_stats; Fig16_units", "Notes 10 and 11"],
        ["Fig 16G, Supp Fig 12", "Do sort purity or library construction order Module post-IP%?",
         "As 16F for each library field kept; a field orders Module post-IP% only if four pre-set conditions hold within donor",
         "17", "Stats_16G.csv", "Fig16G_conditions; Fig16G_stats; Fig16_units", "Notes 10, 11 and 14"],
    ]
    write_table(wb, "Overview", ["Panel", "Question", "Quantity", "Section", "Source file", "Sheets", "See"], ov,
                wrap=True, widths={"Panel": 10, "Question": 40, "Quantity": 60, "Section": 9,
                                   "Source file": 32, "Sheets": 36, "See": 18})

    # HKG8 within-donor Module rs over the 16G fields, for Note 14
    hk_g = G[(G.arm == "HKG8") & (G.level == "within donor") & (G.quantity == "module")]
    n_hk = int(hk_g["estimate"].notna().sum())
    n_hk_sig = int((hk_g["fdr"] < 0.05).sum())
    notes = [
        ["1", "Arms", "M4 = Mic13_PVM3, the masked primary arm. PVM_3 = PVM3_all, the unmasked reference. PVM_1 = PVM1_all, the second supertype. HKG8 = the housekeeping control construct (ACTB with the cHKG8 genes). NeoC and NonNeoC are the two region sets; 16A, 16B, 16F and 16G pool them."],
        ["2", "Resampling and p values", f"Intervals in 16B to 16G are 95% percentile intervals from a donor bootstrap with B = {B_:,}; 16E also gives the 90% interval. The bootstrap p is min(1, 2(k+1)/(B+1)); its smallest value, 2/(B+1) = {2 / (B_ + 1):.4e}, is a bound. 16A uses a permutation null instead: donor labels shuffled within region {R_:,} times, p = (k+1)/(R+1), smallest value 1/(R+1) = {1 / (R_ + 1):.4e}. In Stats_16A the lo and hi columns hold the null median and the null 95th percentile, not an interval. A BH value computed from floor p values is also a bound. 16C and 16D carry no p."],
        ["3", "Multiple testing", "fdr is Benjamini-Hochberg within the group named in the family column. 16A: within arm over the five quantities. 16B: within arm over the three quantities. 16E: over the arms. 16F and 16G: within arm and level, the rs rows (Module, %MT, cell count and depth over the fields) as one family and the contrasts as another. The builder recomputes BH within every family (check B-06)."],
        ["4", "Floors and untestable rows", f"A unit is one donor in one region on one arm with at least {int(cfg['floors']['cells_per_unit'])} cells of the arm. A test needs {min_d} donors. The within-donor level keeps donors with at least {min_u} units with a finite value. 16D and 16E need {int(cfg['floors']['cells_per_group_per_donor'])} cells per donor (per branch in 16E); 16C keeps regions with at least {int(cfg['floors']['cells_per_group_per_donor'])} shared cells. A row below a floor has an empty estimate, interval, p and fdr and is reported as untestable" + (f" ({', '.join(e_untest)} in 16E)." if e_untest else ".")],
        ["5", "Fig 16A", "Donor share = between-donor sum of squares over total sum of squares of the quantity across the arm's units, donor entered first. The TQR residual is Module post-IP% after regression on PMI, RIN, brain pH and fresh brain weight over the arm's units (tissue quality removed). Contamination is the unit mean DecontX contamination; the matched non-MT share null (ctrl_pct_null) is the unit mean of the 200-draw null of a non-MT gene share matched to the MT genes. In Fig16A_readout the gap column is the donor share minus the null 95th percentile, in percentage points."],
        ["6", "Fig 16B", "Donors with at least two units enter. Each split puts a random half of a donor's units in each half; the donor's mean is taken in each half, and rs is computed across donors between the halves. The estimate is the median rs over 200 splits; the donor bootstrap reuses the same splits. The splits are random draws in R from the step seed, so 16B cannot be recomputed from Fig16_units without the analysis script; the builder checks its unit and donor counts (check B-15)."],
        ["7", "Fig 16C", "Cells present in both the test arm and the PVM_3 reference, in regions with at least 20 shared cells. share_both: Module on both; share_test_only: Module on the test arm only; share_ref_only: Module on PVM_3 only; share_neither: Module on neither. kappa is recomputed from the four shares by the builder (check B-07). Pre-set thresholds, read on the pooled point estimate: M4 above 0.60, HKG8 |kappa| below 0.20. PVM_3 is the reference and PVM_1 shares no cells with it."],
        ["8", "Fig 16D", f"Population: the PVM_3 arm cells (MT-CO3-positive Micro-PVM_3 cells). Per donor with at least 20 cells, rs of MT-CO3 with the cDFG composite (mean log1p CP10K of the 8 DFG genes); the partial rs removes log depth on ranks. The median over {nd16d} donors is read against pre-set marks 0.75 and 0.40. The within_iqr_raw row gives the median with the 25th and 75th percentiles as lo and hi. Pooled rs is over all cells; between-donor rs is over donor means. Fig16D_donors holds the per-donor values."],
        ["9", "Fig 16E", "Per donor and branch (Module post-IP, Module pre-IP) with at least 20 cells, OLS of MT UMIs on total UMIs gives intercept a and slope b. The offset share is median a / (median a + median b x reference depth), reference depth = the arm's median total UMIs over its Module cells. The difference is post-IP minus pre-IP in percentage points, with a paired donor bootstrap. Verdict: within the 10 pp margin when the whole 90% interval lies inside +/-10 pp; non-zero when fdr < 0.05 over the arms (dated note of 2026-09-22 in the pre-set readings). The builder recomputes every verdict (check B-09)."],
        ["10", "Fig 16F and 16G", "Between donors: donor means over units. Within donor: units centred on their donor mean. Quantities: Module post-IP% (module), log1p %MT (mito), log cell count (ncells) and log median UMIs (depth). Contrasts on the same donor resamples: module_minus_mito = |Module rs| - |%MT rs|; module_minus_depth = Module rs - depth rs; arm_minus_hkg8 = Module rs of the arm - Module rs of HKG8. Library fields are kept when no more than 25% of units lack them" + (f"; dropped: {', '.join(f'{FIELD_LAB[f]} ({100 * miss[f]:.1f}% of units missing)' for f in dropped)}." if dropped else ".") + " A 16G field orders Module post-IP% only if, within donor, the Module rs has fdr < 0.05, the depth contrast and the HKG8 contrast intervals exclude zero, and the %MT contrast interval is not wholly below zero (Fig16G_conditions). Fields meeting all four: " + (", ".join(orders) if orders else "none") + "."],
        ["11", "Module post-IP% columns", "post_ip_pct_median gives the median Module post-IP% (per cent of the arm's " + ("supertype cells" if den == "supertype" else "labelled cells") + ") over the units behind the row: the units of that arm with a value for the row's quantity (16A), the units of donors with at least two units (16B), the units kept at the within-donor level or the donor means at the between-donor level (16F, 16G). It is filled on the Module rows (module_post_pct and its TQR residual in 16A and 16B; module and the three contrasts in 16F and 16G) and empty on comparator rows. 16C, 16D and 16E do not use Module post-IP%. The builder rebuilds every unit set from 05_units.csv and checks its n_units and n_donors against the stats row (check B-15)."],
        ["12", "Unit values (Fig16_units)", "One row per unit, with every value 16A, 16F and 16G use. post_ip_pct = 100 x n_module_post_ip_cells / " + ("n_supertype_cells" if den == "supertype" else "n_labelled_cells") + ". To recompute a 16A donor share: take the arm's units with a value, and divide the between-donor sum of squares by the total sum of squares. To recompute a 16F or 16G rs: between donors, average the quantity and the field over each donor's units with both values and take Spearman rs over donors; within donor, keep the units with both values, keep donors with at least 3 such units, subtract each donor's mean from both, and take Spearman rs over units. A row with fewer than 10 donors is untestable. The builder recomputes every 16A share and every observed 16F and 16G rs this way (check B-16)."],
        ["13", "Old values", "The archived run is not compared here (decision D20 dropped, revision 49). Results are reported as they stand."],
        ["14", "Supplementary Figure 12", "Figure 16G draws M4 and PVM_3. Supplementary Figure 12 draws the same analysis for PVM_1 (panels A and B) and HKG8 (panels A and C); its values are the PVM_1 and HKG8 rows of Fig16G_stats. "
         "HKG8 is the control, so it has no arm_minus_hkg8 contrast and no row in Fig16G_conditions; "
         f"{n_hk_sig} of its {n_hk} within-donor Module rs have fdr < 0.05."],
    ]
    write_table(wb, "Notes", ["Note", "Subject", "Detail"], notes, wrap=True,
                widths={"Note": 7, "Subject": 28, "Detail": 120})

    gl = [
        ["column", "panel", "Figure panel the row belongs to."],
        ["column", "arm", "M4, PVM_3, PVM_1 or HKG8 (Note 1)."],
        ["column", "level", "16A: units; 16B: donors; 16C: pooled or a region set; 16D: within donor, pooled or between donors; 16E: branches, post-IP, pre-IP or all Module cells; 16F, 16G: between donors or within donor."],
        ["column", "quantity", "What the row estimates; codes below."],
        ["column", "measure", "16C, 16D: the comparison; 16E: the verdict on the d_offset_share_pp_90 row; 16F, 16G: the library field."],
        ["column", "n_units", "Units behind the row (16A, 16B, 16F, 16G; donors at the between-donor level of 16F and 16G), shared cells (16C), cells (16D)."],
        ["column", "n_donors", "Donors behind the row."],
        ["column", "post_ip_pct_median", "Median Module post-IP% over the units or donors behind the row (Note 11)."],
        ["column", "estimate, lo, hi", "Point estimate and its 95% interval (Note 2); exceptions: 16A (null median and null 95th percentile), 16D within_iqr_raw (quartiles), 16E d_offset_share_pp_90 (90% interval)."],
        ["column", "p, fdr", "p value and Benjamini-Hochberg value within family (Notes 2 and 3)."],
        ["column", "family", "The BH group of the row."],
        ["16A", "module_post_pct, module_resid_tq4, mito_log1p, contamination, ctrl_pct_null", "Module post-IP%, its TQR residual, log1p %MT, unit mean DecontX contamination and the matched non-MT share null (Note 5)."],
        ["16A", "donor_share_pct, null_median_pct, null_95th_pct, over_null_95th_pp", "In Fig16A_readout: donor share, null median and null 95th percentile in per cent, and the gap to the null 95th percentile in percentage points."],
        ["16B", "module_post_pct, module_resid_tq4, mito_log1p", "Split-half rs of that quantity, median over splits (Note 6)."],
        ["16C", "kappa", "Cohen's kappa for Module membership against PVM_3."],
        ["16C", "share_both, share_test_only, share_ref_only, share_neither", "Fractions of shared cells in the four agreement cells (Note 7); per cent in Fig16C_readout."],
        ["16D", "within_median_raw, within_median_partial", "Median over donors of the per-donor rs, raw or with log depth partialled out."],
        ["16D", "within_iqr_raw", "Median and quartiles of the per-donor raw rs."],
        ["16D", "pooled_raw, between_raw", "rs over all cells; rs over donor means."],
        ["16D", "rs_raw, rs_partial, n_cells", "In Fig16D_donors: the donor's rs, raw and partial, and its cells."],
        ["16E", "d_offset_share_pp", "Offset share post-IP minus pre-IP, percentage points, 95% interval."],
        ["16E", "d_offset_share_pp_90", "The same with the 90% interval; measure holds the verdict."],
        ["16E", "offset_share_pct, intercept", "Per branch: offset share in per cent and median intercept a (MT UMIs); n_donors = donors fitted in that branch."],
        ["16E", "reference_depth", "The arm's median total UMIs over its Module cells."],
        ["16F, 16G", "module, mito, ncells, depth", "Spearman rs of Module post-IP%, log1p %MT, log cell count or log median UMIs with the field."],
        ["16F, 16G", "module_minus_mito, module_minus_depth, arm_minus_hkg8", "Contrasts on the same donor resamples (Note 10)."],
        ["16G", "c1_nonzero, c2_vs_depth, c3_vs_hkg8, c4_not_matched_by_mt", "In Fig16G_conditions: the four pre-set conditions (Note 10), as Excel formulas over the same row."],
        ["units", "n_labelled_cells, n_module_post_ip_cells, n_supertype_cells", "The unit's labelled cells of the arm, its Module post-IP cells and its cells of the arm's supertype."],
        ["units", "post_ip_pct, post_ip_pct_resid_tq4", "Module post-IP% (Note 12) and its TQR residual."],
        ["units", "mito_pct, mito_log1p", "Unit mean %MT over the 13 MT protein-coding genes, and log1p of it."],
        ["units", "contamination, ctrl_pct_null", "Unit mean DecontX contamination; unit mean matched non-MT share null."],
        ["units", "log_n_cells, median_umi, log_median_umi", "log of n_labelled_cells; median UMIs per cell and its log."],
        ["units", "lib_<field>", "Median over the unit's cells of the library field (constant within a library)."],
    ]
    write_table(wb, "Glossary", ["Scope", "Term", "Meaning"], gl, wrap=True,
                widths={"Scope": 10, "Term": 46, "Meaning": 100})

    # Fig16A readout
    hA = ["arm", "quantity", "n_units", "n_donors", "post_ip_pct_median", "donor_share_pct",
          "null_median_pct", "null_95th_pct", "over_null_95th_pp", "p", "fdr"]
    cA = {h: get_column_letter(j + 1) for j, h in enumerate(hA)}
    rows = []
    for i, (r, pp) in enumerate(zip(A.itertuples(), post_a), start=2):
        s_, h_ = cA["donor_share_pct"], cA["null_95th_pct"]
        rows.append([r.arm, QLAB[r.quantity], r.n_units, r.n_donors, pp, 100 * r.estimate, 100 * r.lo,
                     100 * r.hi, f'=IF(AND(ISNUMBER({s_}{i}),ISNUMBER({h_}{i})),{s_}{i}-{h_}{i},"")', r.p, r.fdr])
    write_table(wb, "Fig16A_readout", hA, rows, p_cols=("p", "fdr"),
                fmt_cols={"post_ip_pct_median": "0.00", "donor_share_pct": "0.0", "null_median_pct": "0.0",
                          "null_95th_pct": "0.0", "over_null_95th_pp": "0.0"})

    # Fig16C readout
    hC = ["arm", "region_set", "shared_cells", "n_donors", "kappa", "kappa_lo", "kappa_hi", "pct_both",
          "pct_test_only", "pct_ref_only", "pct_neither", "threshold", "threshold_met"]
    cC = {h: get_column_letter(j + 1) for j, h in enumerate(hC)}
    rows, i = [], 2
    for (a, lv), g in C.groupby(["arm", "level"], sort=False):
        k = g[g.quantity == "kappa"].iloc[0]
        sh_ = [100 * g[g.quantity == q]["estimate"].iloc[0] for q in SH]
        thr, met = None, None
        if lv == "pooled":
            kc = f"{cC['kappa']}{i}"
            thr = "kappa > 0.60" if a == "M4" else "|kappa| < 0.20"
            met = f'=IF({kc}>0.6,"yes","no")' if a == "M4" else f'=IF(ABS({kc})<0.2,"yes","no")'
        rows.append([a, lv, k["n_units"], k["n_donors"], k["estimate"], k["lo"], k["hi"]] + sh_ + [thr, met])
        i += 1
    write_table(wb, "Fig16C_readout", hC, rows,
                fmt_cols={"kappa": "0.000", "kappa_lo": "0.000", "kappa_hi": "0.000", "pct_both": "0.0",
                          "pct_test_only": "0.0", "pct_ref_only": "0.0", "pct_neither": "0.0"})

    # Fig16E readout
    hE = ["arm", "n_donors_post", "n_donors_pre", "n_donors_either", "reference_depth_umis", "intercept_post",
          "intercept_pre", "offset_share_post_pct", "offset_share_pre_pct", "difference_pp", "lo_95", "hi_95",
          "p", "fdr", "lo_90", "hi_90", "verdict"]
    rows = []
    for a in ARMS:
        e = E[E.arm == a].set_index(["level", "quantity"])
        d, d9 = e.loc[("branches", "d_offset_share_pp")], e.loc[("branches", "d_offset_share_pp_90")]
        rows.append([a, e.loc[("post-IP", "offset_share_pct"), "n_donors"], e.loc[("pre-IP", "offset_share_pct"), "n_donors"],
                     d["n_donors"], e.loc[("all Module cells", "reference_depth"), "estimate"],
                     e.loc[("post-IP", "intercept"), "estimate"], e.loc[("pre-IP", "intercept"), "estimate"],
                     e.loc[("post-IP", "offset_share_pct"), "estimate"], e.loc[("pre-IP", "offset_share_pct"), "estimate"],
                     d["estimate"], d["lo"], d["hi"], d["p"], d["fdr"], d9["lo"], d9["hi"], d9["measure"]])
    write_table(wb, "Fig16E_readout", hE, rows, p_cols=("p", "fdr"),
                fmt_cols={k: "0.0" for k in ("reference_depth_umis", "intercept_post", "intercept_pre",
                                              "offset_share_post_pct", "offset_share_pre_pct", "difference_pp",
                                              "lo_95", "hi_95", "lo_90", "hi_90")},
                widths={"verdict": 44})

    # Fig16G conditions, formulas over the row
    hG = ["arm", "field", "n_units", "n_donors", "post_ip_pct_median", "module_rs", "module_rs_lo", "module_rs_hi",
          "module_rs_fdr", "depth_contrast_lo", "depth_contrast_hi", "hkg8_contrast_lo", "hkg8_contrast_hi",
          "mito_contrast_hi", "c1_nonzero", "c2_vs_depth", "c3_vs_hkg8", "c4_not_matched_by_mt", "verdict"]
    cG = {h: get_column_letter(j + 1) for j, h in enumerate(hG)}
    rows, i = [], 2
    for (a, f) in cond:
        g = G[(G.arm == a) & (G.level == "within donor") & (G.measure == f)].set_index("quantity")
        m = g.loc["module"]
        c = lambda h: f"{cG[h]}{i}"
        rows.append([a, f, m["n_units"], m["n_donors"], R16FG[(a, f, "within donor", "module")]["post"],
                     m["estimate"], m["lo"], m["hi"], m["fdr"],
                     g.loc["module_minus_depth", "lo"], g.loc["module_minus_depth", "hi"],
                     g.loc["arm_minus_hkg8", "lo"], g.loc["arm_minus_hkg8", "hi"], g.loc["module_minus_mito", "hi"],
                     f"=AND(ISNUMBER({c('module_rs_fdr')}),{c('module_rs_fdr')}<0.05)",
                     f"=AND(ISNUMBER({c('depth_contrast_lo')}),OR({c('depth_contrast_lo')}>0,{c('depth_contrast_hi')}<0))",
                     f"=AND(ISNUMBER({c('hkg8_contrast_lo')}),OR({c('hkg8_contrast_lo')}>0,{c('hkg8_contrast_hi')}<0))",
                     f"=AND(ISNUMBER({c('mito_contrast_hi')}),{c('mito_contrast_hi')}>=0)",
                     f'=IF(NOT(ISNUMBER({c("module_rs")})),"untestable (below the donor floor)",'
                     f'IF(AND({c("c1_nonzero")},{c("c2_vs_depth")},{c("c3_vs_hkg8")},{c("c4_not_matched_by_mt")}),'
                     f'"orders Module post-IP%","does not order Module post-IP%"))'])
        i += 1
    write_table(wb, "Fig16G_conditions", hG, rows, p_cols=("module_rs_fdr",),
                fmt_cols={"post_ip_pct_median": "0.00"}, widths={"verdict": 34})

    # Fig16_units
    hU = [n for _, n in ucols]
    write_table(wb, "Fig16_units", hU, SHEET.itertuples(index=False, name=None),
                fmt_cols={"post_ip_pct": "0.000", "post_ip_pct_resid_tq4": "0.000"})

    # Stats and plot-data sheets, as written by the analysis
    stats_sheet(wb, "Fig16A_stats", A, post_ip=post_a)
    stats_sheet(wb, "Fig16B_stats", Bt, post_ip=post_b)
    stats_sheet(wb, "Fig16C_stats", C)
    stats_sheet(wb, "Fig16D_stats", Dd)
    write_table(wb, "Fig16D_donors", list(Pd.columns), Pd.itertuples(index=False, name=None))
    stats_sheet(wb, "Fig16E_stats", E)
    stats_sheet(wb, "Fig16F_stats", F, post_ip=post_fg(F))
    stats_sheet(wb, "Fig16G_stats", G, post_ip=post_fg(G))

    # sheets in panel order: the three general sheets, then 16A to 16G, then the
    # unit values that 16A, 16B, 16F and 16G share
    order = ["Overview", "Notes", "Glossary",
             "Fig16A_readout", "Fig16A_stats", "Fig16B_stats",
             "Fig16C_readout", "Fig16C_stats", "Fig16D_stats", "Fig16D_donors",
             "Fig16E_readout", "Fig16E_stats", "Fig16F_stats",
             "Fig16G_conditions", "Fig16G_stats", "Fig16_units"]
    if sorted(order) != sorted(wb.sheetnames):
        sys.exit(f"ERROR: sheet order list does not match the sheets written: {wb.sheetnames}")
    wb._sheets = [wb[n] for n in order]
    wb.active = 0

    # every string in the workbook, for B-13
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts += [v for v in row if isinstance(v, str)]
    bad = [t for t in texts if any(d_ in t for d_ in DASHES)]
    L.chk("B-13", not bad, "no en dash or em dash in the workbook text", f"{len(bad)} found")

    L._say("")
    L._say(f"Checks: {L.n['PASS']} PASS, {L.n['FAIL']} FAIL, {L.n['NOTE']} NOTE")
    log_path = out.with_name(out.stem + "_build_log.txt")
    out.parent.mkdir(parents=True, exist_ok=True)
    if L.n["FAIL"]:
        L._say("VERDICT: FAIL | workbook not written")
        log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")
        sys.exit(1)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    L._say(f"Output: {out}")
    L._say("VERDICT: PASS")
    log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
