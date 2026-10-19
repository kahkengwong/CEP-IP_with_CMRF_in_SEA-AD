#!/usr/bin/env python3
"""
build_SuppTable_S17.py

Builds SuppTable_Fig19.xlsx (Figure 19 of the Part 12 rerun, Battery 5: do the
eleven severely affected donors drive the results?) from the tables written by
Part12_Fig19_Analysis.r (sections 26 to 30), the figure log written by
Part12_Fig19_Plots.r (section 31), and the Battery 5 config
(Part12_Fig19_config.json). finalize_supp_tables.py turns this workbook into
Supplementary Table S17.

Usage:
    python build_SuppTable_S17.py
    python build_SuppTable_S17.py --root /data/Part12_rerun
    python build_SuppTable_S17.py --root <folder> --out <folder>/SuppTable_Fig19.xlsx

Without --root the script uses $SEAAD_ROOT/Part12_rerun, or Part12_rerun in the
current folder when SEAAD_ROOT is not set. --root is searched recursively; every
input file must occur exactly once under it. Without --out the workbook goes to
<root>/11_Fig19/SuppTable_Fig19.xlsx (the folder is created if it does not
exist). A build log is written beside the workbook.

The script computes no new statistic. Before it writes anything it checks:
  B-01  the five section logs and the figure log end with VERDICT: PASS and 0 FAIL
  B-02  every input file is listed in the outputs of the log of the section that wrote it
  B-03  the five section logs name one analysis script, one config version, one
        labels version, one Battery 5 config with one version and one Battery 5
        readings file; the config here has that config version, the Battery 5
        config here is the one named and has that version, and the readings file
        named exists
  B-04  every row carries the run ID of its log, and every stats row the config and
        labels version of its log
  B-05  row counts equal the counts the logs report
  B-06  every fdr in the nine Figure 19 stats files equals BH recomputed within its family
  B-07  every readout row's without-eleven estimate and FDR equal its stats row, its
        SA change equals s x (e_without - e_all), and its class follows from its
        columns by the pre-set rule
  B-08  the class counts per panel and arm equal the section logs
  B-09  every without-eleven verdict of 17C, 17D and 18A follows from the Figure 19
        stats rows (FDRs, VIF, fitted)
  B-10  every influence change equals s x (e_minus_donor - e_all), and every
        influence summary row follows from the influence rows
  B-11  the SA donor table has the configured number of SA donors, every draw
        removes k donors of its frame's pool and no SA donor, and each frame's
        draw count matches its null status
  B-12  the figure log is PASS and was started after the last analysis section
  B-13  no text written to the workbook contains an en dash or an em dash
Any FAIL stops the build and no workbook is written.

Not carried, because of size: the random-removal estimates (Null_19_*.csv, up to
480 MB). The readout sheets carry their quantiles; the files are in the deposit.
Not recomputed: every statistic that needs the unit table, the joined histology
table, the count matrix or R's random streams (the without-eleven estimates, the
bootstrap and permutation p values, the random-removal quantiles).

Formatting: Arial throughout, bold header row, frozen top row. No fills, no
borders, no colours. The class column of the readout sheets and the verdict
columns of the verdict sheets are Excel formulas; the file is flagged to
recalculate on open.
"""

import argparse
import datetime as dt
import json
import os
import pathlib
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
SECTIONS = [
    ("26", "19_SADonors", ["SA_donors_19.csv", "Frames_19.csv", "NullDraws_19.csv"]),
    ("27", "19_UnitPanels", ["Stats_19_16A.csv", "Stats_19_16B.csv", "Stats_19_16F.csv", "Readout_19_Units.csv"]),
    ("28", "19_Histology", ["Stats_19_17A.csv", "Stats_19_17B.csv", "Stats_19_17C.csv", "Verdicts_19_17C.csv",
                            "Readout_19_Histology.csv"]),
    ("29", "19_Clinical", ["Stats_19_15A.csv", "Stats_19_17D.csv", "Stats_19_18A.csv", "Verdicts_19_17D.csv",
                           "Verdicts_19_18A.csv", "Readout_19_Clinical.csv"]),
    ("30", "19_Influence", ["Influence_19.csv", "Influence_19_summary.csv"]),
]
FIG_LOG = "Part12_Fig19_log.txt"
B5_CONFIG = "Part12_Fig19_config.json"
ARMS = ["M4", "PVM_3", "PVM_1", "HKG8"]
PANELS = ["15A", "16A", "16B", "16F", "17A", "17B", "17C", "17D", "18A"]
PROV_COLS = ["run_id", "config_version", "labels_version"]
TOL = 1e-9
DASHES = ("\u2013", "\u2014")
CHAR_COLS = ("panel", "arm", "level", "quantity", "measure", "family", "run_id", "labels_version",
             "donor_key", "donor_id", "dataset", "frame", "null_status", "removed", "row_id", "class",
             "arms", "verdict_all", "verdict_wo", "predictor", "outcome", "sa_ranks", "severely_affected",
             "high_adnc", "tested_all", "tested_wo", "null_ok", "fitted")

FONT = Font(name="Arial", size=10)
FONT_B = Font(name="Arial", size=10, bold=True)
FMT_P = "0.00E+00"


# ---------------------------------------------------------------------------
# Check ledger and shared helpers (as build_SuppTable_S16.py)
# ---------------------------------------------------------------------------
class Ledger:
    def __init__(self):
        self.lines, self.n = [], {"PASS": 0, "FAIL": 0, "NOTE": 0}

    def chk(self, cid, ok, what, detail=""):
        tag = "PASS" if ok else "FAIL"
        self.n[tag] += 1
        self._say(f"[{tag}] {cid} {what} | {detail}")

    def note(self, what, detail=""):
        self.n["NOTE"] += 1
        self._say(f"[NOTE] {what} | {detail}")

    def _say(self, s):
        self.lines.append(s)
        print(s)


def find(root, name):
    hits = [p for p in root.rglob(name) if p.is_file()]
    if len(hits) != 1:
        sys.exit(f"ERROR: expected exactly one {name} under {root}, found {len(hits)}")
    return hits[0]


_NUM = re.compile(r"([+-]?)(\d*)(?:\.(\d*))?(?:[eE]([+-]?\d+))?$")


def r_strtod(s):
    """One number as R's read.csv parses it (R_strtod5): the digits are
    accumulated in long double, scaled by a power of ten built by repeated
    squaring in long double, then rounded to double. pandas' own parser and
    Python's float() both differ from read.csv in the last bit for some values,
    which changes the order of values that nearly tie after centring."""
    s = s.strip()
    if s in ("", "NA"):
        return np.nan
    low = s.lower()
    if low in ("inf", "+inf"):
        return np.inf
    if low == "-inf":
        return -np.inf
    if low == "nan":
        return np.nan
    m = _NUM.match(s)
    if not m or not (m.group(2) or m.group(3)):
        raise ValueError(s)
    sign, ip, fp, ex = m.group(1), m.group(2) or "", m.group(3) or "", m.group(4)
    digits = ip + fp
    if len(digits.lstrip("0")) >= 18:
        return float(s)
    expn = -len(fp) + (int(ex) if ex else 0)
    ld = np.longdouble
    ans = ld(0)
    for ch in digits:
        ans = ld(10) * ans + ld(ord(ch) - 48)
    if expn != 0 and ans != 0:
        n, p10, fac = abs(expn), ld(10), ld(1)
        while n:
            if n & 1:
                fac *= p10
            n >>= 1
            p10 *= p10
        ans = ans / fac if expn < 0 else ans * fac
    v = float(ans)
    return -v if sign == "-" else v


def r_mean(v):
    """R's mean() of a double vector: long double sum in order, divided by n,
    plus the long double correction pass, rounded to double."""
    ld = np.longdouble
    n = len(v)
    s = ld(0)
    for a in v:
        s += ld(a)
    s /= n
    if np.isfinite(float(s)):
        t = ld(0)
        for a in v:
            t += ld(a) - s
        s += t / n
    return float(s)


def read_csv_like_r(path, char_cols=()):
    """Read a CSV with R read.csv missing-value rules, so complete-case filters
    match the analysis: "NA" is missing in every column; an empty field is
    missing in a numeric column and an empty string in a character column."""
    d = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[])
    for c in d.columns:
        v = d[c].where(d[c] != "NA")
        if c in char_cols:
            d[c] = v
            continue
        try:
            d[c] = pd.Series([np.nan if x is None or (isinstance(x, float) and np.isnan(x)) else r_strtod(x)
                              for x in v], index=d.index, dtype=float)
        except ValueError:
            d[c] = v
    return d


def as_bool(s):
    return s.map({"TRUE": True, "FALSE": False})


# ---------------------------------------------------------------------------
# Recomputation (section 19 and section 20 rules)
def ranks(x):
    return pd.Series(np.asarray(x, float)).rank().values


def parse_log(path):
    """Reads a Part 12 section log. Logs written by older copies of the R
    scripts add a field after the file size, after the config versions and
    after the Battery 5 readings path; the patterns stop before it, so both
    forms give the same result."""
    txt = path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")
    g = lambda pat: (re.search(pat, txt, re.M).group(1) if re.search(pat, txt, re.M) else None)
    outs = set(re.findall(r"^\s+\S*/([^/\s|]+) \| [\d.]+ MB", txt, re.M))
    inputs = {lab.strip(): p for lab, p in
              re.findall(r"^\s+input (.+?)\s{2,}(\S.*?) \| [\d.]+ MB", txt, re.M)}
    return {
        "path": path,
        "script": g(r"^Script: (\S+)"),
        "run_id": g(r"^Run ID: (\S+)"),
        "start": g(r"^Start: (.+)$"),
        "config_version": g(r"^Config: .*? \| version (\S+)"),
        "labels_version": g(r"^Labels version: (\S+)"),
        "b5_config": g(r"^Battery 5 config: (\S+)"),
        "b5_version": g(r"^Battery 5 config: .*? \| version (\S+)"),
        "b5_readings": g(r"^Battery 5 readings: (\S+)"),
        "checks": g(r"^Checks: (.+?) \|"),
        "n_fail": int(g(r"^Checks: \d+ PASS, (\d+) FAIL") or -1),
        "verdict": g(r"^VERDICT: (\S+)"),
        "outputs": outs,
        "inputs": inputs,
        "rows_line": g(r"^\[PASS\] \S+ stats (?:and plot )?files? written(?: and read back)? \| (.+)$"),
        "text": txt,
    }


def bh(p):
    """Benjamini-Hochberg as R p.adjust(method = "BH"), NaN left as NaN."""
    p = np.asarray(p, dtype=float)
    out = np.full(p.shape, np.nan)
    ok = np.isfinite(p)
    m = ok.sum()
    if m == 0:
        return out
    pv = p[ok]
    order = np.argsort(pv)[::-1]
    rk = np.arange(m, 0, -1)
    q = np.minimum.accumulate(pv[order] * m / rk)
    res = np.empty(m)
    res[order] = np.minimum(q, 1.0)
    out[ok] = res
    return out


def close(a, b, tol=TOL):
    a, b = np.asarray(a, float), np.asarray(b, float)
    both_nan = np.isnan(a) & np.isnan(b)
    both_inf = np.isinf(a) & np.isinf(b) & (np.sign(a) == np.sign(b))
    with np.errstate(invalid="ignore"):
        return bool(np.all(both_nan | both_inf | (np.abs(a - b) <= tol)))


class Worst:
    """Largest difference over many compared values, and the count that differ."""
    def __init__(self):
        self.worst, self.at, self.n, self.bad = 0.0, "", 0, 0

    def add(self, v, ref, where, tol=TOL):
        self.n += 1
        v, ref = float(v), float(ref)
        if np.isnan(v) or np.isnan(ref) or np.isinf(v) or np.isinf(ref):
            if not close([v], [ref]):
                self.bad += 1
                self.at = self.at or where
            return
        if abs(v - ref) > self.worst:
            self.worst, self.at = abs(v - ref), where
        if abs(v - ref) >= tol:
            self.bad += 1

    def count(self, ok, where):
        self.n += 1
        if not ok:
            self.bad += 1
            self.at = self.at or where

    def detail(self):
        return (f"{self.n} values | {self.bad} differ | largest difference {self.worst:.1e}"
                + (f" ({self.at})" if self.at else ""))


def clean(v):
    """Excel-safe cell value: NaN and Inf to empty, numpy scalars to Python."""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, (float, np.floating)) and not np.isfinite(v):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    return v


def write_table(wb, name, header, rows, widths=None, p_cols=(), fmt_cols=None, wrap=False):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append([clean(v) for v in r])
    for c in ws[1]:
        c.font = FONT_B
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            if wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    fmt_cols = dict(fmt_cols or {})
    for col in p_cols:
        fmt_cols[col] = FMT_P
    for col, fmt in fmt_cols.items():
        j = header.index(col) + 1
        for i in range(2, ws.max_row + 1):
            ws.cell(i, j).number_format = fmt
    for j, h in enumerate(header, start=1):
        if widths and h in widths:
            w = widths[h]
        else:
            vals = [len(str(h))] + [len(str(ws.cell(i, j).value)) for i in range(2, min(ws.max_row, 200) + 1)
                                    if ws.cell(i, j).value is not None
                                    and not str(ws.cell(i, j).value).startswith("=")]
            w = min(max(vals) + 2, 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def stats_sheet(wb, name, df):
    """A stats file as a sheet, provenance columns left out."""
    cols = [c for c in df.columns if c not in PROV_COLS]
    write_table(wb, name, cols, df[cols].itertuples(index=False, name=None), p_cols=("p", "fdr"))



# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# Battery 5 rules (as Part12_Fig19_Analysis.r)
# ---------------------------------------------------------------------------
def b5_change(e_removed, e):
    e_removed, e = np.asarray(e_removed, float), np.asarray(e, float)
    s = np.sign(e)
    s = np.where(np.isfinite(s) & (s != 0), s, np.nan)
    return s * (e_removed - e)


def b5_class(R, alpha):
    ta, tw, nok = R["tested_all"].astype(bool), R["tested_wo"].astype(bool), R["null_ok"].astype(bool)
    fa, fw = R["fdr_all"].astype(float), R["fdr_wo"].astype(float)
    sa = ta & np.isfinite(fa) & (fa < alpha)
    sw = tw & np.isfinite(fw) & (fw < alpha)
    same = np.sign(R["est_all"].astype(float)) == np.sign(R["est_wo"].astype(float))
    d, q = R["delta_sa"].astype(float), R["null_q_tail"].astype(float)
    beyond = np.isfinite(d) & np.isfinite(q) & (d < q)
    out = []
    for i in range(len(R)):
        if not ta.iloc[i]:
            k = "Untestable"
        elif sa.iloc[i] and not tw.iloc[i]:
            k = "Untestable without the eleven"
        elif sa.iloc[i] and sw.iloc[i] and same.iloc[i]:
            k = "Holds"
        elif sa.iloc[i] and sw.iloc[i]:
            k = "Reversed"
        elif sa.iloc[i] and not nok.iloc[i]:
            k = "Lost, null untestable"
        elif sa.iloc[i] and beyond.iloc[i]:
            k = "Lost, beyond random removal"
        elif sa.iloc[i]:
            k = "Lost, within random removal"
        elif sw.iloc[i]:
            k = "Gained"
        else:
            k = "Neither"
        out.append(k)
    return out


def verdict2(fdr_a, fdr_b, vif, fitted, alpha, vr, labels):
    """p12_17_verdict / p12_18_verdict: labels = (a only, b only)."""
    sa = np.isfinite(fdr_a) & (fdr_a < alpha)
    sb = np.isfinite(fdr_b) & (fdr_b < alpha)
    out = []
    for i in range(len(fdr_a)):
        if not fitted[i]:
            out.append("Untestable")
        elif not (np.isfinite(vif[i]) and vif[i] < vr):
            out.append("Unreadable")
        elif sa[i] and not sb[i]:
            out.append(labels[0])
        elif sb[i] and not sa[i]:
            out.append(labels[1])
        elif sa[i] and sb[i]:
            out.append("Both")
        else:
            out.append("Neither")
    return out


def class_formula(c, alpha):
    sa = f"AND({c('tested_all')},ISNUMBER({c('fdr_all')}),{c('fdr_all')}<{alpha})"
    sw = f"AND({c('tested_wo')},ISNUMBER({c('fdr_wo')}),{c('fdr_wo')}<{alpha})"
    by = f"AND(ISNUMBER({c('delta_sa')}),ISNUMBER({c('null_q_tail')}),{c('delta_sa')}<{c('null_q_tail')})"
    return (f'=IF(NOT({c("tested_all")}),"Untestable",IF(AND({sa},NOT({c("tested_wo")})),"Untestable without the eleven",'
            f'IF(AND({sa},{sw},SIGN({c("est_all")})=SIGN({c("est_wo")})),"Holds",IF(AND({sa},{sw}),"Reversed",'
            f'IF(AND({sa},NOT({c("null_ok")})),"Lost, null untestable",IF(AND({sa},{by}),"Lost, beyond random removal",'
            f'IF({sa},"Lost, within random removal",IF({sw},"Gained","Neither"))))))))')


def verdict_formula(c, fa, fb, alpha, vr, la, lb):
    sa = f"AND(ISNUMBER({c(fa)}),{c(fa)}<{alpha})"
    sb = f"AND(ISNUMBER({c(fb)}),{c(fb)}<{alpha})"
    return (f'=IF(NOT({c("fitted_wo")}),"Untestable",IF(NOT(AND(ISNUMBER({c("vif_wo")}),{c("vif_wo")}<{vr:g})),"Unreadable",'
            f'IF(AND({sa},NOT({sb})),"{la}",IF(AND({sb},NOT({sa})),"{lb}",IF(AND({sa},{sb}),"Both","Neither")))))')


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    default_root = pathlib.Path(os.environ.get("SEAAD_ROOT") or os.getcwd()) / "Part12_rerun"
    ap.add_argument("--root", default=str(default_root),
                    help="the Part12_rerun folder (default: $SEAAD_ROOT/Part12_rerun, else ./Part12_rerun)")
    ap.add_argument("--out", default=None, help="output xlsx (default: <root>/11_Fig19/SuppTable_Fig19.xlsx)")
    args = ap.parse_args()
    root = pathlib.Path(args.root).resolve()
    if not root.is_dir():
        sys.exit(f"ERROR: root folder not found: {root}")
    out = pathlib.Path(args.out).resolve() if args.out else root / "11_Fig19" / "SuppTable_Fig19.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    L = Ledger()
    L._say("build_SuppTable_S17.py")
    L._say(f"Start: {dt.datetime.now():%Y-%m-%d %H:%M:%S} | root {root}")
    L._say("")

    # ---- configs -------------------------------------------------------------------
    cfg_path = find(root, "Part12_01_config.json")
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    b5_path = find(root, B5_CONFIG)
    B5 = json.loads(b5_path.read_text(encoding="utf-8"))
    alpha = float(cfg["inference"]["alpha"]); vr = float(cfg["floors"]["vif_readable"])
    KLEV = list(B5["classes"])
    DATASETS = list(cfg["datasets"].keys())
    OUTS = list(cfg["clinical"]["outcomes"].keys())
    NPAN = len(cfg["figure18"]["panels"])
    R_NULL = int(B5["null"]["draws"]); TAIL = float(B5["null"]["tail"])
    PERM_R = int(cfg["inference"]["permutation_R"]); B_ = int(cfg["inference"]["bootstrap_B"])

    # ---- logs and their outputs ----------------------------------------------------
    L._say("SECTION A -- logs and outputs")
    logs, files, D = {}, {}, {}
    for step, folder, names in SECTIONS:
        lg = parse_log(find(root, f"Part12_{folder}_log.txt"))
        logs[step] = lg
        L.chk(f"B-01.{step}", lg["verdict"] == "PASS" and lg["n_fail"] == 0,
              f"section {step} log is PASS with 0 FAIL", f"{lg['checks']} | {lg['verdict']}")
        for nm in names:
            p = find(root, nm)
            files[nm] = (step, p)
            L.chk(f"B-02.{nm}", nm in lg["outputs"], f"{nm} is listed in the outputs of the section {step} log",
                  lg["path"].name)
            D[nm] = read_csv_like_r(p, char_cols=CHAR_COLS)
    one = lambda k: sorted({str(lg[k]) for lg in logs.values()})
    named = lambda k: pathlib.PurePosixPath(one(k)[0]).name if len(one(k)) == 1 else None
    rd_name = named("b5_readings")
    rd_hits = [p for p in root.rglob(rd_name) if p.is_file()] if rd_name and rd_name != "None" else []
    L.chk("B-03", all(len(one(k)) == 1 for k in ("script", "config_version", "labels_version", "b5_config",
                                                  "b5_version", "b5_readings"))
          and one("config_version") == [str(cfg.get("config_version"))]
          and named("b5_config") == b5_path.name and one("b5_version") == [str(B5.get("battery5_version"))]
          and len(rd_hits) == 1,
          "one analysis script, config version, labels version, Battery 5 config and readings file across sections "
          "26 to 30; the configs here have those versions and the readings file exists",
          f"{one('script')} | config {one('config_version')} | B5 config {one('b5_config')} version "
          f"{one('b5_version')} | B5 readings {one('b5_readings')} ({len(rd_hits)} found)")
    for nm, (step, _) in files.items():
        d, lg = D[nm], logs[step]
        ok = set(d["run_id"]) == {lg["run_id"]}
        if nm.startswith("Stats_"):
            ok = (ok and set(d["config_version"].astype(int).astype(str)) == {lg["config_version"]}
                  and set(d["labels_version"]) == {lg["labels_version"]})
        L.chk(f"B-04.{nm}", ok, f"{nm} rows carry the log's run ID"
              + (", config and labels version" if nm.startswith("Stats_") else ""), lg["run_id"])

    SA, FR, DR = D["SA_donors_19.csv"], D["Frames_19.csv"], D["NullDraws_19.csv"]
    SA["severely_affected"] = as_bool(SA["severely_affected"]); SA["high_adnc"] = as_bool(SA["high_adnc"])
    RD = {"Units": D["Readout_19_Units.csv"], "Histology": D["Readout_19_Histology.csv"],
          "Clinical": D["Readout_19_Clinical.csv"]}
    for R in RD.values():
        for c in ("tested_all", "tested_wo", "null_ok"):
            R[c] = as_bool(R[c])
    ST = {p: D[f"Stats_19_{p}.csv"] for p in PANELS}
    V17C, V17D, V18A = D["Verdicts_19_17C.csv"], D["Verdicts_19_17D.csv"], D["Verdicts_19_18A.csv"]
    IN, IS = D["Influence_19.csv"], D["Influence_19_summary.csv"]
    IN["severely_affected"] = as_bool(IN["severely_affected"]); IN["high_adnc"] = as_bool(IN["high_adnc"])
    L._say("")

    # ---- row counts ------------------------------------------------------------------
    L._say("SECTION B -- row counts")
    tx = {k: logs[k]["text"] for k in logs}
    m26 = re.search(r"tables written \| (\d+) donors \| (\d+) frames \| (\d+) draws", tx["26"])
    m27 = re.search(r"stats and readout written \| 16A (\d+) \| 16B (\d+) \| 16F (\d+) \| readout (\d+) rows", tx["27"])
    m28 = re.search(r"stats and readout written \| 17A (\d+) \| 17B (\d+) \| 17C (\d+) \| readout (\d+) rows", tx["28"])
    m28v = re.search(r"17C verdict tables are in the same model order \| (\d+) models", tx["28"])
    m29 = re.search(r"stats and readout written \| 15A (\d+) \| 17D (\d+) \| 18A (\d+) \| readout (\d+) rows", tx["29"])
    m30 = re.search(r"tables written \| influence (\d+) rows \| summary (\d+) rows", tx["30"])
    got = [(m26, [len(SA), len(FR), len(DR)]),
           (m27, [len(ST["16A"]), len(ST["16B"]), len(ST["16F"]), len(RD["Units"])]),
           (m28, [len(ST["17A"]), len(ST["17B"]), len(ST["17C"]), len(RD["Histology"])]),
           (m28v, [len(V17C)]),
           (m29, [len(ST["15A"]), len(ST["17D"]), len(ST["18A"]), len(RD["Clinical"])]),
           (m30, [len(IN), len(IS)])]
    ok = all(m is not None and [int(x) for x in m.groups()] == v for m, v in got)
    na, nd = len(ARMS), len(DATASETS)
    ok_design = len(V17D) == na * nd * len(OUTS) and len(V18A) == na * nd * NPAN * 2 * len(OUTS)
    L.chk("B-05", ok and ok_design, "row counts equal the logs, and the 17D and 18A verdict tables have the design's rows",
          " | ".join(f"{[int(x) for x in m.groups()] if m else None} vs {v}" for m, v in got)
          + f" | 17D {len(V17D)} | 18A {len(V18A)}")
    L._say("")

    # ---- BH ------------------------------------------------------------------------------
    L._say("SECTION C -- FDR and the readouts")
    for p, S in ST.items():
        w = Worst()
        for f, g in S.groupby("family", sort=False):
            q = bh(g["p"].values)
            for v, r, i in zip(q, g["fdr"].values, g.index):
                w.add(v, r, f"{f} row {i}")
        L.chk(f"B-06.{p}", w.bad == 0, f"Stats_19_{p} fdr equals BH within family", w.detail())

    # readouts against stats, change and class
    for nm, R in RD.items():
        w = Worst(); wc = Worst()
        for i, r in R.iterrows():
            S = ST[r["panel"]]
            j = S.index[(S["arm"] == r["arm"]) & (S["level"] == r["level"]) & (S["quantity"] == r["quantity"])
                        & (S["measure"] == r["measure"])]
            if len(j) != 1:
                w.count(False, f"{r['row_id']} not in Stats_19_{r['panel']}")
                continue
            s = S.loc[j[0]]
            w.add(r["est_wo"], s["estimate"], r["row_id"]); w.add(r["fdr_wo"], s["fdr"], r["row_id"])
            w.add(r["n_donors_wo"], s["n_donors"], r["row_id"])
            w.count(bool(r["tested_wo"]) == bool(np.isfinite(s["p"])), f"{r['row_id']} tested_wo")
        dl = b5_change(R["est_wo"], R["est_all"])
        for v, ref, rid in zip(dl, R["delta_sa"], R["row_id"]):
            wc.add(v, ref, rid)
        L.chk(f"B-07.{nm}.stats", w.bad == 0, f"{nm} readout: without-eleven estimate, FDR and donors equal the stats files",
              w.detail())
        L.chk(f"B-07.{nm}.change", wc.bad == 0, f"{nm} readout: SA change equals s x (e_without - e_all)", wc.detail())
        k = b5_class(R, alpha)
        nbad = sum(a != b for a, b in zip(k, R["class"]))
        L.chk(f"B-07.{nm}.class", nbad == 0 and set(R["class"]) <= set(KLEV),
              f"{nm} readout: every class follows from its columns by the pre-set rule",
              f"{len(R)} rows | {nbad} differ | classes {sorted(set(R['class']))}")

    # class counts against the logs
    logc = {}
    for step in ("27", "28", "29"):
        txt = tx[step]
        i0 = txt.find("-- the reading")
        for line in txt[i0:].split("\n"):
            m = re.match(r"^    (\S+) (\S+)\s{2,}(.+)$", line)
            if m and m.group(1) in PANELS:
                for item in m.group(3).split(" | "):
                    cl, n = item.rsplit(" ", 1)
                    logc[(m.group(1), m.group(2), cl)] = int(n)
    AL = pd.concat(RD.values(), ignore_index=True)
    mine = AL.groupby(["panel", "arm", "class"]).size().to_dict()
    L.chk("B-08", mine == logc, "class counts per panel and arm equal the section logs",
          f"{len(mine)} panel, arm and class cells | log {len(logc)}")
    L._say("")

    # ---- verdicts ------------------------------------------------------------------------
    L._say("SECTION D -- verdicts without the eleven")
    # 17C
    S = ST["17C"]
    key = {(a, lv, ms, q): i for i, (a, lv, ms, q) in enumerate(zip(S["arm"], S["level"], S["measure"], S["quantity"]))}
    g = lambda a, lv, ms, q, col: S[col].iloc[key[(a, lv, ms, q)]] if (a, lv, ms, q) in key else np.nan
    V17C["fdr_module_wo"] = [g(a, lv, ms, "beta_module", "fdr") for a, lv, ms in zip(V17C["arm"], V17C["level"], V17C["measure"])]
    V17C["fdr_mito_wo"] = [g(a, lv, ms, "beta_mito", "fdr") for a, lv, ms in zip(V17C["arm"], V17C["level"], V17C["measure"])]
    V17C["vif_st"] = [g(a, lv, ms, "vif", "estimate") for a, lv, ms in zip(V17C["arm"], V17C["level"], V17C["measure"])]
    V17C["fitted_wo"] = [bool(np.isfinite(g(a, lv, ms, "beta_module", "estimate")) and np.isfinite(g(a, lv, ms, "beta_mito", "estimate")))
                         for a, lv, ms in zip(V17C["arm"], V17C["level"], V17C["measure"])]
    vv = verdict2(V17C["fdr_module_wo"].values.astype(float), V17C["fdr_mito_wo"].values.astype(float),
                  V17C["vif_wo"].values.astype(float), V17C["fitted_wo"].values, alpha, vr, ("Module only", "%MT only"))
    L.chk("B-09.17C", list(vv) == list(V17C["verdict_wo"]) and close(V17C["vif_st"], V17C["vif_wo"]),
          "17C verdicts without the eleven follow from Stats_19_17C", f"{len(V17C)} models | {sum(a != b for a, b in zip(vv, V17C['verdict_wo']))} differ")
    # 17D
    S = ST["17D"]
    key = {(a, lv, ms, q): i for i, (a, lv, ms, q) in enumerate(zip(S["arm"], S["level"], S["measure"], S["quantity"]))}
    g = lambda a, lv, ms, q, col: S[col].iloc[key[(a, lv, ms, q)]] if (a, lv, ms, q) in key else np.nan
    z = list(zip(V17D["arm"], V17D["dataset"], V17D["outcome"]))
    V17D["fdr_module_wo"] = [g(a, ds, o, "beta_module", "fdr") for a, ds, o in z]
    V17D["fdr_mito_wo"] = [g(a, ds, o, "beta_mito", "fdr") for a, ds, o in z]
    V17D["fitted_wo"] = [bool(np.isfinite(g(a, ds, o, "beta_module", "estimate")) and np.isfinite(g(a, ds, o, "beta_mito", "estimate")))
                         for a, ds, o in z]
    vst = [g(a, ds, o, "vif", "estimate") for a, ds, o in z]
    vv = verdict2(V17D["fdr_module_wo"].values.astype(float), V17D["fdr_mito_wo"].values.astype(float),
                  V17D["vif_wo"].values.astype(float), V17D["fitted_wo"].values, alpha, vr, ("Module only", "%MT only"))
    L.chk("B-09.17D", list(vv) == list(V17D["verdict_wo"]) and close(vst, V17D["vif_wo"]),
          "17D verdicts without the eleven follow from Stats_19_17D", f"{len(V17D)} models | {sum(a != b for a, b in zip(vv, V17D['verdict_wo']))} differ")
    # 18A
    S = ST["18A"]
    key = {(a, lv, ms, q): i for i, (a, lv, ms, q) in enumerate(zip(S["arm"], S["level"], S["measure"], S["quantity"]))}
    g = lambda a, lv, ms, q, col: S[col].iloc[key[(a, lv, ms, q)]] if (a, lv, ms, q) in key else np.nan
    z = [(a, ds, f"{p}|{pk}|{o}") for a, ds, p, pk, o in zip(V18A["arm"], V18A["dataset"], V18A["predictor"],
                                                             V18A["panel"], V18A["outcome"])]
    V18A["fdr_second_wo"] = [g(a, ds, ms, "beta_second", "fdr") for a, ds, ms in z]
    V18A["fdr_panel_wo"] = [g(a, ds, ms, "beta_panel", "fdr") for a, ds, ms in z]
    V18A["fitted_wo"] = [bool(np.isfinite(g(a, ds, ms, "beta_second", "estimate")) and np.isfinite(g(a, ds, ms, "beta_panel", "estimate")))
                         for a, ds, ms in z]
    vst = [g(a, ds, ms, "vif", "estimate") for a, ds, ms in z]
    vv = verdict2(V18A["fdr_second_wo"].values.astype(float), V18A["fdr_panel_wo"].values.astype(float),
                  V18A["vif_wo"].values.astype(float), V18A["fitted_wo"].values, alpha, vr, ("second only", "NE panel only"))
    L.chk("B-09.18A", list(vv) == list(V18A["verdict_wo"]) and close(vst, V18A["vif_wo"]),
          "18A verdicts without the eleven follow from Stats_19_18A", f"{len(V18A)} models | {sum(a != b for a, b in zip(vv, V18A['verdict_wo']))} differ")
    L._say("")

    # ---- influence -----------------------------------------------------------------------
    L._say("SECTION E -- influence")
    w = Worst()
    for v, ref, rid in zip(b5_change(IN["est_minus_donor"], IN["est_all"]), IN["change"], IN["row_id"]):
        w.add(v, ref, rid)
    ## Ranks: R ranked the full-precision changes; the CSV keeps 15 significant
    ## digits, so two changes that differ below that read back as a tie. A
    ## stored rank is accepted when it lies within the rank range of the values
    ## within 1e-12 of its own change.
    n_tie = 0
    for rid, gI in IN.groupby("row_id"):
        ch, st = gI["change"].values.astype(float), gI["rank_by_change"].values.astype(float)
        fin = np.isfinite(ch)
        for i in range(len(ch)):
            if not fin[i]:
                w.count(not np.isfinite(st[i]), f"{rid} rank of a missing change")
                continue
            lo = 1 + int(np.sum(ch[fin] < ch[i] - 1e-12))
            hi = int(np.sum(ch[fin] <= ch[i] + 1e-12))
            ok = lo <= st[i] <= hi
            if ok and hi > lo:
                n_tie += 1
            w.count(ok, f"{rid} rank")
    L.chk("B-10.rows", w.bad == 0, "every influence change and rank follows from its row",
          w.detail() + f" | {n_tie} ranks inside a near-tie (changes within 1e-12)")
    w = Worst()
    UR = RD["Units"].set_index("row_id")
    for rid, gI in IN.groupby("row_id"):
        s = IS[IS["row_id"] == rid]
        if len(s) != 1:
            w.count(False, rid); continue
        s = s.iloc[0]
        sa = gI["severely_affected"].values; hi = gI["high_adnc"].values & ~sa
        k = int(sa.sum())
        rks = sorted(int(x) for x in gI["rank_by_change"][sa] if np.isfinite(x))
        w.count(int(s["n_donors"]) == len(gI) and int(s["n_sa"]) == k, f"{rid} counts")
        w.count((s["sa_ranks"] if isinstance(s["sa_ranks"], str) else "") == ";".join(str(x) for x in rks), f"{rid} ranks")
        med = lambda v: float(np.nanmedian(v)) if len(v) and np.isfinite(v).any() else np.nan
        if k:
            w.add(s["sa_in_top_k"], sum(x <= k for x in rks), rid)
            w.add(s["median_change_sa"], med(gI["change"].values[sa]), rid)
            w.add(s["sum_change_sa"], float(np.nansum(gI["change"].values[sa])), rid, tol=1e-12)
        w.add(s["median_change_high_non_sa"], med(gI["change"].values[hi]) if hi.any() else np.nan, rid)
        rest = ~sa & ~hi
        w.add(s["median_change_other"], med(gI["change"].values[rest]) if rest.any() else np.nan, rid)
        w.add(s["change_sa_together"], UR.loc[rid, "delta_sa"] if rid in UR.index else np.nan, rid)
    L.chk("B-10.summary", w.bad == 0 and len(IS) == IN["row_id"].nunique(),
          "every influence summary row follows from the influence rows and the unit readout", w.detail())
    L._say("")

    # ---- SA donors, frames and draws -----------------------------------------------------
    L._say("SECTION F -- SA donors, frames and draws")
    sa_keys = set(SA["donor_key"][SA["severely_affected"]])
    pool = set(SA["donor_key"][SA["high_adnc"] & ~SA["severely_affected"]])
    L.chk("B-11.sa", len(sa_keys) == int(B5["sa"]["expected_n"]), "the SA donor table has the configured number of SA donors",
          f"{len(sa_keys)} of {B5['sa']['expected_n']}")
    bad = 0
    for fr in FR.itertuples():
        d = DR[DR["frame"] == fr.frame]
        want = R_NULL if fr.null_status == "testable" else 0
        if len(d) != want:
            bad += 1
            continue
        for rm in d["removed"]:
            ks = rm.split(";")
            if len(ks) != int(fr.k) or set(ks) & sa_keys or not set(ks) <= pool:
                bad += 1
    L.chk("B-11.draws", bad == 0, "every draw removes k donors of the random-removal pool and no SA donor; draw counts follow the null status",
          f"{len(FR)} frames | {len(DR)} draws | {bad} problems")
    L._say("")

    # ---- figure log ----------------------------------------------------------------------
    fl = parse_log(find(root, FIG_LOG))
    last = max(dt.datetime.strptime(lg["start"], "%Y-%m-%d %H:%M:%S") for lg in logs.values())
    L.chk("B-12", fl["verdict"] == "PASS" and fl["n_fail"] == 0 and
          dt.datetime.strptime(fl["start"], "%Y-%m-%d %H:%M:%S") >= last,
          "figure log is PASS and started after the last analysis section", f"{fl['checks']} | {fl['start']}")
    if L.n["FAIL"]:
        L._say("")
        L._say(f"Checks: {L.n['PASS']} PASS, {L.n['FAIL']} FAIL, {L.n['NOTE']} NOTE")
        L._say("VERDICT: FAIL | workbook not written")
        out.with_name(out.stem + "_build_log.txt").write_text("\n".join(L.lines) + "\n", encoding="utf-8")
        sys.exit(1)

    # ---- workbook ------------------------------------------------------------------------
    wb = Workbook(); wb.remove(wb.active)
    ov = [
        ["Fig 19A", "Is the donor structure carried by the eleven severely affected donors?",
         "16A donor share of variance, 16B split-half rs, 16F between-donor rs with the intronic rate, per arm, with all donors and without the eleven",
         "27", "Readout_19_Units.csv; Stats_19_16A/16B/16F.csv", "Fig19A_readout; Fig19A_16A_stats; Fig19A_16B_stats; Fig19A_16F_stats", "Notes 1 to 6"],
        ["Fig 19B", "Which results hold without the eleven?",
         "Class of every retested row, counted per panel and arm (all nine panels)",
         "27 to 29", "Readout_19_Units/Histology/Clinical.csv; Stats_19_*.csv",
         "Fig19B_counts; Fig19A_readout; Fig19B_readout_histology; Fig19B_readout_clinical; the stats sheets", "Notes 1 to 7"],
        ["Fig 19C", "Is the clinical association carried by the eleven?",
         "FDR of the 15A unadjusted Module post-IP% coefficient (cov4, the Figure 12A model) per arm, region set and outcome, with all donors and without the eleven",
         "29", "Readout_19_Clinical.csv; Stats_19_15A.csv", "Fig19C_clinical; Fig19B_readout_clinical; Fig19B_15A_stats", "Notes 1 to 5 and 8"],
        ["Verdicts", "Do the Battery 3 and 4 verdicts change without the eleven?",
         "17C within donor, 17D and 18A verdicts with all donors and without the eleven (the pre-set 19C reading; no figure in the manuscript)",
         "28, 29", "Verdicts_19_17C/17D/18A.csv", "Verdicts_17C; Verdicts_17D; Verdicts_18A", "Note 9"],
        ["Influence", "Which donors move the unit-level estimates most?",
         "Leave-one-donor-out change of every 16A, 16B and 16F between-donor estimate; no test",
         "30", "Influence_19.csv; Influence_19_summary.csv", "Influence; Influence_summary", "Note 10"],
        ["Donors", "Which donors are removed?",
         "The eleven SA donors, the frames, and the random-removal draws",
         "26", "SA_donors_19.csv; Frames_19.csv; NullDraws_19.csv", "SA_donors; Frames; Null_draws", "Notes 2 and 3"],
    ]
    write_table(wb, "Overview", ["Panel", "Question", "Quantity", "Section", "Source file", "Sheets", "See"], ov,
                wrap=True, widths={"Panel": 11, "Question": 40, "Quantity": 60, "Section": 9, "Source file": 34,
                                   "Sheets": 40, "See": 14})
    kpool = ", ".join(f"{r.frame} {int(r.k)}" for r in FR.itertuples() if r.frame.startswith("U|"))
    notes = [
        ["1", "Arms", "M4 = Mic13_PVM3, the masked primary arm. PVM_3 = PVM3_all, the unmasked reference. PVM_1 = PVM1_all, the second supertype. HKG8 = the housekeeping control construct."],
        ["2", "The eleven", f"The SEA-AD column 'Severely Affected Donor' of obs_dataset.csv, every arm and region set: exactly {B5['sa']['expected_n']} donors flagged, constant within donor, every one with ADNC_rank {B5['high_adnc']['value']} (High). SA_donors lists every donor of the donor table with its SA flag, whether it is in the random-removal pool (High ADNC, not SA) and the arms it appears in. Other donor metadata are not carried."],
        ["3", "Frames and random removal", f"A frame is the donor set a statistic is drawn from: U|arm (unit table; 16A, 16B, 16F), J19|arm (joined to histology or CPS; 17A, 17B), JHU|arm (joined to histology; 17C), DN|arm|region set (donor table; 15A, 17D, 18A). k is the number of SA donors in the frame (unit frames: {kpool}). Random removal: k donors drawn without replacement from the frame's pool, {R_NULL:,} times, made once in section 26 (Null_draws, donors named by donor_id). A contrast between two arms uses draw r of each arm's frame."],
        ["4", "Without the eleven", f"Each statistic is recomputed on the frame less its SA donors with the method, floors and multiple-testing families of the all-donor run and a new seed. Cell labels, the inflection point, units, Module post-IP%, %MT, the 18A panel genes and donor panel scores are held at the all-donor values. Inference: 16A {PERM_R:,} within-region shuffles; 16B, 16F, 17A, 17B, 17C donor bootstrap, B = {B_:,}; 15A, 17D, 18A OLS p. fdr is Benjamini-Hochberg within the family column (families prefixed '19|')."],
        ["5", "Row classes", f"Significant: FDR < {alpha:g}; the all-donor FDR is the Battery 1 to 4 value. Change = s x (e_without - e_all), s = sign(e_all); negative is attenuation. Classes: Holds (significant with and without, same sign); Reversed; Lost, beyond random removal (significant with, not without, change below the {TAIL:.0%} quantile of the random-removal changes); Lost, within random removal; Lost, null untestable; Untestable without the eleven; Gained; Neither; Untestable. In the readout sheets the class column is an Excel formula over the row (check B-07 in the build log compares it with R). p_null = (1 + random changes at or below the SA change) / (1 + finite random changes), descriptive."],
        ["6", "Random-removal columns", f"null_n: finite random-removal estimates of the row (fewer than {R_NULL:,} where a draw left a model below its donor floor, as on the 18A HKG8 NonNeoC cognition rows). null_q_tail, null_q50, null_q95: quantiles of the random-removal changes; null_est_lo, null_est_hi: the {TAIL:.0%} and 95% quantiles on the estimate scale. The random-removal estimates themselves (Null_19_*.csv) are in the deposit; they are too large for this workbook."],
        ["7", "Fig 19B", "Counts of each class per retested panel and arm (Fig19B_counts). 15A: the unadjusted and contamination-adjusted Module coefficients. 16A, 16B: every quantity. 16F: between donors, intronic field, four rs and three contrasts. 17A: within donor. 17B: between donors. 17C: within donor, both coefficients. 17D and 18A: both coefficients."],
        ["8", "Fig 19C", "The 15A unadjusted rows of Fig19B_readout_clinical: outcome ~ cov4 + donor Module post-IP%, standardised, per arm, region set and outcome; BH within arm and region set over the nine outcomes, as in 15A."],
        ["9", "Verdicts", f"Per model, the verdict with all donors (Battery 3 and 4 FDRs) and without the eleven, by the 17C, 17D and 18A rules: Untestable, Unreadable (VIF >= {vr:g}), one predictor only, Both or Neither. The without-eleven verdict is an Excel formula over the FDR and VIF columns of the same row; the all-donor FDRs are in SuppTable_Fig17 and SuppTable_Fig18. The pre-set 19C reading is answered by these sheets; the verdict figure is not in the manuscript."],
        ["10", "Influence", "Every donor of the unit table removed in turn, from every arm, and the 16A, 16B and 16F between-donor estimates recomputed. change = s x (e_minus_donor - e_all); rank_by_change 1 = most attenuating. Influence_summary: per row, the SA donors in the frame, their ranks, sa_in_top_k (SA donors among the k most attenuating), the medians of the change for SA, other High-ADNC and remaining donors, the sum of the SA donors' single changes and the change on removing all of them together (19A)."],
        ["11", "Pre-specification", "Part12_Fig19_config.json and Part12_Preset_Readings_Battery5.md (dated 2026-09-25, written before any statistic was computed without the eleven). Section 26 checked that every Battery 1 to 4 log it reads is PASS and that its analysis file is in Part12_rerun."],
    ]
    write_table(wb, "Notes", ["Note", "Subject", "Detail"], notes, wrap=True, widths={"Note": 7, "Subject": 26, "Detail": 120})

    gl = [
        ["column", "panel", "Source panel of the Battery 1 to 4 figures the row retests; in the stats sheets prefixed '19_'."],
        ["column", "arm, level, quantity, measure", "As in the Battery 1 to 4 stats files: level is the region set (15A, 17D, 18A), units or donors (16A, 16B), within donor or between donors (16F, 17A to 17C)."],
        ["column", "frame, k_removed", "The donor frame of the row and its number of SA donors (Note 3)."],
        ["column", "est_all, fdr_all, n_donors_all", "All-donor estimate, FDR and donors (Battery 1 to 4 stats files)."],
        ["column", "est_wo, lo_wo, hi_wo, fdr_wo, n_donors_wo", "Without the eleven: estimate, 95% interval where one exists, FDR, donors."],
        ["column", "tested_all, tested_wo, null_ok", "Whether the row has a test with all donors and without the eleven; whether its frame's random-removal pool holds at least k donors."],
        ["column", "delta_sa, null_n, null_q_tail, null_q50, null_q95, null_est_lo, null_est_hi, p_null", "SA change and the random-removal comparison (Notes 5 and 6)."],
        ["column", "class", "Row class by the pre-set rule (Note 5); Excel formula in the readout sheets."],
        ["column", "estimate, lo, hi, p, fdr, family", "In the stats sheets: as in the Battery 1 to 4 stats files."],
        ["verdicts", "verdict_all, verdict_wo, vif_all, vif_wo, n_donors_all, n_donors_wo", "Verdict, VIF and donors with all donors and without the eleven (Note 9)."],
        ["influence", "est_minus_donor, change, rank_by_change", "Estimate without one donor, its change and rank (Note 10)."],
        ["donors", "severely_affected, random_removal_pool, arms", "SA flag; High ADNC and not SA (the pool); arms the donor appears in."],
        ["frames", "n_donors, k, n_pool, null_status", "Donors in the frame, SA donors among them, pool size, whether random removal is testable."],
    ]
    write_table(wb, "Glossary", ["Scope", "Term", "Meaning"], gl, wrap=True, widths={"Scope": 10, "Term": 50, "Meaning": 100})

    key2id = dict(zip(SA["donor_key"], SA["donor_id"]))
    write_table(wb, "SA_donors", ["donor_id", "severely_affected", "random_removal_pool", "arms"],
                ([r.donor_id, bool(r.severely_affected), bool(r.high_adnc and not r.severely_affected), r.arms]
                 for r in SA.itertuples()))
    fc = ["frame", "n_donors", "k", "n_pool", "null_status"]
    write_table(wb, "Frames", fc, FR[fc].itertuples(index=False, name=None))
    write_table(wb, "Null_draws", ["frame", "draw", "removed_donor_ids"],
                ([r.frame, int(r.draw), ";".join(key2id.get(k, k) for k in r.removed.split(";"))] for r in DR.itertuples()))

    RC = ["panel", "arm", "level", "quantity", "measure", "frame", "k_removed", "est_all", "fdr_all", "n_donors_all",
          "est_wo", "lo_wo", "hi_wo", "fdr_wo", "n_donors_wo", "tested_all", "tested_wo", "null_ok", "delta_sa",
          "null_n", "null_q_tail", "null_q50", "null_q95", "null_est_lo", "null_est_hi", "p_null", "class"]
    cR = {h: get_column_letter(j + 1) for j, h in enumerate(RC)}

    def readout_sheet(name, R):
        rows = []
        for i, d in enumerate(R.to_dict("records"), start=2):
            c = lambda h: f"{cR[h]}{i}"
            rows.append([bool(d[h]) if h in ("tested_all", "tested_wo", "null_ok") else d[h] for h in RC[:-1]]
                        + [class_formula(c, alpha)])
        write_table(wb, name, RC, rows, p_cols=("fdr_all", "fdr_wo", "p_null"),
                    fmt_cols={k: "0.0000" for k in ("est_all", "est_wo", "lo_wo", "hi_wo", "delta_sa", "null_q_tail",
                                                    "null_q50", "null_q95", "null_est_lo", "null_est_hi")},
                    widths={"class": 28})

    readout_sheet("Fig19A_readout", RD["Units"])
    for p in ("16A", "16B", "16F"):
        stats_sheet(wb, f"Fig19A_{p}_stats", ST[p])
    CT = AL.groupby(["panel", "arm", "class"]).size().unstack(fill_value=0)
    CT = CT.reindex(columns=[k for k in KLEV if k in CT.columns])
    CT = CT.reindex(pd.MultiIndex.from_product([PANELS, ARMS]), fill_value=0)
    write_table(wb, "Fig19B_counts", ["panel", "arm"] + list(CT.columns) + ["total"],
                ([p, a] + [int(x) for x in CT.loc[(p, a)].values] + [int(CT.loc[(p, a)].sum())] for p, a in CT.index))
    readout_sheet("Fig19B_readout_histology", RD["Histology"])
    readout_sheet("Fig19B_readout_clinical", RD["Clinical"])
    for p in ("15A", "17A", "17B", "17C", "17D", "18A"):
        stats_sheet(wb, f"Fig19B_{p}_stats", ST[p])
    C = RD["Clinical"]
    C = C[(C["panel"] == "15A") & (C["quantity"] == "module_beta_unadjusted")]
    hc = ["arm", "region_set", "outcome", "n_donors_all", "est_all", "fdr_all", "n_donors_wo", "est_wo", "fdr_wo", "p_null", "class"]
    write_table(wb, "Fig19C_clinical", hc,
                ([r[h] for h in ("arm", "level", "measure", "n_donors_all", "est_all", "fdr_all", "n_donors_wo",
                                 "est_wo", "fdr_wo", "p_null", "class")] for r in C.to_dict("records")),
                p_cols=("fdr_all", "fdr_wo", "p_null"), fmt_cols={"est_all": "0.0000", "est_wo": "0.0000"}, widths={"class": 28})

    def verdict_sheet(name, V, idc, fa, fb, la, lb):
        hv = idc + ["n_donors_all", "vif_all", "verdict_all", "n_donors_wo", "fitted_wo", "vif_wo", fa, fb, "verdict_wo"]
        cv = {h: get_column_letter(j + 1) for j, h in enumerate(hv)}
        rows = []
        for i, d in enumerate(V.to_dict("records"), start=2):
            c = lambda h: f"{cv[h]}{i}"
            rows.append([d[h] for h in idc] + [d["n_donors_all"], d["vif_all"], d["verdict_all"], d["n_donors_wo"],
                                               bool(d["fitted_wo"]), d["vif_wo"], d[fa], d[fb],
                                               verdict_formula(c, fa, fb, alpha, vr, la, lb)])
        write_table(wb, name, hv, rows, p_cols=(fa, fb), fmt_cols={"vif_all": "0.000", "vif_wo": "0.000"},
                    widths={"verdict_all": 16, "verdict_wo": 16})

    verdict_sheet("Verdicts_17C", V17C, ["arm", "level", "measure"], "fdr_module_wo", "fdr_mito_wo", "Module only", "%MT only")
    verdict_sheet("Verdicts_17D", V17D, ["arm", "dataset", "outcome"], "fdr_module_wo", "fdr_mito_wo", "Module only", "%MT only")
    verdict_sheet("Verdicts_18A", V18A, ["arm", "dataset", "predictor", "panel", "outcome"], "fdr_second_wo", "fdr_panel_wo",
                  "second only", "NE panel only")

    ic = ["panel", "arm", "level", "quantity", "measure", "donor_id", "severely_affected", "high_adnc", "est_all",
          "est_minus_donor", "change", "rank_by_change"]
    IN["donor_id"] = IN["donor_key"].map(key2id)
    write_table(wb, "Influence", ic, ([bool(v) if h in ("severely_affected", "high_adnc") else v
                                       for h, v in zip(ic, row)] for row in IN[ic].itertuples(index=False, name=None)),
                fmt_cols={k: "0.00000" for k in ("est_all", "est_minus_donor", "change")})
    sc = [c for c in IS.columns if c not in ("run_id",)]
    write_table(wb, "Influence_summary", sc, IS[sc].itertuples(index=False, name=None))

    order = ["Overview", "Notes", "Glossary", "SA_donors", "Frames", "Null_draws",
             "Fig19A_readout", "Fig19A_16A_stats", "Fig19A_16B_stats", "Fig19A_16F_stats",
             "Fig19B_counts", "Fig19B_readout_histology", "Fig19B_readout_clinical",
             "Fig19B_15A_stats", "Fig19B_17A_stats", "Fig19B_17B_stats", "Fig19B_17C_stats", "Fig19B_17D_stats",
             "Fig19B_18A_stats", "Fig19C_clinical", "Verdicts_17C", "Verdicts_17D", "Verdicts_18A",
             "Influence", "Influence_summary"]
    if sorted(order) != sorted(wb.sheetnames):
        sys.exit(f"ERROR: sheet order list does not match the sheets written: {wb.sheetnames}")
    wb._sheets = [wb[n] for n in order]
    wb.active = 0

    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts += [v for v in row if isinstance(v, str)]
    bad = [t for t in texts if any(d_ in t for d_ in DASHES)]
    L.chk("B-13", not bad, "no en dash or em dash in the workbook text", f"{len(bad)} found")

    L._say("")
    L._say(f"Checks: {L.n['PASS']} PASS, {L.n['FAIL']} FAIL, {L.n['NOTE']} NOTE")
    log_path = out.with_name(out.stem + "_build_log.txt")
    if L.n["FAIL"]:
        L._say("VERDICT: FAIL | workbook not written")
        log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")
        sys.exit(1)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    L._say(f"Output: {out}")
    L._say("VERDICT: PASS")
    log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
