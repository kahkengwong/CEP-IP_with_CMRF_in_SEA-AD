#!/usr/bin/env python3
"""
build_SuppTable_S15.py

Builds SuppTable_Fig17.xlsx (Figure 17 and Supplementary Figure 13 of the
Part 12 rerun) from the stats and plot-data files written by
Part12_Fig17_Analysis.r (sections 19 and 20), the figure log written by
Part12_Fig17_Plots.r (section 21), the config (Part12_01_config.json) and the
unit and donor tables (05_units.csv, 05_donors.csv) with the step 05 log.
finalize_supp_tables.py turns this workbook into Supplementary Table S15.

Usage:
    python build_SuppTable_S15.py
    python build_SuppTable_S15.py --root /data/Part12_rerun

--root defaults to $SEAAD_ROOT/Part12_rerun when SEAAD_ROOT is set, otherwise
to Part12_rerun in the current folder.

--root is searched recursively; every input file must occur exactly once under
it, so the script works on the Part12_rerun folder tree or on a flat folder
holding copies of the files. Without --out the workbook goes to
<root>/09_Fig17/SuppTable_Fig17.xlsx. A build log is written beside the workbook.

Numbers are parsed and averaged exactly as R's read.csv and mean() do, so the
values recomputed here match the analysis where values nearly tie.

The script computes no new statistic. Before it writes anything it checks:
  B-01  both section logs end with VERDICT: PASS and 0 FAIL
  B-02  every input file is listed in the outputs of the log of the section that wrote it
  B-03  the two analysis logs name one script, one config version, one labels version
        and the same neuropathology and CPS workbook paths
  B-04  every stats, plot-data and joined-unit row carries the run ID of its log,
        and every stats row the config and labels version of its log
  B-05  stats row counts equal the counts the logs report; plot-data and joined-unit
        row counts equal the design (arms x levels x quantities x families, and so on)
  B-06  every fdr equals BH recomputed within its family
  B-07  Plot_17C_verdicts agrees with Stats_17C (coefficients, VIF, marginal rs, FDRs),
        and every verdict and sign flip follows from them
  B-08  Plot_17D_verdicts agrees with Stats_17D, every verdict and sign flip follows
        from them, and every 17D interval is centred on its estimate
  B-09  every S12 family median, family difference and tested count follows from
        the per-measure rs
  B-10  every 17A row equals the S12 per-measure row of the same arm, level,
        quantity and measure
  B-11  the 17C verdict counts, the 17D verdicts and the testable task counts equal
        the section logs
  B-12  the figure log is PASS and was started after the last analysis section
  B-13  no text written to the workbook contains an en dash or an em dash
  B-14  the config exists and has the config version of the section logs, the unit
        table and the donor table are listed in the outputs of their log, which is
        PASS, and both tables carry the labels version of the section logs
  B-15  Joined_17AB_units.csv holds every unit of the unit table once, with the
        unit table's values of the five quantities
  B-16  every observed 17A, 17B and S12 per-measure rs, every 17C coefficient, VIF
        and marginal rs, and every 17E histology-unit |r|, recomputed from the
        Fig17_units values, equals its stats row, with the same n_units and
        n_donors (difference below 1e-9)
  B-17  Fig17_units holds every unit of the unit table
  B-18  every 17D coefficient, single-predictor coefficient and VIF, and every 17E
        clinical-grain |r|, recomputed from the Fig17D_donors values, equals its
        stats or plot-data row, with the same donor count
Any FAIL stops the build and no workbook is written.

Formatting: Arial throughout, bold header row, frozen top row. No fills, no
borders, no colours. The direction and significance columns of Fig17A_readout,
the verdict columns of Fig17C_verdicts and Fig17D_verdicts and the position
column of Fig17E_readout are Excel formulas; the file is flagged to recalculate
on open.
"""

import argparse
import datetime as dt
import json
import os
import pathlib
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
SECTIONS = [
    # (step, folder, files written by that section)
    ("19", "17AB_Histology", ["Stats_17A.csv", "Stats_17B.csv", "Stats_S12.csv", "Stats_S12_measures.csv",
                              "Plot_S12_counts.csv", "Joined_17AB_units.csv"]),
    ("20", "17CE_ModuleVsMT", ["Stats_17C.csv", "Stats_17D.csv", "Stats_17E.csv", "Plot_17C_verdicts.csv",
                               "Plot_17D_verdicts.csv"]),
]
FIG_LOG = "Part12_Fig17_log.txt"
CONFIG = "Part12_01_config.json"
UNITS = "05_units.csv"
DONORS = "05_donors.csv"
UNITS_LOG = "*05_Units*_log.txt"
ARMS = ["M4", "PVM_3", "PVM_1", "HKG8"]
PROV_COLS = ["run_id", "config_version", "labels_version"]
TOL = 1e-9
DASHES = ("\u2013", "\u2014")
QS = ["module_post_pct", "mito_log1p", "cyto_share", "contamination", "log_n_cells"]
LEVELS = ["within donor", "between donors"]
VK = ["Module only", "%MT only", "Both", "Neither", "Unreadable", "Untestable"]
CHAR_COLS = ("panel", "arm", "level", "quantity", "measure", "family", "run_id", "labels_version",
             "donor_key", "donor_id", "dataset", "region", "stain_family", "block", "verdict", "outcome",
             "class", "fitted", "module_sign_flip", "mito_sign_flip", "joined_histology", "joined_cps")
## Fig17_units: unit-table name -> sheet name for the unit columns.
UNIT_RENAME = {"dataset": "region_set", "module_post_pct": "post_ip_pct"}

FONT = Font(name="Arial", size=10)
FONT_B = Font(name="Arial", size=10, bold=True)
FMT_P = "0.00E+00"


# ---------------------------------------------------------------------------
# Check ledger
# ---------------------------------------------------------------------------
class Ledger:
    def __init__(self):
        self.lines, self.n = [], {"PASS": 0, "FAIL": 0, "NOTE": 0}

    def chk(self, cid, ok, what, detail=""):
        tag = "PASS" if ok else "FAIL"
        self.n[tag] += 1
        self._say(f"[{tag}] {cid} {what} | {detail}")

    def note(self, what, detail=""):
        self.n["NOTE"] += 1
        self._say(f"[NOTE] {what} | {detail}")

    def _say(self, s):
        self.lines.append(s)
        print(s)


def find(root, name):
    hits = [p for p in root.rglob(name) if p.is_file()]
    if len(hits) != 1:
        sys.exit(f"ERROR: expected exactly one {name} under {root}, found {len(hits)}")
    return hits[0]


_NUM = re.compile(r"([+-]?)(\d*)(?:\.(\d*))?(?:[eE]([+-]?\d+))?$")


def r_strtod(s):
    """One number as R's read.csv parses it (R_strtod5): the digits are
    accumulated in long double, scaled by a power of ten built by repeated
    squaring in long double, then rounded to double. pandas' own parser and
    Python's float() both differ from read.csv in the last bit for some values,
    which changes the order of values that nearly tie after centring."""
    s = s.strip()
    if s in ("", "NA"):
        return np.nan
    low = s.lower()
    if low in ("inf", "+inf"):
        return np.inf
    if low == "-inf":
        return -np.inf
    if low == "nan":
        return np.nan
    m = _NUM.match(s)
    if not m or not (m.group(2) or m.group(3)):
        raise ValueError(s)
    sign, ip, fp, ex = m.group(1), m.group(2) or "", m.group(3) or "", m.group(4)
    digits = ip + fp
    if len(digits.lstrip("0")) >= 18:
        return float(s)
    expn = -len(fp) + (int(ex) if ex else 0)
    ld = np.longdouble
    ans = ld(0)
    for ch in digits:
        ans = ld(10) * ans + ld(ord(ch) - 48)
    if expn != 0 and ans != 0:
        n, p10, fac = abs(expn), ld(10), ld(1)
        while n:
            if n & 1:
                fac *= p10
            n >>= 1
            p10 *= p10
        ans = ans / fac if expn < 0 else ans * fac
    v = float(ans)
    return -v if sign == "-" else v


def r_mean(v):
    """R's mean() of a double vector: long double sum in order, divided by n,
    plus the long double correction pass, rounded to double."""
    ld = np.longdouble
    n = len(v)
    s = ld(0)
    for a in v:
        s += ld(a)
    s /= n
    if np.isfinite(float(s)):
        t = ld(0)
        for a in v:
            t += ld(a) - s
        s += t / n
    return float(s)


def read_csv_like_r(path, char_cols=()):
    """Read a CSV with R read.csv missing-value rules, so complete-case filters
    match the analysis: "NA" is missing in every column; an empty field is
    missing in a numeric column and an empty string in a character column."""
    d = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[])
    for c in d.columns:
        v = d[c].where(d[c] != "NA")
        if c in char_cols:
            d[c] = v
            continue
        try:
            d[c] = pd.Series([np.nan if x is None or (isinstance(x, float) and np.isnan(x)) else r_strtod(x)
                              for x in v], index=d.index, dtype=float)
        except ValueError:
            d[c] = v
    return d


def as_bool(s):
    return s.map({"TRUE": True, "FALSE": False})


# ---------------------------------------------------------------------------
# Recomputation (section 19 and section 20 rules)
# ---------------------------------------------------------------------------
def ranks(x):
    return pd.Series(np.asarray(x, float)).rank().values


def rs(x, y):
    """Spearman rs as p12_rs(): Pearson correlation of average ranks; NaN when
    fewer than 3 values or when either side has no spread."""
    if len(x) < 3:
        return np.nan
    rx, ry = ranks(x), ranks(y)
    if not (np.std(rx, ddof=1) > 0 and np.std(ry, ddof=1) > 0):
        return np.nan
    return float(np.corrcoef(rx, ry)[0, 1])


def duel(y, x, z):
    """The 17C coefficients as p12_17_stat(kind = "duel"): standardised
    coefficients of rank(y) on rank(x) and rank(z) from the three rank
    correlations."""
    if len(y) < 3:
        return (np.nan, np.nan)
    rx, ry, rz = ranks(x), ranks(y), ranks(z)
    if not (np.std(rx, ddof=1) > 0 and np.std(ry, ddof=1) > 0 and np.std(rz, ddof=1) > 0):
        return (np.nan, np.nan)
    ryx = np.corrcoef(ry, rx)[0, 1]
    ryz = np.corrcoef(ry, rz)[0, 1]
    rxz = np.corrcoef(rx, rz)[0, 1]
    den = 1 - rxz ** 2
    if not np.isfinite(den) or den < 1e-12:
        return (np.nan, np.nan)
    return ((ryx - ryz * rxz) / den, (ryz - ryx * rxz) / den)


def centre_within(v, donor, min_n):
    """p12_centre_within(): each donor's mean removed, for donors with at least
    min_n finite values; NaN otherwise."""
    v = np.asarray(v, float)
    out = np.full(len(v), np.nan)
    donor = np.asarray(donor)
    for u in pd.unique(donor):
        i = np.where((donor == u) & np.isfinite(v))[0]
        if len(i) >= min_n:
            out[i] = v[i] - r_mean(v[i])
    return out


def task(u, cols, level, min_u):
    """p12_17_task(): the values entering one statistic. Units finite in every
    column; between donors, donor means over those units; within donor, each
    column centred on the donor, donors with at least min_u units. Returns the
    columns, the donor of each value, and n_units as the analysis counts it."""
    X = [u[c].astype(float).values for c in cols]
    ok = np.all([np.isfinite(x) for x in X], axis=0)
    d_all = u["donor_key"].values[ok]
    X = [x[ok] for x in X]
    n_units = int(ok.sum())
    if level == "between donors":
        dk = np.array(sorted(pd.unique(d_all)))
        X = [np.array([r_mean(x[d_all == k]) for k in dk]) for x in X]
        d = dk
    else:
        X = [centre_within(x, d_all, min_u) for x in X]
        k = np.all([np.isfinite(x) for x in X], axis=0) if len(d_all) else np.zeros(0, bool)
        X = [x[k] for x in X]
        d = d_all[k]
        n_units = int(k.sum())
    return X, d, n_units


def recompute_units(J, meas, cps, glob, min_d, min_u):
    """Every unit-level statistic of sections 19 and 20 from the joined units.
    Keys: ("rs", arm, level, quantity, measure), ("duel", arm, level, measure),
    ("e_hu", arm, level)."""
    out = {}
    for a in ARMS:
        ua = J[J["arm"] == a]
        for lv in LEVELS:
            for q in QS:
                ms = list(meas) + (list(cps) if lv == "between donors" else [c for c in cps if c not in glob])
                for m in ms:
                    X, d, nu = task(ua, [q, m], lv, min_u)
                    nd = len(pd.unique(d))
                    out[("rs", a, lv, q, m)] = {"n_units": nu, "n_donors": nd,
                                                "estimate": rs(X[0], X[1]) if nd >= min_d else np.nan}
            for m in meas:
                X, d, nu = task(ua, [m, "module_post_pct", "mito_log1p"], lv, min_u)
                nd = len(pd.unique(d))
                y, x, z = X
                testable = nd >= min_d
                bm, bt = duel(y, x, z) if testable else (np.nan, np.nan)
                r12 = rs(x, z) if testable else np.nan
                vif = (1 / (1 - r12 ** 2) if abs(r12) < 1 else np.inf) if np.isfinite(r12) else np.nan
                out[("duel", a, lv, m)] = {"n_units": nu, "n_donors": nd, "beta_module": bm, "beta_mito": bt,
                                           "vif": vif, "rs_module_marginal": rs(x, y) if testable else np.nan,
                                           "rs_mito_marginal": rs(z, y) if testable else np.nan}
        uh = ua[ua["joined_histology"]]
        for lv, key in (("within donor", "within donor HU"), ("between donors", "between donors HU")):
            X, d, nu = task(uh, ["module_post_pct", "mito_log1p"], lv, min_u)
            nd = len(pd.unique(d))
            out[("e_hu", a, key)] = {"n_units": nu, "n_donors": nd,
                                     "estimate": abs(rs(X[0], X[1])) if nd >= min_d else np.nan}
    return out


def scale(v):
    return (v - r_mean(v)) / np.std(v, ddof=1)


def fit_duel(d, y, covs, min_n):
    """fit_duel() of section 20 without the p values: the two standardised
    coefficients of outcome ~ cov4 + Module post-IP% + %MT, each predictor's
    coefficient when fitted alone with cov4, and the VIF of the two
    predictors. Complete cases as R's complete.cases()."""
    cols = list(dict.fromkeys([y, "module_post_pct", "mito_log1p"] + list(covs)))
    M = d[cols].astype(float)
    d = d[~M.isna().any(axis=1).values]
    n = len(d)
    out = {"n": n, "fitted": False}
    if n < min_n:
        return out
    yv, mv, tv = (d[c].astype(float).values for c in (y, "module_post_pct", "mito_log1p"))
    if not (np.std(yv, ddof=1) > 0 and np.std(mv, ddof=1) > 0 and np.std(tv, ddof=1) > 0):
        return out
    Y, Mm, T = scale(yv), scale(mv), scale(tv)
    base = np.column_stack([np.ones(n)] + [d[c].astype(float).values for c in covs])

    def ols(X):
        if np.linalg.matrix_rank(X) < X.shape[1]:
            return None
        return np.linalg.lstsq(X, Y, rcond=None)[0]

    b2 = ols(np.column_stack([base, Mm, T]))
    if b2 is None:
        return out
    bm1 = ols(np.column_stack([base, Mm]))
    bt1 = ols(np.column_stack([base, T]))
    r = np.corrcoef(Mm, T)[0, 1]
    vif = 1 / (1 - r ** 2) if abs(r) < 1 else np.inf
    return {"n": n, "fitted": True, "b_m": b2[-2], "b_t": b2[-1],
            "single_m": bm1[-1] if bm1 is not None else np.nan,
            "single_t": bt1[-1] if bt1 is not None else np.nan, "vif": max(vif, 1.0)}


def verdict(fdr_m, fdr_t, vif, fitted, alpha, vr):
    """p12_17_verdict() for one row."""
    if not fitted:
        return "Untestable"
    if not (np.isfinite(vif) and vif < vr):
        return "Unreadable"
    sm = bool(np.isfinite(fdr_m) and fdr_m < alpha)
    st = bool(np.isfinite(fdr_t) and fdr_t < alpha)
    return "Module only" if sm and not st else "%MT only" if st and not sm else "Both" if sm and st else "Neither"


def parse_log(path):
    txt = path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")
    g = lambda pat: (re.search(pat, txt, re.M).group(1) if re.search(pat, txt, re.M) else None)
    # Output lines: "  <folder>/<file> | <size> MB"; input lines: "  input <label>  <path> | <size> MB".
    # Anything after the size is ignored.
    outs = set(re.findall(r"^\s+\S*/([^/\s|]+) \| [\d.]+ MB", txt, re.M))
    inputs = {lab.strip(): p for lab, p in
              re.findall(r"^\s+input (.+?)\s{2,}(\S.*?) \| [\d.]+ MB", txt, re.M)}
    return {
        "path": path,
        "script": g(r"^Script: (\S+)"),
        "run_id": g(r"^Run ID: (\S+)"),
        "start": g(r"^Start: (.+)$"),
        "config_version": g(r"^Config: .* \| version (\S+?)(?: \||\s*$)"),
        "labels_version": g(r"^Labels version: (\S+)"),
        "checks": g(r"^Checks: (.+?) \|"),
        "n_fail": int(g(r"^Checks: \d+ PASS, (\d+) FAIL") or -1),
        "verdict": g(r"^VERDICT: (\S+)"),
        "outputs": outs,
        "inputs": inputs,
        "rows_line": g(r"^\[PASS\] \S+ stats (?:and plot )?files? written(?: and read back)? \| (.+)$"),
        "text": txt,
    }


def bh(p):
    """Benjamini-Hochberg as R p.adjust(method = "BH"), NaN left as NaN."""
    p = np.asarray(p, dtype=float)
    out = np.full(p.shape, np.nan)
    ok = np.isfinite(p)
    m = ok.sum()
    if m == 0:
        return out
    pv = p[ok]
    order = np.argsort(pv)[::-1]
    rk = np.arange(m, 0, -1)
    q = np.minimum.accumulate(pv[order] * m / rk)
    res = np.empty(m)
    res[order] = np.minimum(q, 1.0)
    out[ok] = res
    return out


def close(a, b, tol=TOL):
    a, b = np.asarray(a, float), np.asarray(b, float)
    both_nan = np.isnan(a) & np.isnan(b)
    both_inf = np.isinf(a) & np.isinf(b) & (np.sign(a) == np.sign(b))
    with np.errstate(invalid="ignore"):
        return bool(np.all(both_nan | both_inf | (np.abs(a - b) <= tol)))


class Worst:
    """Largest difference over many compared values, and the count that differ."""
    def __init__(self):
        self.worst, self.at, self.n, self.bad = 0.0, "", 0, 0

    def add(self, v, ref, where, tol=TOL):
        self.n += 1
        v, ref = float(v), float(ref)
        if np.isnan(v) or np.isnan(ref) or np.isinf(v) or np.isinf(ref):
            if not close([v], [ref]):
                self.bad += 1
                self.at = self.at or where
            return
        if abs(v - ref) > self.worst:
            self.worst, self.at = abs(v - ref), where
        if abs(v - ref) >= tol:
            self.bad += 1

    def count(self, ok, where):
        self.n += 1
        if not ok:
            self.bad += 1
            self.at = self.at or where

    def detail(self):
        return (f"{self.n} values | {self.bad} differ | largest difference {self.worst:.1e}"
                + (f" ({self.at})" if self.at else ""))


def clean(v):
    """Excel-safe cell value: NaN and Inf to empty, numpy scalars to Python."""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, (float, np.floating)) and not np.isfinite(v):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    return v


def write_table(wb, name, header, rows, widths=None, p_cols=(), fmt_cols=None, wrap=False):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append([clean(v) for v in r])
    for c in ws[1]:
        c.font = FONT_B
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            if wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    fmt_cols = dict(fmt_cols or {})
    for col in p_cols:
        fmt_cols[col] = FMT_P
    for col, fmt in fmt_cols.items():
        j = header.index(col) + 1
        for i in range(2, ws.max_row + 1):
            ws.cell(i, j).number_format = fmt
    for j, h in enumerate(header, start=1):
        if widths and h in widths:
            w = widths[h]
        else:
            vals = [len(str(h))] + [len(str(ws.cell(i, j).value)) for i in range(2, min(ws.max_row, 200) + 1)
                                    if ws.cell(i, j).value is not None
                                    and not str(ws.cell(i, j).value).startswith("=")]
            w = min(max(vals) + 2, 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def stats_sheet(wb, name, df):
    """A stats file as a sheet, provenance columns left out."""
    cols = [c for c in df.columns if c not in PROV_COLS]
    write_table(wb, name, cols, df[cols].itertuples(index=False, name=None), p_cols=("p", "fdr"))


def verdict_formula(c, i, alpha, vr):
    """Excel formula of p12_17_verdict() over the columns fitted, vif, fdr_module
    and fdr_mito of row i."""
    fm, ft, v, f = (f"{c[k]}{i}" for k in ("fdr_module", "fdr_mito", "vif", "fitted"))
    sm = f"AND(ISNUMBER({fm}),{fm}<{alpha})"
    st = f"AND(ISNUMBER({ft}),{ft}<{alpha})"
    return (f'=IF(NOT({f}),"Untestable",IF(NOT(AND(ISNUMBER({v}),{v}<{vr})),"Unreadable",'
            f'IF(AND({sm},NOT({st})),"Module only",IF(AND({st},NOT({sm})),"%MT only",'
            f'IF(AND({sm},{st}),"Both","Neither")))))')


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    default_root = pathlib.Path(os.environ.get("SEAAD_ROOT") or ".") / "Part12_rerun"
    ap.add_argument("--root", default=str(default_root),
                    help="folder holding the Figure 17 outputs and logs (default: $SEAAD_ROOT/Part12_rerun, "
                         "or ./Part12_rerun)")
    ap.add_argument("--out", default=None, help="output xlsx (default: <root>/09_Fig17/SuppTable_Fig17.xlsx)")
    args = ap.parse_args()
    root = pathlib.Path(args.root).resolve()
    out = pathlib.Path(args.out).resolve() if args.out else root / "09_Fig17" / "SuppTable_Fig17.xlsx"
    L = Ledger()
    L._say("build_SuppTable_S15.py")
    L._say(f"Start: {dt.datetime.now():%Y-%m-%d %H:%M:%S} | root {root}")
    L._say("")

    # ---- config ----------------------------------------------------------------------
    cfg_path = find(root, CONFIG)
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    HI = cfg["histology"]
    MEAS = [m for f in HI["families"] for m in f["measures"]]
    FAM_OF = {m: f["key"] for f in HI["families"] for m in f["measures"]}
    FAM_LAB = {f["key"]: f["label"] for f in HI["families"]}
    FAM_N = {f["key"]: len(f["measures"]) for f in HI["families"]}
    CPS = [c["column"] for c in HI["cps"]]
    CPS_LAB = {c["column"]: c["label"] for c in HI["cps"]}
    GLOB = [c["column"] for c in HI["cps"] if c["scope"] == "global"]
    LOCAL = [c for c in CPS if c not in GLOB]
    IBA = [z["measure"] for z in HI["iba1_17A"]]
    IBA_LAB = {z["measure"]: z["label"] for z in HI["iba1_17A"]}
    SIGN = {z["measure"]: int(z["predicted_sign"]) for z in HI["iba1_17A"]}
    QLAB = HI["quantities"]
    COVS = list(cfg["clinical"]["cov4"].values())
    OUT = {k: o["column"] for k, o in cfg["clinical"]["outcomes"].items()}
    DATASETS = list(cfg["datasets"].keys())
    min_d = int(cfg["floors"]["donors_per_test"])
    min_u = int(cfg["floors"]["finite_units_within_donor"])
    alpha = float(cfg["inference"]["alpha"])
    vr = float(cfg["floors"]["vif_readable"])
    B_ = int(cfg["inference"]["bootstrap_B"])

    # ---- logs and input files ---------------------------------------------------------
    L._say("SECTION A -- logs and inputs")
    logs, files, D = {}, {}, {}
    for step, folder, names in SECTIONS:
        lg = parse_log(find(root, f"Part12_{folder}_log.txt"))
        logs[step] = lg
        L.chk(f"B-01.{step}", lg["verdict"] == "PASS" and lg["n_fail"] == 0,
              f"section {step} log is PASS with 0 FAIL", f"{lg['checks']} | {lg['verdict']}")
        for nm in names:
            p = find(root, nm)
            files[nm] = (step, p)
            L.chk(f"B-02.{nm}", nm in lg["outputs"],
                  f"{nm} is listed in the outputs of the section {step} log", lg["path"].name)
            D[nm] = read_csv_like_r(p, char_cols=CHAR_COLS)
    one = lambda k: sorted({str(lg[k]) for lg in logs.values()})
    wbk = lambda lg: lg["inputs"]
    same_wb = (wbk(logs["19"]) == wbk(logs["20"]) and len(wbk(logs["19"])) == 2)
    L.chk("B-03", len(one("script")) == 1 and len(one("config_version")) == 1
          and len(one("labels_version")) == 1 and same_wb,
          "one analysis script, config, labels version and pair of workbooks across sections 19 and 20",
          f"{one('script')} | config {one('config_version')} | labels {one('labels_version')}"
          f" | workbooks {', '.join(f'{k} {pathlib.PurePath(v).name}' for k, v in wbk(logs['19']).items())}")

    # ---- row provenance -------------------------------------------------------------
    for nm, (step, _) in files.items():
        d, lg = D[nm], logs[step]
        ok = set(d["run_id"]) == {lg["run_id"]}
        if nm.startswith("Stats_"):
            ok = (ok and set(d["config_version"].astype(int).astype(str)) == {lg["config_version"]}
                  and set(d["labels_version"]) == {lg["labels_version"]})
        L.chk(f"B-04.{nm}", ok, f"{nm} rows carry the log's run ID"
              + (", config and labels version" if nm.startswith("Stats_") else ""),
              f"{lg['run_id']} | {lg['config_version']} | {lg['labels_version']}")

    # ---- row counts -------------------------------------------------------------------
    key = {"17A": "Stats_17A.csv", "17B": "Stats_17B.csv", "S12": "Stats_S12.csv",
           "S12 measures": "Stats_S12_measures.csv", "17C": "Stats_17C.csv", "17D": "Stats_17D.csv",
           "17E": "Stats_17E.csv"}
    for step, folder, names in SECTIONS:
        rl = logs[step]["rows_line"] or ""
        got = dict(re.findall(r"(17A|17B|S12 measures|S12|17C|17D|17E) (\d+)", rl))
        want = [n for n in names if n.startswith("Stats_")]
        ok = sorted(key[k] for k in got) == sorted(want) and all(int(v) == len(D[key[k]]) for k, v in got.items())
        L.chk(f"B-05.{step}", ok, f"section {step} stats row counts equal the log", rl)
    nf = len(HI["families"])
    design = {"Plot_S12_counts.csv": len(ARMS) * len(LEVELS) * len(QS) * nf,
              "Plot_17C_verdicts.csv": len(ARMS) * len(LEVELS) * len(MEAS),
              "Plot_17D_verdicts.csv": len(ARMS) * len(DATASETS) * len(OUT)}
    L.chk("B-05.plots", all(len(D[k]) == v for k, v in design.items()),
          "plot-data row counts equal the design", " | ".join(f"{k} {len(D[k])} of {v}" for k, v in design.items()))

    # ---- BH within family -----------------------------------------------------------
    for nm in [n for n in D if n.startswith("Stats_")]:
        d = D[nm]
        ok = all(close(bh(g["p"].values), g["fdr"].values, 1e-12) for _, g in d.groupby("family", sort=False))
        L.chk(f"B-06.{nm}", ok, f"{nm} fdr equals BH within family", f"{d['family'].nunique()} families")
    L._say("")

    # ---- panel consistency -------------------------------------------------------------
    L._say("SECTION B -- panel consistency")
    A, Bt, S, SM, SC = (D[n] for n in ("Stats_17A.csv", "Stats_17B.csv", "Stats_S12.csv",
                                       "Stats_S12_measures.csv", "Plot_S12_counts.csv"))
    C, Dd, E, CV, DV = (D[n] for n in ("Stats_17C.csv", "Stats_17D.csv", "Stats_17E.csv",
                                       "Plot_17C_verdicts.csv", "Plot_17D_verdicts.csv"))
    DV = DV.rename(columns={"class": "oclass"})
    for X in (CV, DV):
        for c in ("fitted", "module_sign_flip", "mito_sign_flip"):
            X[c] = as_bool(X[c])

    # 17C
    Ci = C.set_index(["arm", "level", "measure", "quantity"])
    W = Worst()
    for r in CV.itertuples():
        g = lambda q, col="estimate": Ci.loc[(r.arm, r.level, r.measure, q), col]
        W.add(r.beta_module, g("beta_module"), f"17C {r.arm} {r.level} {r.measure} beta_module")
        W.add(r.beta_mito, g("beta_mito"), f"17C {r.arm} {r.level} {r.measure} beta_mito")
        W.add(r.vif, g("vif"), f"17C {r.arm} {r.level} {r.measure} vif")
        W.add(r.rs_module, g("rs_module_marginal"), f"17C {r.arm} {r.level} {r.measure} rs_module")
        W.add(r.rs_mito, g("rs_mito_marginal"), f"17C {r.arm} {r.level} {r.measure} rs_mito")
        W.add(r.fdr_module, g("beta_module", "fdr"), f"17C {r.arm} {r.level} {r.measure} fdr_module")
        W.add(r.fdr_mito, g("beta_mito", "fdr"), f"17C {r.arm} {r.level} {r.measure} fdr_mito")
        W.count(r.fitted == bool(np.isfinite(r.beta_module) and np.isfinite(r.beta_mito)),
                f"17C {r.arm} {r.level} {r.measure} fitted")
        W.count(r.verdict == verdict(r.fdr_module, r.fdr_mito, r.vif, r.fitted, alpha, vr),
                f"17C {r.arm} {r.level} {r.measure} verdict")
        W.count(r.module_sign_flip == bool(np.isfinite(r.beta_module) and np.isfinite(r.rs_module)
                                           and r.beta_module * r.rs_module < 0), f"17C {r.measure} flip")
        W.count(r.mito_sign_flip == bool(np.isfinite(r.beta_mito) and np.isfinite(r.rs_mito)
                                         and r.beta_mito * r.rs_mito < 0), f"17C {r.measure} flip")
    L.chk("B-07", W.bad == 0 and W.worst < TOL,
          "Plot_17C_verdicts agrees with Stats_17C and every verdict and sign flip follows from it", W.detail())

    # 17D
    Di = Dd.set_index(["arm", "level", "measure", "quantity"])
    W = Worst()
    for r in DV.itertuples():
        g = lambda q, col="estimate": Di.loc[(r.arm, r.dataset, r.outcome, q), col]
        W.add(r.beta_module, g("beta_module"), f"17D {r.arm} {r.dataset} {r.outcome} beta_module")
        W.add(r.beta_mito, g("beta_mito"), f"17D {r.arm} {r.dataset} {r.outcome} beta_mito")
        W.add(r.vif, g("vif"), f"17D {r.arm} {r.dataset} {r.outcome} vif")
        W.add(r.fdr_module, g("beta_module", "fdr"), f"17D {r.arm} {r.dataset} {r.outcome} fdr_module")
        W.add(r.fdr_mito, g("beta_mito", "fdr"), f"17D {r.arm} {r.dataset} {r.outcome} fdr_mito")
        W.add(r.n_donors, g("beta_module", "n_donors"), f"17D {r.arm} {r.dataset} {r.outcome} n")
        W.count(r.verdict == verdict(r.fdr_module, r.fdr_mito, r.vif, r.fitted, alpha, vr),
                f"17D {r.arm} {r.dataset} {r.outcome} verdict")
        W.count(r.module_sign_flip == bool(np.isfinite(r.beta_module) and np.isfinite(r.single_module)
                                           and r.beta_module * r.single_module < 0), f"17D {r.outcome} flip")
        W.count(r.mito_sign_flip == bool(np.isfinite(r.beta_mito) and np.isfinite(r.single_mito)
                                         and r.beta_mito * r.single_mito < 0), f"17D {r.outcome} flip")
    beta = Dd[Dd.quantity.isin(["beta_module", "beta_mito"]) & np.isfinite(Dd.estimate)]
    for r in beta.itertuples():
        W.add((r.lo + r.hi) / 2, r.estimate, f"17D {r.arm} {r.level} {r.measure} {r.quantity} interval centre")
    L.chk("B-08", W.bad == 0 and W.worst < TOL,
          "Plot_17D_verdicts agrees with Stats_17D, every verdict and sign flip follows from it, "
          "and every 17D interval is centred on its estimate", W.detail())

    # S12
    SMi = SM.set_index(["arm", "level", "quantity", "measure"])
    W = Worst()
    for r in S.itertuples():
        ms = [m for m in MEAS if FAM_OF[m] == r.measure]
        if r.quantity == "abs_module_minus_abs_mito":
            vm = np.array([SMi.loc[(r.arm, r.level, "module_post_pct", m), "estimate"] for m in ms])
            vt = np.array([SMi.loc[(r.arm, r.level, "mito_log1p", m), "estimate"] for m in ms])
            k = np.isfinite(vm) & np.isfinite(vt)
            W.add(np.median(np.abs(vm[k]) - np.abs(vt[k])) if k.any() else np.nan, r.estimate,
                  f"S12B {r.arm} {r.level} {r.measure}")
        else:
            v = np.array([SMi.loc[(r.arm, r.level, r.quantity, m), "estimate"] for m in ms])
            v = v[np.isfinite(v)]
            W.add(np.median(v) if len(v) else np.nan, r.estimate, f"S12A {r.arm} {r.level} {r.quantity} {r.measure}")
    for r in SC.itertuples():
        ms = [m for m in MEAS if FAM_OF[m] == r.family]
        v = np.array([SMi.loc[(r.arm, r.level, r.quantity, m), "estimate"] for m in ms])
        W.count(int(np.isfinite(v).sum()) == r.n_tested and len(ms) == r.n_total,
                f"S12 counts {r.arm} {r.level} {r.quantity} {r.family}")
    L.chk("B-09", W.bad == 0 and W.worst < 1e-12,
          "every S12 family median, family difference and tested count follows from the per-measure rs", W.detail())

    # 17A against S12 per-measure
    W = Worst()
    for r in A.itertuples():
        s = SMi.loc[(r.arm, r.level, r.quantity, r.measure)]
        for col in ("estimate", "lo", "hi", "n_units", "n_donors"):
            W.add(getattr(r, col), s[col], f"17A {r.arm} {r.level} {r.quantity} {r.measure} {col}", 1e-15)
    L.chk("B-10", W.bad == 0, "every 17A row equals the S12 per-measure row of the same arm, level, quantity and measure",
          W.detail())

    # against the logs
    t20 = logs["20"]["text"]
    pat = (r"^\s+(\S+)\s+(within donor|between donors)\s+Module only (\d+), %MT only (\d+), Both (\d+), "
           r"Neither (\d+), Unreadable (\d+), Untestable (\d+) \| %MT only with a Module sign flip (\d+) \| donors (\d+)")
    logc = {(m[0], m[1]): tuple(int(x) for x in m[2:]) for m in re.findall(pat, t20, re.M)}
    mine = {}
    for (a, lv), g in CV.groupby(["arm", "level"], sort=False):
        cnt = tuple(int((g.verdict == k).sum()) for k in VK)
        mine[(a, lv)] = cnt + (int(((g.verdict == "%MT only") & g.module_sign_flip).sum()), int(g.n_donors.max()))
    sec_e = t20.split("SECTION E", 1)[-1].split("SECTION F", 1)[0]
    logd = {}
    for a, ds, rest in re.findall(r"^\s+(\S+)\s+(" + "|".join(DATASETS) + r")\s+(.+)$", sec_e, re.M):
        for item in rest.split(" | "):
            o, v = item.rsplit(": ", 1)
            logd[(a, ds, o)] = v.strip()
    mined = {(r.arm, r.dataset, r.outcome): r.verdict for r in DV.itertuples()}
    tk19 = re.search(r"\[NOTE\] tasks \| (\d+) \((\d+) histology, (\d+) CPS\) \| (\d+) testable", logs["19"]["text"])
    tk20 = re.search(r"\[NOTE\] tasks \| (\d+) 17C models, (\d+) 17E rows \| testable (\d+)", t20)
    n19 = int(np.isfinite(SM.estimate).sum() + np.isfinite(Bt.estimate).sum())
    n20 = int(CV.fitted.sum() + np.isfinite(E[E.quantity == "abs_r"].estimate).sum())
    ok = (logc == mine and len(logc) == len(ARMS) * len(LEVELS) and logd == mined and len(logd) == len(DV)
          and tk19 is not None and int(tk19.group(4)) == n19 and int(tk19.group(2)) == len(SM)
          and int(tk19.group(3)) == len(Bt) and tk20 is not None and int(tk20.group(3)) == n20
          and int(tk20.group(1)) == len(CV) and int(tk20.group(2)) == int((E.quantity == "abs_r").sum()))
    L.chk("B-11", ok, "the 17C verdict counts, the 17D verdicts and the testable task counts equal the section logs",
          f"17C {len(logc)} arm-level lines | 17D {len(logd)} verdicts | testable {n19} (section 19), {n20} (section 20)")

    # ---- figure log ----------------------------------------------------------------
    fl = parse_log(find(root, FIG_LOG))
    last = max(dt.datetime.strptime(lg["start"], "%Y-%m-%d %H:%M:%S") for lg in logs.values())
    fstart = dt.datetime.strptime(fl["start"], "%Y-%m-%d %H:%M:%S")
    L.chk("B-12", fl["verdict"] == "PASS" and fl["n_fail"] == 0 and fstart >= last,
          "figure log is PASS and started after the last analysis section",
          f"{fl['script']} | {fl['start']}")
    L._say("")

    # ---- config, unit and donor tables -------------------------------------------------
    L._say("SECTION C -- unit and donor tables")
    L.chk("B-14.config", cfg_path.is_file() and str(cfg.get("config_version")) == one("config_version")[0],
          f"{CONFIG} exists and has the config version of the section logs",
          f"version {cfg.get('config_version')}")
    up_log = parse_log(find(root, UNITS_LOG)) if len([p for p in root.rglob(UNITS_LOG) if p.is_file()]) == 1 \
        else sys.exit(f"ERROR: expected exactly one file matching {UNITS_LOG} under {root}")
    tabs = {}
    for nm in (UNITS, DONORS):
        p = find(root, nm)
        L.chk(f"B-14.{nm}", nm in up_log["outputs"] and up_log["verdict"] == "PASS",
              f"{nm} is listed in the outputs of its log, which is PASS", up_log["path"].name)
        tabs[nm] = (p, read_csv_like_r(p, char_cols=CHAR_COLS + ("libs",)))
    U, DN = tabs[UNITS][1], tabs[DONORS][1]
    L.chk("B-14.labels", set(U["labels_version"]) == set(one("labels_version"))
          and set(DN["labels_version"]) == set(one("labels_version")),
          "unit and donor tables carry the labels version of the section logs",
          ", ".join(sorted(set(U["labels_version"]) | set(DN["labels_version"]))))

    # Joined units against the unit table
    J = D["Joined_17AB_units.csv"].copy()
    for c in ("joined_histology", "joined_cps"):
        J[c] = as_bool(J[c])
    kcols = ["arm", "dataset", "donor_key", "region"]
    jk = J[kcols].astype(str).agg("|".join, axis=1)
    uk = U[kcols].astype(str).agg("|".join, axis=1)
    ok = (not jk.duplicated().any() and set(jk) == set(uk) and len(J) == len(U)
          and list(jk) == list(uk))
    if ok:
        for q in QS:
            ok &= close(J[q].values, U[q].values, 0.0)
    L.chk("B-15", ok, "Joined_17AB_units.csv holds every unit of the unit table once, in its order, "
          "with its values of the five quantities", f"{len(J)} joined | {len(U)} in the unit table")

    # ---- Fig17_units: the values the recomputation reads -------------------------------
    ucols = ["arm", "dataset", "donor_key", "region"] + QS + ["joined_histology", "joined_cps"] + MEAS + CPS
    J["_a"] = J["arm"].map({a: i for i, a in enumerate(ARMS)})
    ## Stable sort by arm only: within an arm the rows keep the unit-table order,
    ## which is the order R averaged them in.
    SHEET = J[J["_a"].notna()].sort_values("_a", kind="stable").reset_index(drop=True)[ucols].copy()
    RU = recompute_units(SHEET, MEAS, CPS, GLOB, min_d, min_u)
    W = Worst()
    for S_ in (A, Bt, SM):
        for r in S_.itertuples():
            x = RU[("rs", r.arm, r.level, r.quantity, r.measure)]
            where = f"{r.panel} {r.arm} {r.level} {r.quantity} {r.measure}"
            W.add(x["estimate"], r.estimate, where)
            W.count(x["n_units"] == r.n_units and x["n_donors"] == r.n_donors, where + " counts")
    for r in C.itertuples():
        x = RU[("duel", r.arm, r.level, r.measure)]
        where = f"17C {r.arm} {r.level} {r.measure} {r.quantity}"
        W.add(x[r.quantity], r.estimate, where)
        W.count(x["n_units"] == r.n_units and x["n_donors"] == r.n_donors, where + " counts")
    for r in E[E.level.str.endswith("HU") & (E.quantity == "abs_r")].itertuples():
        x = RU[("e_hu", r.arm, r.level)]
        where = f"17E {r.arm} {r.level}"
        W.add(x["estimate"], r.estimate, where)
        W.count(x["n_units"] == r.n_units and x["n_donors"] == r.n_donors, where + " counts")
    for r in E[E.level.str.endswith("HU") & (E.quantity == "vif")].itertuples():
        x = RU[("e_hu", r.arm, r.level)]["estimate"]
        W.add((1 / (1 - x ** 2) if x < 1 else np.inf) if np.isfinite(x) else np.nan, r.estimate,
              f"17E {r.arm} {r.level} vif")
    L.chk("B-16", W.bad == 0 and W.worst < TOL,
          "every 17A, 17B, S12, 17C and 17E histology-unit statistic recomputed from the Fig17_units values "
          "equals its stats row", W.detail())
    L.chk("B-17", len(SHEET) == len(U), "Fig17_units holds every unit of the unit table", f"{len(SHEET)} units")

    # ---- Fig17D_donors: the donor frame 17D and the 17E clinical rows read -------------
    dcols = ["arm", "dataset", "donor_key", "donor_id"] + list(dict.fromkeys(COVS + list(OUT.values()))) \
        + ["module_post_pct", "mito_log1p"]
    miss = [c for c in dcols if c not in DN.columns]
    if miss:
        sys.exit(f"ERROR: {DONORS} lacks {', '.join(miss)}")
    DN["_a"] = DN["arm"].map({a: i for i, a in enumerate(ARMS)})
    DSHEET = DN[DN["_a"].notna()].sort_values("_a", kind="stable").reset_index(drop=True)[dcols].copy()
    W = Worst()
    DVi = DV.set_index(["arm", "dataset", "outcome"])
    for a in ARMS:
        for ds in DATASETS:
            d = DSHEET[(DSHEET.arm == a) & (DSHEET.dataset == ds)]
            for k, y in OUT.items():
                cv = [c for c in COVS if c != y]
                f = fit_duel(d, y, cv, min_d)
                v = DVi.loc[(a, ds, k)]
                where = f"17D {a} {ds} {k}"
                W.count(f["fitted"] == v["fitted"] and f["n"] == v["n_donors"], where + " fitted and n")
                if f["fitted"]:
                    for mine_k, col in (("b_m", "beta_module"), ("b_t", "beta_mito"), ("single_m", "single_module"),
                                        ("single_t", "single_mito"), ("vif", "vif")):
                        W.add(f[mine_k], v[col], f"{where} {col}")
            dc = d[~d[COVS].astype(float).isna().any(axis=1).values]
            ok_ = np.isfinite(dc["module_post_pct"].astype(float).values) & np.isfinite(dc["mito_log1p"].astype(float).values)
            x, yv = dc["module_post_pct"].astype(float).values[ok_], dc["mito_log1p"].astype(float).values[ok_]
            nd = len(pd.unique(dc["donor_key"].values[ok_]))
            est = abs(float(np.corrcoef(x, yv)[0, 1])) if nd >= min_d else np.nan
            e = E[(E.arm == a) & (E.level == "between donors CG") & (E.measure == ds)].set_index("quantity")
            W.add(est, e.loc["abs_r", "estimate"], f"17E {a} {ds} |r|")
            W.count(int(ok_.sum()) == e.loc["abs_r", "n_units"] and nd == e.loc["abs_r", "n_donors"],
                    f"17E {a} {ds} counts")
    L.chk("B-18", W.bad == 0 and W.worst < TOL,
          "every 17D coefficient, single-predictor coefficient and VIF and every 17E clinical-grain |r| "
          "recomputed from the Fig17D_donors values equals its row", W.detail())
    L._say("")

    # ---- workbook ------------------------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)
    ns = cfg["region_map"]["no_sequencing"]
    npath = dict(logs["19"]["inputs"])
    n_mt = CV[(CV.arm != "HKG8") & (CV.verdict == "%MT only")]
    r_line = float(np.sqrt(1 - 1 / vr))

    ov = [
        ["Fig 17A", "Does microglial shape track Module post-IP%, or just %MT?",
         "Within-donor Spearman rs of Module post-IP%, log1p %MT and the cytoplasmic share with the 6 IBA1 measures",
         "19", "Stats_17A.csv", "Fig17A_readout; Fig17A_stats; Fig17_units", "Notes 2 to 7 and 13"],
        ["Fig 17B", "Does nearby pathology track Module post-IP%, or just %MT?",
         "Spearman rs of Module post-IP%, log1p %MT, the cytoplasmic share, contamination and log cell count with the CPS scores, within donor (local scores) and between donors (all six)",
         "19", "Stats_17B.csv", "Fig17B_stats; Fig17_units", "Notes 2 to 6, 8 and 13"],
        ["Fig 17C", "Can Module post-IP% and %MT be separated on histology?",
         "rank(measure) ~ rank(Module post-IP%) + rank(%MT), standardised, no covariates, for each of the 28 stain measures, within donor and between donors",
         "20", "Stats_17C.csv; Plot_17C_verdicts.csv", "Fig17C_verdicts; Fig17C_stats; Fig17_units", "Notes 2 to 6, 9 and 13"],
        ["Fig 17D", "Can Module post-IP% and %MT be separated on the clinical outcomes?",
         "outcome ~ cov4 + Module post-IP% + %MT, standardised, OLS on the donor table, region sets fitted separately, for the 9 outcomes",
         "20", "Stats_17D.csv; Plot_17D_verdicts.csv", "Fig17D_verdicts; Fig17D_stats; Fig17D_donors", "Notes 3, 4 and 10"],
        ["Fig 17E", "Why are so many models unreadable?",
         "|r| between Module post-IP% and %MT and the VIF, in the frame each 17C and 17D model uses",
         "20", "Stats_17E.csv", "Fig17E_readout; Fig17E_stats; Fig17_units; Fig17D_donors", "Notes 3 and 11"],
        ["Supp Fig 13", "How do the stain families track each quantity?",
         "Panel A: median rs over each stain family's measures; panel B: family median of |Module rs| - |%MT rs|",
         "19", "Stats_S12.csv; Stats_S12_measures.csv; Plot_S12_counts.csv",
         "FigS13_stats; FigS13_measures; FigS13_counts; Fig17_units", "Notes 2 to 6, 12, 13 and 15"],
    ]
    write_table(wb, "Overview", ["Panel", "Question", "Quantity", "Section", "Source file", "Sheets", "See"], ov,
                wrap=True, widths={"Panel": 11, "Question": 40, "Quantity": 60, "Section": 9,
                                   "Source file": 34, "Sheets": 38, "See": 22})

    fam_txt = ", ".join(f"{FAM_LAB[f['key']]} ({len(f['measures'])})" for f in HI["families"])
    notes = [
        ["1", "Arms", "M4 = Mic13_PVM3, the masked primary arm. PVM_3 = PVM3_all, the unmasked reference. PVM_1 = PVM1_all, the second supertype. HKG8 = the housekeeping control construct (ACTB with the cHKG8 genes). NeoC and NonNeoC are the two region sets. 17A, 17B, 17C, Supplementary Figure 13 and the histology-unit rows of 17E pool them within donor; 17D and the clinical-grain rows of 17E fit them separately."],
        ["2", "Histology and CPS", f"Neuropathology: {pathlib.PurePath(npath.get('neuropathology table', '')).name}, sheet '{cfg['inputs']['neuropath_table']['sheet']}', one row per donor and region, grey matter only. The 28 stain measures are its measure columns less the four hematoxylin nucleus measures (a counterstain). Stain families: {fam_txt}. CPS (continuous pseudo-progression score): {pathlib.PurePath(npath.get('CPS workbook', '')).name}, sheet '{cfg['inputs']['cps_workbook']['sheet']}'; the three local scores are region-resolved and the three global scores are one value per donor. Region names are mapped to the snRNA-seq regions by the config map; {' and '.join(ns)} have no sequencing and are left out, and LEC units have no histology or CPS row."],
        ["3", "Resampling and p values", f"Intervals are 95% percentile intervals from a donor bootstrap with B = {B_:,}: donors are resampled with replacement, each bringing all its units. Donor indices are drawn once per arm and frame, so every rs of an arm in a frame uses the same resamples. The bootstrap p is min(1, 2(k+1)/(B+1)), where k is the number of resamples on the far side of zero; its smallest value, 2/(B+1) = {2 / (B_ + 1):.2e}, is a bound, and a BH value computed from floor p values is also a bound. 17D is the exception: its p values and 95% intervals are from OLS, as in the Figure 12 clinical model. 17E and the per-measure rows of Supplementary Figure 13 carry no test."],
        ["4", "Multiple testing", "fdr is Benjamini-Hochberg within the group named in the family column. 17A: within arm, level and quantity over the 6 IBA1 measures. 17B: within arm, level and quantity over the scores of that level (3 within donor, 6 between donors). 17C: within arm, level and predictor over the 28 stain measures. 17D: within arm, region set and predictor over the 9 outcomes. Supplementary Figure 13 panel A: within arm, level and quantity over the 7 stain families; panel B: within arm and level over the 7 families. Rows whose family ends in 'descriptive' carry no p. The builder recomputes BH within every family (check B-06)."],
        ["5", "Floors and untestable rows", f"A unit is one donor in one region on one arm with at least {int(cfg['floors']['cells_per_unit'])} cells of the arm. A test needs {min_d} donors in the observed data; a bootstrap resample needs 3 distinct donors. The within-donor level keeps donors with at least {min_u} units with a finite value on every variable of the row. A row below a floor has an empty estimate, interval, p and fdr and is drawn as untestable."],
        ["6", "Levels", "Within donor: every variable of the row is centred on the donor's mean over its units, and rs is taken over the pooled centred values. Between donors: one value per donor, the mean over its units; a global CPS score is the donor's value. n_units counts the units entering the row (at the between-donor level, the units averaged); n_donors counts the donors."],
        ["7", "Fig 17A", "Pre-set reading: the activation account predicts a positive rs with activated IBA1+ cells per area and a negative rs with inactivated IBA1+ cells per area, process area per cell and process length per cell. IBA1+ cells per area and IBA1+ area (%) are abundance measures with no sign set in advance. Fig17A_readout gives the within-donor rows as drawn, with the predicted sign; its direction, significance and same-sign columns are Excel formulas. The between-donor rows are in Fig17A_stats and are not drawn."],
        ["8", "Fig 17B", "A positive rs supports the account that Module post-IP% follows pathology. Within donor uses the three local scores, since the global scores are constant within donor; between donors uses all six. The figure draws Module post-IP%, %MT, the cytoplasmic share and log cell count; contamination is in Fig17B_stats only."],
        ["9", "Fig 17C", f"Within donor, every variable is centred on the donor before ranking; between donors, donor means are ranked. The two standardised coefficients come from the three rank correlations, and each bootstrap resample refits them. VIF = 1/(1 - r^2), r the rank correlation of the two predictors on the full sample; a model with VIF of {vr:g} or more is unreadable. Verdict per row: Untestable (not fitted), Unreadable, otherwise Module only, %MT only, Both or Neither from the two FDRs at {alpha:g}. A sign flip is a joint coefficient whose sign differs from that predictor's marginal rs (rs_module_marginal, rs_mito_marginal). On the real arms {len(n_mt)} rows are %MT only, {int(n_mt.module_sign_flip.sum())} of them with the Module coefficient reversed. In Fig17C_verdicts the verdict column is an Excel formula over the same row; the builder checks the analysis's verdicts against the same rule (check B-07)."],
        ["10", "Fig 17D", f"The 15A clinical model with both predictors: outcome ~ {' + '.join(COVS)} + Module post-IP% + %MT, fitted by OLS on the donor table of each arm and region set, outcome and the two predictors standardised. APOE e4 dose leaves the covariates when it is the outcome. Complete cases only; a model needs {min_d} donors. VIF from the two predictors on the model's donors. Verdicts as 17C. A sign flip is a joint coefficient whose sign differs from the coefficient of that predictor fitted alone with cov4 (single_module, single_mito). Fig17D_donors holds the donor values; the builder refits every model from them and checks the coefficients, single-predictor coefficients, VIF and donor counts (check B-18). The OLS p values are not recomputed."],
        ["11", "Fig 17E", f"|r| between Module post-IP% and log1p %MT and its VIF, in the frame each model uses: within donor and between donors on the histology units (Spearman, region sets pooled, as 17C; level codes 'within donor HU' and 'between donors HU'), and between donors on the donor table with cov4 complete (Pearson, region sets split, as 17D; level code 'between donors CG', region set in the measure column). The VIF {vr:g} line lies at |r| = {r_line:.3f}. Fig17E_readout gives the position of each row against the line as an Excel formula. Descriptive, no test."],
        ["12", "Supplementary Figure 13", "Panel A (quantity = the unit quantity): per arm, level, quantity and stain family, the median rs over the family's tested measures, on the same resamples as the per-measure rs. Panel B (quantity = abs_module_minus_abs_mito): the family median of |Module rs| - |%MT rs| over measures tested for both; negative means %MT tracks the family more strongly. On family rows n_units and n_donors are the largest over the family's tested measures. FigS13_counts gives the tested and total measures per family; FigS13_measures the per-measure rs with intervals and no test. The figure draws the within-donor rows for Module post-IP%, %MT, contamination and log cell count."],
        ["13", "Unit values (Fig17_units)", f"One row per unit, with the five unit quantities, the 28 stain measures and the six CPS scores joined by donor and region. post_ip_pct = Module post-IP%, per cent of the arm's {'supertype cells' if cfg['quantities']['module_post_pct_denominator'] == 'supertype' else 'labelled cells'}. To recompute a 17A, 17B or Supplementary Figure 13 rs: take the arm's units with a value on both variables; between donors, average each over the donor's units and take Spearman rs over donors; within donor, keep donors with at least {min_u} such units, subtract each donor's mean from both, and take Spearman rs over units. 17C uses the same frames with three variables. The 17E histology-unit rows use the units with joined_histology TRUE. A row with fewer than {min_d} donors is untestable. The builder recomputes every such statistic this way (check B-16)."],
        ["14", "Old values", "The archived run is not compared here (decision D20 dropped, revision 49). Results are reported as they stand."],
        ["15", "Supplementary Figure 13 file names", "Supplementary Figure 13 was numbered S12 when the analysis ran. The source files Stats_S12.csv, Stats_S12_measures.csv and Plot_S12_counts.csv, and the S12 labels in the build log checks, keep that name, as the logs name them. In FigS13_stats and FigS13_measures the panel column reads S13 and the family column starts with S13; every other value is as in the source files."],
    ]
    write_table(wb, "Notes", ["Note", "Subject", "Detail"], notes, wrap=True,
                widths={"Note": 7, "Subject": 28, "Detail": 120})

    gl = [
        ["column", "panel", "Figure panel the row belongs to."],
        ["column", "arm", "M4, PVM_3, PVM_1 or HKG8 (Note 1)."],
        ["column", "level", "17A, 17B, 17C, S13: within donor or between donors; 17D: the region set; 17E: within donor HU, between donors HU or between donors CG (Note 11)."],
        ["column", "quantity", "What the row estimates; codes below."],
        ["column", "measure", "17A, 17C, S13 per-measure: the stain measure; 17B: the CPS score; 17D: the outcome; 17E: pooled or the region set; S13 family rows: the stain family key."],
        ["column", "n_units", "Units behind the row (Note 6); empty in 17D."],
        ["column", "n_donors", "Donors behind the row."],
        ["column", "estimate, lo, hi", "Point estimate and its 95% interval (Note 3)."],
        ["column", "p, fdr", "p value and Benjamini-Hochberg value within family (Notes 3 and 4)."],
        ["column", "family", "The BH group of the row."],
        ["17A, 17B, S13", "module_post_pct, mito_log1p, cyto_share, contamination, log_n_cells", "Spearman rs of Module post-IP%, log1p %MT, the cytoplasmic share, unit mean DecontX contamination or log cell count with the measure or score."],
        ["17A", "predicted_sign", "In Fig17A_readout: +1, -1, or 0 for no sign set in advance (Note 7)."],
        ["17C", "beta_module, beta_mito", "Standardised rank coefficients of Module post-IP% and %MT in the joint model (Note 9)."],
        ["17C", "vif, rs_module_marginal, rs_mito_marginal", "VIF of the two ranked predictors; each predictor's rs with the measure alone. Descriptive rows."],
        ["17D", "beta_module, beta_mito, vif", "Standardised OLS coefficients of Module post-IP% and %MT with cov4; VIF of the two predictors (Note 10)."],
        ["17C, 17D", "fitted, verdict, module_sign_flip, mito_sign_flip", "In the verdict sheets: whether the model was fitted, its verdict (Excel formula) and the sign flips (Notes 9 and 10)."],
        ["17D", "single_module, single_mito", "In Fig17D_verdicts: the coefficient of each predictor fitted alone with cov4."],
        ["17E", "abs_r, vif", "|r| of Module post-IP% with %MT, and 1/(1 - r^2) (Note 11)."],
        ["S13", "abs_module_minus_abs_mito", "Panel B: family median of |Module rs| - |%MT rs| (Note 12)."],
        ["S13", "n_tested, n_total", "In FigS13_counts: measures of the family tested, and in the family."],
        ["units", "region_set, donor_key, region", "The unit's region set, donor and region."],
        ["units", "post_ip_pct, mito_log1p, cyto_share, contamination, log_n_cells", "Unit Module post-IP% (Note 13), log1p unit mean %MT, unit mean cytoplasmic share, unit mean DecontX contamination, log of the unit's labelled cells."],
        ["units", "joined_histology, joined_cps", "Whether the unit has a neuropathology row and a CPS row."],
        ["units", "stain measures and CPS scores", "The neuropathology and CPS values of the unit's donor and region, by the workbook column names."],
        ["donors", "cov4 and outcome columns", f"The donor-table columns of the 17D models: cov4 {', '.join(COVS)}; outcomes {', '.join(OUT.values())}."],
        ["donors", "module_post_pct, mito_log1p", "Donor Module post-IP% and donor log1p %MT, as Figure 12."],
    ]
    write_table(wb, "Glossary", ["Scope", "Term", "Meaning"], gl, wrap=True,
                widths={"Scope": 14, "Term": 46, "Meaning": 100})

    # Fig17A readout: within donor, as drawn
    hA = ["arm", "measure", "predicted_sign", "n_units", "n_donors", "module_rs", "module_lo", "module_hi",
          "module_fdr", "mito_rs", "mito_fdr", "cyto_rs", "cyto_fdr", "module_significant", "module_direction",
          "mito_same_sign"]
    cA = {h: get_column_letter(j + 1) for j, h in enumerate(hA)}
    Ai = A.set_index(["arm", "level", "quantity", "measure"])
    rows, i = [], 2
    for a in ARMS:
        for m in IBA:
            g = lambda q: Ai.loc[(a, "within donor", q, m)]
            mo, mi, cy = g("module_post_pct"), g("mito_log1p"), g("cyto_share")
            c = lambda h: f"{cA[h]}{i}"
            rows.append([a, IBA_LAB[m], SIGN[m], mo["n_units"], mo["n_donors"], mo["estimate"], mo["lo"], mo["hi"],
                         mo["fdr"], mi["estimate"], mi["fdr"], cy["estimate"], cy["fdr"],
                         f'=IF(ISNUMBER({c("module_fdr")}),IF({c("module_fdr")}<{alpha},"yes","no"),"")',
                         f'=IF(NOT(ISNUMBER({c("module_rs")})),"untestable",IF({c("predicted_sign")}=0,"no sign set",'
                         f'IF({c("module_rs")}*{c("predicted_sign")}>0,"predicted","opposite")))',
                         f'=IF(AND(ISNUMBER({c("module_rs")}),ISNUMBER({c("mito_rs")})),'
                         f'IF(SIGN({c("module_rs")})=SIGN({c("mito_rs")}),"yes","no"),"")'])
            i += 1
    write_table(wb, "Fig17A_readout", hA, rows, p_cols=("module_fdr", "mito_fdr", "cyto_fdr"),
                fmt_cols={k: "0.000" for k in ("module_rs", "module_lo", "module_hi", "mito_rs", "cyto_rs")},
                widths={"measure": 40})

    # Fig17C and Fig17D verdicts, verdict as a formula over the row
    hC = ["arm", "level", "measure", "stain_family", "block", "n_units", "n_donors", "fitted", "vif",
          "beta_module", "beta_mito", "rs_module", "rs_mito", "fdr_module", "fdr_mito", "module_sign_flip",
          "mito_sign_flip", "verdict"]
    cC = {h: get_column_letter(j + 1) for j, h in enumerate(hC)}
    rows = []
    for i, r in enumerate(CV.itertuples(), start=2):
        rows.append([r.arm, r.level, r.measure, FAM_LAB[r.stain_family], r.block, r.n_units, r.n_donors,
                     bool(r.fitted), r.vif, r.beta_module, r.beta_mito, r.rs_module, r.rs_mito, r.fdr_module,
                     r.fdr_mito, bool(r.module_sign_flip), bool(r.mito_sign_flip), verdict_formula(cC, i, alpha, vr)])
    write_table(wb, "Fig17C_verdicts", hC, rows, p_cols=("fdr_module", "fdr_mito"),
                fmt_cols={k: "0.000" for k in ("vif", "beta_module", "beta_mito", "rs_module", "rs_mito")},
                widths={"measure": 40, "verdict": 14})

    hD = ["arm", "region_set", "outcome", "class", "n_donors", "fitted", "vif", "beta_module", "beta_mito",
          "single_module", "single_mito", "fdr_module", "fdr_mito", "module_sign_flip", "mito_sign_flip", "verdict"]
    cD = {h: get_column_letter(j + 1) for j, h in enumerate(hD)}
    rows = []
    for i, r in enumerate(DV.itertuples(), start=2):
        rows.append([r.arm, r.dataset, r.outcome, r.oclass,
                     r.n_donors, bool(r.fitted), r.vif, r.beta_module, r.beta_mito, r.single_module,
                     r.single_mito, r.fdr_module, r.fdr_mito, bool(r.module_sign_flip), bool(r.mito_sign_flip),
                     verdict_formula(cD, i, alpha, vr)])
    write_table(wb, "Fig17D_verdicts", hD, rows, p_cols=("fdr_module", "fdr_mito"),
                fmt_cols={k: "0.000" for k in ("vif", "beta_module", "beta_mito", "single_module", "single_mito")},
                widths={"verdict": 14})

    # Fig17E readout
    hE = ["arm", "level", "region_set", "n_units", "n_donors", "abs_r", "lo", "hi", "vif", "position"]
    cE = {h: get_column_letter(j + 1) for j, h in enumerate(hE)}
    Ei = E.set_index(["arm", "level", "measure", "quantity"])
    rows, i = [], 2
    for a in ARMS:
        for lv, ms in (("within donor HU", ["pooled"]), ("between donors HU", ["pooled"]),
                       ("between donors CG", DATASETS)):
            for m in ms:
                r, v = Ei.loc[(a, lv, m, "abs_r")], Ei.loc[(a, lv, m, "vif")]
                vc = f"{cE['vif']}{i}"
                rows.append([a, lv, m, r["n_units"], r["n_donors"], r["estimate"], r["lo"], r["hi"],
                             v["estimate"] if np.isfinite(v["estimate"]) else None,
                             f'=IF(ISNUMBER({vc}),IF({vc}>={vr:g},"at or above the VIF {vr:g} line",'
                             f'"below the VIF {vr:g} line"),"untestable")'])
                i += 1
    write_table(wb, "Fig17E_readout", hE, rows,
                fmt_cols={k: "0.000" for k in ("abs_r", "lo", "hi")} | {"vif": "0.0"}, widths={"position": 30})

    # Fig17_units and Fig17D_donors
    hU = [UNIT_RENAME.get(c, c) for c in ucols]
    write_table(wb, "Fig17_units", hU, SHEET.itertuples(index=False, name=None),
                fmt_cols={"post_ip_pct": "0.000"})
    hDn = [UNIT_RENAME.get(c, c) for c in dcols]
    write_table(wb, "Fig17D_donors", hDn, DSHEET.itertuples(index=False, name=None),
                fmt_cols={"post_ip_pct": "0.000"})

    # Stats and plot-data sheets, as written by the analysis
    stats_sheet(wb, "Fig17A_stats", A)
    stats_sheet(wb, "Fig17B_stats", Bt)
    stats_sheet(wb, "Fig17C_stats", C)
    stats_sheet(wb, "Fig17D_stats", Dd)
    stats_sheet(wb, "Fig17E_stats", E)
    # Supplementary Figure 13 was S12 at analysis time; relabel panel and family on the sheets only.
    # The checks above ran on the source files as read.
    def as_s13(df):
        d = df.copy()
        d["panel"] = d["panel"].replace({"S12": "S13"})
        d["family"] = d["family"].where(d["family"].isna(),
                                        d["family"].astype(str).str.replace(r"^S12", "S13", regex=True))
        return d
    stats_sheet(wb, "FigS13_stats", as_s13(S))
    stats_sheet(wb, "FigS13_measures", as_s13(SM))
    scol = [c for c in SC.columns if c != "run_id"]
    write_table(wb, "FigS13_counts", scol, SC[scol].itertuples(index=False, name=None))

    # sheets in panel order: the three general sheets, then 17A to 17E, then
    # Supplementary Figure 13, then the unit and donor values
    order = ["Overview", "Notes", "Glossary",
             "Fig17A_readout", "Fig17A_stats", "Fig17B_stats",
             "Fig17C_verdicts", "Fig17C_stats", "Fig17D_verdicts", "Fig17D_stats",
             "Fig17E_readout", "Fig17E_stats",
             "FigS13_stats", "FigS13_measures", "FigS13_counts",
             "Fig17_units", "Fig17D_donors"]
    if sorted(order) != sorted(wb.sheetnames):
        sys.exit(f"ERROR: sheet order list does not match the sheets written: {wb.sheetnames}")
    wb._sheets = [wb[n] for n in order]
    wb.active = 0

    # every string in the workbook, for B-13
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts += [v for v in row if isinstance(v, str)]
    bad = [t for t in texts if any(d_ in t for d_ in DASHES)]
    L.chk("B-13", not bad, "no en dash or em dash in the workbook text", f"{len(bad)} found")

    L._say("")
    L._say(f"Checks: {L.n['PASS']} PASS, {L.n['FAIL']} FAIL, {L.n['NOTE']} NOTE")
    log_path = out.with_name(out.stem + "_build_log.txt")
    out.parent.mkdir(parents=True, exist_ok=True)
    if L.n["FAIL"]:
        L._say("VERDICT: FAIL | workbook not written")
        log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")
        sys.exit(1)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    L._say(f"Output: {out}")
    L._say("VERDICT: PASS")
    log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
