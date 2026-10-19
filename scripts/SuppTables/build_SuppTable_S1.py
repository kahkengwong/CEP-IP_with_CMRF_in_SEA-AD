#!/usr/bin/env python3
"""
build_SuppTable_S1.py

Builds the data sheets of Supplementary Table S1 (cell selection: choice of
the 3.0% cutoff and of the M4 scorer, in each region set) from the tables
written by Part 2 and Part 3 of the Python pipeline. Sheets A to F and J are
the arm-level Part 2 tables (cutoff rule, marker dilution, mixture fits,
bootstrap stability, arm matching, scorer overlap and sweep summary), sheet I
is the Part 2 scorer comparison at the primary 3.0% cutoff, and sheets G and
H are the Part 3 composition tables of the M4 selected set and of all six
scorers. Where a Part 2 table spans all six scorers, only the M4_M1andM2 rows
are kept, except in sheet C (the three base scorers) and sheet I (all six).
Each sheet keeps the published columns in the published order; an arm column
(NeoC or NonNeoC) is added in front where the source has none, and the arms
are stacked NeoC first. Values are copied as written by the pipeline:
booleans are written as the text TRUE or FALSE, missing values as blank
cells, and en or em dashes in text as a plain hyphen. finalize_supp_tables.py
adds the Contents sheet and saves the published file.

Usage:
    python build_SuppTable_S1.py
    python build_SuppTable_S1.py --root /data/SEA-AD
Root: --root, else $SEAAD_ROOT, else the current folder.

Inputs (<arm dir> is [1]_NeoC/Part2_NeoC_M4 or [2]_Non-NeoC/Part2_NonNeoC_M4;
a file may also carry the arm prefix, for example NeoC_S2b_cutoff_rule_verdict.csv):
    <arm dir>/S2b_cutoff_rule_verdict.csv     sheet A  (Part2_<arm>_CP10K_code.py)
    <arm dir>/S2c_marker_dilution.csv         sheet B  (Part2_<arm>_CP10K_code.py)
    <arm dir>/S2d_mixture.csv                 sheet C  (Part2_<arm>_CP10K_code.py)
    <arm dir>/S2e_bootstrap_stability.csv     sheet D  (Part2_<arm>_CP10K_code.py)
    <arm dir>/S2f_arm_matching.csv            sheet E  (Part2_<arm>_CP10K_code.py;
                                              written only by the arm run second)
    <arm dir>/S2g_scorer_overlap.csv          sheet F  (Part2_<arm>_CP10K_code.py)
    <arm dir>/top_3.0pct/<arm>_S1_supertype_enrichment_M4.csv
                                              sheet G  (Part3_M4Composition_code.py)
    <arm dir>/top_3.0pct/<arm>_S7_scorer_composition.csv
                                              sheet H  (Part3_M4Composition_code.py)
    <arm dir>/top_3.0pct/S6_method_comparison.csv
                                              sheet I  (Part2_<arm>_CP10K_code.py)
    <arm dir>/S7_sweep_summary.csv            sheet J  (Part2_<arm>_CP10K_code.py)
Output:
    <root>/Manuscript/Supp_Tables/build/SuppTable_S1_CellSelection.xlsx
        sheets A_cutoff_rule_verdict, B_marker_dilution, C_mixture,
        D_bootstrap_stability, E_arm_matching, F_scorer_overlap,
        G_selection_by_supertype, H_composition_by_scorer,
        I_method_comparison, J_sweep_summary
    <root>/Manuscript/Supp_Tables/build/build_SuppTable_S1_log.txt
        inputs, checks and row counts, ending with VERDICT: PASS or
        VERDICT: FAIL. On FAIL no workbook is written.

Formatting: Calibri 11 (openpyxl default), bold header row, frozen top row,
no fills, no borders. finalize_supp_tables.py applies the published fonts.
"""

import argparse
import datetime as dt
import math
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

ARM_FOLDER = {"NeoC": "[1]_NeoC", "NonNeoC": "[2]_Non-NeoC"}
M4 = "M4_M1andM2"
PRIMARY_DIR = "top_3.0pct"
OUT_NAME = "SuppTable_S1_CellSelection.xlsx"
LOG_NAME = "build_SuppTable_S1_log.txt"

# Sheet specifications, in the published order.
#   stem    source file name as Part 2 or Part 3 writes it
#   sub     subfolder of the arm's Part 2 folder ("" for the folder itself)
#   prefix  True: the arm prefix is part of the written name (Part 3)
#   filter  (column, value) rows kept, or None for all rows
#   arm     True: add the arm column in front; False: the source has the
#           columns already (sheet H has its own arm column, sheet E has
#           this_arm and other_arm)
#   cols    published columns, in order (without the added arm column)
#   n_arm   rows per arm in the published table, used only in a check
SHEETS = [
    dict(name="A_cutoff_rule_verdict", stem="S2b_cutoff_rule_verdict.csv", sub="",
         prefix=False, filter=("method", M4), arm=True, n_arm=101,
         cols=["top_pct", "n_top_total", "pass_raw", "pass_guarded",
               "n_violators_raw", "n_violators_guarded", "violators_raw",
               "violators_guarded", "margin_cells_raw", "margin_cells_guarded",
               "enriched_are_top2"]),
    dict(name="B_marker_dilution", stem="S2c_marker_dilution.csv", sub="",
         prefix=False, filter=("method", M4), arm=True, n_arm=101,
         cols=["top_pct", "n_cum", "n_band",
               "cum_fold_heldout_up", "marg_fold_heldout_up", "cum_det_heldout_up",
               "cum_fold_inpanel_up", "marg_fold_inpanel_up", "cum_det_inpanel_up",
               "cum_fold_homeostatic", "marg_fold_homeostatic", "cum_det_homeostatic"]),
    dict(name="C_mixture", stem="S2d_mixture.csv", sub="",
         prefix=False, filter=None, arm=True, n_arm=3,
         cols=["method", "n_used", "bic_k1", "bic_k2", "dbic_1_minus_2", "best_k",
               "w_upper", "separation", "bimodality_coef", "null_dbic_median",
               "null_dbic_max", "null_sep_median", "p_dbic_vs_null", "p_sep_vs_null",
               "verdict"]),
    dict(name="D_bootstrap_stability", stem="S2e_bootstrap_stability.csv", sub="",
         prefix=False, filter=("method", M4), arm=True, n_arm=21,
         cols=["top_pct", "n_ref", "implied_pct_mean", "implied_pct_sd",
               "implied_pct_cv", "mean_jaccard_vs_full", "core_n", "core_frac"]),
    dict(name="E_arm_matching", stem="S2f_arm_matching.csv", sub="",
         prefix=False, filter=None, arm=False, n_arm=None,
         cols=["criterion", "this_arm", "this_top_pct", "this_n", "this_purity",
               "other_arm", "other_top_pct", "other_n", "other_purity"]),
    dict(name="F_scorer_overlap", stem="S2g_scorer_overlap.csv", sub="",
         prefix=False, filter=("intersection", M4), arm=True, n_arm=5,
         cols=["parent_a", "parent_b", "spearman_a_b", "top_pct",
               "n_parent_a_at_pct", "n_intersection", "retention_vs_parent_a",
               "effective_pct_of_arm", "contained_in_topk_A", "contained_in_topk_B"]),
    dict(name="G_selection_by_supertype", stem="S1_supertype_enrichment_M4.csv",
         sub=PRIMARY_DIR, prefix=True, filter=None, arm=True, n_arm=8,
         cols=["supertype", "n_total", "n_top", "n_donors_top", "pct_top",
               "enrichment_fold", "boot_ci_lo", "boot_ci_hi"]),
    dict(name="H_composition_by_scorer", stem="S7_scorer_composition.csv",
         sub=PRIMARY_DIR, prefix=True, filter=None, arm=False, n_arm=48,
         cols=["arm", "scorer", "supertype", "n_total", "n_top", "pct_top",
               "enrichment_fold"]),
    dict(name="I_method_comparison", stem="S6_method_comparison.csv",
         sub=PRIMARY_DIR, prefix=False, filter=None, arm=True, n_arm=6,
         cols=["method", "label", "n_BH_sig", "n_raw_sig", "mean_winner_rho",
               "mean_abs_rho_pos", "median_abs_rho_pos", "n_correct_dir",
               "n_pos_tests"]),
    dict(name="J_sweep_summary", stem="S7_sweep_summary.csv", sub="",
         prefix=False, filter=None, arm=True, n_arm=5,
         cols=["top_pct", "is_primary", "n_top_M1", "N_M4_TOP", "M4_thr_M1",
               "M4_thr_M2", "pct_PVM3", "pct_PVM4", "worst_other", "worst_other_pct",
               "worst_other_n", "worst_other_gd", "worst_other_gd_pct",
               "rule_pass_raw", "rule_pass_guarded", "margin_cells_raw",
               "margin_cells_gd", "donor_rho_ADNC", "donor_rho_Braak",
               "best_method", "mean_winner_rho"]),
]

# Reference values used only in checks.
M4_TOTAL = {"NeoC": 3412, "NonNeoC": 2704}   # M4 selected nuclei at 3.0% (Figure 1B)
E_ROWS = 3                                    # sheet E: the NonNeoC run's three criteria


##### Log #####
class Log:
    def __init__(self):
        self.lines, self.fails = [], 0

    def say(self, s=""):
        self.lines.append(str(s))
        print(s, flush=True)

    def chk(self, ok, what):
        self.say(f"  [{'PASS' if ok else 'FAIL'}] {what}")
        if not ok:
            self.fails += 1
        return bool(ok)

    def write(self, path):
        self.say("")
        self.say(f"VERDICT: {'PASS' if self.fails == 0 else 'FAIL'}")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.lines) + "\n")


##### Helpers #####
def find_source(root, arm, spec):
    """Path of the source file for one arm, or None. The name the pipeline
    writes is tried first, then the same name with or without the arm prefix."""
    d = os.path.join(root, ARM_FOLDER[arm], f"Part2_{arm}_M4", spec["sub"])
    plain, pref = spec["stem"], f"{arm}_{spec['stem']}"
    for name in ([pref, plain] if spec["prefix"] else [plain, pref]):
        p = os.path.join(d, name)
        if os.path.isfile(p):
            return p
    return None


def cell_value(v):
    """The value as the published sheets hold it: booleans as the text TRUE or
    FALSE, missing values blank, whole numbers as integers, en and em dashes in
    text as a hyphen."""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        v = float(v)
        if not math.isfinite(v):
            return None
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, str):
        if v == "":
            return None
        return v.replace(" \u2014 ", " - ").replace("\u2014", "-").replace("\u2013", "-")
    if pd.isna(v):
        return None
    return v


def count_dashes(df):
    n = 0
    for c in df.columns:
        if df[c].dtype == object:
            n += int(df[c].map(lambda v: isinstance(v, str)
                               and ("\u2014" in v or "\u2013" in v)).sum())
    return n


def build_sheet(L, root, spec):
    """Returns the sheet as a DataFrame with the published columns, or None."""
    L.say("")
    L.say(spec["name"])
    parts = []
    for arm in ARM_FOLDER:
        path = find_source(root, arm, spec)
        if path is None:
            if spec["name"] == "E_arm_matching":
                L.say(f"  {arm}: no {spec['stem']} (written only by the arm run second)")
                continue
            L.chk(False, f"{arm}: source {spec['stem']} found under "
                         f"{ARM_FOLDER[arm]}/Part2_{arm}_M4/{spec['sub']}")
            return None
        L.say(f"  {arm}: {os.path.relpath(path, root)}")
        df = pd.read_csv(path, float_precision="round_trip", keep_default_na=True)
        n_read = len(df)
        if spec["filter"]:
            col, val = spec["filter"]
            if not L.chk(col in df.columns, f"{arm}: column {col} present"):
                return None
            df = df[df[col] == val]
        missing = [c for c in spec["cols"] if c not in df.columns]
        if not L.chk(not missing, f"{arm}: published columns present in the source"
                                  + (f" (missing {missing})" if missing else "")):
            return None
        df = df[spec["cols"]].reset_index(drop=True)
        if spec["arm"]:
            df.insert(0, "arm", arm)
        L.say(f"    {n_read} rows read, {len(df)} kept")
        if spec["n_arm"] is not None:
            L.chk(len(df) == spec["n_arm"], f"{arm}: {len(df)} rows (published {spec['n_arm']})")
        parts.append(df)
    if not parts:
        L.chk(False, "at least one source file found")
        return None
    # Sheet E comes from the arm run second (NonNeoC). A rerun in a folder that
    # still holds the other arm's S2f_half file from an earlier run also gives a
    # NeoC file; the NonNeoC rows are kept, as in the published sheet.
    if spec["name"] == "E_arm_matching" and len(parts) > 1:
        keep = [d for d in parts if "this_arm" in d.columns
                and set(d["this_arm"].astype(str)) == {"NonNeoC"}]
        if keep:
            L.say(f"  both arms wrote {spec['stem']} (a rerun in a used folder); "
                  f"the NonNeoC rows are kept")
            parts = keep
    out = pd.concat(parts, ignore_index=True)
    nd = count_dashes(out)
    if nd:
        L.say(f"  {nd} text cells with an en or em dash, written with a hyphen")
    return out


def frames_differ(a, b):
    a = a.reset_index(drop=True).astype(str)
    b = b.reset_index(drop=True).astype(str)
    return a.shape != b.shape or not (a.values == b.values).all()


def write_workbook(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name, df in sheets:
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = bold
        for rec in df.itertuples(index=False):
            ws.append([cell_value(v) for v in rec])
        ws.freeze_panes = "A2"
    wb.save(path)


##### Main #####
def main():
    ap = argparse.ArgumentParser(description="Build Supplementary Table S1 (cell selection).")
    ap.add_argument("--root", default=os.environ.get("SEAAD_ROOT", os.getcwd()),
                    help="SEA-AD analysis folder (default $SEAAD_ROOT, else the current folder)")
    args = ap.parse_args()
    root = os.path.abspath(args.root)
    out_dir = os.path.join(root, "Manuscript", "Supp_Tables", "build")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, OUT_NAME)

    L = Log()
    L.say("build_SuppTable_S1.py")
    L.say(f"run   : {dt.datetime.now().isoformat(timespec='seconds')}")
    L.say(f"root  : {root}")
    L.say(f"output: {out_path}")

    built = []
    for spec in SHEETS:
        df = build_sheet(L, root, spec)
        if df is None:
            continue
        built.append((spec["name"], df))

    L.say("")
    L.say("Cross-sheet checks")
    got = dict(built)
    if "G_selection_by_supertype" in got:
        g = got["G_selection_by_supertype"]
        tot = {a: int(g.loc[g.arm == a, "n_top"].sum()) for a in ARM_FOLDER}
        L.chk(tot == M4_TOTAL, f"G: M4 selected totals {tot} (expected {M4_TOTAL})")
    if "J_sweep_summary" in got:
        j = got["J_sweep_summary"]
        prim = j[j["is_primary"].map(lambda v: str(v).upper() == "TRUE")]
        L.chk(len(prim) == 2 and all(np.isclose(prim["top_pct"].astype(float), 3.0)),
              "J: one primary row per arm, at 3.0%")
        if "G_selection_by_supertype" in got:
            m4n = {a: int(prim.loc[prim.arm == a, "N_M4_TOP"].iloc[0]) for a in prim.arm}
            L.chk(m4n == M4_TOTAL, f"J: N_M4_TOP at the primary cutoff {m4n} equals the G totals")
    if "E_arm_matching" in got:
        e = got["E_arm_matching"]
        if len(e) != E_ROWS:
            L.say(f"  note: E has {len(e)} rows; the published sheet has {E_ROWS} "
                  f"(only the NonNeoC run wrote S2f_arm_matching.csv)")
        else:
            L.say(f"  E: {len(e)} rows (this_arm {sorted(set(e.this_arm))})")
    for name, df in built:
        if "arm" in df.columns and set(df.arm) >= set(ARM_FOLDER):
            L.chk(frames_differ(df[df.arm == "NeoC"].drop(columns="arm"),
                                df[df.arm == "NonNeoC"].drop(columns="arm")),
                  f"{name}: NeoC and NonNeoC rows are not identical")

    L.chk(len(built) == len(SHEETS), f"{len(built)} of {len(SHEETS)} sheets built")
    L.say("")
    L.say("Sheets")
    for name, df in built:
        L.say(f"  {name:<26} {len(df):>4} rows x {df.shape[1]:>2} columns")
    if L.fails == 0:
        write_workbook(out_path, built)
        L.say(f"  saved {out_path}")
    else:
        L.say("  workbook not written (checks failed)")
    L.write(os.path.join(out_dir, LOG_NAME))
    sys.exit(0 if L.fails == 0 else 1)


if __name__ == "__main__":
    main()
