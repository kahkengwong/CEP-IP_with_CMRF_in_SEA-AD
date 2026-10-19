#!/usr/bin/env python3
"""
build_SuppTable_S2.py

Builds the data sheets of Supplementary Table S2 (Mic.13 validation) from the
workbooks written by Part 5 and Part 5b of the Python pipeline. Sheet A is the
summary of the differential expression of M4-selected against non-selected
cells (sheet 00_summary of the Part 5 master workbook). Sheets B and C are the
hypergeometric recovery tests of the Mic.13 markers and the per-gene Mic.13
table (sheets hypergeometric_tests and mic13_gene_detail of the Part 5b
workbook). For each sheet the two region sets are stacked, NeoC first, with
an arm column (NeoC or NonNeoC) in front; the other columns, their names,
their order and the row order are those of the source sheet. Values are
copied as stored in the source workbooks; empty text cells are written as
blank cells. finalize_supp_tables.py adds the Contents sheet, writes ">=" as
"≥" in the column name of sheet A and in the two robust-DEG test names of
sheet B, and saves the published file.

Usage:
    python build_SuppTable_S2.py
    python build_SuppTable_S2.py --root /data/SEA-AD
Root: --root, else $SEAAD_ROOT, else the current folder.

Inputs (<arm folder> is [1]_NeoC or [2]_Non-NeoC):
    <arm folder>/Part5_<arm>_DEGs/DEG_master_results.xlsx, sheet 00_summary
        (Part5_<arm>_DEGs.py; <arm>-DEG_master_results.xlsx is also accepted)
    <arm folder>/Part5b_<arm>_Validation/Part5b_<arm>_Mic13_validation.xlsx,
        sheets hypergeometric_tests and mic13_gene_detail (Part5b_<arm>_Validation.py)
Output:
    <root>/Manuscript/Supp_Tables/build/SuppTable_S2_Mic13Validation.xlsx
        sheets A_deg_summary, B_hypergeometric, C_mic13_gene_detail
    <root>/Manuscript/Supp_Tables/build/build_SuppTable_S2_log.txt
        inputs, checks and row counts, ending with VERDICT: PASS or
        VERDICT: FAIL. On FAIL no workbook is written.

Formatting: Calibri 11 (openpyxl default), bold header row, frozen top row,
no fills, no borders. finalize_supp_tables.py applies the published fonts.
"""

import argparse
import datetime as dt
import math
import os
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font

ARM_FOLDER = {"NeoC": "[1]_NeoC", "NonNeoC": "[2]_Non-NeoC"}
OUT_NAME = "SuppTable_S2_Mic13Validation.xlsx"
LOG_NAME = "build_SuppTable_S2_log.txt"

# Source columns, as Part 5 and Part 5b write them (checked, not renamed).
COLS = {
    "A_deg_summary": ["Supertype", "Region", "Mode", "N fg (M4-top)", "N bg (non-M4-top)",
                      "Genes tested", "N sig (p_adj<0.05, |log2FC|>=0.25, pct_fg>=0.15)",
                      "N up (M4-top)", "N down (M4-top)", "Status"],
    "B_hypergeometric": ["Test", "Reference set", "DEG list", "Universe", "Overlap",
                         "Expected", "Fold enrichment", "p-value", "Sig", "Overlap genes"],
    "C_mic13_gene_detail": ["Gene", "pct_exp (Mic.13)", "mean_exp (Mic.13)", "In scorer?",
                            "Status", "Sig DEG?", "Direction", "N regions sig", "Regions",
                            "log2FC range"],
}

# Reference values used only in checks (published table and Part 5b design:
# 68 Mic.13 markers held out of the scoring panel, 27 in it, 12 DAM2 held out).
ROWS_A = {"NeoC": 8, "NonNeoC": 6}     # regions plus the pooled PVM_2_3 and PVM_4 rows
ROWS_B = {"NeoC": 12, "NonNeoC": 10}   # one row per region, plus six pooled tests
ROWS_C = {"NeoC": 95, "NonNeoC": 95}
N_HELD, N_IN, N_DAM2 = 68, 27, 12


##### Log #####
class Log:
    def __init__(self):
        self.lines, self.fails = [], 0

    def say(self, s=""):
        self.lines.append(str(s))
        print(s, flush=True)

    def chk(self, ok, what):
        self.say(f"  [{'PASS' if ok else 'FAIL'}] {what}")
        if not ok:
            self.fails += 1
        return bool(ok)

    def write(self, path):
        self.say("")
        self.say(f"VERDICT: {'PASS' if self.fails == 0 else 'FAIL'}")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.lines) + "\n")


##### Helpers #####
def first(paths):
    return next((p for p in paths if os.path.isfile(p)), None)


def source(root, arm, sheet):
    if sheet == "A_deg_summary":
        d = os.path.join(root, ARM_FOLDER[arm], f"Part5_{arm}_DEGs")
        return first([os.path.join(d, "DEG_master_results.xlsx"),
                      os.path.join(d, f"{arm}-DEG_master_results.xlsx")]), "00_summary"
    d = os.path.join(root, ARM_FOLDER[arm], f"Part5b_{arm}_Validation")
    tab = "hypergeometric_tests" if sheet == "B_hypergeometric" else "mic13_gene_detail"
    return first([os.path.join(d, f"Part5b_{arm}_Mic13_validation.xlsx")]), tab


def read_tab(path, tab):
    """Header and rows of one sheet, values as stored; trailing empty rows dropped."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if tab not in wb.sheetnames:
            return None, None
        rows = [list(r) for r in wb[tab].iter_rows(values_only=True)]
    finally:
        wb.close()
    hdr = list(rows[0])
    while hdr and hdr[-1] is None:
        hdr.pop()
    body = [r[:len(hdr)] for r in rows[1:]]
    while body and all(v is None or v == "" for v in body[-1]):
        body.pop()
    return hdr, body


def cell_value(v):
    if isinstance(v, str) and v == "":
        return None
    if isinstance(v, float) and not math.isfinite(v):
        return None
    return v


def build_sheet(L, root, name, n_rows):
    L.say("")
    L.say(name)
    hdr_out, rows_out = None, []
    for arm in ARM_FOLDER:
        path, tab = source(root, arm, name)
        if not L.chk(path is not None, f"{arm}: source workbook found for [{tab}]"):
            return None, None
        L.say(f"  {arm}: {os.path.relpath(path, root)} [{tab}]")
        hdr, body = read_tab(path, tab)
        if not L.chk(hdr is not None, f"{arm}: sheet {tab} present"):
            return None, None
        if not L.chk(hdr == COLS[name], f"{arm}: columns as written by the pipeline"
                     + ("" if hdr == COLS[name] else f" (found {hdr})")):
            return None, None
        L.chk(len(body) == n_rows[arm], f"{arm}: {len(body)} rows (published {n_rows[arm]})")
        hdr_out = ["arm"] + hdr
        rows_out += [[arm] + [cell_value(v) for v in r] for r in body]
    return hdr_out, rows_out


def write_workbook(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name, hdr, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(hdr)
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
    wb.save(path)


##### Main #####
def main():
    ap = argparse.ArgumentParser(description="Build Supplementary Table S2 (Mic.13 validation).")
    ap.add_argument("--root", default=os.environ.get("SEAAD_ROOT", os.getcwd()),
                    help="SEA-AD analysis folder (default $SEAAD_ROOT, else the current folder)")
    args = ap.parse_args()
    root = os.path.abspath(args.root)
    out_dir = os.path.join(root, "Manuscript", "Supp_Tables", "build")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, OUT_NAME)

    L = Log()
    L.say("build_SuppTable_S2.py")
    L.say(f"run   : {dt.datetime.now().isoformat(timespec='seconds')}")
    L.say(f"root  : {root}")
    L.say(f"output: {out_path}")

    built = []
    for name, n_rows in (("A_deg_summary", ROWS_A), ("B_hypergeometric", ROWS_B),
                         ("C_mic13_gene_detail", ROWS_C)):
        hdr, rows = build_sheet(L, root, name, n_rows)
        if hdr is not None:
            built.append((name, hdr, rows))

    L.say("")
    L.say("Content checks")
    got = {n: (h, r) for n, h, r in built}
    if "B_hypergeometric" in got:
        h, rows = got["B_hypergeometric"]
        it, ir = h.index("Test"), h.index("Reference set")
        for arm in ARM_FOLDER:
            sub = [r for r in rows if r[0] == arm]
            held = [r[ir] for r in sub if str(r[it]).startswith("Mic.13 held-out")]
            ins = [r[ir] for r in sub if "IN-SCORER" in str(r[it])]
            dam = [r[ir] for r in sub if str(r[it]).startswith("DAM2 held-out")]
            L.chk(held and all(v == N_HELD for v in held),
                  f"B {arm}: {len(held)} held-out tests, each on {N_HELD} markers")
            L.chk(ins == [N_IN], f"B {arm}: one IN-SCORER test on {N_IN} markers")
            L.chk(dam and all(v == N_DAM2 for v in dam),
                  f"B {arm}: {len(dam)} DAM2 tests, each on {N_DAM2} genes")
    if "C_mic13_gene_detail" in got:
        h, rows = got["C_mic13_gene_detail"]
        js = h.index("Status")
        for arm in ARM_FOLDER:
            sub = [r for r in rows if r[0] == arm]
            n_held = sum(r[js] == "HELD-OUT" for r in sub)
            n_in = sum(str(r[js]).startswith("IN-SCORER") for r in sub)
            L.chk(n_held == N_HELD, f"C {arm}: {n_held} held-out Mic.13 markers")
            L.chk(n_in == N_IN, f"C {arm}: {n_in} Mic.13 markers in the scoring panel")
    for name, h, rows in built:
        neo = [r[1:] for r in rows if r[0] == "NeoC"]
        non = [r[1:] for r in rows if r[0] == "NonNeoC"]
        L.chk(neo != non, f"{name}: NeoC and NonNeoC rows are not identical")

    L.chk(len(built) == 3, f"{len(built)} of 3 sheets built")
    L.say("")
    L.say("Sheets")
    for name, h, rows in built:
        L.say(f"  {name:<22} {len(rows):>4} rows x {len(h):>2} columns")
    if L.fails == 0:
        write_workbook(out_path, built)
        L.say(f"  saved {out_path}")
    else:
        L.say("  workbook not written (checks failed)")
    L.write(os.path.join(out_dir, LOG_NAME))
    sys.exit(0 if L.fails == 0 else 1)


if __name__ == "__main__":
    main()
