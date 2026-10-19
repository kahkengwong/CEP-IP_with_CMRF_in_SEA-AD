#!/usr/bin/env python3
"""
build_SuppTable_S5.py

Builds the 16 data sheets of Supplementary Table S5 (pre-checks of the immune
object) from the outputs of the Part 1 scripts in Part1_PreChecks/. Thirteen
sheets are columns of a Part 1 CSV, in the published column order and in the
row order of the CSV. Three changes are made to the CSV tables: F_per_library
keeps the libraries with at least 40 analyzed nuclei (the libraries that carry
a per-library rho) and replaces the library names with the codes L001, L002,
... in the order of the CSV; K_supertype_specificity adds the depth-matched
columns of S1e_supertype_depth_matched.csv to S1e_per_supertype.csv by
supertype; O_donor_module_rho takes four columns of S1f_donor_table.csv.
Sheets G_ambient_mediation, H_ambient_index_quartiles and J_region_LORO are
read from S1e_techqc_report.txt, because Part 1e prints these values and does
not save them to a CSV; they have the precision printed there. Numbers are
written with 15 significant digits (whole numbers as integers) and logical
values as the text TRUE or FALSE, as in the published table.
finalize_supp_tables.py adds the Contents sheet, applies the published
wording of two cells and saves the published file.

Usage:
    python build_SuppTable_S5.py
    python build_SuppTable_S5.py --root /data/SEA-AD
Root: --root, else $SEAAD_ROOT, else the current folder.

Inputs (all in <root>/Part1_PreChecks/):
    S1b_mito_gene_inventory.csv      Part1b_MT-audit_code.py       -> A
    S1g_trend_fits.csv               Part1g_TrendReconcile_code.py -> B
    S1c_within_donor_rho.csv         Part1c_WithinDonor_code.py    -> C
    S1d_errata_specimen_audit.csv    Part1d_Errata_code.py         -> D
    S1e_matched_panel_null.csv       Part1e_TechQC_code.py         -> E
    S1e_per_library.csv              Part1e_TechQC_code.py         -> F
    S1e_techqc_report.txt            Part1e_TechQC_code.py         -> G, H, J
    S1e_per_region.csv               Part1e_TechQC_code.py         -> I
    S1e_per_supertype.csv and
    S1e_supertype_depth_matched.csv  Part1e_TechQC_code.py         -> K
    S1e_segmented_screen.csv         Part1e_TechQC_code.py         -> L
    S1e_median_specialness.csv       Part1e_TechQC_code.py         -> M
    S1f_ordinal_levels.csv           Part1f_Ordinals_code.py       -> N
    S1f_donor_table.csv              Part1f_Ordinals_code.py       -> O
    S1f_shape_tests.csv              Part1f_Ordinals_code.py       -> P
Output:
    <root>/Manuscript/Supp_Tables/build/SuppTable_S5_PreChecks.xlsx
        sheets A_mito_gene_inventory, B_trend_fits, C_within_donor_rho,
        D_errata_specimen_audit, E_matched_panel_null, F_per_library,
        G_ambient_mediation, H_ambient_index_quartiles, I_per_region,
        J_region_LORO, K_supertype_specificity, L_segmented_screen,
        M_median_specialness, N_ordinal_levels, O_donor_module_rho,
        P_shape_tests
    <root>/Manuscript/Supp_Tables/build/build_SuppTable_S5_log.txt
        checks and row counts, ending with VERDICT: PASS or VERDICT: FAIL.
        On FAIL no workbook is written.

Formatting: Calibri 11 (openpyxl default), bold header row, frozen top row,
no fills, no borders.
"""

import argparse
import datetime as dt
import math
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

MIN_LIB_CELLS = 40          # MIN_CELLS_GROUP of Part1e_TechQC_code.py
ARMS = ("NeoC", "NonNeoC")

# Published column order of each CSV-based sheet: (sheet, CSV, columns).
CSV_SHEETS = [
    ("A_mito_gene_inventory", "S1b_mito_gene_inventory.csv",
     ["gene", "in_module1", "is_mito", "detect_pct", "detect_pct_analysed", "mean_counts",
      "mean_counts_analysed", "ratio_analysed_over_all", "max_counts", "mean_norm", "max_norm",
      "rho_with_GOI", "pred_rho_analysed", "resid_analysed"]),
    ("B_trend_fits", "S1g_trend_fits.csv",
     ["abundance_axis", "slope", "intercept", "r2", "xr", "yr", "pred_ref", "resid", "worst",
      "worst_gene", "ratio", "n_genes"]),
    ("C_within_donor_rho", "S1c_within_donor_rho.csv",
     ["donor", "n_cells", "rho_within", "rho_within_depth"]),
    ("D_errata_specimen_audit", "S1d_errata_specimen_audit.csv",
     ["specimen", "fix", "n_nuclei", "donor_state", "tissue_state"]),
    ("E_matched_panel_null", "S1e_matched_panel_null.csv", ["null_rho"]),
    ("I_per_region", "S1e_per_region.csv",
     ["region", "arm", "n_analysed", "rho", "p", "n_donors", "n_libs", "median_depth",
      "median_mito"]),
    ("L_segmented_screen", "S1e_segmented_screen.csv",
     ["supertype", "n", "breakpoint", "rss_drop", "rss_drop_vs_quad", "slope_lo", "slope_hi",
      "slope_ratio"]),
    ("M_median_specialness", "S1e_median_specialness.csv",
     ["supertype", "n", "median_GOI", "q_preferred", "x_preferred", "rank_of_median",
      "n_candidates", "n_offered", "q_eval_lo", "q_eval_hi", "rss_penalty",
      "rss_penalty_central", "rank_at_u_NeoC", "rank_at_u_NonNeoC", "penalty_at_u_NeoC",
      "penalty_at_u_NonNeoC", "dslope_at_median"]),
    ("N_ordinal_levels", "S1f_ordinal_levels.csv",
     ["ordinal", "column", "level", "naive_digit_parse", "declared_map", "is_categorical"]),
    ("O_donor_module_rho", "S1f_donor_table.csv",
     ["donor", "n_nuclei", "module_n_libs", "module_rho"]),
    ("P_shape_tests", "S1f_shape_tests.csv",
     ["response", "axis", "n", "rho", "p", "r2_linear", "r2_step", "r2_gain", "best_cut",
      "step_delta", "n_below", "n_above", "rho_adjusted", "p_adjusted", "n_adjusted",
      "adjusted_for"]),
]
F_COLS = ["n_analysed", "rho", "amb_index", "median_depth", "median_mito"]
K_COLS_ST = ["supertype", "n_total", "n_goi_pos", "rho", "p", "goi_detect_pct",
             "median_depth", "median_mito"]
K_COLS_DM = ["strata_used", "rho_depth_matched", "min_stratum_n", "shift"]

SHEET_ORDER = ["A_mito_gene_inventory", "B_trend_fits", "C_within_donor_rho",
               "D_errata_specimen_audit", "E_matched_panel_null", "F_per_library",
               "G_ambient_mediation", "H_ambient_index_quartiles", "I_per_region",
               "J_region_LORO", "K_supertype_specificity", "L_segmented_screen",
               "M_median_specialness", "N_ordinal_levels", "O_donor_module_rho",
               "P_shape_tests"]

# Report lines of the library-level ambient test (Part 1e, F3): report key,
# test, adjustment.
MEDIATION = [
    ("ambient index vs library mito%", "Ambient index vs library mitochondrial %", None),
    ("ambient index vs library depth", "Ambient index vs library median depth", None),
    ("library mito% vs module strength", "Library mitochondrial % vs module rho", None),
    ("library depth vs module strength", "Library median depth vs module rho", None),
    ("ambient -> module | mito%", "Ambient index vs module rho", "mitochondrial %"),
    ("ambient -> module | depth", "Ambient index vs module rho", "median depth"),
    ("ambient -> module | mito% + depth", "Ambient index vs module rho",
     "mitochondrial % and median depth"),
    ("mito%   -> module | ambient", "Library mitochondrial % vs module rho", "ambient index"),
]
NUM = r"([+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)"


##### Log #####
class Log:
    def __init__(self):
        self.lines, self.fails = [], 0

    def say(self, s=""):
        self.lines.append(str(s))
        print(s, flush=True)

    def chk(self, ok, what):
        self.say(f"  [{'PASS' if ok else 'FAIL'}] {what}")
        if not ok:
            self.fails += 1
        return bool(ok)

    def write(self, path):
        self.say("")
        self.say(f"VERDICT: {'PASS' if self.fails == 0 else 'FAIL'}")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.lines) + "\n")


##### Helpers #####
def cell_value(v):
    """Missing values blank, numbers with 15 significant digits (whole numbers
    as integers), logical values as the text TRUE or FALSE."""
    if v is None:
        return None
    if isinstance(v, (bool, np.bool_)):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, np.integer)):
        return int(v)
    if isinstance(v, (float, np.floating)):
        v = float(v)
        if not math.isfinite(v):
            return None
        v = float(f"{v:.15g}")
        return int(v) if v.is_integer() and abs(v) < 1e15 else v
    if isinstance(v, str):
        return v
    if pd.isna(v):
        return None
    return v


def read_csv(path):
    # Identifier columns stay text; everything else is parsed by pandas.
    return pd.read_csv(path, dtype={"donor": str, "specimen": str, "library": str})


def select(df, cols, log, label):
    missing = [c for c in cols if c not in df.columns]
    log.chk(not missing, f"{label}: published columns present in the source"
            + (f" (missing {missing})" if missing else ""))
    return df[[c for c in cols if c in df.columns]].reset_index(drop=True)


##### Report parsing (S1e_techqc_report.txt) #####
def parse_report(txt, log):
    def one(pattern, what):
        m = re.findall(pattern, txt, flags=re.M)
        log.chk(len(m) == 1, f"report: one line for {what} (found {len(m)})")
        return m[0] if m else None

    # G: library-level ambient test and its mediation (Part 1e, F3)
    med = []
    raw = one(r"rho\( intronic:exonic , module strength \) = " + NUM + r"\s+p=" + NUM
              + r"\s+n=(\d+)", "the ambient index vs module rho test")
    n_lib = int(raw[2]) if raw else None
    if raw:
        med.append({"test": "Ambient index vs module rho", "partialled_on": None,
                    "rho": float(raw[0]), "p": float(raw[1]), "n_libraries": n_lib})
    for key, test, adj in MEDIATION:
        m = one(r"^\s+" + re.escape(key) + r"\s+" + NUM + r"\s+p=" + NUM + r"\s*$", key)
        if m:
            med.append({"test": test, "partialled_on": adj, "rho": float(m[0]),
                        "p": float(m[1]), "n_libraries": n_lib})
    med = pd.DataFrame(med, columns=["test", "partialled_on", "rho", "p", "n_libraries"])

    # H: module rho by quartile of the nucleus-level ambient index (F2b)
    q = re.findall(r"^\s+(Q[1-4])\s+([\d,]+)\s+" + NUM + r"\s+" + NUM + r"\s+\[" + NUM
                   + r", " + NUM + r"\]\s*$", txt, flags=re.M)
    log.chk(len(q) == 4 and [r[0] for r in q] == ["Q1", "Q2", "Q3", "Q4"],
            f"report: four ambient index quartile rows, Q1 to Q4 (found {len(q)})")
    quart = pd.DataFrame([{"quartile": a, "n_nuclei": int(b.replace(",", "")),
                           "median_ambient_index_pct": float(c), "rho": float(d),
                           "ci95_lo": float(e), "ci95_hi": float(f)}
                          for a, b, c, d, e, f in q],
                         columns=["quartile", "n_nuclei", "median_ambient_index_pct", "rho",
                                  "ci95_lo", "ci95_hi"])

    # J: leave-one-region-out per region set (F10). Each block runs from the
    # "<arm>: pooled rho" line to the next blank line.
    loro = []
    for arm in ARMS:
        hdr = re.findall(r"^  " + arm + r": pooled rho " + NUM + r"\s+n=([\d,]+)\s*\n((?:[ \t]+\S.*\n)+)",
                         txt, flags=re.M)
        log.chk(len(hdr) == 1, f"report: one leave-one-region-out block for {arm} (found {len(hdr)})")
        if not hdr:
            continue
        pooled, _, block = hdr[0]
        for reg, rho, n, share in re.findall(
                r"without (\S+)\s+rho=" + NUM + r"\s+n=\s*([\d,]+)\s+\(that region was\s+"
                + NUM + r"% of the arm\)", block):
            loro.append({"arm": arm, "region_removed": reg,
                         "region_share_of_arm_pct": float(share),
                         "n_remaining": int(n.replace(",", "")),
                         "rho_without_region": float(rho),
                         "rho_arm_pooled": float(pooled),
                         "change": round(float(rho) - float(pooled), 4)})
    loro = pd.DataFrame(loro, columns=["arm", "region_removed", "region_share_of_arm_pct",
                                       "n_remaining", "rho_without_region", "rho_arm_pooled",
                                       "change"])

    libs = one(r"libraries with >=(\d+) analysed cells: (\d+) of (\d+)",
               "the number of libraries with a per-library rho")
    return med, quart, loro, libs


##### Main #####
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[1])
    ap.add_argument("--root", default=os.environ.get("SEAAD_ROOT", os.getcwd()))
    args = ap.parse_args()
    root = os.path.abspath(args.root)
    p1 = os.path.join(root, "Part1_PreChecks")
    out_dir = os.path.join(root, "Manuscript", "Supp_Tables", "build")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "SuppTable_S5_PreChecks.xlsx")
    log_path = os.path.join(out_dir, "build_SuppTable_S5_log.txt")
    if os.path.exists(out):
        os.remove(out)

    log = Log()
    log.say("build_SuppTable_S5.py")
    log.say(f"started {dt.datetime.now():%Y-%m-%d %H:%M:%S}")
    log.say(f"root   : {root}")
    log.say(f"output : {out}")
    log.say("")

    need = sorted({f for _, f, _ in CSV_SHEETS} | {
        "S1e_per_library.csv", "S1e_per_supertype.csv", "S1e_supertype_depth_matched.csv",
        "S1e_techqc_report.txt"})
    log.say("Inputs")
    missing = [f for f in need if not os.path.isfile(os.path.join(p1, f))]
    for f in need:
        log.say(f"  {'found  ' if f not in missing else 'MISSING'}  Part1_PreChecks/{f}")
    if missing:
        log.chk(False, f"{len(missing)} input(s) missing")
        log.write(log_path)
        sys.exit(1)

    sheets = {}

    # Sheets copied column by column from one CSV
    log.say("")
    log.say("CSV sheets")
    src = {}
    for name, fname, cols in CSV_SHEETS:
        df = read_csv(os.path.join(p1, fname))
        src[name] = df
        sheets[name] = select(df, cols, log, f"{name} <- {fname}")
        log.say(f"    {len(sheets[name])} rows, {len(cols)} of {df.shape[1]} source columns")

    # F: libraries with a per-library rho, coded L001, L002, ...
    log.say("")
    log.say("F_per_library <- S1e_per_library.csv")
    lib = read_csv(os.path.join(p1, "S1e_per_library.csv"))
    keep = lib["n_analysed"] >= MIN_LIB_CELLS
    log.chk(bool((keep == lib["rho"].notna()).all()),
            f"the libraries with >={MIN_LIB_CELLS} analyzed nuclei are those with a rho")
    f = select(lib.loc[keep], F_COLS, log, "F_per_library")
    f.insert(0, "lib_index", [f"L{i:03d}" for i in range(1, len(f) + 1)])
    sheets["F_per_library"] = f
    log.say(f"    {len(f)} of {len(lib)} libraries kept, coded L001 to L{len(f):03d}")

    # K: per-supertype rho with the depth-matched columns added by supertype
    log.say("")
    log.say("K_supertype_specificity <- S1e_per_supertype.csv + S1e_supertype_depth_matched.csv")
    st = read_csv(os.path.join(p1, "S1e_per_supertype.csv"))
    dm = read_csv(os.path.join(p1, "S1e_supertype_depth_matched.csv"))
    log.chk(dm["supertype"].is_unique and set(dm["supertype"]) <= set(st["supertype"]),
            "depth-matched supertypes are unique and all in S1e_per_supertype.csv")
    both = st.merge(dm, on="supertype", how="left", suffixes=("", "_dm"))
    chk = both.dropna(subset=["rho_raw"])
    log.chk(bool(np.allclose(chk["rho"], chk["rho_raw"], rtol=0, atol=1e-12)),
            "rho_raw of the depth-matched file equals rho of the per-supertype file")
    sheets["K_supertype_specificity"] = select(both, K_COLS_ST + K_COLS_DM, log,
                                               "K_supertype_specificity")
    no_dm = both.loc[both["rho_depth_matched"].isna(), "supertype"].tolist()
    log.say(f"    {len(both)} supertypes; without a depth-matched estimate: {no_dm}")

    # G, H, J: from the Part 1e report
    log.say("")
    log.say("G_ambient_mediation, H_ambient_index_quartiles, J_region_LORO <- S1e_techqc_report.txt")
    with open(os.path.join(p1, "S1e_techqc_report.txt"), encoding="utf-8") as fh:
        txt = fh.read()
    med, quart, loro, libs = parse_report(txt, log)
    sheets["G_ambient_mediation"] = med
    sheets["H_ambient_index_quartiles"] = quart
    sheets["J_region_LORO"] = loro
    log.chk(len(med) == 1 + len(MEDIATION), f"G: {len(med)} rows (expected {1 + len(MEDIATION)})")
    if libs:
        log.chk(int(libs[0]) == MIN_LIB_CELLS and int(libs[1]) == len(f)
                and int(libs[2]) == len(lib),
                f"report: {libs[1]} of {libs[2]} libraries with >={libs[0]} analysed cells, "
                f"as in sheet F ({len(f)} of {len(lib)})")
    log.chk(med["n_libraries"].nunique() == 1 and int(med["n_libraries"].iloc[0]) == len(f),
            "G: n_libraries equals the number of libraries in sheet F")
    reg = sheets["I_per_region"]
    for arm in ARMS:
        a = loro[loro["arm"] == arm]
        want = set(reg.loc[reg["arm"] == arm, "region"])
        log.chk(set(a["region_removed"]) == want and len(a) == len(want),
                f"J: {arm} leaves out each of its {len(want)} regions once")
        n_arm = int(reg.loc[reg["arm"] == arm, "n_analysed"].sum())
        r_n = reg.set_index("region")["n_analysed"]
        ok = all(int(r.n_remaining) == n_arm - int(r_n[r.region_removed]) for r in a.itertuples())
        log.chk(ok, f"J: {arm} n_remaining equals the arm's analyzed nuclei ({n_arm:,}) "
                    f"minus those of the region left out (sheet I)")
    q_n = int(quart["n_nuclei"].sum())
    log.chk(q_n == int(reg["n_analysed"].sum()),
            f"H: the quartiles hold all {q_n:,} analyzed nuclei (sheet I total "
            f"{int(reg['n_analysed'].sum()):,})")

    # Consistency checks between sheets
    log.say("")
    log.say("Cross-sheet checks")
    o = sheets["O_donor_module_rho"]
    with_rho = int(o["module_rho"].notna().sum())
    log.chk(int(o["module_n_libs"].fillna(0).sum()) == len(f),
            f"O: module_n_libs sums to the {len(f)} libraries of sheet F")
    log.say(f"    O: {len(o)} donors, {with_rho} with a module rho, {len(o) - with_rho} blank")
    a = sheets["A_mito_gene_inventory"]
    b = sheets["B_trend_fits"]
    actb = a[a["gene"] == "ACTB"]
    if len(actb) == 1 and "analysed" in set(b["abundance_axis"]):
        ref = b[b["abundance_axis"] == "analysed"].iloc[0]
        log.chk(abs(float(actb["pred_rho_analysed"].iloc[0]) - float(ref["pred_ref"])) < 1e-9,
                "A and B agree on the predicted ACTB rho on the analyzed axis")
    log.chk(list(b["abundance_axis"]) == ["all_PVM3", "analysed"],
            "B: rows all_PVM3 and analysed (finalize writes 'analyzed')")

    # Write
    log.say("")
    log.say("Sheets")
    for name in SHEET_ORDER:
        df = sheets[name]
        has_dash = any(isinstance(v, str) and (chr(0x2013) in v or chr(0x2014) in v)
                       for v in list(df.columns) + df.astype(object).values.ravel().tolist())
        log.chk(not has_dash, f"{name}: {len(df)} rows x {df.shape[1]} columns, no en or em dash")
    if log.fails:
        log.say("")
        log.say("Checks failed; no workbook written.")
        log.write(log_path)
        sys.exit(1)

    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name in SHEET_ORDER:
        df = sheets[name]
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = bold
        for rec in df.astype(object).itertuples(index=False):
            ws.append([cell_value(v) for v in rec])
        ws.freeze_panes = "A2"
    wb.save(out)
    log.say("")
    log.say(f"saved {out}")
    log.write(log_path)


if __name__ == "__main__":
    main()
