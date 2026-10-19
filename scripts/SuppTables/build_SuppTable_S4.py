#!/usr/bin/env python3
"""
build_SuppTable_S4.py

Builds the data sheets of Supplementary Table S4 (GOI-DFG screens) from the
Part 6 and Part 6b outputs of both region sets. Sheets A to C cover the
regional screens of M4-selected cells (Part 6): the population tally, the DFG
counts per population and the 50 highest-ranked genes over all populations,
each copied as written by Part 6 with a leading arm column. Sheet D summarises
the pooled screen of each supertype without the M4 selection (Part 6b): the
number of genes with DFG partners, the top gene with its DFG count, max rs and
max tau, and the rank and DFG count of MT-CO3, taken from each population's
MasterDFG ranking. Sheet E is the Part 6b population tally with the column
cells_screened (the smaller of subsample_n and n_cells_after_donor_filter)
added after subsample_n. finalize_supp_tables.py adds the Contents sheet and
saves the published file.

Usage:
    python build_SuppTable_S4.py
    python build_SuppTable_S4.py --root /data/SEA-AD
Root: --root, else $SEAAD_ROOT, else the current folder.

Inputs (<arm> is NeoC in [1]_NeoC and NonNeoC in [2]_Non-NeoC):
    <root>/<arm folder>/Part6_<arm>_GOI-DFG/Part6_population_tally.csv
    <root>/<arm folder>/Part6_<arm>_GOI-DFG/Part6_cross_population_summary.xlsx
        sheets 01_dfg_counts and 02_top_GOIs_all_populations
        (both written by scripts/Python/Part6_<arm>_GOI-DFG_code.py)
    <root>/<arm folder>/Part6b_<arm>_Pooled_GOI-DFG/Part6b_population_tally.csv
    <root>/<arm folder>/Part6b_<arm>_Pooled_GOI-DFG/Part6b_cross_population_summary.xlsx
        sheet 01_dfg_counts (used to check sheet D)
    <root>/<arm folder>/Part6b_<arm>_Pooled_GOI-DFG/<population>/<population>_MasterDFG_Ranking.xlsx
        sheet Master_DFG_Ranking; Part 6b writes no ranking for a population
        without DFG hits
        (all three written by scripts/Python/Part6b_<arm>_Pooled_GOI-DFG.py)
Output:
    <root>/Manuscript/Supp_Tables/build/SuppTable_S4_GOI_DFG_screen.xlsx
        sheets A_population_tally, B_dfg_counts, C_top_GOIs,
        D_pooled_cross_population, E_pooled_tally
    <root>/Manuscript/Supp_Tables/build/build_SuppTable_S4_log.txt
        checks and row counts, ending with VERDICT: PASS or VERDICT: FAIL.
        On FAIL no workbook is written.

Formatting: bold header row, frozen top row, no fills, no borders.
"""

import argparse
import math
import os
import sys

import pandas as pd
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ARMS = {"NeoC": "[1]_NeoC", "NonNeoC": "[2]_Non-NeoC"}
GOI = "MT-CO3"
OUT_NAME = "SuppTable_S4_GOI_DFG_screen.xlsx"
LOG_NAME = "build_SuppTable_S4_log.txt"
D_COLS = ["arm", "Population", "GOIs with DFGs", "Top GOI", "Top DFG count", "Top max Sp",
          "Top max Kd", "MT-CO3 rank", "MT-CO3 DFGs", "Status"]


class Log:
    def __init__(self):
        self.lines, self.fails = [], 0

    def say(self, s=""):
        print(s)
        self.lines.append(str(s))

    def check(self, ok, what):
        self.say(f"  [{'PASS' if ok else 'FAIL'}] {what}")
        if not ok:
            self.fails += 1
        return ok

    def write(self, path):
        self.say("")
        self.say(f"VERDICT: {'PASS' if self.fails == 0 else 'FAIL'}")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(self.lines) + "\n")


def p6(root, arm, name):
    return os.path.join(root, ARMS[arm], f"Part6_{arm}_GOI-DFG", name)


def p6b(root, arm, *name):
    return os.path.join(root, ARMS[arm], f"Part6b_{arm}_Pooled_GOI-DFG", *name)


def py(v):
    """numpy scalar to Python; NaN to None."""
    if v is None or v is pd.NA:
        return None
    if hasattr(v, "item"):
        v = v.item()
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def read_csv(path, arm):
    d = pd.read_csv(path)
    d.insert(0, "arm", arm)
    return d.columns.tolist(), [[py(v) for v in r] for r in d.itertuples(index=False)]


def read_sheet(path, sheet, arm):
    """A sheet as written by the pipeline: header row, then data rows; cell
    values are kept as stored (no type conversion)."""
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    while hdr and hdr[-1] is None:
        hdr.pop()
    data = [list(r[:len(hdr)]) for r in rows[1:] if any(v is not None for v in r[:len(hdr)])]
    return ["arm"] + hdr, [[arm] + r for r in data]


def pooled_summary(root, arm, tally):
    """Sheet D rows of one arm, one per Part 6b population in tally order."""
    out = []
    for rec in tally.itertuples(index=False):
        pop = rec.population
        path = p6b(root, arm, pop, f"{pop}_MasterDFG_Ranking.xlsx")
        if rec.status != "OK":
            out.append([arm, pop, 0, "-", 0, 0, 0, "-", 0, "skipped"])
            continue
        if not os.path.isfile(path):
            # Part 6b writes no ranking when no GOI has a DFG partner
            out.append([arm, pop, 0, "-", 0, 0, 0, "-", 0, "completed_zero_hits"])
            continue
        m = pd.read_excel(path, sheet_name="Master_DFG_Ranking")
        m = m[m["DFG_Total"] >= 1].sort_values("Rank").reset_index(drop=True)
        if m.empty:
            out.append([arm, pop, 0, "-", 0, 0, 0, "-", 0, "completed_zero_hits"])
            continue
        t = m.iloc[0]
        co3 = m[m["GOI"] == GOI]
        out.append([arm, pop, len(m), t["GOI"], py(t["DFG_Total"]), py(t["max_Spearman"]),
                    py(t["max_Kendall"]),
                    py(co3["Rank"].iloc[0]) if len(co3) else "absent",
                    py(co3["DFG_Total"].iloc[0]) if len(co3) else 0, "completed"])
    return out


def add_sheet(wb, name, header, rows):
    bold = Font(bold=True)
    w = wb.create_sheet(name)
    w.append(header)
    for c in w[1]:
        c.font = bold
    for r in rows:
        w.append(r)
    for j, col in enumerate(header, 1):
        w.column_dimensions[get_column_letter(j)].width = max(10, min(32, len(str(col)) + 2))
    w.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--root", default=os.environ.get("SEAAD_ROOT", os.getcwd()))
    root = os.path.abspath(ap.parse_args().root)
    out_dir = os.path.join(root, "Manuscript", "Supp_Tables", "build")
    os.makedirs(out_dir, exist_ok=True)
    out_path, log_path = os.path.join(out_dir, OUT_NAME), os.path.join(out_dir, LOG_NAME)
    L = Log()
    L.say("build_SuppTable_S4.py")
    L.say(f"root: {root}")

    need = []
    for arm in ARMS:
        need += [p6(root, arm, "Part6_population_tally.csv"),
                 p6(root, arm, "Part6_cross_population_summary.xlsx"),
                 p6b(root, arm, "Part6b_population_tally.csv"),
                 p6b(root, arm, "Part6b_cross_population_summary.xlsx")]
    for v in need:
        L.say(f"  {'found  ' if os.path.isfile(v) else 'MISSING'}  {os.path.relpath(v, root)}")
    if not L.check(all(os.path.isfile(v) for v in need), "all inputs present"):
        L.write(log_path)
        sys.exit(1)

    sheets = {k: (None, []) for k in ("A", "B", "C", "D", "E")}

    def extend(key, header, rows):
        h0, r0 = sheets[key]
        if h0 is not None and h0 != header:
            L.check(False, f"sheet {key}: the two arms have different columns ({h0} / {header})")
        sheets[key] = (header, r0 + rows)

    for arm in ARMS:
        L.say("")
        L.say(f"{arm}")
        hA, rA = read_csv(p6(root, arm, "Part6_population_tally.csv"), arm)
        hB, rB = read_sheet(p6(root, arm, "Part6_cross_population_summary.xlsx"), "01_dfg_counts", arm)
        hC, rC = read_sheet(p6(root, arm, "Part6_cross_population_summary.xlsx"),
                            "02_top_GOIs_all_populations", arm)
        L.check([r[1] for r in rA] == [r[1] for r in rB],
                "Part 6 tally and 01_dfg_counts list the same populations in the same order")
        L.check(len(rC) <= 50, f"02_top_GOIs_all_populations holds {len(rC)} rows (at most 50)")

        tally6b = pd.read_csv(p6b(root, arm, "Part6b_population_tally.csv"))
        hE = ["arm"] + tally6b.columns.tolist()
        k = hE.index("subsample_n") + 1
        hE.insert(k, "cells_screened")
        rE = []
        for rec in tally6b.itertuples(index=False):
            r = [arm] + [py(v) for v in rec]
            r.insert(k, min(int(rec.subsample_n), int(rec.n_cells_after_donor_filter)))
            rE.append(r)

        rD = pooled_summary(root, arm, tally6b)
        _, c6b = read_sheet(p6b(root, arm, "Part6b_cross_population_summary.xlsx"), "01_dfg_counts", arm)
        ref = {r[1]: r for r in c6b}           # arm, Population, Supertype, Region, screened, with_DFG, max, top
        ok = all(r[1] in ref and r[2] == ref[r[1]][5] and r[4] == ref[r[1]][6] for r in rD)
        L.check(ok, "sheet D: genes with DFGs and top DFG count agree with Part 6b 01_dfg_counts")
        for r in rD:
            L.say(f"    {r[1]:16s} {r[2]:>4} {str(r[3]):10s} {r[4]:>3} {r[5]:>7} {r[6]:>7} "
                  f"{str(r[7]):>7} {r[8]:>3} {r[9]}")

        extend("A", hA, rA)
        extend("B", hB, rB)
        extend("C", hC, rC)
        extend("D", D_COLS, rD)
        extend("E", hE, rE)

    names = {"A": "A_population_tally", "B": "B_dfg_counts", "C": "C_top_GOIs",
             "D": "D_pooled_cross_population", "E": "E_pooled_tally"}
    L.say("")
    for key, name in names.items():
        h, r = sheets[key]
        L.say(f"  {name}: {len(r)} rows x {len(h)} columns")
    if L.fails:
        L.say("")
        L.say("Checks failed; workbook not written.")
        L.write(log_path)
        sys.exit(1)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for key, name in names.items():
        add_sheet(wb, name, *sheets[key])
    wb.save(out_path)
    L.say(f"  saved: {os.path.relpath(out_path, root)}")
    L.write(log_path)


if __name__ == "__main__":
    main()
