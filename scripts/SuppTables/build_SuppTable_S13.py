#!/usr/bin/env python3
"""
build_SuppTable_S13.py

Builds SuppTable_Fig15.xlsx (Figure 15 and Supplementary Figure 11 of the
Part 12 rerun) from the stats and plot-data files written by
Part12_Fig15_Analysis.r (sections 07 to 11) and the figure log written by
Part12_Fig15_Plots.r (section 12). finalize_supp_tables.py turns this workbook
into Supplementary Table S13.

Usage:
    python build_SuppTable_S13.py
    python build_SuppTable_S13.py --root /data/Part12_rerun

--root defaults to $SEAAD_ROOT/Part12_rerun when SEAAD_ROOT is set, otherwise
to Part12_rerun in the current folder. --out defaults to
<root>/07_Fig15/SuppTable_Fig15.xlsx.

--root is searched recursively; every input file must occur exactly once under
it, so the script works on the Part12_rerun folder tree or on a flat folder
holding copies of the files. Besides the Figure 15 outputs and logs it reads
the config (Part12_01_config.json), the donor table (05_donors.csv), the cell
labels (04_cell_labels.csv.gz) and the logs of steps 04 and 05, for the
Module post-IP% columns of the 15A and 15B sheets (readouts, Fig15B_removed and
the two stats sheets). Each value is the median Module post-IP% over the donors
of that row.

The script computes no new statistic. Before it writes anything it checks:
  B-01  every section log ends with VERDICT: PASS and 0 FAIL
  B-02  every input file is listed in the outputs of the log of the section that wrote it
  B-03  the five analysis logs name one script, one config version and one labels version
  B-04  every stats row carries the run ID, config version and labels version of its log
  B-05  stats row counts equal the counts the logs report
  B-06  every fdr equals BH recomputed within its family
  B-07  15A retention equals dR2 adjusted / dR2 on the same donors
  B-08  15D plot data equal the tertile rs and the T3 - T1 estimate in Stats_15D
  B-09  15E median of the per-sample log ratios equals the Stats_15E estimate
  B-10  15F per-gene and primary medians and the quartile row equal the plot data
  B-11  15G per-gene and primary estimates equal the plot data
  B-12  the figure log is PASS and was started after the last analysis section
  B-13  no text written to the workbook contains an en dash or an em dash
  B-14  the config exists and has the config version of the section logs; the donor
        table and the cell labels are listed in the outputs of their logs, which are
        PASS, and carry the labels version of the section logs
  B-15  the donor sets rebuilt for the post-IP% columns have the n_donors of every
        15A and 15B row (unadjusted, same-donor, adjusted and doublet models, and the
        removed-share units), and the removed shares equal Stats_15B
  B-16  every 15A and 15B model refitted from the donor values of Fig15AB_donors
        gives the n_donors, beta and dR2 of its stats row (difference below 1e-8)
  B-17  Fig15AB_donors holds every donor of the arm and region-set units
Any FAIL stops the build and no workbook is written. A build log is written
beside the workbook.

Formatting: Arial throughout, bold header row, frozen top row. No fills, no
borders, no colours. The status and retention columns of the 15A and 15B
readout sheets are Excel formulas; the file is flagged to recalculate on open.
"""

import argparse
import datetime as dt
import json
import os
import pathlib
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Inputs
# ---------------------------------------------------------------------------
SECTIONS = [
    # (step, folder, files written by that section)
    ("07", "15AB_ContamDoublets", ["Stats_15A.csv", "Stats_15B.csv"]),
    ("08", "15CD_DecontXStrata", ["Stats_15C.csv", "Stats_15D.csv", "Plot_15D_tertiles.csv"]),
    ("09", "15EG_GerritsAmbient", ["Stats_15E.csv", "Stats_15F.csv", "Stats_15G.csv",
                                   "Plot_15E_samples.csv", "Plot_15F_samples.csv",
                                   "Plot_15G_tertiles.csv"]),
    ("10", "15H_DepthThinning", ["Stats_15H.csv"]),
    ("11", "S10_Cargo", ["Stats_S10.csv"]),
]
FIG_LOG = "Part12_Fig15_log.txt"
CONFIG = "Part12_01_config.json"
UPSTREAM = [("05_donors.csv", "*05_Units*_log.txt"), ("04_cell_labels.csv.gz", "*04_Labels*_log.txt")]
ARMS = ["M4", "PVM_3", "PVM_1", "HKG8"]
REGION_SETS = ["NeoC", "NonNeoC"]
PROV_COLS = ["run_id", "config_version", "labels_version"]
TOL = 1e-9
DASHES = ("\u2013", "\u2014")

FONT = Font(name="Arial", size=10)
FONT_B = Font(name="Arial", size=10, bold=True)
FMT_P = "0.00E+00"


# ---------------------------------------------------------------------------
# Check ledger
# ---------------------------------------------------------------------------
class Ledger:
    def __init__(self):
        self.lines, self.n = [], {"PASS": 0, "FAIL": 0, "NOTE": 0}

    def chk(self, cid, ok, what, detail=""):
        tag = "PASS" if ok else "FAIL"
        self.n[tag] += 1
        self._say(f"[{tag}] {cid} {what} | {detail}")

    def note(self, what, detail=""):
        self.n["NOTE"] += 1
        self._say(f"[NOTE] {what} | {detail}")

    def _say(self, s):
        self.lines.append(s)
        print(s)


def find(root, name):
    hits = [p for p in root.rglob(name) if p.is_file()]
    if len(hits) != 1:
        sys.exit(f"ERROR: expected exactly one {name} under {root}, found {len(hits)}")
    return hits[0]


def find_glob(root, pattern):
    hits = [p for p in root.rglob(pattern) if p.is_file()]
    if len(hits) != 1:
        sys.exit(f"ERROR: expected exactly one file matching {pattern} under {root}, found {len(hits)}")
    return hits[0]


def read_csv_like_r(path, char_cols=(), usecols=None):
    """Read a CSV with R read.csv missing-value rules, so complete-case filters
    match the analysis: "NA" is missing in every column; an empty field is
    missing in a numeric column and an empty string in a character column."""
    d = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[], usecols=usecols)
    for c in d.columns:
        v = d[c].where(d[c] != "NA")
        if c in char_cols:
            d[c] = v
            continue
        num = pd.to_numeric(v.where(v != ""), errors="coerce")
        bad = v.notna() & (v != "") & num.isna() & ~v.str.lower().isin(["nan"])
        d[c] = v if bad.any() else num
    return d


def post_ip_donor_sets(DN, LB, cfg, arms, region_sets):
    """Rebuild the donor set and Module post-IP% of every 15A and 15B model as
    section 07 of Part12_Fig15_Analysis.r does, and return the median
    Module post-IP% over each set with its size.

    15A: unadjusted = complete cases on outcome, module_post_pct and cov4 (APOE
    e4 dose dropped when it is the outcome); same donors = donors with a finite
    contam_super_nonmodule, then the unadjusted complete cases; adjusted = the
    unadjusted columns plus contam_super_nonmodule. 15B: Module post-IP% recomputed per donor after
    removing Module post-IP nuclei with doublet score >= cut from numerator and
    denominator, then the unadjusted donor set on that column."""
    cl = cfg["clinical"]
    ## cov4 is either a list of donor-table columns or a named list {name: column};
    ## R's unlist() gives the columns in both cases, so the values are used here.
    covs = list(cl["cov4"].values()) if isinstance(cl["cov4"], dict) else list(cl["cov4"])
    outcomes = {k: v["column"] for k, v in cl["outcomes"].items()}
    cuts = [float(c) for c in cfg["doublet_cuts"]]
    den_col = "n_super" if cfg["quantities"]["module_post_pct_denominator"] == "supertype" else "n_cells"

    def cc(d, cols):
        return d[d[list(dict.fromkeys(cols))].notna().all(axis=1)]

    med = lambda x: float(np.median(x)) if len(x) else np.nan
    A, Bm, share, unit = {}, {}, {}, {}
    DX = DN.copy()  # donor table plus the per-cut removed counts and post-IP%
    for a in arms:
        for ds in region_sets:
            d = DN[(DN["arm"] == a) & (DN["dataset"] == ds)]
            unit[(a, ds)] = med(d["module_post_pct"].dropna().values)
            for k, y in outcomes.items():
                cv = [c for c in covs if c != y]
                u = cc(d, [y, "module_post_pct"] + cv)
                dd = d[np.isfinite(d["contam_super_nonmodule"].astype(float))]
                us = cc(dd, [y, "module_post_pct"] + cv)
                ad = cc(d, [y, "module_post_pct"] + cv + ["contam_super_nonmodule"])
                A[(a, ds, k)] = {"n": len(u), "med": med(u["module_post_pct"].values),
                                 "n_same": len(us), "med_same": med(us["module_post_pct"].values),
                                 "n_adj": len(ad), "med_adj": med(ad["module_post_pct"].values)}
    lb = LB[(LB["group"] == "Module post-IP") & np.isfinite(LB["doublet_score"].astype(float))]
    for cut in cuts:
        tag = f"{cut:.2f}"
        rm = lb[lb["doublet_score"] >= cut].groupby(["arm", "dataset", "donor_key"]).size()
        r = pd.Series([rm.get((a, ds, dk), 0) for a, ds, dk in
                       zip(DN["arm"], DN["dataset"], DN["donor_key"])], index=DN.index, dtype=float)
        with np.errstate(divide="ignore", invalid="ignore"):
            col = 100 * (DN["n_post"] - r) / (DN[den_col] - r)
        col = col.where(~np.isnan(col))
        DX[f"n_removed_{tag}"] = r
        DX[f"post_ip_pct_{tag}"] = col
        nd = DN.assign(module_nodbl=col)
        for a in arms:
            for ds in region_sets:
                i = (nd["arm"] == a) & (nd["dataset"] == ds)
                d = nd[i]
                share[(a, ds, tag)] = (r[i].sum() / max(1, DN.loc[i, "n_post"].sum()), int(i.sum()))
                unit[(a, ds, tag)] = med(d["module_nodbl"].dropna().values)
                for k, y in outcomes.items():
                    cv = [c for c in covs if c != y]
                    u = cc(d, [y, "module_nodbl"] + cv)
                    Bm[(a, ds, tag, k)] = {"n": len(u), "med": med(u["module_nodbl"].values)}
    return A, Bm, share, unit, DX


def ols_refit(d, y, x, covs, min_n, extra=()):
    """The section 07 model in numpy: complete cases on y, x, covariates and extra;
    outcome and x standardised (sd with n - 1, as R scale()); covariates entered
    as they are, a text covariate as treatment dummies. Returns n, the beta of x
    and dR2 = R2(covariates + x) - R2(covariates). The p value and interval need
    the t distribution and are not recomputed here."""
    cols = list(dict.fromkeys([y, x] + list(covs) + list(extra)))
    d = d[d[cols].notna().all(axis=1)]
    n = len(d)
    if n < min_n:
        return n, np.nan, np.nan
    yv, xv = d[y].astype(float).values, d[x].astype(float).values
    if not (np.std(yv, ddof=1) > 0 and np.std(xv, ddof=1) > 0):
        return n, np.nan, np.nan
    yz = (yv - yv.mean()) / np.std(yv, ddof=1)
    xz = (xv - xv.mean()) / np.std(xv, ddof=1)
    parts = [np.ones((n, 1))]
    for c in list(dict.fromkeys([c for c in list(covs) + list(extra) if c not in (y, x)])):
        v = d[c]
        if pd.api.types.is_numeric_dtype(v):
            parts.append(v.astype(float).values[:, None])
        else:
            parts.append(pd.get_dummies(v.astype(str), drop_first=True).astype(float).values)
    X0 = np.hstack(parts)
    X1 = np.hstack([X0, xz[:, None]])
    tss = np.sum((yz - yz.mean()) ** 2)
    r2 = []
    for X in (X0, X1):
        b, *_ = np.linalg.lstsq(X, yz, rcond=None)
        r2.append(1 - np.sum((yz - X @ b) ** 2) / tss)
        beta = b[-1]
    return n, float(beta), float(r2[1] - r2[0])


def parse_log(path):
    txt = path.read_text(encoding="utf-8", errors="replace")
    g = lambda pat: (re.search(pat, txt, re.M).group(1) if re.search(pat, txt, re.M) else None)
    # Output lines: "  <folder>/<file> | <size> MB"; anything after the size is ignored.
    outs = set(re.findall(r"^\s+\S*/([^/\s|]+) \| [\d.]+ MB", txt, re.M))
    return {
        "path": path,
        "script": g(r"^Script: (\S+)"),
        "run_id": g(r"^Run ID: (\S+)"),
        "start": g(r"^Start: (.+)$"),
        "config_version": g(r"^Config: .* \| version (\S+?)(?: \||\s*$)"),
        "labels_version": g(r"^Labels version: (\S+)"),
        "checks": g(r"^Checks: (.+?) \|"),
        "n_fail": int(g(r"^Checks: \d+ PASS, (\d+) FAIL") or -1),
        "verdict": g(r"^VERDICT: (\S+)"),
        "outputs": outs,
        "rows_line": g(r"stats files? written and read back \| (.+)$"),
        "text": txt,
    }


def bh(p):
    """Benjamini-Hochberg as R p.adjust(method = "BH"), NaN left as NaN."""
    p = np.asarray(p, dtype=float)
    out = np.full(p.shape, np.nan)
    ok = np.isfinite(p)
    m = ok.sum()
    if m == 0:
        return out
    pv = p[ok]
    order = np.argsort(pv)[::-1]
    ranks = np.arange(m, 0, -1)
    q = np.minimum.accumulate(pv[order] * m / ranks)
    res = np.empty(m)
    res[order] = np.minimum(q, 1.0)
    out[ok] = res
    return out


def close(a, b, tol=TOL):
    a, b = np.asarray(a, float), np.asarray(b, float)
    both_nan = np.isnan(a) & np.isnan(b)
    return bool(np.all(both_nan | (np.abs(a - b) <= tol)))


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------
def clean(v):
    """Excel-safe cell value: NaN to empty, numpy scalars to Python."""
    if v is None:
        return None
    if isinstance(v, (float, np.floating)) and not np.isfinite(v):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    return v


def write_table(wb, name, header, rows, widths=None, p_cols=(), fmt_cols=None, wrap=False):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append([clean(v) for v in r])
    for c in ws[1]:
        c.font = FONT_B
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            if wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    fmt_cols = dict(fmt_cols or {})
    for col in p_cols:
        fmt_cols[col] = FMT_P
    for col, fmt in fmt_cols.items():
        j = header.index(col) + 1
        for i in range(2, ws.max_row + 1):
            ws.cell(i, j).number_format = fmt
    for j, h in enumerate(header, start=1):
        if widths and h in widths:
            w = widths[h]
        else:
            vals = [len(str(h))] + [len(str(ws.cell(i, j).value)) for i in range(2, min(ws.max_row, 200) + 1)
                                    if ws.cell(i, j).value is not None
                                    and not str(ws.cell(i, j).value).startswith("=")]
            w = min(max(vals) + 2, 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def stats_sheet(wb, name, df, post_ip=None):
    """A stats file as a sheet, provenance columns left out. post_ip, if given, is
    one value per row, written as post_ip_pct_median after n_donors."""
    cols = [c for c in df.columns if c not in PROV_COLS]
    d = df[cols].copy()
    fmt = {}
    if post_ip is not None:
        d.insert(cols.index("n_donors") + 1, "post_ip_pct_median", list(post_ip))
        fmt = {"post_ip_pct_median": "0.00"}
    write_table(wb, name, list(d.columns), d.itertuples(index=False, name=None),
                p_cols=("p", "fdr"), fmt_cols=fmt)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    default_root = pathlib.Path(os.environ.get("SEAAD_ROOT") or ".") / "Part12_rerun"
    ap.add_argument("--root", default=str(default_root),
                    help="folder holding the Figure 15 outputs and logs (default: $SEAAD_ROOT/Part12_rerun, "
                         "or ./Part12_rerun)")
    ap.add_argument("--out", default=None, help="output xlsx (default: <root>/07_Fig15/SuppTable_Fig15.xlsx)")
    args = ap.parse_args()
    root = pathlib.Path(args.root).resolve()
    out = pathlib.Path(args.out).resolve() if args.out else root / "07_Fig15" / "SuppTable_Fig15.xlsx"
    L = Ledger()
    L._say("build_SuppTable_S13.py")
    L._say(f"Start: {dt.datetime.now():%Y-%m-%d %H:%M:%S} | root {root}")
    L._say("")

    # ---- logs and input files ---------------------------------------------------------
    L._say("SECTION A -- logs and inputs")
    logs, files, D = {}, {}, {}
    for step, folder, names in SECTIONS:
        lg = parse_log(find(root, f"Part12_{folder}_log.txt"))
        logs[step] = lg
        L.chk(f"B-01.{step}", lg["verdict"] == "PASS" and lg["n_fail"] == 0,
              f"section {step} log is PASS with 0 FAIL", f"{lg['checks']} | {lg['verdict']}")
        for nm in names:
            p = find(root, nm)
            files[nm] = (step, p)
            L.chk(f"B-02.{nm}", nm in lg["outputs"],
                  f"{nm} is listed in the outputs of the section {step} log", lg["path"].name)
            D[nm] = pd.read_csv(p)
    one = lambda k: sorted({lg[k] for lg in logs.values()})
    L.chk("B-03", len(one("script")) == 1 and len(one("config_version")) == 1
          and len(one("labels_version")) == 1,
          "one analysis script, config and labels version across sections 07 to 11",
          f"{one('script')} | config {one('config_version')} | labels {one('labels_version')}")

    # ---- row provenance -------------------------------------------------------------
    for nm, (step, _) in files.items():
        if not nm.startswith("Stats_"):
            continue
        d, lg = D[nm], logs[step]
        ok = (set(d["run_id"]) == {lg["run_id"]}
              and set(d["config_version"].astype(str)) == {lg["config_version"]}
              and set(d["labels_version"]) == {lg["labels_version"]})
        L.chk(f"B-04.{nm}", ok, f"{nm} rows carry the log's run ID, config and labels version",
              f"{lg['run_id']} | {lg['config_version']} | {lg['labels_version']}")

    # ---- row counts against the logs ----------------------------------------------------
    for step, folder, names in SECTIONS:
        rl = logs[step]["rows_line"] or ""
        stats = [n for n in names if n.startswith("Stats_")]
        pairs = dict(re.findall(r"(\w+) (\d+) rows", rl))
        if pairs:
            ok = all(int(pairs.get(n[6:-4], -1)) == len(D[n]) for n in stats)
        else:
            m = re.search(r"(\d+) rows", rl)
            ok = bool(m) and int(m.group(1)) == sum(len(D[n]) for n in stats)
        L.chk(f"B-05.{step}", ok, f"section {step} stats row counts equal the log", rl)

    # ---- BH within family -----------------------------------------------------------
    for nm in [n for n in D if n.startswith("Stats_")]:
        d = D[nm]
        ok = True
        for fam, g in d.groupby("family", sort=False):
            if not close(bh(g["p"].values), g["fdr"].values, 1e-12):
                ok = False
        L.chk(f"B-06.{nm}", ok, f"{nm} fdr equals BH within family", f"{d['family'].nunique()} families")
    L._say("")

    # ---- 15A --------------------------------------------------------------------
    L._say("SECTION B -- panel consistency")
    A = D["Stats_15A.csv"]
    key = ["arm", "level", "measure"]
    q = {k: A[A.quantity == k].set_index(key).sort_index() for k in A.quantity.unique()}
    # rows in arm and region-set order, outcomes in the order the stats file lists them
    first = list(A[A.quantity == "module_beta_unadjusted"][key].itertuples(index=False, name=None))
    idx = [k for a in ARMS for lv in REGION_SETS for k in first if k[0] == a and k[1] == lv]
    ret = q["dR2_contam_adjusted"]["estimate"] / q["dR2_unadjusted_same_donors"]["estimate"]
    ret[~(q["dR2_unadjusted_same_donors"]["estimate"] > 0)] = np.nan
    L.chk("B-07", close(ret.values, q["retention"]["estimate"].reindex(ret.index).values),
          "15A retention equals dR2 adjusted / dR2 on the same donors", f"{len(ret)} rows")

    # ---- 15D --------------------------------------------------------------------
    Dd, Dp = D["Stats_15D.csv"], D["Plot_15D_tertiles.csv"]
    ok = True
    for a in ARMS:
        s = Dp[Dp.arm == a].set_index("tertile")["rs"]
        for t in (1, 2, 3):
            r = Dd[(Dd.arm == a) & (Dd.quantity == f"rs_T{t}")]["estimate"]
            ok &= len(r) == 1 and close(r.values, [s[t]])
        e = Dd[(Dd.arm == a) & (Dd.quantity == "rs_T3_minus_T1")]["estimate"]
        ok &= len(e) == 1 and close(e.values, [s[3] - s[1]], 1e-12)
    L.chk("B-08", ok, "15D plot data equal the tertile rs and T3 - T1 in Stats_15D", "4 arms")

    # ---- 15E --------------------------------------------------------------------
    E, Pe = D["Stats_15E.csv"], D["Plot_15E_samples.csv"]
    L.chk("B-09", close([np.median(Pe["log_ratio"])], E["estimate"].values),
          "15E median per-sample log ratio equals the stats estimate",
          f"{np.median(Pe['log_ratio']):.6f}")

    # ---- 15F --------------------------------------------------------------------
    F, Pf = D["Stats_15F.csv"], D["Plot_15F_samples.csv"]
    ok = True
    for col, qn in (("rs_raw", "rs_raw"), ("rs_partial", "rs_partial")):
        per_gene = Pf.groupby("partner")[col].median()
        for g, v in per_gene.items():
            r = F[(F.quantity == qn) & (F.measure == g)]["estimate"]
            ok &= len(r) == 1 and close(r.values, [v])
        per_sample = Pf.groupby("sample")[col].median()
        r = F[(F.quantity == qn) & (F.measure == "median over 8 partners")]["estimate"]
        ok &= len(r) == 1 and close(r.values, [np.median(per_sample)])
    ps = Pf.groupby("sample")["rs_raw"].median().values
    qr = F[F.quantity == "rs_raw_per_sample_quartiles"].iloc[0]
    ok &= close([qr["lo"], qr["estimate"], qr["hi"]], np.percentile(ps, [25, 50, 75]))
    L.chk("B-10", ok, "15F medians and quartile row equal the plot data", f"{Pf['sample'].nunique()} samples")

    # ---- 15G --------------------------------------------------------------------
    G, Pg = D["Stats_15G.csv"], D["Plot_15G_tertiles.csv"]
    w = Pg.pivot_table(index=["sample", "partner"], columns="tertile", values="rs").reset_index()
    w["d"] = w[3] - w[1]
    ok = True
    for g, v in w.groupby("partner")["d"].median().items():
        r = G[G.measure == g]["estimate"]
        ok &= len(r) == 1 and close(r.values, [v])
    r = G[G.measure == "median over 8 partners"]["estimate"]
    ok &= len(r) == 1 and close(r.values, [w.groupby("sample")["d"].median().median()])
    L.chk("B-11", ok, "15G estimates equal the plot data", f"{w['sample'].nunique()} samples")

    # ---- figure log ----------------------------------------------------------------
    fl = parse_log(find(root, FIG_LOG))
    last = max(dt.datetime.strptime(lg["start"], "%Y-%m-%d %H:%M:%S") for lg in logs.values())
    fstart = dt.datetime.strptime(fl["start"], "%Y-%m-%d %H:%M:%S")
    L.chk("B-12", fl["verdict"] == "PASS" and fl["n_fail"] == 0 and fstart >= last,
          "figure log is PASS and started after the last analysis section",
          f"{fl['script']} | {fl['start']}")
    L._say("")

    # ---- upstream inputs for the Module post-IP% columns --------------------------------
    cfg_path = find(root, CONFIG)
    cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
    L.chk("B-14.config", cfg_path.is_file() and str(cfg.get("config_version")) == one("config_version")[0],
          f"{CONFIG} exists and has the config version of the section logs",
          f"version {cfg.get('config_version')}")
    up = {}
    for nm, pat in UPSTREAM:
        p, lg = find(root, nm), parse_log(find_glob(root, pat))
        up[nm] = (p, lg)
        L.chk(f"B-14.{nm}", nm in lg["outputs"] and lg["verdict"] == "PASS",
              f"{nm} is listed in the outputs of its log, which is PASS", lg["path"].name)
    DN = read_csv_like_r(up["05_donors.csv"][0], char_cols=("donor_key", "donor_id"))
    LB = read_csv_like_r(up["04_cell_labels.csv.gz"][0], char_cols=("barcode", "donor_key"),
                         usecols=["arm", "dataset", "donor_key", "group", "doublet_score", "labels_version"])
    L.chk("B-14.labels", set(DN["labels_version"]) == set(LB["labels_version"]) == set(one("labels_version")),
          "donor table and cell labels carry the labels version of the section logs",
          ", ".join(sorted(set(DN["labels_version"]))))
    PA, PB, PSH, PU, DX = post_ip_donor_sets(DN, LB, cfg, ARMS, REGION_SETS)

    # ---- derived tallies for the Notes ------------------------------------------------
    ub, ab = q["module_beta_unadjusted"], q["module_beta_contam_adjusted"]
    sig = lambda s: s.fillna(1.0) < 0.05
    fin = q["dR2_unadjusted_same_donors"]["estimate"].notna() & q["dR2_contam_adjusted"]["estimate"].notna()
    sb = sig(ub["fdr"]).reindex(idx)[fin.reindex(idx)]
    sa = sig(ab["fdr"]).reindex(idx)[fin.reindex(idx)]
    st = pd.Series(np.where(sb & sa, "both", np.where(sb, "lost", np.where(sa, "gained", "never"))),
                   index=sb.index)
    n_st = st.value_counts().to_dict()
    lost = st[st == "lost"].index
    gained = st[st == "gained"].index
    lost_txt = "; ".join(f"{a} {lv} {m} ({ret[(a, lv, m)]:.2f})" for a, lv, m in lost)
    # Pre-set reading: a lost row with retention below 0.5 is attributed in part to
    # contamination; with retention of 0.5 or more it is borderline.
    lost_low = [k for k in lost if ret[k] < 0.5]
    lost_border = [k for k in lost if ret[k] >= 0.5]
    row_txt = lambda ks: ", ".join(" ".join(k) for k in ks) if ks else "none"
    gained_txt = "; ".join(f"{a} {lv} {m}" for a, lv, m in gained)
    cog = ["Cognitive status", "MMSE score", "CASI score"]
    m4c = [(lv, m) for lv in REGION_SETS for m in cog]
    m4_ok = all(sb[("M4", lv, m)] and sa[("M4", lv, m)] for lv, m in m4c)
    m4_min = min(m4c, key=lambda k: ret[("M4",) + k])
    donors_cov = {(a, lv): (int(ub.loc[(a, lv)]["n_donors"].max()), int(ab.loc[(a, lv)]["n_donors"].max()))
                  for a in ARMS for lv in REGION_SETS}

    B = D["Stats_15B.csv"]
    cuts = sorted({c.split("_")[-1] for c in B.quantity if c.startswith("module_beta_doublet_")})
    b_lost = {}
    for c in cuts:
        bb = B[B.quantity == f"module_beta_doublet_{c}"].set_index(key)["fdr"].reindex(ub.index)
        b_lost[c] = [k for k in idx if sig(ub["fdr"])[k] and not sig(bb)[k]]
    shares = B[B.quantity.str.startswith("removed_share_")]["estimate"]

    # the rebuilt donor sets must reproduce every n_donors in the stats files
    ok_a = all(PA[k]["n"] == ub.loc[k]["n_donors"]
               and PA[k]["n_same"] == q["dR2_unadjusted_same_donors"].loc[k]["n_donors"]
               and PA[k]["n_adj"] == ab.loc[k]["n_donors"] for k in idx)
    L.chk("B-15.15A", ok_a, "rebuilt 15A donor sets have the stats n_donors",
          f"{len(idx)} rows, unadjusted, same-donor and adjusted")
    ok_b, n_b = True, 0
    for c in cuts:
        bb = B[B.quantity == f"module_beta_doublet_{c}"].set_index(key)["n_donors"]
        for k in idx:
            ok_b &= PB[(k[0], k[1], c, k[2])]["n"] == bb[k]
            n_b += 1
        sh = B[B.quantity == f"removed_share_{c}"].set_index(["arm", "level"])
        for a in ARMS:
            for lv in REGION_SETS:
                ok_b &= close([PSH[(a, lv, c)][0]], [sh.loc[(a, lv)]["estimate"]], 1e-12) \
                    and PSH[(a, lv, c)][1] == sh.loc[(a, lv)]["n_donors"]
    L.chk("B-15.15B", ok_b, "rebuilt 15B donor sets and removed shares equal Stats_15B",
          f"{n_b} rows over {len(cuts)} thresholds")

    # the donor sheet must reproduce the models: refit all of them from its values
    cl = cfg["clinical"]
    cov_map = dict(cl["cov4"]) if isinstance(cl["cov4"], dict) else {c: c for c in cl["cov4"]}
    out_map = {k: v["column"] for k, v in cl["outcomes"].items()}
    min_n = int(cfg["floors"]["donors_per_test"])
    den_col = "n_super" if cfg["quantities"]["module_post_pct_denominator"] == "supertype" else "n_cells"
    out_cols = list(dict.fromkeys(out_map.values()))
    cov_cols = [c for c in dict.fromkeys(cov_map.values()) if c not in out_cols]
    # donor-sheet header of each source column: outcome label, else covariate name
    hdr_of = {v: k for k, v in out_map.items()}
    hdr_of.update({v: k for k, v in cov_map.items() if v not in out_cols})
    shared = [hdr_of[c] for c in cov_map.values() if c in out_cols]
    worst, n_fit = 0.0, 0
    def cmp(n, b, r2, row_b, row_r):
        nonlocal worst, n_fit
        n_fit += 1
        ok = n == row_r["n_donors"]
        for v, ref in ((b, row_b["estimate"] if row_b is not None else None), (r2, row_r["estimate"])):
            if ref is None:
                continue
            if np.isnan(v) and pd.isna(ref):
                continue
            if np.isnan(v) or pd.isna(ref):
                return False
            worst = max(worst, abs(v - ref))
        return ok
    ok_r = True
    for (a, lv, m) in idx:
        d = DX[(DX["arm"] == a) & (DX["dataset"] == lv)]
        y = out_map[m]
        cv = [c for c in cov_map.values() if c != y]
        n, b, r2 = ols_refit(d, y, "module_post_pct", cv, min_n)
        ok_r &= cmp(n, b, r2, ub.loc[(a, lv, m)], q["dR2_unadjusted"].loc[(a, lv, m)])
        dd = d[np.isfinite(d["contam_super_nonmodule"].astype(float))]
        n, b, r2 = ols_refit(dd, y, "module_post_pct", cv, min_n)
        ok_r &= cmp(n, b, r2, None, q["dR2_unadjusted_same_donors"].loc[(a, lv, m)])
        n, b, r2 = ols_refit(d, y, "module_post_pct", cv, min_n, extra=["contam_super_nonmodule"])
        ok_r &= cmp(n, b, r2, ab.loc[(a, lv, m)], q["dR2_contam_adjusted"].loc[(a, lv, m)])
        for c in cuts:
            n, b, r2 = ols_refit(d, y, f"post_ip_pct_{c}", cv, min_n)
            bb = B[(B.quantity == f"module_beta_doublet_{c}")].set_index(key).loc[(a, lv, m)]
            br = B[(B.quantity == f"dR2_doublet_{c}")].set_index(key).loc[(a, lv, m)]
            ok_r &= cmp(n, b, r2, bb, br)
    ok_r &= worst < 1e-8
    L.chk("B-16", ok_r, "every 15A and 15B model refitted from the Fig15AB_donors values gives the stats n_donors, beta and dR2",
          f"{n_fit} models | largest difference {worst:.1e}")

    # ---- workbook ------------------------------------------------------------------
    wb = Workbook()
    wb.remove(wb.active)
    texts = []  # every string written, for B-13

    # Overview
    ov = [
        ["Fig 15A", "Does contamination explain the clinical association?",
         "Module post-IP% in the Figure 12A model (cov4 OLS, 9 outcomes, region sets separate) before and after adding donor-mean DecontX contamination of the arm's non-Module supertype cells",
         "07", "Stats_15A.csv", "Fig15A_readout; Fig15A_stats; Fig15AB_donors", "Notes 5, 6, 15 and 16"],
        ["Fig 15B", "Do doublets drive the clinical association?",
         "The same model after removing Module post-IP nuclei with a doublet score at or above 0.15, 0.20 and 0.30",
         "07", "Stats_15B.csv", "Fig15B_readout; Fig15B_removed; Fig15B_stats; Fig15AB_donors", "Notes 7, 15 and 16"],
        ["Fig 15C", "Are Module post-IP cells more contaminated?",
         "Median over donors of the within-donor difference in group median DecontX contamination; three contrasts",
         "08", "Stats_15C.csv", "Fig15C_stats", "Note 4"],
        ["Fig 15D", "Does co-expression rise with contamination?",
         "Spearman rs of the GOI with the composite per within-donor contamination tertile; T3 minus T1",
         "08", "Stats_15D.csv; Plot_15D_tertiles.csv", "Fig15D_stats; Fig15D_tertiles", None],
        ["Fig 15E", "Is the ambient pool more mitochondrial than the nuclei? (Gerrits)",
         "Per sample, log of ambient MT fraction over nuclear MT fraction; median over samples",
         "09", "Stats_15E.csv; Plot_15E_samples.csv", "Fig15E_stats; Fig15E_samples", "Note 8"],
        ["Fig 15F", "Does co-expression survive removing ambient RNA and depth? (Gerrits)",
         "Per sample, median over the 8 DFG genes of rs with MT-CO3, raw and with soup index and log depth partialled out",
         "09", "Stats_15F.csv; Plot_15F_samples.csv", "Fig15F_stats; Fig15F_samples", "Notes 8 and 9"],
        ["Fig 15G", "Does co-expression change with ambient load? (Gerrits)",
         "Per sample, median over the 8 DFG genes of rs in the top minus the bottom soup-index tertile",
         "09", "Stats_15G.csv; Plot_15G_tertiles.csv", "Fig15G_stats; Fig15G_tertiles", "Notes 8 and 10"],
        ["Fig 15H", "Do depth and mitochondrial content explain the gap to Gerrits? (simulation)",
         "Median rs over 20 binomial thinning repeats per target depth, with and without matching the Gerrits MT fraction",
         "10", "Stats_15H.csv", "Fig15H_stats", "Note 11"],
        ["Supp Fig 11", "Is the extra material the cell's own cytoplasm or another lineage's?",
         "Median over donors of the within-donor difference, Module post-IP minus Module pre-IP, in four metrics",
         "11", "Stats_S10.csv", "SuppFig11_stats", "Notes 12 and 18"],
    ]
    hdr = ["Panel", "Question", "Quantity", "Section", "Source file", "Sheets", "See"]
    write_table(wb, "Overview", hdr, ov, wrap=True,
                widths={"Panel": 11, "Question": 40, "Quantity": 60, "Section": 9,
                        "Source file": 32, "Sheets": 34, "See": 20})
    texts += [str(v) for r in ov for v in r if v]

    # Notes
    fmt_share = lambda c: (f"{100 * B[B.quantity == f'removed_share_{c}']['estimate'].min():.1f}% to "
                           f"{100 * B[B.quantity == f'removed_share_{c}']['estimate'].max():.1f}%")
    n15 = D["Stats_15H.csv"]
    hk = n15[n15.quantity == "rs_housekeeping"]
    hk_same = close(hk[hk.level == "depth"].set_index("measure")["estimate"].sort_index().values,
                    hk[hk.level == "depth+mito"].set_index("measure")["estimate"].sort_index().values, 1e-12)
    notes = [
        ["1", "Arms", "M4 = Mic13_PVM3, the masked primary arm. PVM_3 = PVM3_all, the unmasked reference. PVM_1 = PVM1_all, the second supertype. HKG8 = the housekeeping control construct (ACTB with the cHKG8 genes). NeoC and NonNeoC are the two region sets."],
        ["2", "Resampling and p values", "Intervals are 95% percentile intervals from a donor bootstrap with B = 10,000. The bootstrap p is min(1, 2(k+1)/(B+1)); its smallest value, 2/(B+1) = 1.9998e-04, is a bound, and a BH value computed from floor p values is also a bound. The figures write such values as \"FDR<0.001\". The clinical models in 15A and 15B keep the OLS p and the OLS 95% confidence interval."],
        ["3", "Multiple testing", "fdr is Benjamini-Hochberg within the group named in the family column of each stats sheet. The builder recomputes BH within every family and checks it against the stats file (check B-06)."],
        ["4", "Floors and untestable rows", "A donor enters a within-donor contrast with at least 20 cells in each group; a test needs at least 10 donors. A test below the floor has an empty estimate, interval, p and fdr, and is reported as untestable. HKG8 is untestable in 15C and Supplementary Figure 11. In 15C the group-median rows count the donors with at least 20 cells in that group."],
        ["5", "Fig 15A models", "Outcome and Module post-IP% are standardised within each model's donors, so the coefficient is a standardised beta. dR2 is R2(cov4 + Module) minus R2(cov4); after adjustment it is R2(cov4 + contamination + Module) minus R2(cov4 + contamination). APOE e4 dose leaves the covariate block when it is the outcome. Retention is dR2 after adjustment divided by dR2 before adjustment on the same donors (the donors with a contamination value); it is empty when that dR2 is not above zero. Donors per region set before/after adjustment: " + "; ".join(f"{a} {lv} {b}/{c}" for (a, lv), (b, c) in donors_cov.items()) + "."],
        ["6", "Fig 15A result", f"Over {len(st)} rows: {n_st.get('both', 0)} significant before and after, {n_st.get('lost', 0)} lost, {n_st.get('gained', 0)} gained, {n_st.get('never', 0)} never significant. Lost rows with retention: {lost_txt}. Gained: {gained_txt}. The six M4 cognition rows (cognitive status, MMSE, CASI in each region set) are {'all' if m4_ok else 'not all'} significant before and after; the lowest retention among them is M4 {m4_min[0]} {m4_min[1]} at {ret[('M4',) + m4_min]:.2f}. Under the pre-set reading, lost rows with retention below 0.5 are attributed in part to contamination ({len(lost_low)}: {row_txt(lost_low)}) and lost rows with retention of 0.5 or more are borderline ({len(lost_border)}: {row_txt(lost_border)}). On PVM_1 NonNeoC the contamination covariate is missing for {donors_cov[('PVM_1', 'NonNeoC')][0] - donors_cov[('PVM_1', 'NonNeoC')][1]} of the unit's donors."],
        ["7", "Fig 15B", "Module post-IP nuclei with a doublet score at or above the threshold are removed from the numerator and the denominator of Module post-IP%, and the 15A unadjusted model is refitted. Significant rows lost, by threshold: " + "; ".join(f"{c}: {len(b_lost[c])}" + (f" ({', '.join(' '.join(k) for k in b_lost[c])})" if b_lost[c] else "") for c in cuts) + ". Share of Module post-IP nuclei removed per arm and region set: " + "; ".join(f"{c}: {fmt_share(c)}" for c in cuts) + " (sheet Fig15B_removed)."],
        ["8", "Gerrits data (15E to 15H)", "GSE148822: 36 samples from 18 donors, two cortical regions each, all NEUNneg OLIG2neg sorted nuclei. The sample is the unit; the donor bootstrap resamples donors and keeps a donor's two samples together. Nuclei are called by emptyDrops (lower 100 UMIs, FDR 0.01, fixed seed); the ambient pool is the barcodes with 1 to 100 UMIs. 15F and 15G use the MT-CO3-positive called nuclei."],
        ["9", "Fig 15F", "Partial rs removes the soup index and log depth on ranks. partial_over_raw is the retained share and carries no test. The rs_raw_per_sample_quartiles row (estimate = median, lo and hi = 25th and 75th percentiles of the per-sample raw values) is the Gerrits reference range used by 15H."],
        ["10", "Fig 15G", "Tertiles of the soup index are cut within each sample. Depth is not adjusted in this panel; 15F is the panel that removes depth."],
        ["11", "Fig 15H", "Simulation, no test. Population: the PVM_3 arm cells, all MT-CO3-positive. Each cell's counts are binomially thinned to the target depth; a cell already below the target is not thinned. The 'depth+mito' condition also thins the 13 MT genes to the Gerrits MT fraction. lo and hi are the smallest and largest value over the 20 repeats. Targets 1504 (the Gerrits median depth) and 1500 are both run. " + ("The housekeeping values are identical with and without the MT match, because the housekeeping genes are not MT genes and the normalisation depth is not changed by the match." if hk_same else "The housekeeping values differ between the two conditions.")],
        ["12", "Supplementary Figure 11 units", "Own-ER, ribosomal and NEAT1 shares are per cent of non-MT UMIs; the other-lineage index is per cent of total UMIs, under the circularity guard of the cell's own dataset. Differences are in percentage points."],
        ["13", "Column name \"partner\"", "In the Gerrits sheets the column \"partner\" names a DFG gene, one of the eight genes paired with MT-CO3. The figures and text use \"DFG genes\". The column name is kept so the sheets match the source files."],
        ["14", "Formulas", "In Fig15A_readout the retention and status columns, and in Fig15B_readout the dR2_change and status columns, are Excel formulas over the values in the same row. The builder checks that retention equals the value in Stats_15A (check B-07)."],
        ["15", "Module post-IP% columns", "The post_ip_pct_median columns give the median Module post-IP% (per cent of the arm's " + ("supertype cells" if cfg["quantities"]["module_post_pct_denominator"] == "supertype" else "cells") + ", per donor) over the donors behind that row. They appear only on the sheets that use Module post-IP%: Fig15A_readout, Fig15B_readout, Fig15B_removed, Fig15A_stats and Fig15B_stats. Each median uses the donor set of its row: the unadjusted, same-donor or contamination-adjusted model in 15A, the model after doublet removal in 15B, and for the removed-share rows every donor of the arm and region set with a value. The values are rebuilt from 05_donors.csv and 04_cell_labels.csv.gz with the donor filters of section 07, and the builder checks that every rebuilt donor set has the n_donors of its stats row and that the removed shares equal Stats_15B (check B-15). Rows of one arm and region set share nearly the same donors, so their medians are close. In 15B the value after removal is recomputed per donor with the removed nuclei taken out of the numerator and the denominator."],
        ["16", "Donor values (Fig15AB_donors)", "One row per arm, region set and donor, with every value the 15A and 15B models use, so each model can be refitted from this sheet. Outcomes: " + ", ".join(out_map) + ". Covariates: " + ", ".join(hdr_of[c] for c in cov_map.values()) + (f" ({', '.join(shared)} is both a covariate and an outcome and appears once, under its outcome name)" if shared else "") + ". post_ip_pct is Module post-IP% = 100 x n_module_post_ip_cells / " + ("n_supertype_cells" if den_col == "n_super" else "n_cells") + ". contamination_mean is the donor's mean DecontX contamination over the non-Module cells of the arm's supertype in that region set, as written by step 05; an empty cell means the donor has no value and is left out of the adjusted model. post_ip_pct_<t> is Module post-IP% after removing the n_removed_<t> Module post-IP nuclei with doublet score at or above t from the numerator and the denominator. To refit a model: keep the donors of one arm and region set with no empty cell in the outcome, the Module post-IP% column and the covariates (and contamination_mean for the adjusted model); drop a covariate from the covariates when it is the outcome; standardise the outcome and the Module post-IP% column over those donors (mean 0, standard deviation with n - 1); fit OLS with an intercept, the covariates as they are (sex is the Sex_num column: 0 = female, 1 = male), and the standardised Module post-IP%. dR2 is the R2 of that model minus the R2 of the model without Module post-IP%. The builder refits every 15A and 15B model this way and checks n_donors, beta and dR2 against the stats files (check B-16). A text value that is empty in the source is written as (empty), because R treats it as a level of its own."],
        ["17", "Old values", "The archived run is not compared here (decision D20 dropped, revision 49). Results are reported as they stand."],
        ["18", "Supplementary Figure 11 file names", "Supplementary Figure 11 was numbered S10 when the analysis ran. The source file Stats_S10.csv, section 11 (S10_Cargo) and the FigS10 figure files keep that name, as the logs name them. In SuppFig11_stats the panel column reads S11 and the family column starts with S11; every other value is as in Stats_S10.csv."],
    ]
    write_table(wb, "Notes", ["Note", "Subject", "Detail"], notes, wrap=True,
                widths={"Note": 7, "Subject": 30, "Detail": 120})
    texts += [str(v) for r in notes for v in r]

    # Glossary
    gl = [
        ["column", "panel", "Figure panel the row belongs to."],
        ["column", "arm", "M4, PVM_3, PVM_1 or HKG8 (Note 1); Gerrits for the external data; PVM_3 for the 15H simulation."],
        ["column", "level", "Region set (NeoC, NonNeoC), 'pooled region sets', 'sample' (Gerrits), or the 15H condition (observed, depth, depth+mito, Gerrits)."],
        ["column", "quantity", "What the row estimates; codes below."],
        ["column", "measure", "Outcome (15A, 15B), gene or summary (15F, 15G), target depth (15H) or metric (Supplementary Figure 11)."],
        ["column", "n_units", "Samples (Gerrits) or cells (15H) behind the row."],
        ["column", "n_donors", "Donors behind the row."],
        ["column", "estimate, lo, hi", "Point estimate and its 95% interval (Note 2); for 15H, lo and hi are the range over the repeats."],
        ["column", "p, fdr", "p value and Benjamini-Hochberg value within family (Notes 2 and 3)."],
        ["column", "family", "The BH group of the row."],
        ["15A", "module_beta_unadjusted", "Standardised beta of Module post-IP% in the cov4 OLS model, all donors with complete data."],
        ["15A", "module_beta_contam_adjusted", "The same beta with donor-mean non-Module supertype DecontX contamination added."],
        ["15A", "dR2_unadjusted", "dR2 of Module post-IP% before adjustment, all donors."],
        ["15A", "dR2_unadjusted_same_donors", "dR2 before adjustment on the donors with a contamination value."],
        ["15A", "dR2_contam_adjusted", "dR2 of Module post-IP% over cov4 plus contamination."],
        ["15A", "retention", "dR2_contam_adjusted / dR2_unadjusted_same_donors."],
        ["15A, 15B", "post_ip_pct_median", "Median Module post-IP% (per cent) over the donors of the model or unit in that row (Note 15). In Fig15B_stats, the removed-share rows give the median before removal over the unit's donors."],
        ["15A", "post_ip_pct_median_same", "In Fig15A_readout, the same over the donors of the same-donor unadjusted model."],
        ["15A", "post_ip_pct_median_adjusted", "In Fig15A_readout, the same over the donors of the contamination-adjusted model."],
        ["15B", "post_ip_pct_median_all_nuclei", "In Fig15B_readout, the median before doublet removal over the donors of the all-nuclei model; in Fig15B_removed, over the unit's donors."],
        ["15B", "post_ip_pct_median_removed", "In Fig15B_removed, the median over the unit's donors after removing nuclei at or above the threshold."],
        ["15B", "removed_share_<t>", "Share (fraction) of the unit's Module post-IP nuclei with doublet score at or above t; n_donors = donors in the unit."],
        ["15B", "module_beta_doublet_<t>, dR2_doublet_<t>", "Beta and dR2 of Module post-IP% recomputed without those nuclei."],
        ["15C", "donor_median_post, _pre, _non", "Median over donors of the donor's median DecontX contamination in Module post-IP, Module pre-IP or non-Module cells."],
        ["15C", "post_minus_pre, post_minus_nonmodule, pre_minus_nonmodule", "Median over donors of the within-donor difference of group medians."],
        ["15D", "rs_T1, rs_T2, rs_T3", "Spearman rs of the GOI (MT-CO3; ACTB on HKG8) with the composite in the Module cells of that contamination tertile, tertiles cut within donor, cells pooled per tertile."],
        ["15D", "rs_T3_minus_T1", "rs in the most minus the least contaminated tertile."],
        ["15E", "log_ambient_over_nuclear_mito", "Median over samples of log(ambient MT fraction / nuclear MT fraction)."],
        ["15F", "rs_raw, rs_partial", "Median over samples of rs(MT-CO3, DFG gene), raw or partial; per gene, or the per-sample median over the 8 genes."],
        ["15F", "partial_over_raw", "Median over samples of partial / raw."],
        ["15F", "rs_raw_per_sample_quartiles", "Quartiles of the per-sample raw values (Note 9)."],
        ["15G", "rs_top_minus_bottom", "Median over samples of rs in the top minus the bottom soup-index tertile."],
        ["15H", "rs_median_over_partners", "Median over repeats of the median over the 8 DFG genes of rs with MT-CO3."],
        ["15H", "rs_housekeeping", "The same for ACTB with the cHKG8 genes, on the same cells and draws."],
        ["15H", "median_depth, median_mito_fraction", "Median UMIs per cell and MT fraction after thinning."],
        ["15H", "reference_depth, reference_mito_fraction, reference_rs_raw_iqr", "Gerrits median depth, median MT fraction and the per-sample raw rs quartiles."],
        ["S11", "donor_median_post, donor_median_pre", "Median over donors of the donor's median metric in Module post-IP or Module pre-IP cells."],
        ["donors", "arm, region_set, donor_id, donor_key", "The unit and the donor. donor_id is the SEA-AD donor identifier; donor_key is the key used inside the pipeline."],
        ["donors", "outcome columns", "One column per clinical outcome, named as in the 15A and 15B sheets (Note 16)."],
        ["donors", "covariate columns", "Age, sex and PMI as in the donor table (sex: 0 = female, 1 = male, from Sex_num of the Part 10h Donor_Table); APOE e4 dose is the outcome column of that name."],
        ["donors", "n_module_post_ip_cells, n_supertype_cells", "The donor's Module post-IP cells and cells of the arm's supertype in that region set."],
        ["donors", "post_ip_pct", "Module post-IP%, per cent (Note 16)."],
        ["donors", "contamination_mean", "Donor mean DecontX contamination of the non-Module supertype cells; the 15A adjustment covariate."],
        ["donors", "n_removed_<t>, post_ip_pct_<t>", "Module post-IP nuclei with doublet score at or above t, and Module post-IP% after removing them (15B)."],
        ["S11", "post_minus_pre", "Median over donors of the within-donor difference, percentage points."],
    ]
    write_table(wb, "Glossary", ["Scope", "Term", "Meaning"], gl, wrap=True,
                widths={"Scope": 10, "Term": 44, "Meaning": 100})
    texts += [str(v) for r in gl for v in r]

    # Fig15A readout (one row per arm, region set, outcome; formulas for retention and status)
    h15a = ["arm", "region_set", "outcome",
            "n_donors", "post_ip_pct_median", "beta", "beta_lo", "beta_hi", "p", "fdr", "dR2",
            "n_donors_same", "post_ip_pct_median_same", "dR2_same_donors",
            "n_donors_adjusted", "post_ip_pct_median_adjusted", "beta_adjusted", "beta_adjusted_lo",
            "beta_adjusted_hi", "p_adjusted", "fdr_adjusted", "dR2_adjusted", "retention", "status"]
    cA = {h: get_column_letter(j + 1) for j, h in enumerate(h15a)}
    rows = []
    for i, k in enumerate(idx, start=2):
        u, s_, a_ = ub.loc[k], q["dR2_unadjusted_same_donors"].loc[k], ab.loc[k]
        L_, S_, I_, R_ = cA["dR2_same_donors"], cA["dR2_adjusted"], cA["fdr"], cA["fdr_adjusted"]
        rows.append([k[0], k[1], k[2],
                     u["n_donors"], PA[k]["med"], u["estimate"], u["lo"], u["hi"], u["p"], u["fdr"],
                     q["dR2_unadjusted"].loc[k]["estimate"],
                     s_["n_donors"], PA[k]["med_same"], s_["estimate"],
                     a_["n_donors"], PA[k]["med_adj"], a_["estimate"], a_["lo"], a_["hi"], a_["p"],
                     a_["fdr"], q["dR2_contam_adjusted"].loc[k]["estimate"],
                     f'=IF(AND(ISNUMBER({L_}{i}),ISNUMBER({S_}{i})),IF({L_}{i}>0,{S_}{i}/{L_}{i},""),"")',
                     f'=IF(AND(ISNUMBER({I_}{i}),{I_}{i}<0.05),IF(AND(ISNUMBER({R_}{i}),{R_}{i}<0.05),"significant before and after","lost"),IF(AND(ISNUMBER({R_}{i}),{R_}{i}<0.05),"gained","never significant"))'])
    write_table(wb, "Fig15A_readout", h15a, rows, p_cols=("p", "fdr", "p_adjusted", "fdr_adjusted"),
                fmt_cols={"retention": "0.000", "post_ip_pct_median": "0.00",
                          "post_ip_pct_median_same": "0.00", "post_ip_pct_median_adjusted": "0.00"},
                widths={"retention": 11, "status": 28})

    # Fig15B readout
    h15b = ["arm", "region_set", "threshold", "outcome",
            "n_donors_all_nuclei", "post_ip_pct_median_all_nuclei", "dR2_all_nuclei", "fdr_all_nuclei",
            "n_donors", "post_ip_pct_median", "beta", "beta_lo", "beta_hi", "p", "fdr", "dR2",
            "dR2_change", "status"]
    cB = {h: get_column_letter(j + 1) for j, h in enumerate(h15b)}
    rows = []
    i = 2
    for c in cuts:
        bq = {x: B[B.quantity == f"{x}_{c}"].set_index(key) for x in ("module_beta_doublet", "dR2_doublet")}
        for k in idx:
            u, b_, d_ = ub.loc[k], bq["module_beta_doublet"].loc[k], bq["dR2_doublet"].loc[k]
            F_, N_, G_, M_ = cB["dR2_all_nuclei"], cB["dR2"], cB["fdr_all_nuclei"], cB["fdr"]
            rows.append([k[0], k[1], float(c), k[2],
                         u["n_donors"], PA[k]["med"], q["dR2_unadjusted"].loc[k]["estimate"], u["fdr"],
                         b_["n_donors"], PB[(k[0], k[1], c, k[2])]["med"],
                         b_["estimate"], b_["lo"], b_["hi"], b_["p"], b_["fdr"],
                         d_["estimate"],
                         f'=IF(AND(ISNUMBER({F_}{i}),ISNUMBER({N_}{i})),{N_}{i}-{F_}{i},"")',
                         f'=IF(AND(ISNUMBER({G_}{i}),{G_}{i}<0.05),IF(AND(ISNUMBER({M_}{i}),{M_}{i}<0.05),"significant before and after","lost"),IF(AND(ISNUMBER({M_}{i}),{M_}{i}<0.05),"gained","never significant"))'])
            i += 1
    write_table(wb, "Fig15B_readout", h15b, rows, p_cols=("fdr_all_nuclei", "p", "fdr"),
                fmt_cols={"threshold": "0.00", "post_ip_pct_median_all_nuclei": "0.00",
                          "post_ip_pct_median": "0.00"},
                widths={"status": 28})

    sh = B[B.quantity.str.startswith("removed_share_")]
    rows = [[r.arm, r.level, float(r.quantity.split("_")[-1]), r.n_donors, r.estimate,
             PU[(r.arm, r.level)], PU[(r.arm, r.level, r.quantity.split("_")[-1])]]
            for r in sh.itertuples()]
    write_table(wb, "Fig15B_removed", ["arm", "region_set", "threshold", "n_donors", "removed_share",
                                       "post_ip_pct_median_all_nuclei", "post_ip_pct_median_removed"],
                rows, fmt_cols={"threshold": "0.00", "removed_share": "0.0%",
                                "post_ip_pct_median_all_nuclei": "0.00",
                                "post_ip_pct_median_removed": "0.00"})

    # Donor values behind 15A and 15B, one row per arm, region set and donor
    dcols = (["donor_id", "donor_key"] + out_cols + cov_cols + ["n_post", den_col, "module_post_pct",
             "contam_super_nonmodule"] + [x for c in cuts for x in (f"n_removed_{c}", f"post_ip_pct_{c}")])
    dsheet = DX[DX["arm"].isin(ARMS) & DX["dataset"].isin(REGION_SETS)].copy()
    dsheet["_a"] = dsheet["arm"].map({a: i for i, a in enumerate(ARMS)})
    dsheet["_r"] = dsheet["dataset"].map({r: i for i, r in enumerate(REGION_SETS)})
    dsheet = dsheet.sort_values(["_a", "_r", "donor_id"])
    names = {"n_post": "n_module_post_ip_cells",
             den_col: "n_supertype_cells" if den_col == "n_super" else "n_cells",
             "module_post_pct": "post_ip_pct", "contam_super_nonmodule": "contamination_mean"}
    dh = ["arm", "region_set"] + [hdr_of.get(c, names.get(c, c)) for c in dcols]
    # An empty text value is a level of its own in R (read.csv keeps "" in a text
    # column), but an empty Excel cell reads back as missing. Such values are
    # written as "(empty)" so a refit from the sheet keeps those donors.
    n_empty = 0
    for c in dcols:
        if not pd.api.types.is_numeric_dtype(dsheet[c]):
            e = dsheet[c].notna() & (dsheet[c].astype(str) == "")
            n_empty += int(e.sum())
            dsheet.loc[e, c] = "(empty)"
    if n_empty:
        L.note("Fig15AB_donors empty text values written as (empty)", f"{n_empty} cells")
    rows = [[r["arm"], r["dataset"]] + [r[c] for c in dcols] for _, r in dsheet.iterrows()]
    write_table(wb, "Fig15AB_donors", dh, rows,
                fmt_cols={h: "0.000" for h in dh if h.startswith("post_ip_pct")})
    L.chk("B-17", len(dsheet) == sum(PSH[(a, lv, cuts[0])][1] for a in ARMS for lv in REGION_SETS),
          "Fig15AB_donors holds every donor of the removed-share units", f"{len(dsheet)} rows")

    # Stats and plot-data sheets, as written by the analysis
    # post-IP% per stats row, over the donors of the model or unit behind the row
    SET_A = {"module_beta_unadjusted": "med", "dR2_unadjusted": "med",
             "dR2_unadjusted_same_donors": "med_same", "module_beta_contam_adjusted": "med_adj",
             "dR2_contam_adjusted": "med_adj", "retention": "med_adj"}
    pa_col = [PA[(r.arm, r.level, r.measure)][SET_A[r.quantity]] for r in A.itertuples()]
    pb_col = []
    for r in B.itertuples():
        c = r.quantity.split("_")[-1]
        if r.quantity.startswith("removed_share_"):
            pb_col.append(PU[(r.arm, r.level)])
        else:
            pb_col.append(PB[(r.arm, r.level, c, r.measure)]["med"])
    stats_sheet(wb, "Fig15A_stats", A, post_ip=pa_col)
    stats_sheet(wb, "Fig15B_stats", B, post_ip=pb_col)
    stats_sheet(wb, "Fig15C_stats", D["Stats_15C.csv"])
    stats_sheet(wb, "Fig15D_stats", Dd)
    write_table(wb, "Fig15D_tertiles", list(Dp.columns), Dp.itertuples(index=False, name=None))
    stats_sheet(wb, "Fig15E_stats", E)
    write_table(wb, "Fig15E_samples", list(Pe.columns), Pe.itertuples(index=False, name=None))
    stats_sheet(wb, "Fig15F_stats", F)
    write_table(wb, "Fig15F_samples", list(Pf.columns), Pf.itertuples(index=False, name=None))
    stats_sheet(wb, "Fig15G_stats", G)
    write_table(wb, "Fig15G_tertiles", list(Pg.columns), Pg.itertuples(index=False, name=None))
    stats_sheet(wb, "Fig15H_stats", n15)
    # Supplementary Figure 11 was S10 at analysis time; relabel panel and family on the sheet only.
    # The checks above ran on the source file as read.
    s11 = D["Stats_S10.csv"].copy()
    s11["panel"] = s11["panel"].replace({"S10": "S11"})
    s11["family"] = s11["family"].where(s11["family"].isna(),
                                        s11["family"].astype(str).str.replace(r"^S10\|", "S11|", regex=True))
    stats_sheet(wb, "SuppFig11_stats", s11)

    # every string in the workbook, for B-13
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts += [v for v in row if isinstance(v, str)]
    bad = [t for t in texts if any(d in t for d in DASHES)]
    L.chk("B-13", not bad, "no en dash or em dash in the workbook text", f"{len(bad)} found")

    L._say("")
    L._say(f"Checks: {L.n['PASS']} PASS, {L.n['FAIL']} FAIL, {L.n['NOTE']} NOTE")
    log_path = out.with_name(out.stem + "_build_log.txt")
    out.parent.mkdir(parents=True, exist_ok=True)
    if L.n["FAIL"]:
        L._say("VERDICT: FAIL | workbook not written")
        log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")
        sys.exit(1)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    L._say(f"Output: {out}")
    L._say("VERDICT: PASS")
    log_path.write_text("\n".join(L.lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
