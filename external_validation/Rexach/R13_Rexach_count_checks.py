# R13_Rexach_count_checks.py
# Count checks of external validation 2 (Rexach et al. 2024, insula), run after MapMyCells and
# before R14. Counts only: no correlation, screen, GAM, IP, or association is computed here.
#
# Checks:
#   A. MapMyCells files: all six present (parts 1 to 3, Deep Generative and SEA-AD Hierarchical
#      Mapping), each covering exactly the nuclei of its R12 part, with no duplicates or extras.
#      The algorithm and codebase version are taken from each csv header.
#   B. Taxonomy names: Micro-PVM_3-SEAAD and Micro-PVM_1 written exactly as in the SEA-AD release.
#   C. Agreement: Deep Generative against Hierarchical Mapping at subclass and
#      supertype, within Microglia-PVM, and against the authors' cell type; the Sst Chodl count.
#   D. Micro-PVM_3 nuclei per donor, by diagnosis, with donors below 5, 10, 20, and 50 nuclei. The
#      arm is Micro-PVM_3 by Deep Generative Mapping.
#   E. Depth of the PVM_3 arm (Micro-PVM_3 with MT-CO3 > 0): median UMIs, genes,
#      MT UMIs, and MT-CO3 UMIs per nucleus, %MT range, and the HKG8 arm count
#      (PVM_3 arm nuclei with ACTB > 0).
#   F. Export for the R steps: the raw counts of every nucleus that is Micro-PVM_3 in either mapping
#      (all genes, genes x nuclei), with a nucleus table that flags each mapping.
#
# Usage (WSL, the same venv as R11 and R12):
#   cd /mnt/c/Users/Wong/Desktop/SEA-AD/External_Dataset_2
#   ~/venv_synapse/bin/python R13_Rexach_count_checks.py
# Reads, from the script's folder: the CELLxGENE .h5ad, R12_mapmycells_input/R12_Rexach_nuclei.csv.gz,
# and the MapMyCells csv files in R13_mapmycells_output/.
# Outputs, in the folder R13_count_checks:
#   R13_Rexach_count_checks.xlsx   sheets Files, Taxonomy_names, Agreement, Microglia_DGM_vs_HM,
#                                  Author_vs_DGM, Donors, Arm_depth, Settings, and Session
#   R13_Rexach_MicroPVM3_counts.mtx.gz   raw counts, genes x nuclei (MatrixMarket)
#   R13_Rexach_MicroPVM3_genes.csv.gz    gene_id and symbol, in matrix row order
#   R13_Rexach_MicroPVM3_nuclei.csv.gz   one row per nucleus, in matrix column order
#   R13_Rexach_log.txt                   console output, with progress every 30 s

import glob
import gzip
import importlib.metadata as im
import os
import platform
import re
import sys
import time

import anndata as ad
import h5py
import numpy as np
import pandas as pd
import scipy.io
import scipy.sparse as sp

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    sys.exit("Install openpyxl first: ~/venv_synapse/bin/pip install -q openpyxl")

# The data folder: REXACH_ROOT when set (the container sets it to /data), otherwise the script folder.
HERE = os.environ.get("REXACH_ROOT", os.path.dirname(os.path.abspath(__file__)))
MAP_DIR = os.path.join(HERE, "R13_mapmycells_output")
NUC_CSV = os.path.join(HERE, "R12_mapmycells_input", "R12_Rexach_nuclei.csv.gz")
OUT = os.path.join(HERE, "R13_count_checks")
os.makedirs(OUT, exist_ok=True)
SET = {"arm_supertype": "Micro-PVM_3-SEAAD", "pvm1_supertype": "Micro-PVM_1",
       "microglia_subclass": "Microglia-PVM", "author_microglia": "microglial cell",
       "goi": "MT-CO3", "hkg8_goi": "ACTB", "donor_floor_section3": 5,
       "chunk_rows": 2000, "progress_s": 30}
MT13 = ["MT-ND1", "MT-ND2", "MT-CO1", "MT-CO2", "MT-ATP8", "MT-ATP6", "MT-CO3", "MT-ND3",
        "MT-ND4L", "MT-ND4", "MT-ND5", "MT-ND6", "MT-CYB"]
ALGS = {"DeepGenerativeMapping": "DGM", "SEA-ADHierarchicalMapping": "HM"}

log = open(os.path.join(OUT, "R13_Rexach_log.txt"), "w", encoding="utf-8")
pd.set_option("display.width", 220)
pd.set_option("display.max_rows", 300)
pd.set_option("display.max_columns", 30)


def say(*parts):
    s = " ".join(str(p) for p in parts)
    print(s, flush=True)
    log.write(s + "\n")
    log.flush()


def stop(msg):
    say("STOPPED:", msg)
    log.close()
    sys.exit(1)


t0 = time.time()
started = time.strftime("%Y-%m-%d %H:%M:%S")
say("R13_Rexach_count_checks.py started", started)

h5ads = glob.glob(os.path.join(HERE, "*.h5ad"))
if len(h5ads) != 1:
    stop("expected one .h5ad in %s, found %d" % (HERE, len(h5ads)))
H5AD = h5ads[0]
if not os.path.exists(NUC_CSV):
    stop("R12 output not found: " + NUC_CSV)
nuc = pd.read_csv(NUC_CSV)
kept = nuc[nuc["pass_floor"]].copy()
say("R12 nuclei: %d insula | %d kept by the floor" % (len(nuc), len(kept)))

# ---- A. MapMyCells files ------------------------------------------------------------------
say("\n== A. MapMyCells files")
pat = re.compile(r"part(\d+)_.*_(DeepGenerativeMapping|SEA-ADHierarchicalMapping)_UTC_\d+\.csv$")
files = {}
for p in sorted(glob.glob(os.path.join(MAP_DIR, "*.csv"))):
    m = pat.search(os.path.basename(p))
    if m:
        key = (int(m.group(1)), ALGS[m.group(2)])
        if key in files:
            stop("two csv files for part %d %s" % key)
        files[key] = p
parts = sorted(kept["part"].unique())
missing = [(p, a) for p in parts for a in ("DGM", "HM") if (p, a) not in files]
if missing:
    stop("missing MapMyCells csv files: " + ", ".join("part%d %s" % k for k in missing))


def read_mapping(path):
    with open(path, encoding="utf-8") as fh:
        header = [ln.strip() for ln in fh if ln.startswith("#")]
    d = pd.read_csv(path, comment="#")
    alg = next((h for h in header if h.startswith("# algorithm")), "")
    return d, alg


maps = {"DGM": [], "HM": []}
file_rows = []
for (p, a), path in sorted(files.items()):
    d, alg = read_mapping(path)
    exp = set(kept.loc[kept["part"] == p, "cell_id"])
    got = d["cell_id"].astype(str)
    row = {"part": p, "algorithm": a, "file": os.path.basename(path), "header": alg,
           "nuclei": len(d), "expected": len(exp), "duplicates": int(got.duplicated().sum()),
           "missing": len(exp - set(got)), "extra": len(set(got) - exp)}
    row["ok"] = row["nuclei"] == row["expected"] and row["duplicates"] == 0 \
        and row["missing"] == 0 and row["extra"] == 0
    file_rows.append(row)
    say("part %d %s: %d nuclei (expected %d) | duplicates %d | missing %d | extra %d | %s" %
        (p, a, row["nuclei"], row["expected"], row["duplicates"], row["missing"], row["extra"],
         "OK" if row["ok"] else "NOT OK"))
    say("   ", alg)
    maps[a].append(d)
fr = pd.DataFrame(file_rows)
if not fr["ok"].all():
    stop("a MapMyCells file does not match its R12 part (see above)")


def levels(d):
    """class, subclass, supertype name and probability columns of a MapMyCells csv."""
    out = {}
    for lv in ("class", "subclass", "supertype"):
        name = lv + "_name"
        probs = [c for c in d.columns if c.startswith(lv + "_") and c.endswith("probability")]
        if name not in d.columns:
            stop("column %s not found; columns are %s" % (name, list(d.columns)))
        out[lv] = (name, probs[0] if probs else None)
    return out


mp = {}
for a in ("DGM", "HM"):
    d = pd.concat(maps[a], ignore_index=True)
    lv = levels(d)
    cols = {"cell_id": d["cell_id"].astype(str)}
    for k, (name, prob) in lv.items():
        cols["%s_%s" % (a, k)] = d[name].astype(str)
        if prob:
            cols["%s_%s_prob" % (a, k)] = pd.to_numeric(d[prob], errors="coerce")
    mp[a] = pd.DataFrame(cols)
    say("%s probability columns used: %s" % (a, ", ".join(p for _, p in lv.values() if p)))
cells = kept.merge(mp["DGM"], on="cell_id", how="left").merge(mp["HM"], on="cell_id", how="left")
if cells["DGM_supertype"].isna().any() or cells["HM_supertype"].isna().any():
    stop("some kept nuclei have no mapping")

# ---- B. Taxonomy names -----------------------------------------------------------------------
say("\n== B. Taxonomy names")
tn = []
for a in ("DGM", "HM"):
    st = set(cells[a + "_supertype"])
    for nm in (SET["arm_supertype"], SET["pvm1_supertype"]):
        tn.append({"algorithm": a, "name": nm, "present": nm in st})
    say("%s Microglia-PVM supertypes: %s" % (a, ", ".join(sorted(
        cells.loc[cells[a + "_subclass"] == SET["microglia_subclass"], a + "_supertype"].unique()))))
tn = pd.DataFrame(tn)
say(tn.to_string(index=False))

# ---- C. Agreement ------------------------------------------------------------------------------
say("\n== C. Agreement")
agr = []
for lv in ("class", "subclass", "supertype"):
    same = (cells["DGM_" + lv] == cells["HM_" + lv])
    agr.append({"comparison": "DGM vs HM, all kept nuclei", "level": lv, "nuclei": len(cells),
                "agree": int(same.sum()), "pct_agree": round(100 * same.mean(), 2)})
mgd = cells["DGM_subclass"] == SET["microglia_subclass"]
same = cells.loc[mgd, "DGM_supertype"] == cells.loc[mgd, "HM_supertype"]
agr.append({"comparison": "DGM vs HM, DGM Microglia-PVM nuclei", "level": "supertype",
            "nuclei": int(mgd.sum()), "agree": int(same.sum()), "pct_agree": round(100 * same.mean(), 2)})
arm_d = cells["DGM_supertype"] == SET["arm_supertype"]
arm_h = cells["HM_supertype"] == SET["arm_supertype"]
both = arm_d & arm_h
agr.append({"comparison": "Micro-PVM_3: DGM nuclei also Micro-PVM_3 in HM", "level": "supertype",
            "nuclei": int(arm_d.sum()), "agree": int(both.sum()),
            "pct_agree": round(100 * both.sum() / max(int(arm_d.sum()), 1), 2)})
agr.append({"comparison": "Micro-PVM_3: HM nuclei also Micro-PVM_3 in DGM", "level": "supertype",
            "nuclei": int(arm_h.sum()), "agree": int(both.sum()),
            "pct_agree": round(100 * both.sum() / max(int(arm_h.sum()), 1), 2)})
am = cells["cell_type"] == SET["author_microglia"]
for a in ("DGM", "HM"):
    x = cells[a + "_subclass"] == SET["microglia_subclass"]
    agr.append({"comparison": "%s Microglia-PVM nuclei that the authors call microglial cell" % a,
                "level": "subclass", "nuclei": int(x.sum()), "agree": int((x & am).sum()),
                "pct_agree": round(100 * (x & am).sum() / max(int(x.sum()), 1), 2)})
    agr.append({"comparison": "Authors' microglial cell nuclei that %s calls Microglia-PVM" % a,
                "level": "subclass", "nuclei": int(am.sum()), "agree": int((x & am).sum()),
                "pct_agree": round(100 * (x & am).sum() / max(int(am.sum()), 1), 2)})
    y = cells[a + "_supertype"] == SET["arm_supertype"]
    agr.append({"comparison": "%s Micro-PVM_3 nuclei that the authors call microglial cell" % a,
                "level": "supertype", "nuclei": int(y.sum()), "agree": int((y & am).sum()),
                "pct_agree": round(100 * (y & am).sum() / max(int(y.sum()), 1), 2)})
agr = pd.DataFrame(agr)
say(agr.to_string(index=False))
if "DGM_supertype_prob" in cells.columns:
    pr = cells.loc[arm_d, "DGM_supertype_prob"]
    say("DGM Micro-PVM_3 supertype probability: median %.3f | below 0.5: %d of %d" %
        (pr.median(), int((pr < 0.5).sum()), len(pr)))
for a in ("DGM", "HM"):
    sc = cells[a + "_subclass"] == "Sst Chodl"
    say("%s Sst Chodl nuclei: %d | authors' labels for them: %s" % (a, int(sc.sum()), "; ".join(
        "%s %d" % kv for kv in cells.loc[sc, "cell_type"].value_counts().head(4).items())))
sc_d = cells["DGM_subclass"] == "Sst Chodl"
say("DGM Sst Chodl nuclei, HM subclass: %s" % "; ".join(
    "%s %d" % kv for kv in cells.loc[sc_d, "HM_subclass"].value_counts().head(5).items()))
say("DGM Sst Chodl nuclei: median genes %.0f against %.0f for all kept nuclei" %
    (cells.loc[sc_d, "n_genes"].median(), cells["n_genes"].median()))

xt_mg = pd.crosstab(cells.loc[mgd | (cells["HM_subclass"] == SET["microglia_subclass"]), "DGM_supertype"],
                    cells.loc[mgd | (cells["HM_subclass"] == SET["microglia_subclass"]), "HM_supertype"])
xt_mg = xt_mg.loc[:, xt_mg.sum() > 0]
xt_auth = pd.crosstab(cells["cell_type"], cells["DGM_subclass"])
say("\nMicroglia-PVM supertypes, DGM (rows) against HM (columns):")
say(xt_mg.to_string())

# ---- D. Micro-PVM_3 per donor ------------------------------------------------------------------
say("\n== D. Micro-PVM_3 nuclei per donor (arm: Deep Generative Mapping)")
dn = cells.groupby("donor_id").agg(disease=("disease", "first"), sex=("sex", "first"),
                                   nuclei_kept=("cell_id", "size")).reset_index()
for col, mask in (("microglia_pvm_dgm", mgd), ("micro_pvm3_dgm", arm_d), ("micro_pvm3_hm", arm_h),
                  ("micro_pvm3_both", both), ("micro_pvm1_dgm", cells["DGM_supertype"] == SET["pvm1_supertype"]),
                  ("author_microglia", am)):
    dn[col] = dn["donor_id"].map(cells[mask].groupby("donor_id").size()).fillna(0).astype(int)
# ---- E. Read counts for the union of Micro-PVM_3 nuclei (both mappings) ---------------------------
f = h5py.File(H5AD, "r")
A = ad.read_h5ad(H5AD, backed="r")
counts_path = "raw/X" if "raw/X" in f else "X"
var = A.raw.var if (counts_path == "raw/X" and A.raw is not None) else A.var
gene_ids = np.asarray(var.index.astype(str).tolist(), dtype=object)
symbols = np.asarray(var["feature_name"].astype(str).tolist(), dtype=object)
row_of = pd.Series(np.arange(A.n_obs), index=pd.Index(A.obs.index.astype(str).tolist(), dtype=object))
g = f[counts_path]
indptr = g["indptr"][:]
ng_total = int(g.attrs["shape"][1])

sel = cells[arm_d | arm_h].copy()
sel["row"] = row_of.reindex(sel["cell_id"].tolist()).values
if sel["row"].isna().any():
    stop("some Micro-PVM_3 cell IDs are not in the h5ad")
sel = sel.sort_values("row").reset_index(drop=True)
rows = sel["row"].astype(np.int64).values
say("\nReading counts of %d nuclei that are Micro-PVM_3 in either mapping" % len(rows))
blocks, last = [], time.time()
for s in range(0, len(rows), SET["chunk_rows"]):
    r = rows[s:s + SET["chunk_rows"]]
    breaks = np.where(np.diff(r) != 1)[0] + 1
    dparts, iparts, lens = [], [], []
    for run in np.split(r, breaks):
        a0, b0 = indptr[run[0]], indptr[run[-1] + 1]
        dparts.append(g["data"][a0:b0]); iparts.append(g["indices"][a0:b0])
        lens.append(np.diff(indptr[run[0]:run[-1] + 2]))
    m = sp.csr_matrix((np.concatenate(dparts).astype(np.int32), np.concatenate(iparts),
                       np.concatenate([[0], np.cumsum(np.concatenate(lens))])), shape=(len(r), ng_total))
    m.eliminate_zeros()
    blocks.append(m)
    if time.time() - last >= SET["progress_s"]:
        say("  %d of %d nuclei read | %.0f s" % (min(s + SET["chunk_rows"], len(rows)), len(rows), time.time() - t0))
        last = time.time()
f.close()
M = sp.vstack(blocks, format="csr")
gi = {s: i for i, s in enumerate(symbols)}
i_mt = [gi[x] for x in MT13]
sel["total_umis_check"] = np.asarray(M.sum(axis=1)).ravel()
if not np.array_equal(sel["total_umis_check"].values, sel["total_umis"].values):
    stop("total UMIs read here differ from R12")
sel["mt_co3_umis"] = M[:, gi[SET["goi"]]].toarray().ravel()
sel["actb_umis"] = M[:, gi[SET["hkg8_goi"]]].toarray().ravel()
sel["is_pvm3_dgm"] = sel["DGM_supertype"] == SET["arm_supertype"]
sel["is_pvm3_hm"] = sel["HM_supertype"] == SET["arm_supertype"]
sel["in_pvm3_arm"] = sel["is_pvm3_dgm"] & (sel["mt_co3_umis"] > 0)
sel["in_hkg8_arm"] = sel["in_pvm3_arm"] & (sel["actb_umis"] > 0)

arm = sel[sel["in_pvm3_arm"]]
dn["pvm3_arm"] = dn["donor_id"].map(arm.groupby("donor_id").size()).fillna(0).astype(int)
dn["hkg8_arm"] = dn["donor_id"].map(sel[sel["in_hkg8_arm"]].groupby("donor_id").size()).fillna(0).astype(int)
dn["pct_pvm3_mtco3_pos"] = (100 * dn["pvm3_arm"] / dn["micro_pvm3_dgm"].replace(0, np.nan)).round(1)
dn["section3_eligible"] = dn["micro_pvm3_dgm"] >= SET["donor_floor_section3"]
dn = dn.sort_values(["disease", "donor_id"])
say(dn.to_string(index=False))
say("\nMicro-PVM_3 (DGM) per donor by diagnosis:")
say(dn.groupby("disease")["micro_pvm3_dgm"].agg(["count", "sum", "median", "min", "max"]).to_string())
for fl in (5, 10, 20, 50):
    below = dn.loc[dn["micro_pvm3_dgm"] < fl, "donor_id"].tolist()
    say("Donors with fewer than %d Micro-PVM_3 nuclei: %d %s" % (fl, len(below), ", ".join(below)))
el = dn[dn["section3_eligible"]]
say("Section 3 donors (>= %d Micro-PVM_3 nuclei): %d | by diagnosis: %s" % (
    SET["donor_floor_section3"], len(el), "; ".join("%s %d" % kv for kv in el["disease"].value_counts().items())))

say("\n== E. Depth of the PVM_3 arm (Micro-PVM_3 by DGM with MT-CO3 > 0)")
n_pvm3 = int(sel["is_pvm3_dgm"].sum())
say("Micro-PVM_3 (DGM): %d | with MT-CO3 > 0 (PVM_3 arm): %d (%.1f%%) | HKG8 arm (also ACTB > 0): %d" %
    (n_pvm3, len(arm), 100 * len(arm) / max(n_pvm3, 1), int(sel["in_hkg8_arm"].sum())))
dep = pd.DataFrame({
    "measure": ["total UMIs", "genes detected", "MT UMIs", "MT-CO3 UMIs", "%MT"],
    "median": [arm["total_umis"].median(), arm["n_genes"].median(), arm["mt_umis"].median(),
               arm["mt_co3_umis"].median(), arm["pct_mt"].median()],
    "p10": [arm[c].quantile(0.1) for c in ("total_umis", "n_genes", "mt_umis", "mt_co3_umis", "pct_mt")],
    "p90": [arm[c].quantile(0.9) for c in ("total_umis", "n_genes", "mt_umis", "mt_co3_umis", "pct_mt")],
    "max": [arm[c].max() for c in ("total_umis", "n_genes", "mt_umis", "mt_co3_umis", "pct_mt")]})
dep[["median", "p10", "p90", "max"]] = dep[["median", "p10", "p90", "max"]].round(3)
say(dep.to_string(index=False))
say("PVM_3 arm nuclei with %%MT above 8: %d" % int((arm["pct_mt"] > 8).sum()))
say("MT-CO3 UMIs in the PVM_3 arm, share of nuclei with 1, 2, 3, and 4 or more: %s" % ", ".join(
    "%.1f%%" % (100 * v) for v in [(arm["mt_co3_umis"] == 1).mean(), (arm["mt_co3_umis"] == 2).mean(),
                                   (arm["mt_co3_umis"] == 3).mean(), (arm["mt_co3_umis"] >= 4).mean()]))

# ---- F. Export --------------------------------------------------------------------------------------
say("\n== F. Export of the Micro-PVM_3 nuclei (either mapping)")
with gzip.open(os.path.join(OUT, "tmp_R13_Rexach_MicroPVM3_counts.mtx.gz"), "wb") as fh:
    scipy.io.mmwrite(fh, M.T.tocoo(), field="integer")
os.replace(os.path.join(OUT, "tmp_R13_Rexach_MicroPVM3_counts.mtx.gz"),
           os.path.join(OUT, "R13_Rexach_MicroPVM3_counts.mtx.gz"))
pd.DataFrame({"gene_id": gene_ids, "symbol": symbols}).to_csv(
    os.path.join(OUT, "R13_Rexach_MicroPVM3_genes.csv.gz"), index=False, compression="gzip")
keep_cols = ["cell_id", "donor_id", "library_id", "disease", "sex", "development_stage", "cell_type",
             "ct_subcluster", "total_umis", "n_genes", "mt_umis", "pct_mt", "mt_co3_umis", "actb_umis",
             "DGM_subclass", "DGM_supertype", "DGM_supertype_prob", "HM_subclass", "HM_supertype",
             "HM_supertype_prob", "is_pvm3_dgm", "is_pvm3_hm", "in_pvm3_arm", "in_hkg8_arm"]
sel[[c for c in keep_cols if c in sel.columns]].to_csv(
    os.path.join(OUT, "R13_Rexach_MicroPVM3_nuclei.csv.gz"), index=False, compression="gzip")
say("Matrix: %d genes x %d nuclei | genes with duplicate symbols: %d" %
    (M.shape[1], M.shape[0], int(pd.Series(symbols).duplicated().sum())))

settings = pd.DataFrame({"setting": list(SET.keys()) + ["counts", "input"],
                         "value": [str(v) for v in SET.values()] + [counts_path, os.path.basename(H5AD)]})
pk = ["anndata", "h5py", "numpy", "pandas", "scipy", "openpyxl"]
session = pd.DataFrame({"item": ["python", "platform", "started", "finished"] + pk,
                        "value": [sys.version.split()[0], platform.platform(), started,
                                  time.strftime("%Y-%m-%d %H:%M:%S")] + [im.version(x) for x in pk]})


def write_xlsx(path, sheets):
    """Plain workbook: bold header and a frozen top row, no fills or borders."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, df in sheets:
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                       (v.item() if hasattr(v, "item") else v) for v in r])
        ws.freeze_panes = "A2"
    wb.save(path + ".tmp")
    os.replace(path + ".tmp", path)


write_xlsx(os.path.join(OUT, "R13_Rexach_count_checks.xlsx"), [
    ("Files", fr), ("Taxonomy_names", tn), ("Agreement", agr),
    ("Microglia_DGM_vs_HM", xt_mg.reset_index()), ("Author_vs_DGM", xt_auth.reset_index()),
    ("Donors", dn), ("Arm_depth", dep), ("Settings", settings), ("Session", session)])
say("\nDone in %.0f s" % (time.time() - t0))
log.close()
