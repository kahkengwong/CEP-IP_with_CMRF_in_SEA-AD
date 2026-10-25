# R12_Rexach_mapping_input.py
# Prepares the Rexach et al. 2024 insula nuclei (CELLxGENE dataset aa87f914-07e0-48aa-b915-8de906c95baf)
# for MapMyCells, and records the per-nucleus QC used by the count checks.
# No correlation or other statistic of the readings is computed here.
#
# Steps:
#   1. Insula nuclei (tissue "insular cortex") and their raw counts (raw/X, integer counts).
#   2. Per nucleus: total UMIs, genes detected (counts > 0), MT UMIs (13 MT genes), and %MT.
#   3. The 500-gene floor of the readings (Section 1). The authors' QC is kept as published, so no
#      %MT cut and no doublet removal are added here.
#   4. Donor checks: one insula library per donor (library_id against donor_id), and the age format
#      in development_stage.
#   5. MapMyCells input: the floor-passing nuclei as gzip h5ad files, split by donor into parts below
#      the 2 GB upload limit of the MapMyCells web app. X holds the raw integer counts, genes are
#      Ensembl IDs (var index), and cells are the CELLxGENE cell IDs (obs index), both written as
#      plain strings that the MapMyCells server can read.
#
# Usage (WSL, the same venv as R11; openpyxl is needed for the summary workbook):
#   ~/venv_synapse/bin/pip install -q openpyxl
#   cd /mnt/c/Users/Wong/Desktop/SEA-AD/External_Dataset_2
#   ~/venv_synapse/bin/python R12_Rexach_mapping_input.py
# Without an argument, it uses the only .h5ad file in the script's folder.
# Outputs, in the folder R12_mapmycells_input next to the .h5ad:
#   R12_Rexach_MapMyCells_part<k>.h5ad   MapMyCells input, one file per part
#   R12_Rexach_nuclei.csv.gz             one row per insula nucleus, with its QC and part
#   R12_Rexach_summary.xlsx              sheets Donors, Floor_by_cell_type, Libraries, Parts,
#                                        Settings, and Session
#   R12_Rexach_log.txt                   console output, with progress every 30 s

import glob
import importlib.metadata as im
import os
import platform
import re
import sys
import time

import anndata as ad
import h5py
import numpy as np
import pandas as pd
import scipy.sparse as sp

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    sys.exit("Install openpyxl first: ~/venv_synapse/bin/pip install -q openpyxl")

SET = {
    "tissue": "insular cortex",
    "min_genes": 500,          # readings Section 1
    "part_cells": 80000,       # nuclei per MapMyCells part (whole donors are kept together)
    "max_bytes": 1_950_000_000,  # below the 2 GB MapMyCells upload limit
    "chunk_rows": 10000,       # rows read from the file at a time
    "progress_s": 30,
}
MT13 = ["MT-ND1", "MT-ND2", "MT-CO1", "MT-CO2", "MT-ATP8", "MT-ATP6", "MT-CO3", "MT-ND3",
        "MT-ND4L", "MT-ND4", "MT-ND5", "MT-ND6", "MT-CYB"]
MICROGLIA = "microglial cell"

if len(sys.argv) > 1:
    H5AD = sys.argv[1]
else:
    here = os.path.dirname(os.path.abspath(__file__))
    found = glob.glob(os.path.join(here, "*.h5ad"))
    if len(found) != 1:
        sys.exit("Give the .h5ad path as the argument (found %d .h5ad files here)." % len(found))
    H5AD = found[0]
OUT = os.path.join(os.path.dirname(os.path.abspath(H5AD)), "R12_mapmycells_input")
os.makedirs(OUT, exist_ok=True)
log = open(os.path.join(OUT, "R12_Rexach_log.txt"), "w", encoding="utf-8")
pd.set_option("display.width", 200)
pd.set_option("display.max_rows", 200)
pd.set_option("display.max_columns", 30)


def say(*parts):
    s = " ".join(str(p) for p in parts)
    print(s, flush=True)
    log.write(s + "\n")
    log.flush()


t0 = time.time()
started = time.strftime("%Y-%m-%d %H:%M:%S")
say("R12_Rexach_mapping_input.py started", started)
say("File:", H5AD, "|", round(os.path.getsize(H5AD) / 1e9, 2), "GB")
say("Output folder:", OUT)

# ---- 1. Insula nuclei and raw counts -----------------------------------------------------
f = h5py.File(H5AD, "r")
A = ad.read_h5ad(H5AD, backed="r")  # obs and var in memory; matrices stay on disk
obs = A.obs
counts_path = "raw/X" if "raw/X" in f else "X"
g = f[counts_path]
if g.attrs.get("encoding-type") != "csr_matrix":
    sys.exit("STOPPED: %s is %s; this script reads csr_matrix only." %
             (counts_path, g.attrs.get("encoding-type")))
var = A.raw.var if (counts_path == "raw/X" and A.raw is not None) else A.var
gene_ids = np.asarray(var.index.astype(str).tolist(), dtype=object)
symbols = np.asarray(var["feature_name"].astype(str).tolist(), dtype=object)
n_genes_total = int(g.attrs["shape"][1])
if len(gene_ids) != n_genes_total:
    sys.exit("STOPPED: %d genes in var, %d columns in %s." % (len(gene_ids), n_genes_total, counts_path))
say("\n== 1. Counts and genes")
say("Counts from:", counts_path, "| genes:", n_genes_total)
say("Gene IDs starting with ENSG: %d of %d" % (sum(s.startswith("ENSG") for s in gene_ids),
                                              len(gene_ids)))
if len(set(gene_ids)) != len(gene_ids):
    sys.exit("STOPPED: the gene IDs are not unique.")
i_mt = np.where(np.isin(symbols, MT13))[0]
say("MT genes found:", len(i_mt), "of 13")
if len(i_mt) != 13:
    sys.exit("STOPPED: expected 13 MT genes.")
is_mt = np.zeros(n_genes_total, dtype=bool)
is_mt[i_mt] = True

ins = (obs["tissue"].astype(str) == SET["tissue"]).values
rows_ins = np.where(ins)[0]
cell_ids = np.asarray(obs.index.astype(str).tolist(), dtype=object)
say("Insula nuclei:", len(rows_ins))
if len(set(cell_ids[rows_ins])) != len(rows_ins):
    sys.exit("STOPPED: the insula cell IDs are not unique.")

oi = pd.DataFrame({
    "row": rows_ins,
    "cell_id": cell_ids[rows_ins],
    "donor_id": obs["donor_id"].astype(str).to_numpy(dtype=object)[rows_ins],
    "library_id": obs["library_id"].astype(str).to_numpy(dtype=object)[rows_ins],
    "disease": obs["disease"].astype(str).to_numpy(dtype=object)[rows_ins],
    "sex": obs["sex"].astype(str).to_numpy(dtype=object)[rows_ins],
    "development_stage": obs["development_stage"].astype(str).to_numpy(dtype=object)[rows_ins],
    "cell_type": obs["cell_type"].astype(str).to_numpy(dtype=object)[rows_ins],
    "ct_subcluster": obs["ct_subcluster"].astype(str).to_numpy(dtype=object)[rows_ins]
    if "ct_subcluster" in obs.columns else "",
})

# ---- Donor checks (before the counts are read, so a problem shows early) ---------------------
say("\n== 2. Donors and libraries")
lib = oi.groupby("library_id").agg(n_donors=("donor_id", "nunique"),
                                   donors=("donor_id", lambda s: ";".join(sorted(s.unique()))),
                                   nuclei=("cell_id", "size")).reset_index()
per_donor_libs = oi.groupby("donor_id")["library_id"].nunique()
one_to_one = bool((lib["n_donors"] == 1).all() and (per_donor_libs == 1).all())
say("Insula libraries:", lib.shape[0], "| donors:", per_donor_libs.shape[0])
say("One library per donor and one donor per library:", one_to_one)
if not one_to_one:
    say("  Donors with more than one library:",
        ", ".join(per_donor_libs[per_donor_libs > 1].index) or "none")
    say("  Libraries with more than one donor:",
        ", ".join(lib.loc[lib["n_donors"] > 1, "library_id"]) or "none")

stages = oi.groupby("donor_id")["development_stage"].agg(lambda s: ";".join(sorted(s.unique())))


def parse_age(stage):
    m = re.match(r"^(\d+)-year-old", stage)
    return int(m.group(1)) if m else np.nan


ages = stages.map(parse_age)
say("development_stage values:", ", ".join(sorted(oi["development_stage"].unique())))
say("Donors with an age in years:", int(ages.notna().sum()), "of", len(ages))
if ages.isna().any():
    say("  WARNING: no age in years for:", ", ".join(ages[ages.isna()].index),
        "(their stage is a range; age at death for them is taken from Table S1)")

# ---- Parts: whole donors, in donor order, up to part_cells nuclei each ----------------------
donor_n = oi.groupby("donor_id").size().sort_index()
part_of, part, n_in = {}, 1, 0
for d, n in donor_n.items():
    if n_in > 0 and n_in + n > SET["part_cells"]:
        part, n_in = part + 1, 0
    part_of[d] = part
    n_in += n
oi["part"] = oi["donor_id"].map(part_of)
n_parts = part
say("MapMyCells parts:", n_parts, "| up to", SET["part_cells"], "nuclei each, whole donors kept together")

# ---- 2, 3, and 5. Read counts per part, QC, floor, and write the MapMyCells input ------------
indptr = g["indptr"][:]
data_ds, ind_ds = g["data"], g["indices"]


def read_block(rows):
    """Rows of the CSR count matrix as a CSR matrix with int32 counts (rows must be sorted)."""
    breaks = np.where(np.diff(rows) != 1)[0] + 1
    dparts, iparts, lens = [], [], []
    for run in np.split(rows, breaks):
        a, b = indptr[run[0]], indptr[run[-1] + 1]
        dparts.append(data_ds[a:b])
        iparts.append(ind_ds[a:b])
        lens.append(np.diff(indptr[run[0]:run[-1] + 2]))
    data = np.concatenate(dparts)
    if data.size and not np.all(np.equal(np.mod(data, 1), 0)):
        sys.exit("STOPPED: non-integer values in %s." % counts_path)
    m = sp.csr_matrix((data.astype(np.int32), np.concatenate(iparts),
                       np.concatenate([[0], np.cumsum(np.concatenate(lens))])),
                      shape=(len(rows), n_genes_total))
    m.eliminate_zeros()
    return m


say("\n== 3. Counts, QC, floor, and MapMyCells files")
qc_parts, part_rows = [], []
last = time.time()
for p in range(1, n_parts + 1):
    sub = oi[oi["part"] == p].sort_values("row")
    rows_p = sub["row"].values
    blocks, tot_l, ng_l, mt_l = [], [], [], []
    for s in range(0, len(rows_p), SET["chunk_rows"]):
        m = read_block(rows_p[s:s + SET["chunk_rows"]])
        tot = np.asarray(m.sum(axis=1)).ravel()
        ng = np.diff(m.indptr)
        mt = np.asarray(m[:, is_mt].sum(axis=1)).ravel()
        blocks.append(m[ng >= SET["min_genes"]])
        tot_l.append(tot); ng_l.append(ng); mt_l.append(mt)
        if time.time() - last >= SET["progress_s"]:
            say("  part %d: %d of %d nuclei read | %.0f s" %
                (p, min(s + SET["chunk_rows"], len(rows_p)), len(rows_p), time.time() - t0))
            last = time.time()
    sub = sub.copy()
    sub["total_umis"] = np.concatenate(tot_l)
    sub["n_genes"] = np.concatenate(ng_l)
    sub["mt_umis"] = np.concatenate(mt_l)
    sub["pct_mt"] = 100 * sub["mt_umis"] / np.maximum(sub["total_umis"], 1)
    sub["pass_floor"] = sub["n_genes"] >= SET["min_genes"]
    qc_parts.append(sub)

    kept = sub[sub["pass_floor"]]
    X = sp.vstack(blocks, format="csr")
    # The cell and gene IDs are written as plain (object) strings with no other obs or var
    # columns. With pandas 3, anndata otherwise writes them as "nullable-string-array", which the
    # older anndata on the MapMyCells server cannot read (the first upload failed on this).
    adata = ad.AnnData(X=X, obs=pd.DataFrame(index=pd.Index(list(kept["cell_id"]), dtype=object)),
                       var=pd.DataFrame(index=pd.Index(list(gene_ids), dtype=object)))
    name = "R12_Rexach_MapMyCells_part%d.h5ad" % p
    tmp = os.path.join(OUT, "tmp_" + name)
    say("  part %d: writing %s (%d nuclei, %d donors) | %.0f s" %
        (p, name, X.shape[0], kept["donor_id"].nunique(), time.time() - t0))
    adata.write_h5ad(tmp, compression="gzip")
    with h5py.File(tmp, "r") as chk:
        enc = {k: chk[k + "/" + chk[k].attrs["_index"]].attrs.get("encoding-type") for k in ("obs", "var")}
    if any(v != "string-array" for v in enc.values()):
        sys.exit("STOPPED: part %d has index encodings %s; MapMyCells needs string-array." % (p, enc))
    os.replace(tmp, os.path.join(OUT, name))
    size = os.path.getsize(os.path.join(OUT, name))
    say("  part %d: %.2f GB" % (p, size / 1e9))
    if size > SET["max_bytes"]:
        say("  WARNING: part %d is above the upload limit. Lower part_cells in SET and run again." % p)
    part_rows.append({"part": p, "file": name, "donors": ";".join(sorted(kept["donor_id"].unique())),
                      "n_donors": kept["donor_id"].nunique(), "nuclei": X.shape[0],
                      "bytes": size, "below_limit": size <= SET["max_bytes"]})
    del adata, X, blocks
f.close()

qc = pd.concat(qc_parts).sort_values("row")
qc.drop(columns=["row"]).to_csv(os.path.join(OUT, "R12_Rexach_nuclei.csv.gz"), index=False,
                                compression="gzip")

# ---- Summaries ---------------------------------------------------------------------------------
say("\n== 4. Summary")
k = qc[qc["pass_floor"]]
say("Insula nuclei: %d | removed by the %d-gene floor: %d (%.1f%%) | kept: %d" %
    (len(qc), SET["min_genes"], int((~qc["pass_floor"]).sum()),
     100 * (~qc["pass_floor"]).mean(), len(k)))
mg = qc["cell_type"] == MICROGLIA
say("Authors' microglia: %d | removed by the floor: %d (%.1f%%) | kept: %d" %
    (int(mg.sum()), int((mg & ~qc["pass_floor"]).sum()),
     100 * (mg & ~qc["pass_floor"]).sum() / max(int(mg.sum()), 1), int((mg & qc["pass_floor"]).sum())))
say("Kept nuclei: median UMIs %.0f | median genes %.0f | median %%MT %.2f | max %%MT %.2f" %
    (k["total_umis"].median(), k["n_genes"].median(), k["pct_mt"].median(), k["pct_mt"].max()))

fl = qc.groupby("cell_type").agg(nuclei=("cell_id", "size"),
                                 removed=("pass_floor", lambda s: int((~s).sum()))).reset_index()
fl["pct_removed"] = (100 * fl["removed"] / fl["nuclei"]).round(2)
fl = fl.sort_values("nuclei", ascending=False)
say("\nFloor by the authors' cell type:")
say(fl.to_string(index=False))

dn = qc.groupby("donor_id").agg(
    disease=("disease", "first"), sex=("sex", "first"),
    library_ids=("library_id", lambda s: ";".join(sorted(s.unique()))),
    n_libraries=("library_id", "nunique"), part=("part", "first"),
    insula_nuclei=("cell_id", "size"),
    nuclei_kept=("pass_floor", "sum"),
    author_microglia=("cell_type", lambda s: int((s == MICROGLIA).sum()))).reset_index()
dn["development_stage"] = dn["donor_id"].map(stages)
dn["age_years"] = dn["donor_id"].map(ages)
dn["pct_removed"] = (100 * (1 - dn["nuclei_kept"] / dn["insula_nuclei"])).round(2)
mk = qc[mg & qc["pass_floor"]].groupby("donor_id")
dn["author_microglia_kept"] = dn["donor_id"].map(mk.size()).fillna(0).astype(int)
kd = k.groupby("donor_id")
dn["median_umis_kept"] = dn["donor_id"].map(kd["total_umis"].median())
dn["median_genes_kept"] = dn["donor_id"].map(kd["n_genes"].median())
dn["median_pct_mt_kept"] = dn["donor_id"].map(kd["pct_mt"].median()).round(3)
dn = dn[["donor_id", "disease", "sex", "development_stage", "age_years", "library_ids", "n_libraries",
         "insula_nuclei", "nuclei_kept", "pct_removed", "author_microglia", "author_microglia_kept",
         "median_umis_kept", "median_genes_kept", "median_pct_mt_kept", "part"]]
dn = dn.sort_values(["disease", "donor_id"])
say("\nDonors:")
say(dn.drop(columns=["library_ids", "development_stage"]).to_string(index=False))

pr = pd.DataFrame(part_rows)
settings = pd.DataFrame({"setting": list(SET.keys()) + ["counts", "MT genes", "input"],
                         "value": [str(v) for v in SET.values()] +
                                  [counts_path, ", ".join(MT13), os.path.basename(H5AD)]})
pk = ["anndata", "h5py", "numpy", "pandas", "scipy", "openpyxl"]
session = pd.DataFrame({"item": ["python", "platform", "started", "finished"] + pk,
                        "value": [sys.version.split()[0], platform.platform(), started,
                                  time.strftime("%Y-%m-%d %H:%M:%S")] + [im.version(x) for x in pk]})


def write_xlsx(path, sheets):
    """Plain workbook: bold header and a frozen top row, no fills or borders."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, df in sheets:
        ws = wb.create_sheet(name)
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                       (v.item() if hasattr(v, "item") else v) for v in r])
        ws.freeze_panes = "A2"
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


write_xlsx(os.path.join(OUT, "R12_Rexach_summary.xlsx"),
           [("Donors", dn), ("Floor_by_cell_type", fl), ("Libraries", lib), ("Parts", pr),
            ("Settings", settings), ("Session", session)])
say("\nUpload each R12_Rexach_MapMyCells_part<k>.h5ad to MapMyCells (10x Human MTG SEA-AD,"
    " CCN20230505), once with Deep Generative Mapping and once with Hierarchical Mapping.")
say("Done in %.0f s" % (time.time() - t0))
log.close()
