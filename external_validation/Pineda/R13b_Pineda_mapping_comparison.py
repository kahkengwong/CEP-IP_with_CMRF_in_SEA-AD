# R13b_Pineda_mapping_comparison.py
# Compares two MapMyCells runs of the Pineda nuclei, nucleus by nucleus. By default:
#   old: all R02 kept nuclei mapped before the authors' QC (R13_mapmycells_output_allKept/)
#   new: only the nuclei in the authors' published QC (R13_mapmycells_output/)
# A repeat run of the same input gives the run-to-run baseline to compare against.
# The labels of a nucleus can depend on the other nuclei in the same file, and Deep Generative Mapping
# (DGM) can also differ between repeat runs of the same file. This script measures the agreement of
# DGM and Hierarchical Mapping (HM) labels and probabilities between the two runs.
#
# Usage (WSL, the same venv as R12 and R13):
#   cd /mnt/c/Users/Wong/Desktop/SEA-AD/External_Dataset_1
#   ~/venv_synapse/bin/python R13b_Pineda_mapping_comparison.py
# Outputs, in the folder R13b_mapping_comparison:
#   R13b_Pineda_mapping_comparison.xlsx   sheets Files, Agreement, Probability, PVM3_DGM, PVM3_HM,
#                                         Changed_DGM, and Session
#   R13b_Pineda_log.txt                   console output

import glob
import importlib.metadata as im
import os
import platform
import re
import sys
import time

import numpy as np
import pandas as pd

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    sys.exit("Install openpyxl first: ~/venv_synapse/bin/pip install -q openpyxl")

# The data folder: PINEDA_ROOT when set (the container sets it to /data), otherwise the script folder.
HERE = os.environ.get("PINEDA_ROOT", os.path.dirname(os.path.abspath(__file__)))
# Folders can be given on the command line: old folder, new folder, and a label for the output folder.
#   python R13b_Pineda_mapping_comparison.py R13_mapmycells_output_allKept R13_mapmycells_output_allKept_repeat repeat
ARGS = sys.argv[1:]
DIRS = {"old": os.path.join(HERE, ARGS[0] if len(ARGS) > 0 else "R13_mapmycells_output_allKept"),
        "new": os.path.join(HERE, ARGS[1] if len(ARGS) > 1 else "R13_mapmycells_output")}
OUT = os.path.join(HERE, "R13b_mapping_comparison" + ("_" + ARGS[2] if len(ARGS) > 2 else ""))
os.makedirs(OUT, exist_ok=True)
ALGS = {"DeepGenerativeMapping": "DGM", "SEA-ADHierarchicalMapping": "HM"}
ARM = "Micro-PVM_3-SEAAD"
MICRO = "Microglia-PVM"
log = open(os.path.join(OUT, "R13b_Pineda_log.txt"), "w", encoding="utf-8")
pd.set_option("display.width", 220)
pd.set_option("display.max_columns", 40)
pd.set_option("display.max_rows", 200)


def say(*parts):
    s = " ".join(str(p) for p in parts)
    print(s, flush=True)
    log.write(s + "\n")
    log.flush()


def stop(msg):
    say("STOPPED:", msg)
    log.close()
    sys.exit(1)


t0 = time.time()
say("R13b_Pineda_mapping_comparison.py started", time.strftime("%Y-%m-%d %H:%M:%S"))
say("old:", DIRS["old"])
say("new:", DIRS["new"])

# ---- Read both runs ---------------------------------------------------------------------------
pat = re.compile(r"part(\d+)_.*_(DeepGenerativeMapping|SEA-ADHierarchicalMapping)_UTC_\d+\.csv$")
runs, file_rows = {}, []
for run, folder in DIRS.items():
    if not os.path.isdir(folder):
        stop("folder not found: %s" % folder)
    tabs = {"DGM": [], "HM": []}
    for p in sorted(glob.glob(os.path.join(folder, "*.csv"))):
        m = pat.search(os.path.basename(p))
        if not m:
            continue
        a = ALGS[m.group(2)]
        with open(p, encoding="utf-8") as fh:
            header = next((ln.strip() for ln in fh if ln.startswith("# algorithm")), "")
        d = pd.read_csv(p, comment="#")
        tabs[a].append(d)
        file_rows.append({"run": run, "part": int(m.group(1)), "algorithm": a, "file": os.path.basename(p),
                          "nuclei": len(d), "header": header})
    for a in ("DGM", "HM"):
        if not tabs[a]:
            stop("no %s csv files in %s" % (a, folder))
        d = pd.concat(tabs[a], ignore_index=True)
        if d["cell_id"].duplicated().any():
            stop("%s %s: duplicated cell IDs" % (run, a))
        cols = {"cell_id": d["cell_id"].astype(str)}
        for lv in ("class", "subclass", "supertype"):
            cols[lv] = d[lv + "_name"].astype(str)
            probs = [c for c in d.columns if c.startswith(lv + "_") and c.endswith("probability")]
            cols[lv + "_prob"] = pd.to_numeric(d[probs[0]], errors="coerce") if probs else np.nan
        runs[(run, a)] = pd.DataFrame(cols).set_index("cell_id")
    say("%s run: %d DGM files (%d nuclei) | %d HM files (%d nuclei)" %
        (run, len(tabs["DGM"]), len(runs[(run, "DGM")]), len(tabs["HM"]), len(runs[(run, "HM")])))
files = pd.DataFrame(file_rows)
for (run, a), d in runs.items():
    hs = files.loc[(files["run"] == run) & (files["algorithm"] == a), "header"].unique()
    say("  %s %s header: %s" % (run, a, " | ".join(hs)))

# ---- Nucleus-by-nucleus agreement ------------------------------------------------------------------
say("\n== Agreement on the nuclei in both runs")
agree_rows, prob_rows, trans, changed = [], [], {}, None
for a in ("DGM", "HM"):
    old, new = runs[("old", a)], runs[("new", a)]
    common = new.index.intersection(old.index)
    not_in_old = len(new.index.difference(old.index))
    if not_in_old:
        say("WARNING: %d %s nuclei of the new run are not in the old run" % (not_in_old, a))
    o, n = old.loc[common], new.loc[common]
    for lv in ("class", "subclass", "supertype"):
        same = (o[lv] == n[lv])
        agree_rows.append({"algorithm": a, "level": lv, "nuclei_in_both": len(common),
                           "same_label": int(same.sum()), "pct_same": round(100 * same.mean(), 4),
                           "changed": int((~same).sum())})
        po, pn = o[lv + "_prob"], n[lv + "_prob"]
        ok = po.notna() & pn.notna()
        if ok.any():
            dif = (pn[ok] - po[ok]).abs()
            prob_rows.append({"algorithm": a, "level": lv, "nuclei": int(ok.sum()),
                              "max_abs_diff": float(dif.max()), "median_abs_diff": float(dif.median()),
                              "share_identical": round(float((dif == 0).mean()), 6),
                              "pearson_r": float(np.corrcoef(po[ok], pn[ok])[0, 1])
                              if ok.sum() > 2 and po[ok].std() > 0 and pn[ok].std() > 0 else np.nan})
    # Micro-PVM_3 among the Microglia-PVM nuclei: old label against new label.
    mg = (o["subclass"] == MICRO) | (n["subclass"] == MICRO)
    t = pd.crosstab(o.loc[mg, "supertype"].rename("old_supertype"), n.loc[mg, "supertype"].rename("new_supertype"))
    trans[a] = t.reset_index()
    n_old, n_new = int((o["supertype"] == ARM).sum()), int((n["supertype"] == ARM).sum())
    both = int(((o["supertype"] == ARM) & (n["supertype"] == ARM)).sum())
    say("%s: Micro-PVM_3 in the old run %d | in the new run %d | in both %d (of the nuclei in both runs)" %
        (a, n_old, n_new, both))
    agree_rows.append({"algorithm": a, "level": "Micro-PVM_3 (old | new | both)", "nuclei_in_both": len(common),
                       "same_label": both, "pct_same": round(100 * both / max(n_old, n_new, 1), 4),
                       "changed": (n_old - both) + (n_new - both)})
    if a == "DGM":
        ch = ~(o["supertype"] == n["supertype"])
        changed = pd.DataFrame({"cell_id": common[ch.values], "old_supertype": o.loc[ch, "supertype"].values,
                                "new_supertype": n.loc[ch, "supertype"].values,
                                "old_prob": o.loc[ch, "supertype_prob"].values,
                                "new_prob": n.loc[ch, "supertype_prob"].values})
agr = pd.DataFrame(agree_rows)
prb = pd.DataFrame(prob_rows)
say(agr.to_string(index=False))
say("\nProbabilities (new minus old, absolute):")
say(prb.to_string(index=False))
dgm_sup = agr[(agr["algorithm"] == "DGM") & (agr["level"] == "supertype")].iloc[0]
say("\nDGM supertype labels changed for %d of %d nuclei (%.4f%%)." %
    (dgm_sup["changed"], dgm_sup["nuclei_in_both"], 100 - dgm_sup["pct_same"]))
if changed is not None and len(changed):
    say("Changed DGM supertypes, most frequent:")
    say(changed.groupby(["old_supertype", "new_supertype"]).size().sort_values(ascending=False).head(15).to_string())


# ---- Save ----------------------------------------------------------------------------------------
def write_xlsx(path, sheets):
    """Plain workbook: bold header and a frozen top row, no fills or borders."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, df in sheets:
        ws = wb.create_sheet(name)
        ws.append([str(c) for c in df.columns])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                       (v.item() if hasattr(v, "item") else v) for v in r])
        ws.freeze_panes = "A2"
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


pk = ["numpy", "pandas", "openpyxl"]
session = pd.DataFrame({"item": ["python", "platform", "finished"] + pk,
                        "value": [sys.version.split()[0], platform.platform(), time.strftime("%Y-%m-%d %H:%M:%S")] +
                                 [im.version(x) for x in pk]})
write_xlsx(os.path.join(OUT, "R13b_Pineda_mapping_comparison.xlsx"),
           [("Files", files), ("Agreement", agr), ("Probability", prb), ("PVM3_DGM", trans["DGM"]),
            ("PVM3_HM", trans["HM"]), ("Changed_DGM", changed if changed is not None else pd.DataFrame()),
            ("Session", session)])
say("\nWritten to %s | done in %.0f s" % (OUT, time.time() - t0))
log.close()
