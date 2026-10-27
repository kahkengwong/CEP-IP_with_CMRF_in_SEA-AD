# R13c_Pineda_noAuthorQC_arm.py
# Rebuilds the PVM_3 arm of the earlier Pineda run without the authors' QC, for Supplementary Table
# S20 and the Methods sentence that cites it. That run mapped all R02 kept nuclei with MapMyCells
# (R13_mapmycells_output_allKept/); its arm is the Micro-PVM_3 nuclei by Deep Generative Mapping with
# MT-CO3 > 0. The script counts the arm nuclei by %MT band and by donor, and flags whether each nucleus
# is in the authors' published QC. No statistic is computed.
#
# Usage (WSL, the same venv as R12 and R13):
#   cd /mnt/c/Users/Wong/Desktop/SEA-AD/External_Dataset_1
#   ~/venv_synapse/bin/python R13c_Pineda_noAuthorQC_arm.py
# Inputs: R11_nuclei/R11_Pineda_nuclei.csv.gz and R11_Pineda_genes.csv.gz, the Deep Generative Mapping
#   csv files in R13_mapmycells_output_allKept/, counts/<LIB>_raw_feature_bc_matrix.h5, and the
#   authors' metadata in Pineda2024_metadata/.
# Outputs, in the folder R13c_noAuthorQC_arm:
#   R13c_Pineda_noAuthorQC_arm.xlsx   sheets Bands, Donors, Settings, and Session
#   R13c_Pineda_log.txt               console output

import glob
import importlib.metadata as im
import os
import platform
import re
import sys
import time

import h5py
import numpy as np
import pandas as pd

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    sys.exit("Install openpyxl first: ~/venv_synapse/bin/pip install -q openpyxl")

# The data folder: PINEDA_ROOT when set (the container sets it to /data), otherwise the script folder.
HERE = os.environ.get("PINEDA_ROOT", os.path.dirname(os.path.abspath(__file__)))
F_NUC = os.path.join(HERE, "R11_nuclei", "R11_Pineda_nuclei.csv.gz")
F_GENES = os.path.join(HERE, "R11_nuclei", "R11_Pineda_genes.csv.gz")
MAP_DIR = os.path.join(HERE, "R13_mapmycells_output_allKept")
COUNTS = os.path.join(HERE, "counts")
AUTHOR_DIR = os.path.join(HERE, "Pineda2024_metadata")
OUT = os.path.join(HERE, "R13c_noAuthorQC_arm")
os.makedirs(OUT, exist_ok=True)
SET = {"arm_supertype": "Micro-PVM_3-SEAAD", "goi": "MT-CO3", "bands_pct_mt": [0, 1, 5, 10, 20, 100],
       # The earlier run (log of 2026-10-27, first R13 run) had 2,795 arm nuclei, 239 of them above 20% MT.
       "expected_arm": 2795, "expected_above_20": 239}
log = open(os.path.join(OUT, "R13c_Pineda_log.txt"), "w", encoding="utf-8")
pd.set_option("display.width", 200)
pd.set_option("display.max_rows", 200)


def say(*parts):
    s = " ".join(str(p) for p in parts)
    print(s, flush=True)
    log.write(s + "\n")
    log.flush()


def stop(msg):
    say("STOPPED:", msg)
    log.close()
    sys.exit(1)


t0 = time.time()
say("R13c_Pineda_noAuthorQC_arm.py started", time.strftime("%Y-%m-%d %H:%M:%S"))
for p in (F_NUC, F_GENES, MAP_DIR, COUNTS, AUTHOR_DIR):
    if not os.path.exists(p):
        stop("not found: " + p)

# ---- Kept nuclei and their Deep Generative Mapping labels (earlier run) -------------------------
nuc = pd.read_csv(F_NUC, dtype={"donor_id": str, "barcode": str, "cell_id": str})
kept = nuc[nuc["keep"]].copy()
pat = re.compile(r"part(\d+)_.*_DeepGenerativeMapping_UTC_\d+\.csv$")
files = [p for p in sorted(glob.glob(os.path.join(MAP_DIR, "*.csv"))) if pat.search(os.path.basename(p))]
if not files:
    stop("no Deep Generative Mapping csv files in " + MAP_DIR)
dgm = pd.concat([pd.read_csv(p, comment="#", usecols=["cell_id", "supertype_name"]) for p in files],
                ignore_index=True)
dgm["cell_id"] = dgm["cell_id"].astype(str)
if dgm["cell_id"].duplicated().any():
    stop("duplicated cell IDs in the mapping files")
say("Kept nuclei (R02): %d | mapped in the earlier run: %d (%d files)" % (len(kept), len(dgm), len(files)))
if set(dgm["cell_id"]) != set(kept["cell_id"]):
    stop("the earlier mapping does not cover exactly the R02 kept nuclei")
kept = kept.merge(dgm, on="cell_id", how="left")
pvm3 = kept[kept["supertype_name"] == SET["arm_supertype"]].copy()
say("Micro-PVM_3 (DGM) in the earlier run: %d" % len(pvm3))

# ---- Authors' published QC --------------------------------------------------------------------------
meta = {}
for p in glob.glob(os.path.join(AUTHOR_DIR, "**", "*_col_metadata.tsv"), recursive=True):
    meta.setdefault(os.path.basename(p)[:-len("_col_metadata.tsv")], p)
pvm3["bc16"] = pvm3["barcode"].str.split("-").str[0]
pvm3["in_author_qc"] = False
for lib, s in pvm3.groupby("library_id"):
    if lib not in meta:
        stop("authors' metadata missing for " + lib)
    b16 = set(pd.read_csv(meta[lib], sep="\t", dtype=str, usecols=["Barcode"])["Barcode"].str.split("-").str[0])
    pvm3.loc[s.index, "in_author_qc"] = s["bc16"].isin(b16).values

# ---- MT-CO3 UMIs of these nuclei from the raw .h5 ---------------------------------------------------
genes = pd.read_csv(F_GENES)
gi = int(np.where(genes["symbol"].astype(str).values == SET["goi"])[0][0])
pvm3["mt_co3_umis"] = 0
for lib, s in pvm3.groupby("library_id"):
    with h5py.File(os.path.join(COUNTS, lib + "_raw_feature_bc_matrix.h5"), "r") as f:
        g = f["matrix"]
        bc = pd.Index([x.decode() for x in g["barcodes"][:]])
        pos = bc.get_indexer(s["barcode"].tolist())
        if (pos < 0).any():
            stop("%s: barcodes not found in the raw .h5" % lib)
        indptr = g["indptr"][:]
        vals = []
        for p_ in pos:
            idx = g["indices"][indptr[p_]:indptr[p_ + 1]]
            dat = g["data"][indptr[p_]:indptr[p_ + 1]]
            hit = np.where(idx == gi)[0]
            vals.append(int(dat[hit[0]]) if hit.size else 0)
    pvm3.loc[s.index, "mt_co3_umis"] = vals
arm = pvm3[pvm3["mt_co3_umis"] > 0].copy()
say("PVM_3 arm of the earlier run (Micro-PVM_3 with MT-CO3 > 0): %d (expected %d)" %
    (len(arm), SET["expected_arm"]))

# ---- Counts by %MT band and by donor --------------------------------------------------------------------
edges = SET["bands_pct_mt"]
labels = ["<=%g%%" % edges[1]] + ["%g to %g%%" % (a, b) for a, b in zip(edges[1:-2], edges[2:-1])] + [">%g%%" % edges[-2]]
arm["pct_mt_band"] = pd.cut(arm["pct_mt"], edges, labels=labels, include_lowest=True)
bands = arm.groupby("pct_mt_band", observed=False).agg(
    arm_nuclei=("cell_id", "size"), in_author_qc=("in_author_qc", "sum"),
    median_genes=("n_genes", "median"), median_umis=("total_umis", "median")).reset_index()
bands["removed_by_author_qc"] = bands["arm_nuclei"] - bands["in_author_qc"]
say("\nPVM_3 arm nuclei by %MT band:")
say(bands.to_string(index=False))
n20 = int((arm["pct_mt"] > 20).sum())
say("Arm nuclei above 20%% MT: %d (expected %d)" % (n20, SET["expected_above_20"]))
donors = arm.groupby(["donor_id", "disease"]).agg(
    arm_nuclei=("cell_id", "size"), above_5pct_mt=("pct_mt", lambda v: int((v > 5).sum())),
    above_10pct_mt=("pct_mt", lambda v: int((v > 10).sum())), above_20pct_mt=("pct_mt", lambda v: int((v > 20).sum())),
    in_author_qc=("in_author_qc", "sum")).reset_index().sort_values("above_20pct_mt", ascending=False)
top5 = donors.head(5)
say("\nFive donors with the most arm nuclei above 20%% MT: %s | %d of %d (%.0f%%)" %
    (", ".join("%s %d" % (d, n) for d, n in zip(top5["donor_id"], top5["above_20pct_mt"])),
     int(top5["above_20pct_mt"].sum()), n20, 100 * top5["above_20pct_mt"].sum() / max(n20, 1)))
if len(arm) != SET["expected_arm"] or n20 != SET["expected_above_20"]:
    say("WARNING: the counts differ from the earlier run; check that R13_mapmycells_output_allKept holds "
        "the mapping of that run")


# ---- Save ----------------------------------------------------------------------------------------
def write_xlsx(path, sheets):
    """Plain workbook: bold header and a frozen top row, no fills or borders."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, df in sheets:
        ws = wb.create_sheet(name)
        ws.append([str(c) for c in df.columns])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else
                       (v.item() if hasattr(v, "item") else (str(v) if not isinstance(v, (int, float, str, bool)) else v))
                       for v in r])
        ws.freeze_panes = "A2"
    tmp = path + ".tmp"
    wb.save(tmp)
    os.replace(tmp, path)


settings = pd.DataFrame({"setting": ["arm", "mapping", "authors' QC", "%MT bands"],
                         "value": ["Micro-PVM_3 (Deep Generative Mapping) with MT-CO3 > 0, R02 kept nuclei",
                                   "R13_mapmycells_output_allKept (all R02 kept nuclei, mapped before the authors' QC)",
                                   "barcode present in the authors' <LIB>_col_metadata.tsv",
                                   ", ".join(labels)]})
pk = ["h5py", "numpy", "pandas", "openpyxl"]
session = pd.DataFrame({"item": ["python", "platform", "finished"] + pk,
                        "value": [sys.version.split()[0], platform.platform(), time.strftime("%Y-%m-%d %H:%M:%S")] +
                                 [im.version(x) for x in pk]})
write_xlsx(os.path.join(OUT, "R13c_Pineda_noAuthorQC_arm.xlsx"),
           [("Bands", bands), ("Donors", donors), ("Settings", settings), ("Session", session)])
say("\nWritten to %s | done in %.0f s" % (OUT, time.time() - t0))
log.close()
